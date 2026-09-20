# -*- coding: utf-8 -*-
"""RIZE W2 B2 — the rails for the year AFTER the goals are agreed.

The live Chrome run is what proves the pages render; these are the floor.
Every one is written against a failure this half of the module could have and
that nothing at runtime would report:

  * a monthly check-in made twice is one row per person per morning, for ever,
    with a cheerful count in the log;
  * a check-in that is never marked missed makes the compliance figure — the
    one number this whole phase exists for — a lie that always reads 100%;
  * applicability worked out from the wrong date puts somebody who arrived in
    November through a year-end review about a year they were barely in, and
    nothing anywhere says why;
  * a change request carried out on one path and not the other is a change
    that silently does not happen on a database where the route was never
    switched on (R204);
  * a score shown while one key result is unmarked is a figure somebody spends
    an afternoon on and that is about to move;
  * a closed year whose numbers are still a live compute re-answers the day
    somebody corrects a weight in a neighbouring row;
  * an employee who can score their own key results is a self-assessment
    wearing a manager's clothes.

WHY THE FIXTURES ARE NAMED "DEMO" — as B1: everything here is rolled back with
the transaction, and the names are the live convention anyway (ledger rule 9),
because a fixture copied into a live script keeps the name it was written with.
"""

import json

from datetime import date, timedelta

from odoo.exceptions import UserError, ValidationError
from odoo.tests import tagged

from .test_goals import GoalsCase


class YearCase(GoalsCase):
    """B1's world, with one sheet agreed and locked so a year can happen."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.Checkin = cls.env['pb.goal.checkin']
        cls.Review = cls.env['pb.goal.review']
        cls.Change = cls.env['pb.goal.change']
        cls.Audit = cls.env['pb.goal.audit']
        cls.Band = cls.env['pb.goal.band']
        cls.Job = cls.env['pb.goals.automation']
        cls.Numbers = cls.env['pb.goals.analytics']
        cls.Home = cls.env['pb.goals.home']

    def _agreed(self, employee=None, weights=(60.0, 40.0), joined_on=None):
        """A sheet with two weighted goals on it, agreed and locked.

        THE JOINING DATE IS SET AND RESTAMPED, and that is not fixture noise.
        `_join_date` falls back to the day the employee record was MADE, which
        in a test is today — so an employee created this morning "joined"
        three weeks before the half-way point and gets no half-way review at
        all. Every test here that is not ABOUT applicability wants somebody
        who was here for the whole year, and saying so is the difference
        between a fixture and a coincidence.
        """
        sheet = self._sheet(employee)
        sheet.sudo().write(
            {'joined_on': joined_on or (self.cycle.date_start
                                        - timedelta(days=400))})
        sheet._stamp_applicability(force=True)
        for index, weight in enumerate(weights):
            self._goal(sheet, title='DEMO goal %s' % (index + 1),
                       weight=weight, krs=2)
        self._lock(sheet)
        return sheet

    def _carry_out(self, change):
        """Take a change request all the way, WITHOUT driving the engine.

        B1's `_lock` helper and this one are the same idea and the same
        reason. A route IS published on this database, so `_advance_state`
        correctly hands each press to the engine — and the engine correctly
        refuses a test user who is not the seat holder ("This step is not
        waiting for you"). Minting seats for two approvers here would be a
        test of the ENGINE rather than of this module; the route end to end is
        proved live.

        What is being proved is that the CONSEQUENCE hangs off the STATUS
        (R204), so it is identical whether the route drove the request there
        or the record's own ladder did — which is exactly what these two lines
        replay.
        """
        change.sudo().write({'state': 'manager_ok'})
        change._after_approval_transition('approved')
        change.sudo().write({'state': 'approved'})
        change.invalidate_recordset()
        return change

    def _score_all(self, sheet, mark='4'):
        for goal in sheet.with_context(active_test=False).goal_ids:
            for kr in goal.kr_ids:
                kr.sudo().write({'score': mark})
        sheet.invalidate_recordset()
        return sheet


# =========================================================================
#  T2. The monthly conversation
# =========================================================================
@tagged('post_install', '-at_install')
class TestCheckins(YearCase):

    def test_the_job_makes_this_months_check_in_once(self):
        """A JOB THAT RUNS THIRTY TIMES A MONTH CANNOT BE PROTECTED BY A
        PYTHON CHECK ALONE, so this asserts the count after two runs."""
        sheet = self._agreed()
        first = self.Job._make_checkins()
        self.assertGreaterEqual(first, 1)
        rows = self.Checkin.search([('set_id', '=', sheet.id)])
        self.assertEqual(len(rows), 1)
        self.Job._make_checkins()
        self.assertEqual(
            len(self.Checkin.search([('set_id', '=', sheet.id)])), 1,
            'a second run made a second check-in for the same month')

    def test_two_rows_for_one_month_are_refused_by_the_database(self):
        """`_sql_constraints` as a LIST is silently ignored on Odoo 19 — the
        rule simply is not enforced and nothing says so. This proves the
        `models.Constraint` form really reached Postgres."""
        sheet = self._agreed()
        month = date(2026, 6, 1)
        self.Checkin.create({'set_id': sheet.id, 'month': month,
                             'scheduled_date': date(2026, 6, 25)})
        with self.assertRaises(Exception):
            with self.env.cr.savepoint():
                self.Checkin.create({'set_id': sheet.id, 'month': month,
                                     'scheduled_date': date(2026, 6, 25)})


    def test_running_it_twice_reports_one_then_nothing(self):
        """R90 — ONLY THE COUNT LIES, AND THAT IS ENOUGH TO BREAK IT.

        "Did this row get made just now" cannot be asked of `create_date`: a
        row made at six o'clock this morning was made TODAY, so a second run
        on the same day reported one made every time over a table it had not
        touched — the rows identical, the number a fiction, and the job's
        whole claim being that it is idempotent. Found live by pressing "run
        it now" twice.
        """
        self._agreed()
        self.assertEqual(self.Job._make_checkins(), 1)
        self.assertEqual(self.Job._make_checkins(), 0,
                         'the second run reported work it did not do')

    def test_the_review_count_is_honest_on_a_second_run(self):
        sheet = self._agreed()
        near = self.cycle.mid_year_date - timedelta(days=10)
        self.assertEqual(self.Job._make_reviews(near), 1)
        self.assertEqual(self.Job._make_reviews(near), 0)

    def test_a_check_in_belongs_to_a_whole_month(self):
        sheet = self._agreed()
        with self.assertRaises(ValidationError):
            self.Checkin.create({'set_id': sheet.id,
                                 'month': date(2026, 6, 14),
                                 'scheduled_date': date(2026, 6, 25)})

    def test_the_planned_day_comes_from_the_goal_year(self):
        """THE YEAR'S OWN DAY WINS over the product-wide switch, and it is
        capped at the 28th so it is the same day every month."""
        self.cycle.write({'checkin_day': 31})
        sheet = self._agreed()
        row = self.Checkin.ensure_for(sheet, date(2026, 2, 3))
        self.assertEqual(row.scheduled_date.day, 28)
        self.cycle.write({'checkin_day': 10})
        other = self.Checkin.ensure_for(sheet, date(2026, 6, 3))
        self.assertEqual(other.scheduled_date, date(2026, 6, 10))

    def test_a_note_is_required_and_the_snapshot_is_taken(self):
        """A CHECK-IN WITH NOTHING ON IT IS A TICK IN A BOX, and a tick in a
        box cannot be read next March."""
        sheet = self._agreed()
        row = self.Checkin.ensure_for(sheet)
        with self.assertRaises(UserError):
            row.action_checkin_done('   ')
        kr = sheet.goal_ids[0].kr_ids[0]
        kr.write({'progress': 40.0, 'current': 20.0})
        row.action_checkin_done('DEMO the north depot opened.',
                                'DEMO waiting on the second van.')
        self.assertEqual(row.state, 'done')
        self.assertTrue(row.done_at)
        self.assertEqual(row.done_by_id, self.env.user)
        rows = row.snapshot_rows()
        self.assertTrue(rows)
        self.assertTrue(any(abs(r['progress'] - 40.0) < 0.01 for r in rows),
                        'the snapshot did not freeze where the numbers stood')

    def test_a_snapshot_does_not_move_when_the_numbers_do(self):
        """THE SNAPSHOT IS WHAT MAKES IT EVIDENCE. The live figures keep
        moving; the frozen ones do not."""
        sheet = self._agreed()
        row = self.Checkin.ensure_for(sheet)
        kr = sheet.goal_ids[0].kr_ids[0]
        kr.write({'progress': 30.0})
        row.action_checkin_done('DEMO a third of the way.')
        kr.write({'progress': 90.0})
        frozen = [r for r in row.snapshot_rows() if r['kr'] == kr.title]
        self.assertAlmostEqual(frozen[0]['progress'], 30.0, places=1)

    def test_either_side_may_write_it_up_and_a_stranger_may_not(self):
        sheet = self._agreed()
        row = self.Checkin.ensure_for(sheet)
        with self.assertRaises(UserError):
            row.with_user(self.stranger_user).action_checkin_done('DEMO no.')
        row.with_user(self.staff_user).action_checkin_done('DEMO from them.')
        self.assertEqual(row.state, 'done')

    def test_a_month_that_ended_with_nothing_on_it_is_missed(self):
        """MISSED IS A REAL OUTCOME. A row that quietly stays planned for ever
        says nothing at all, and the compliance figure then always reads 100%
        because it has nothing to count against."""
        sheet = self._agreed()
        old = self.Checkin.create({
            'set_id': sheet.id, 'month': date(2026, 5, 1),
            'scheduled_date': date(2026, 5, 25)})
        self.Job._close_missed(date(2026, 7, 4))
        self.assertEqual(old.state, 'missed')

    def test_a_check_in_is_never_made_before_the_goals_were_agreed(self):
        """A conversation about a plan nobody had yet is a row with nothing in
        it."""
        sheet = self._agreed()
        sheet.sudo().write({'locked_at': '2026-08-02 09:00:00'})
        self.assertFalse(
            sheet._checkin_is_due_this_month(date(2026, 6, 1)))
        self.assertTrue(
            sheet._checkin_is_due_this_month(date(2026, 9, 1)))

    def test_a_check_in_is_never_made_after_the_year_ends(self):
        sheet = self._agreed()
        self.assertFalse(
            sheet._checkin_is_due_this_month(date(2027, 6, 1)))

    def test_off_is_a_real_state_and_the_job_says_what_it_would_have_done(self):
        """R54 — A THING THAT IS OFF AND DOES NOT SAY SO IS REPORTED AS
        BROKEN."""
        sheet = self._agreed()
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_goals.checkins', '0')
        made = self.Job._make_checkins()
        self.assertEqual(made, 0)
        # SCOPED TO THIS TEST'S OWN SHEET. `search([])` is a question about
        # the WHOLE DATABASE, and this suite runs against a live demo one —
        # so an assertion that the table is empty is an assertion that passes
        # only until somebody loads a demo, which is a test that will fail for
        # a reason that has nothing to do with the rule it is guarding.
        self.assertFalse(self.Checkin.search([('set_id', '=', sheet.id)]))
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_goals.checkins', '1')

    def test_a_nudge_goes_once_and_to_both_sides(self):
        """BOTH SIDES, because a conversation is two people — and the key is
        written even when there was no address, or somebody with no email is
        due a nudge every morning for the rest of the year."""
        sheet = self._agreed()
        today = date.today()
        row = self.Checkin.create({
            'set_id': sheet.id,
            'month': date(today.year, today.month, 1),
            'scheduled_date': today})
        before = self.env['mail.mail'].search_count([])
        sent = self.Job._checkin_nudges(today)
        self.assertEqual(sent, 1)
        after = self.env['mail.mail'].search_count([])
        self.assertEqual(after - before, 2,
                         'a check-in nudge goes to the employee AND their '
                         'manager')
        self.assertTrue(row._already_nudged('day'))
        self.assertEqual(self.Job._checkin_nudges(today), 0)
        self.env['mail.mail'].search(
            [('id', '>', before)]).write({'state': 'cancel'})


# =========================================================================
#  T3. Who a review is fair to
# =========================================================================
@tagged('post_install', '-at_install')
class TestApplicability(YearCase):

    def _joiner(self, joined_on, name='DEMO Joiner'):
        user = self.env['res.users'].with_context(
            no_reset_password=True).create({
                'name': name,
                'login': '%s@example.com' % name.lower().replace(' ', '.'),
                'email': '%s@example.com' % name.lower().replace(' ', '.'),
                'company_id': self.company.id,
                'company_ids': [(6, 0, [self.company.id])],
                'group_ids': [(6, 0, [self.env.ref('base.group_user').id])],
            })
        employee = self.env['hr.employee'].create({
            'name': name, 'company_id': self.company.id,
            'user_id': user.id, 'parent_id': self.boss.id,
            'work_email': '%s@example.com' % name.lower().replace(' ', '.')})
        sheet = self.Set.create({
            'employee_id': employee.id, 'cycle_id': self.cycle.id,
            'company_id': self.company.id, 'joined_on': joined_on,
            'deadline': date.today() + timedelta(days=14)})
        sheet._stamp_applicability(force=True)
        return sheet

    def test_somebody_who_was_here_all_year_gets_both(self):
        sheet = self._joiner(date(2025, 1, 6), 'DEMO Early')
        self.assertTrue(sheet.applies_mid_year)
        self.assertTrue(sheet.applies_year_end)
        self.assertFalse(sheet.prorated)
        self.assertEqual(sheet.covered_from, self.cycle.date_start)
        self.assertIn('whole of', sheet.applicability_note)

    def test_somebody_who_joined_just_before_the_half_way_point_does_not(self):
        """THREE WHOLE MONTHS, counted from the day their goals start
        covering. August the 15th to October the 1st is one month, not two."""
        sheet = self._joiner(date(2026, 8, 15), 'DEMO Late Summer')
        self.assertFalse(sheet.applies_mid_year)
        self.assertTrue(sheet.applies_year_end)
        self.assertTrue(sheet.prorated)
        self.assertIn('too close to the half-way point',
                      sheet.applicability_note)

    def test_somebody_who_joined_after_the_half_way_point_never_gets_one(self):
        sheet = self._joiner(date(2026, 11, 3), 'DEMO November')
        self.assertFalse(sheet.applies_mid_year)
        self.assertTrue(sheet.applies_year_end,
                        'November to March is five months, which is over the '
                        'shipped minimum of three')

    def test_somebody_who_joined_too_late_to_be_scored_is_carried(self):
        """JANUARY TO MARCH IS TWO MONTHS, under the shipped minimum of
        three, so there is not enough of the year left to score them on."""
        sheet = self._joiner(date(2027, 1, 20), 'DEMO January')
        self.assertFalse(sheet.applies_year_end)
        self.assertTrue(sheet.prorated)
        self.assertIn('carry into next year', sheet.applicability_note)

    def test_the_minimum_is_a_dial_the_company_turns(self):
        """"Fair" is a company's word and not ours, so the November example
        the handover gives is reachable by raising the minimum to six."""
        self.cycle.write({'year_end_min_months': 6})
        sheet = self._joiner(date(2026, 11, 3), 'DEMO November Six')
        self.assertFalse(sheet.applies_year_end)

    def test_it_is_stamped_once_and_restamping_is_a_deliberate_act(self):
        """A computed answer would silently move somebody in or out of a
        review they have already been told about."""
        sheet = self._joiner(date(2025, 1, 6), 'DEMO Stamped')
        self.assertTrue(sheet.applicability_set)
        sheet.write({'joined_on': date(2027, 1, 20)})
        self.assertFalse(sheet._stamp_applicability())
        self.assertTrue(sheet.applies_year_end, 'the stamp moved on its own')
        sheet.restamp_applicability()
        self.assertFalse(sheet.applies_year_end)

    def test_reviews_are_opened_only_when_they_are_nearly_due(self):
        """A LIST OF THINGS DUE IN NINE MONTHS IS A LIST NOBODY READS."""
        sheet = self._agreed()
        sheet._stamp_applicability(force=True)
        far = self.cycle.mid_year_date - timedelta(days=200)
        self.Job._make_reviews(far)
        self.assertFalse(self.Review.search([('set_id', '=', sheet.id)]))
        near = self.cycle.mid_year_date - timedelta(days=10)
        self.Job._make_reviews(near)
        rows = self.Review.search([('set_id', '=', sheet.id)])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.kind, 'mid_year')
        self.Job._make_reviews(near)
        self.assertEqual(
            len(self.Review.search([('set_id', '=', sheet.id)])), 1)

    def test_a_sheet_with_no_year_end_never_gets_a_year_end_review(self):
        sheet = self._joiner(date(2027, 1, 20), 'DEMO No Year End')
        self._goal(sheet, weight=100.0)
        self._lock(sheet)
        self.Job._make_reviews(self.cycle.date_end - timedelta(days=3))
        self.assertFalse(self.Review.search(
            [('set_id', '=', sheet.id), ('kind', '=', 'year_end')]))

    def test_a_review_needs_a_write_up_and_freezes_the_score(self):
        sheet = self._agreed()
        sheet._stamp_applicability(force=True)
        review = sheet._ensure_review('mid_year', self.cycle.mid_year_date)
        with self.assertRaises(UserError):
            review.action_review_done('  ')
        self._score_all(sheet, '4')
        review.action_review_done('DEMO half-way and on track.')
        self.assertEqual(review.state, 'done')
        self.assertAlmostEqual(review.score_at_review, 4.0, places=2)


# =========================================================================
#  T4. Changing goals that have already been agreed
# =========================================================================
@tagged('post_install', '-at_install')
class TestChanges(YearCase):

    def test_a_change_can_only_be_asked_for_after_the_lock(self):
        sheet = self._sheet()
        self._goal(sheet, weight=100.0)
        with self.assertRaises(UserError):
            self.Change.raise_change(sheet.id, 'edit', {'title': 'DEMO x'},
                                     'DEMO why', sheet.goal_ids[0].id)

    def test_a_reason_is_required(self):
        sheet = self._agreed()
        with self.assertRaises(UserError):
            self.Change.raise_change(sheet.id, 'edit',
                                     {'title': 'DEMO new title'}, '   ',
                                     sheet.goal_ids[0].id)

    def test_a_reweight_that_does_not_add_up_is_refused_at_the_door(self):
        """REFUSED WHEN SOMEBODY TYPES, not three days later after two people
        have agreed it."""
        sheet = self._agreed()
        goals = sheet.goal_ids
        with self.assertRaises(UserError):
            self.Change.raise_change(sheet.id, 'reweight', {'weights': {
                goals[0].id: 60, goals[1].id: 30}}, 'DEMO why')
        change = self.Change.raise_change(sheet.id, 'reweight', {'weights': {
            goals[0].id: 70, goals[1].id: 30}}, 'DEMO the north came first.')
        self.assertEqual(change.state, 'draft')

    def test_a_reweight_must_name_every_goal(self):
        sheet = self._agreed()
        goals = sheet.goal_ids
        with self.assertRaises(UserError):
            self.Change.raise_change(sheet.id, 'reweight',
                                     {'weights': {goals[0].id: 100}},
                                     'DEMO why')

    def test_dropping_a_goal_archives_it_and_writes_the_trail(self):
        """ARCHIVED AND NEVER DELETED: history that is deleted cannot be
        produced when somebody asks what the year was about. And the numbers
        do not move, because every compute reads its goals with
        `active_test=False`."""
        sheet = self._agreed()
        goal = sheet.goal_ids[0]
        before_total = sheet.weight_total
        change = self.Change.raise_change(
            sheet.id, 'drop', {}, 'DEMO the region was merged.', goal.id)
        change.action_change_submit()
        self._carry_out(change)
        self.assertEqual(change.state, 'approved')
        self.assertFalse(goal.active)
        self.assertTrue(goal.exists())
        sheet.invalidate_recordset()
        self.assertAlmostEqual(sheet.weight_total, before_total, places=2,
                               msg='archiving a goal moved the arithmetic')
        audit = self.Audit.search([('change_id', '=', change.id)])
        self.assertEqual(len(audit), 1)
        self.assertIn('DEMO goal 1', audit.what)
        self.assertEqual(audit.before().get('title'), 'DEMO goal 1')
        self.assertTrue(audit.after().get('dropped'))
        goal.invalidate_recordset()
        self.assertTrue(goal.changed_on)
        self.assertIn('after it was agreed', goal._changed_words())

    def test_an_edit_is_carried_out_and_the_before_is_kept(self):
        sheet = self._agreed()
        goal = sheet.goal_ids[0]
        change = self.Change.raise_change(
            sheet.id, 'edit', {'title': 'DEMO grow the south region'},
            'DEMO the north was handed over.', goal.id)
        change.action_change_submit()
        self._carry_out(change)
        self.assertEqual(goal.title, 'DEMO grow the south region')
        audit = self.Audit.search([('change_id', '=', change.id)])
        self.assertEqual(audit.before()['title'], 'DEMO goal 1')
        self.assertEqual(audit.after()['title'], 'DEMO grow the south region')

    def test_adding_a_goal_arrives_at_nought_and_locked(self):
        """TWO CHANGES IN ONE REQUEST IS TWO DECISIONS IN ONE PRESS, and the
        second one is the one nobody reads. A new goal arrives at nought and
        the weights are a separate request."""
        sheet = self._agreed()
        change = self.Change.raise_change(sheet.id, 'add', {
            'title': 'DEMO open the second depot',
            'description': 'DEMO why it matters.',
            'date_end': str(self.cycle.date_end),
            'kr_titles': ['DEMO staff trained', 'DEMO deliveries a week'],
        }, 'DEMO the board approved the depot in June.')
        change.action_change_submit()
        self._carry_out(change)
        made = change.goal_id
        self.assertTrue(made)
        self.assertEqual(made.weight, 0.0)
        self.assertTrue(made.locked)
        self.assertEqual(len(made.kr_ids), 2)

    def test_turning_one_down_changes_nothing(self):
        sheet = self._agreed()
        goal = sheet.goal_ids[0]
        change = self.Change.raise_change(
            sheet.id, 'edit', {'title': 'DEMO something else'},
            'DEMO why', goal.id)
        change.action_change_submit()
        change.sudo().write({'state': 'manager_ok'})
        change.sudo().write({'refuse_note':
                             'DEMO the year still needs this one.'})
        change._after_approval_transition('refused')
        change.sudo().write({'state': 'refused'})
        self.assertEqual(change.state, 'refused')
        self.assertEqual(goal.title, 'DEMO goal 1')
        self.assertFalse(self.Audit.search([('change_id', '=', change.id)]))
        self.assertIn('still needs this one', change.refuse_note)


    def test_a_refusal_puts_its_reason_on_the_record(self):
        """THE PERSON WHO HAS TO ACT ON IT READS THEIR OWN GOALS PAGE, not an
        approval inbox. Under a route the engine keeps the reason on the
        REQUEST and the record's own note stays empty — so the page says
        "Turned down" with nothing beside it and the email has nothing to
        quote. Found live on the first refusal."""
        sheet = self._agreed()
        change = self.Change.raise_change(
            sheet.id, 'edit', {'title': 'DEMO x'}, 'DEMO why',
            sheet.goal_ids[0].id)
        change.sudo().write({'state': 'manager_ok'})
        change._approval_reject(None, 'DEMO the goal still stands.')
        self.assertEqual(change.state, 'refused')
        self.assertIn('still stands', change.refuse_note)

    def test_the_consequence_hangs_off_the_status_and_not_the_door(self):
        """R204 — the adapter's `_approval_apply` is the ROUTE'S half and is
        never called on the record's own ladder, so a change carried out
        there alone silently does not happen where no route is published."""
        sheet = self._agreed()
        goal = sheet.goal_ids[0]
        change = self.Change.raise_change(
            sheet.id, 'edit', {'title': 'DEMO reworded'}, 'DEMO why', goal.id)
        change.sudo().write({'state': 'manager_ok'})
        change._after_approval_transition('approved')
        self.assertEqual(goal.title, 'DEMO reworded')
        self.assertTrue(change.applied_at)

    def test_it_is_carried_out_exactly_once(self):
        sheet = self._agreed()
        goal = sheet.goal_ids[0]
        change = self.Change.raise_change(
            sheet.id, 'edit', {'title': 'DEMO once'}, 'DEMO why', goal.id)
        change.sudo().write({'state': 'manager_ok'})
        change._after_approval_transition('approved')
        change._after_approval_transition('approved')
        self.assertEqual(
            len(self.Audit.search([('change_id', '=', change.id)])), 1)

    def test_the_route_is_registered_with_no_date_field(self):
        """R192 — a `date_field` asks who held the seat BACK THEN, and a
        change request's own dates are the goal year's, which are months old
        by construction."""
        from odoo.addons.biz_approval_workflow.models.chain_shim import (
            CHAIN_PROCESS_KEYS)
        spec = CHAIN_PROCESS_KEYS.get('pb.goal.change')
        self.assertTrue(spec)
        self.assertEqual(spec['process_key'], 'goal_change')
        self.assertIsNone(spec['date_field'])
        self.assertEqual(spec['submit_state'], 'submitted')
        self.assertEqual(spec['driven'], ('manager_ok', 'approved'))
        self.assertEqual(spec['reverse_to'], ('draft',))

    def test_the_catalogue_row_this_module_owns_is_shipped(self):
        """R208 — without it `Seed.lay` refuses with ONE INFO LINE, the
        install reports success, and a request sent in never reaches
        anybody's inbox."""
        row = self.env.ref('pb_goals.process_goal_change',
                           raise_if_not_found=False)
        self.assertTrue(row)
        self.assertEqual(row.key, 'goal_change')
        self.assertEqual(row.model_name, 'pb.goal.change')

    def test_the_revision_stamp_is_the_proposal_and_the_reason(self):
        """R202/AM32 — never `write_date`, and never anything an approver is
        meant to change on the way through."""
        sheet = self._agreed()
        change = self.Change.raise_change(
            sheet.id, 'edit', {'title': 'DEMO x'}, 'DEMO why',
            sheet.goal_ids[0].id)
        stamp = change._chain_revision_values()
        self.assertIn('payload', stamp)
        self.assertIn('reason', stamp)
        self.assertNotIn('write_date', stamp)
        self.assertNotIn('state', stamp)


# =========================================================================
#  T5. Scoring
# =========================================================================
@tagged('post_install', '-at_install')
class TestScoring(YearCase):

    def test_nothing_is_scored_until_everything_is(self):
        """A NUMBER THAT IS RIGHT ONLY IF YOU KNOW WHAT IS MISSING FROM IT is
        worse than no number, because a screen cannot say what is missing."""
        sheet = self._agreed()
        goal = sheet.goal_ids[0]
        goal.kr_ids[0].sudo().write({'score': '5'})
        goal.invalidate_recordset()
        self.assertFalse(goal.scored)
        self.assertFalse(sheet.scored)
        goal.kr_ids[1].sudo().write({'score': '3'})
        goal.invalidate_recordset()
        self.assertTrue(goal.scored)
        self.assertAlmostEqual(goal.score, 4.0, places=2)
        sheet.invalidate_recordset()
        self.assertFalse(sheet.scored, 'one goal of two is not a sheet')

    def test_the_sheet_score_is_the_hand_calculation(self):
        """60% x 4.0 + 40% x 3.0 = 3.6, and the band is Strong."""
        sheet = self._agreed(weights=(60.0, 40.0))
        goals = sheet.goal_ids
        for kr in goals[0].kr_ids:
            kr.sudo().write({'score': '4'})
        for kr in goals[1].kr_ids:
            kr.sudo().write({'score': '3'})
        sheet.invalidate_recordset()
        self.assertTrue(sheet.scored)
        self.assertAlmostEqual(sheet.score, 3.6, places=2)
        self.assertEqual(sheet.score_band, 'Strong')

    def test_every_band_is_reachable_and_the_top_one_takes_five(self):
        sheet = self._agreed(weights=(100.0, 0.0))
        cases = (('5', 'Outstanding'), ('4', 'Strong'), ('3', 'Solid'),
                 ('1', 'Needs attention'))
        for mark, band in cases:
            self._score_all(sheet, mark)
            self.assertEqual(sheet.score_band, band,
                             'a score of %s should be %s' % (mark, band))

    def test_a_score_is_not_the_employees_word(self):
        """The employee already said what they thought in April — that is the
        self-rating, and it is a different question asked of a different
        person."""
        sheet = self._agreed()
        kr = sheet.goal_ids[0].kr_ids[0]
        with self.assertRaises(UserError):
            kr.with_user(self.staff_user).write({'score': '5'})
        kr.with_user(self.boss_user).write({'score': '5'})
        self.assertEqual(kr.score, '5')

    def test_the_door_refuses_anybody_but_the_manager_or_hr(self):
        sheet = self._agreed()
        kr = sheet.goal_ids[0].kr_ids[0]
        with self.assertRaises(UserError):
            sheet.with_user(self.staff_user).score_key_results(
                {kr.id: '4'})
        answer = sheet.with_user(self.boss_user).score_key_results(
            {kr.id: '4'})
        self.assertTrue(answer['ok'])
        self.assertEqual(answer['saved'], 1)

    def test_a_partial_save_is_allowed_and_says_what_is_left(self):
        """A manager scores a year over two sittings, and a screen that
        refuses nine marks because the tenth is undecided loses nine marks."""
        sheet = self._agreed()
        kr = sheet.goal_ids[0].kr_ids[0]
        answer = sheet.score_key_results({kr.id: '4'})
        self.assertFalse(answer['scored'])
        self.assertIn('still to score', answer['sentence'])

    def test_a_key_result_from_another_sheet_is_skipped_and_not_shouted_at(self):
        sheet = self._agreed()
        other = self._agreed(employee=self.stranger)
        answer = sheet.score_key_results({other.goal_ids[0].kr_ids[0].id: '5'})
        self.assertEqual(answer['saved'], 0)
        self.assertFalse(other.goal_ids[0].kr_ids[0].score)

    def test_a_bad_mark_is_refused_with_the_list_of_good_ones(self):
        sheet = self._agreed()
        with self.assertRaises(UserError):
            sheet.score_key_results({sheet.goal_ids[0].kr_ids[0].id: '9'})

    def test_marking_a_goal_complete_is_not_a_score(self):
        sheet = self._agreed()
        goal = sheet.goal_ids[0]
        goal.with_user(self.staff_user).action_goal_done()
        self.assertTrue(goal.done_at)
        self.assertFalse(goal.scored)
        sheet.invalidate_recordset()
        self.assertEqual(sheet.goals_done, 1)
        goal.action_goal_reopen()
        self.assertFalse(goal.done_at)

    def test_a_stranger_cannot_mark_somebody_elses_goal_complete(self):
        sheet = self._agreed()
        with self.assertRaises(UserError):
            sheet.goal_ids[0].with_user(
                self.stranger_user).action_goal_done()


# =========================================================================
#  T6. Closing the year
# =========================================================================
@tagged('post_install', '-at_install')
class TestClose(YearCase):

    def test_the_preview_counts_and_names_what_is_missing(self):
        sheet = self._agreed()
        preview = self.Cycle.preview_close(self.cycle.id)
        self.assertTrue(preview['ok'])
        self.assertEqual(preview['live'], 1)
        self.assertEqual(preview['unscored'], 1)
        self.assertFalse(preview['can_close'])
        self.assertIn(sheet.employee_id.name, preview['names'])

    def test_it_refuses_while_anything_is_unscored(self):
        self._agreed()
        with self.assertRaises(UserError):
            self.Cycle.close_cycle(self.cycle.id)

    def test_closing_with_gaps_is_written_on_the_year(self):
        """A year closed with holes in it is a legitimate thing a company
        sometimes has to do; one that nobody wrote down is how a figure gets
        quoted in March that nobody can account for."""
        self._agreed()
        self.Cycle.close_cycle(self.cycle.id, with_gaps=True)
        self.assertEqual(self.cycle.state, 'closed')
        self.assertEqual(self.cycle.closed_with_gaps, 1)
        self.assertTrue(self.cycle.closed_at)

    def test_closing_freezes_the_numbers_and_files_the_goals(self):
        sheet = self._agreed()
        self._score_all(sheet, '4')
        live_score = sheet.score
        self.Cycle.close_cycle(self.cycle.id)
        self.assertEqual(sheet.state, 'closed')
        frozen = sheet.frozen()
        self.assertAlmostEqual(frozen['score'], live_score, places=2)
        self.assertEqual(frozen['band'], sheet.score_band)
        self.assertEqual(len(frozen['goals']), 2)
        self.assertFalse(
            sheet.with_context(active_test=False).goal_ids.filtered('active'),
            'the goals were not filed')

    def test_a_closed_year_keeps_its_arithmetic(self):
        """A LIVE COMPUTE OVER A CLOSED YEAR RE-ANSWERS the day somebody
        corrects a weight in a neighbouring row. Archiving must change
        nothing."""
        sheet = self._agreed()
        self._score_all(sheet, '4')
        before = (sheet.weight_total, sheet.goal_count, sheet.score)
        self.Cycle.close_cycle(self.cycle.id)
        sheet.invalidate_recordset()
        self.assertEqual(
            (sheet.weight_total, sheet.goal_count, sheet.score), before)

    def test_a_carried_sheet_is_not_scored_and_says_why(self):
        sheet = self._agreed()
        sheet.sudo().write({'applies_year_end': False, 'prorated': True})
        self.Cycle.close_cycle(self.cycle.id, with_gaps=True)
        self.assertEqual(sheet.state, 'closed')
        self.assertIn('joined too late', sheet._close_sentence())

    def test_the_year_end_email_goes_once_and_carries_the_band(self):
        sheet = self._agreed()
        self._score_all(sheet, '5')
        before = self.env['mail.mail'].search_count([])
        self.Cycle.close_cycle(self.cycle.id)
        made = self.env['mail.mail'].search([('id', '>', 0)], order='id desc',
                                            limit=5)
        self.assertEqual(self.env['mail.mail'].search_count([]) - before, 1)
        self.assertIn('Outstanding', ' '.join(made.mapped('subject')))
        made.write({'state': 'cancel'})

    def test_closing_can_be_undone(self):
        """NOTHING IS LOST EITHER WAY, and saying so out loud is most of why
        anybody will press it."""
        sheet = self._agreed()
        self._score_all(sheet, '4')
        self.Cycle.close_cycle(self.cycle.id)
        frozen = sheet.frozen_json
        self.cycle.action_reopen_cycle()
        self.assertEqual(self.cycle.state, 'open')
        self.assertEqual(sheet.state, 'locked')
        self.assertTrue(
            sheet.with_context(active_test=False).goal_ids.filtered('active'))
        self.assertEqual(sheet.frozen_json, frozen,
                         'what the year looked like when it was closed is '
                         'still a fact')

    def test_scores_cannot_change_after_the_year_is_closed(self):
        sheet = self._agreed()
        self._score_all(sheet, '4')
        self.Cycle.close_cycle(self.cycle.id)
        with self.assertRaises(UserError):
            sheet.score_key_results({sheet.with_context(
                active_test=False).goal_ids[0].kr_ids[0].id: '5'})

    def test_a_new_year_can_open_once_the_old_one_is_closed(self):
        self._agreed()
        self.Cycle.close_cycle(self.cycle.id, with_gaps=True)
        nxt = self.Cycle.create({
            'name': 'DEMO FY2027', 'company_id': self.company.id,
            'date_start': date(2027, 4, 1), 'date_end': date(2028, 3, 31)})
        nxt.action_open()
        self.assertEqual(nxt.state, 'open')

    def test_closing_is_for_the_hr_team(self):
        self._agreed()
        with self.assertRaises(UserError):
            self.Cycle.with_user(self.staff_user).close_cycle(
                self.cycle.id, with_gaps=True)


# =========================================================================
#  T7. The numbers
# =========================================================================
@tagged('post_install', '-at_install')
class TestNumbers(YearCase):

    def test_an_empty_goal_year_gets_a_sentence_and_not_a_grid_of_noughts(self):
        answer = self.Numbers.get_numbers(self.cycle.id)
        self.assertTrue(answer['allowed'])
        self.assertTrue(answer['empty'])
        self.assertIn('Nobody has a goal sheet', answer['headline'])
        self.assertEqual(answer['tiles'], [])

    def test_the_compliance_figure_equals_a_hand_recomputation(self):
        """THE NUMBER THIS LENS EXISTS FOR. Three conversations owed, two
        had."""
        sheet = self._agreed()
        for month, state in ((5, 'done'), (6, 'done'), (7, 'missed')):
            self.Checkin.create({
                'set_id': sheet.id, 'month': date(2026, month, 1),
                'scheduled_date': date(2026, month, 25), 'state': state})
        answer = self.Numbers.get_numbers(self.cycle.id)
        self.assertFalse(answer['empty'])
        tile = [t for t in answer['tiles'] if t['key'] == 'checkins'][0]
        self.assertEqual(tile['value'], '66.7%')
        manager = answer['by_manager'][0]
        self.assertAlmostEqual(manager['compliance'], 66.7, places=1)
        self.assertEqual(manager['owed'], 3)
        self.assertEqual(manager['done'], 2)
        self.assertEqual(manager['missed'], 1)

    def test_a_planned_conversation_is_not_yet_owed(self):
        """"Owed" is what has already come round, not what is still to come —
        otherwise every manager reads 0% on the first of the month."""
        sheet = self._agreed()
        self.Checkin.create({
            'set_id': sheet.id, 'month': date(2026, 9, 1),
            'scheduled_date': date(2026, 9, 25)})
        answer = self.Numbers.get_numbers(self.cycle.id)
        tile = [t for t in answer['tiles'] if t['key'] == 'checkins'][0]
        self.assertEqual(tile['value'], '—')

    def test_the_bands_come_back_in_band_order(self):
        first = self._agreed()
        second = self._agreed(employee=self.stranger)
        self._score_all(first, '5')
        self._score_all(second, '3')
        answer = self.Numbers.get_numbers(self.cycle.id)
        labels = [row['label'] for row in answer['bands']]
        self.assertEqual(
            labels[:4], ['Outstanding', 'Strong', 'Solid', 'Needs attention'])
        counts = {row['label']: row['count'] for row in answer['bands']}
        self.assertEqual(counts['Outstanding'], 1)
        self.assertEqual(counts['Solid'], 1)

    def test_a_negative_span_is_floored_rather_than_dropped(self):
        """R155 — a span that crosses a timezone boundary is an ARTEFACT, and
        a `>= 0` guard drops exactly the fastest rows."""
        sheet = self._agreed()
        sheet.sudo().write({'submitted_at': '2026-06-10 08:00:00',
                            'manager_ok_at': '2026-06-09 08:00:00'})
        answer = self.Numbers.get_numbers(self.cycle.id)
        rung = [r for r in answer['rungs'] if r['key'] == 'manager'][0]
        self.assertEqual(rung['n'], 1)
        self.assertEqual(rung['days'], 0.0)

    def test_the_filter_lists_come_from_the_data(self):
        """R27 — a filter that matches nothing is a broken promise."""
        self._agreed()
        answer = self.Numbers.get_numbers(self.cycle.id)
        names = [row['name'] for row in answer['lists']['managers']]
        self.assertIn('DEMO Boss', names)

    def test_it_is_for_the_hr_team_and_says_so_rather_than_erroring(self):
        self._agreed()
        answer = self.Numbers.with_user(self.staff_user).get_numbers(
            self.cycle.id)
        self.assertFalse(answer['allowed'])
        self.assertIn('HR team', answer['why'])

    def test_the_spreadsheet_re_reads_and_carries_the_figures(self):
        sheet = self._agreed()
        self._score_all(sheet, '4')
        self.Checkin.create({
            'set_id': sheet.id, 'month': date(2026, 5, 1),
            'scheduled_date': date(2026, 5, 25), 'state': 'done'})
        answer = self.Numbers.export_xlsx(self.cycle.id)
        self.assertTrue(answer['file_b64'])
        self.assertIn('spreadsheetml', answer['mimetype'])
        self.assertIn('DEMO FY2026', answer['filename'])
        import base64
        import io
        import openpyxl
        book = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(answer['file_b64'])))
        self.assertIn('By manager', book.sheetnames)
        page = book['By manager']
        self.assertEqual(page.cell(row=2, column=1).value, 'DEMO Boss')
        self.assertEqual(page.cell(row=2, column=9).value, 100.0)

    def test_the_spreadsheet_raises_rather_than_downloading_a_refusal(self):
        self._agreed()
        from odoo.exceptions import AccessError
        with self.assertRaises(AccessError):
            self.Numbers.with_user(self.staff_user).export_xlsx(self.cycle.id)


# =========================================================================
#  T8. "Goals waiting on you"
# =========================================================================
@tagged('post_install', '-at_install')
class TestHome(YearCase):

    def test_a_manager_sees_the_conversations_they_owe(self):
        sheet = self._agreed()
        row = self.Checkin.ensure_for(sheet)
        row.sudo().write({'scheduled_date': date.today() - timedelta(days=2)})
        answer = self.Home.with_user(self.boss_user).get_home()
        kinds = [r['kind'] for r in answer['rows']]
        self.assertIn('checkin', kinds)
        checkin = [r for r in answer['rows'] if r['kind'] == 'checkin'][0]
        self.assertTrue(checkin['late'])
        self.assertEqual(checkin['id'], row.id)

    def test_the_count_and_the_list_are_the_same_number(self):
        """R80 — a chip that counts one thing over a list that shows another
        is two bugs."""
        sheet = self._agreed()
        row = self.Checkin.ensure_for(sheet)
        row.sudo().write({'scheduled_date': date.today()})
        answer = self.Home.with_user(self.boss_user).get_home()
        counted = sum(section['count'] for section in answer['sections'])
        self.assertEqual(counted, answer['count'])
        self.assertEqual(len(answer['rows']), min(answer['count'], 12))

    def test_an_employee_sees_their_own_unwritten_sheet(self):
        sheet = self._sheet()
        answer = self.Home.with_user(self.staff_user).get_home()
        rows = [r for r in answer['rows'] if r['kind'] == 'mine']
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['id'], sheet.id)

    def test_an_agreed_sheet_is_not_waiting_on_anybody(self):
        self._agreed()
        answer = self.Home.with_user(self.staff_user).get_home()
        self.assertFalse([r for r in answer['rows'] if r['kind'] == 'mine'])

    def test_nothing_waiting_is_a_sentence_and_never_a_blank(self):
        answer = self.Home.with_user(self.stranger_user).get_home()
        self.assertEqual(answer['count'], 0)
        self.assertIn('Nothing about goals', answer['headline'])

    def test_every_row_kind_has_a_door_and_the_windows_carry_views(self):
        """R125 — the client maps over `action.views` unconditionally, and
        the theme shows the failure as a generic dialog with nothing in the
        console."""
        sheet = self._agreed()
        row = self.Checkin.ensure_for(sheet)
        review = sheet._ensure_review('mid_year', self.cycle.mid_year_date)
        for kind, ident in (('checkin', row.id), ('review', review.id),
                            ('mine', sheet.id), ('approve', 1)):
            door = self.Home.open_row(kind, ident)
            self.assertTrue(door.get('type'))
            if door['type'] == 'ir.actions.act_window':
                self.assertIn('views', door)


# =========================================================================
#  T9/T10 wiring, and the employee's own page
# =========================================================================
@tagged('post_install', '-at_install')
class TestMyGoalsYear(YearCase):

    def test_the_page_carries_the_year(self):
        sheet = self._agreed()
        self.Checkin.ensure_for(sheet)
        page = self.My.with_user(self.staff_user).home()
        self.assertTrue(page['ok'])
        self.assertTrue(page['checkin'])
        self.assertIn('month_word', page['checkin'])
        self.assertTrue(page['why_reviews'])
        self.assertTrue(page['may_change'])
        self.assertIn('scores', page)
        self.assertTrue(all('done' in goal for goal in page['goals']))

    def test_the_employee_writes_up_their_own_check_in(self):
        sheet = self._agreed()
        row = self.Checkin.ensure_for(sheet)
        answer = self.My.with_user(self.staff_user).write_up_checkin(
            row.id, 'DEMO the depot is open.')
        self.assertTrue(answer['ok'])
        self.assertEqual(row.state, 'done')

    def test_somebody_elses_check_in_is_not_reachable_from_the_page(self):
        other = self._agreed(employee=self.stranger)
        row = self.Checkin.ensure_for(other)
        with self.assertRaises(UserError):
            self.My.with_user(self.staff_user).write_up_checkin(
                row.id, 'DEMO not mine.')

    def test_the_employee_may_ask_for_a_change_but_never_a_reweight(self):
        sheet = self._agreed()
        with self.assertRaises(UserError):
            self.My.with_user(self.staff_user).ask_for_change(
                'reweight', {}, 'DEMO why')
        answer = self.My.with_user(self.staff_user).ask_for_change(
            'drop', {}, 'DEMO the region was merged.',
            sheet.goal_ids[0].id)
        self.assertTrue(answer['ok'])
        change = self.Change.browse(answer['change_id'])
        self.assertEqual(change.state, 'submitted')

    def test_a_half_scored_sheet_shows_no_number(self):
        sheet = self._agreed()
        sheet.goal_ids[0].kr_ids[0].sudo().write({'score': '5'})
        page = self.My.with_user(self.staff_user).home()
        self.assertFalse(page['scores']['scored'])
        self.assertNotIn('5', page['scores']['sentence'])

    def test_a_past_year_is_read_from_the_frozen_copy(self):
        """NEVER RECOMPUTED. A past year that changes is a past year nobody
        believes."""
        sheet = self._agreed()
        self._score_all(sheet, '4')
        self.Cycle.close_cycle(self.cycle.id)
        page = self.My.with_user(self.staff_user).past_year(sheet.id)
        self.assertTrue(page['ok'])
        self.assertFalse(page['live'])
        self.assertAlmostEqual(page['score'], 4.0, places=2)
        # Move the live data underneath it: the frozen page must not budge.
        sheet.with_context(active_test=False).goal_ids[0].sudo().write(
            {'weight': 10.0})
        again = self.My.with_user(self.staff_user).past_year(sheet.id)
        self.assertAlmostEqual(again['score'], 4.0, places=2)

    def test_a_past_year_that_is_not_mine_is_refused(self):
        other = self._agreed(employee=self.stranger)
        with self.assertRaises(UserError):
            self.My.with_user(self.staff_user).past_year(other.id)


# =========================================================================
#  T10. The board's new tabs and the wiring
# =========================================================================
@tagged('post_install', '-at_install')
class TestBoardYear(YearCase):

    def test_every_tab_answers_with_its_own_payload(self):
        """R188 — one reader per tab, and the payload says which tab it is."""
        sheet = self._agreed()
        self.Checkin.ensure_for(sheet)
        for tab in ('checkins', 'reviews', 'changes', 'scores'):
            answer = self.Board.get_year(self.cycle.id, tab)
            self.assertTrue(answer['allowed'])
            self.assertEqual(answer['tab'], tab)
            self.assertIsInstance(answer['rows'], list)
            self.assertIsInstance(answer['stats'], list)
            self.assertTrue(answer['headline'])

    def test_an_unknown_tab_falls_back_to_the_conversations(self):
        answer = self.Board.get_year(self.cycle.id, 'nonsense')
        self.assertEqual(answer['tab'], 'checkins')

    def test_the_drawer_carries_every_new_section(self):
        sheet = self._agreed()
        self.Checkin.ensure_for(sheet)
        sheet._ensure_review('mid_year', self.cycle.mid_year_date)
        self.Change.raise_change(sheet.id, 'edit', {'title': 'DEMO x'},
                                 'DEMO why', sheet.goal_ids[0].id)
        drawer = self.Board.get_set(sheet.id)
        self.assertTrue(drawer['ok'])
        for key in ('checkins', 'reviews', 'changes', 'audit', 'scores',
                    'applicability', 'score_options'):
            self.assertIn(key, drawer)
        self.assertEqual(len(drawer['checkins']), 1)
        self.assertEqual(len(drawer['reviews']), 1)
        self.assertEqual(len(drawer['changes']), 1)
        self.assertEqual(len(drawer['scores']['goals']), 2)
        self.assertTrue(drawer['applicability']['note'])

    def test_the_scores_tab_puts_the_unscored_first(self):
        first = self._agreed()
        second = self._agreed(employee=self.stranger)
        self._score_all(first, '4')
        rows = self.Board.get_year(self.cycle.id, 'scores')['rows']
        self.assertEqual(rows[0]['id'], second.id)
        self.assertFalse(rows[0]['scored'])

    def test_a_record_argument_survives_the_wire(self):
        """R43 — a recordset arrives over JSON-RPC as a plain integer."""
        sheet = self._agreed()
        row = self.Checkin.ensure_for(sheet)
        self.assertEqual(
            self.Board.open_checkin(row)['res_id'],
            self.Board.open_checkin(row.id)['res_id'])

    def test_every_new_door_carries_views(self):
        sheet = self._agreed()
        row = self.Checkin.ensure_for(sheet)
        change = self.Change.raise_change(
            sheet.id, 'edit', {'title': 'DEMO x'}, 'DEMO why',
            sheet.goal_ids[0].id)
        for door in (self.Board.open_checkin(row.id),
                     self.Board.open_change(change.id),
                     self.Numbers.open_checkins(self.cycle.id),
                     self.Numbers.open_changes(self.cycle.id),
                     row.action_view_set(), change.action_view_set()):
            self.assertIn('views', door)

    def test_the_daily_job_runs_every_leg_and_reports_each_one(self):
        """SIX LEGS, EACH IN ITS OWN SAVEPOINT. The count is what somebody
        compares against the morning's log."""
        self._agreed()
        answer = self.Job._cron_daily()
        for key in ('deadlines', 'checkins_made', 'checkins_missed',
                    'checkin_nudges', 'reviews_made', 'review_nudges'):
            self.assertIn(key, answer)

    def test_the_shipped_bands_are_company_less(self):
        """R8 — a seeded record with a company installs into whichever company
        ran the install, and the standard rule then hides it from everybody
        else. Three companies out of four would score people and have no word
        for the answer."""
        for xmlid in ('band_outstanding', 'band_strong', 'band_solid',
                      'band_attention'):
            band = self.env.ref('pb_goals.%s' % xmlid)
            self.assertFalse(band.company_id)
            self.assertFalse(band.cycle_id)

    def test_a_goal_year_may_carry_its_own_bands(self):
        self.Band.create({
            'name': 'DEMO Met', 'min_score': 3.0, 'max_score': 5.0,
            'cycle_id': self.cycle.id, 'tone': 'good'})
        self.Band.create({
            'name': 'DEMO Not met', 'min_score': 0.0, 'max_score': 3.0,
            'cycle_id': self.cycle.id, 'tone': 'bad'})
        sheet = self._agreed(weights=(100.0, 0.0))
        self._score_all(sheet, '4')
        self.assertEqual(sheet.score_band, 'DEMO Met')

    def test_the_new_switches_read_through_the_fallback(self):
        from odoo.addons.pb_goals.models.goals_common import (
            DEFAULTS, flag, number)
        icp = self.env['ir.config_parameter'].sudo()
        for key in ('pb_goals.checkins', 'pb_goals.checkin_day',
                    'pb_goals.mid_year_min_months',
                    'pb_goals.year_end_min_months', 'pb_goals.yearend_mail',
                    'pb_goals.review_lead_days'):
            self.assertIn(key, DEFAULTS)
            icp.search([('key', '=', key)]).unlink()
        self.assertTrue(flag(self.env, 'pb_goals.checkins'))
        self.assertEqual(number(self.env, 'pb_goals.checkin_day'), 25)
        self.assertEqual(
            number(self.env, 'pb_goals.mid_year_min_months'), 3)
        self.assertEqual(number(self.env, 'pb_goals.review_lead_days'), 30)

    def test_the_month_helpers_do_not_walk_off_the_end_of_a_month(self):
        """The 31st of January plus one month is the 28th of February, and a
        schedule that walks off the end of a month walks the whole year."""
        from odoo.addons.pb_goals.models.checkin import (
            add_months, month_first, month_last, months_between)
        self.assertEqual(add_months(date(2026, 1, 31), 1), date(2026, 2, 28))
        self.assertEqual(month_first(date(2026, 6, 14)), date(2026, 6, 1))
        self.assertEqual(month_last(date(2026, 2, 3)), date(2026, 2, 28))
        self.assertEqual(months_between(date(2026, 8, 15), date(2026, 10, 1)),
                         1)
        self.assertEqual(months_between(date(2026, 8, 1), date(2026, 11, 1)),
                         3)

    def test_the_frozen_copy_is_json_a_screen_can_read(self):
        sheet = self._agreed()
        self._score_all(sheet, '4')
        payload = json.loads(sheet._freeze())
        self.assertIn('goals', payload)
        self.assertIn('score', payload)
        self.assertEqual(len(payload['goals'][0]['krs']), 2)


# =========================================================================
#  The gates B2 adds to B1's
# =========================================================================
@tagged('post_install', '-at_install')
class TestYearGates(YearCase):

    def test_every_b2_model_has_an_access_row_for_every_tier(self):
        """An ACL that names three of four tiers is a model one tier cannot
        read, and the symptom is an empty screen rather than an error."""
        rows = open(
            __file__.rsplit('/', 2)[0] + '/security/ir.model.access.csv',
            encoding='utf-8').read()
        for model in ('pb_goal_checkin', 'pb_goal_review', 'pb_goal_change',
                      'pb_goal_audit', 'pb_goal_band'):
            for tier in ('group_goals_user', 'group_goals_manager',
                         'group_goals_admin'):
                self.assertIn('model_%s,pb_goals.%s' % (model, tier), rows,
                              '%s has no row for %s' % (model, tier))

    def test_every_narrow_rule_ships_with_a_wide_one(self):
        """R60/R65 — SHIP THE PAIR, ALWAYS. `ir.rule` group rules are ORed
        over the rules that APPLY, so a narrow rule shipped alone NARROWS
        anybody who holds both groups."""
        for model in ('pb.goal.checkin', 'pb.goal.review', 'pb.goal.change',
                      'pb.goal.audit'):
            wide = self.env['ir.rule'].search([
                ('model_id.model', '=', model),
                ('domain_force', '=', "[(1, '=', 1)]")])
            self.assertTrue(
                wide, '%s has no "the HR team sees them all" rule' % model)

    def test_the_change_seat_may_write_and_not_only_read(self):
        """R193 — READ-ONLY PRODUCES THE WORST OUTCOME THERE IS: the route
        records the approval and the record does not move, with the reason
        inside the request's own `block_reason` and nothing on any screen."""
        rule = self.env.ref('pb_goals.rule_change_seat')
        self.assertTrue(rule.perm_read)
        self.assertTrue(rule.perm_write)
        self.assertFalse(rule.perm_create)
        self.assertFalse(rule.perm_unlink)


    def test_an_approval_drawer_row_is_a_dict_and_never_a_list(self):
        """THE ONE INBOX SILENTLY DISCARDS A ROW THAT IS NOT A DICT.

        `pb_approval_config/models/inbox_facade.py:598` skips anything that is
        not a dict with `head` / `sub` / `cells`, and it drops the whole
        detail only when there are no CHIPS either — so a consumer that hands
        it lists gets a drawer with its title, its chips and its note all
        present and no table at all, which reads as "there was nothing to
        show" rather than as a mistake. Nothing errors and nothing is logged.

        Found live on the first goal-change request; B1's goal-sheet drawer
        had shipped with the same shape and had been drawing no table since.
        """
        sheet = self._agreed()
        change = self.Change.raise_change(
            sheet.id, 'edit', {'title': 'DEMO reworded'}, 'DEMO why',
            sheet.goal_ids[0].id)
        for record in (sheet, change):
            detail = record._approval_detail(None)
            self.assertTrue(detail['rows'],
                            '%s has no drawer rows at all' % record._name)
            for row in detail['rows']:
                self.assertIsInstance(
                    row, dict,
                    '%s hands the inbox a list, which it discards without '
                    'a word' % record._name)
                self.assertTrue(row.get('head'))
                self.assertIsInstance(row.get('cells'), list)
            self.assertLessEqual(
                len(detail['rows'][0]['cells']), len(detail['columns']),
                '%s has more cells than column headings' % record._name)

    def test_the_manifest_declares_the_two_hubs_it_imports_from(self):
        """A registry you import is a dependency you have, whether or not the
        manifest says so — and the day somebody re-orders the other module's
        list it stops loading."""
        manifest = open(
            __file__.rsplit('/', 2)[0] + '/__manifest__.py',
            encoding='utf-8').read()
        self.assertIn("'pb_insights_hub'", manifest)
        self.assertIn("'pb_home_hub'", manifest)
        self.assertIn("19.0.1.1.0", manifest)

    def test_the_migration_for_this_version_exists(self):
        """`post_init_hook` fires on INSTALL ONLY and never on `-u` (AM70), so
        a new route and a backfill need a migration or the databases that
        already have this module come out of the upgrade without them."""
        import os
        path = __file__.rsplit('/', 2)[0] + '/migrations/19.0.1.1.0'
        self.assertTrue(os.path.isdir(path))
        self.assertTrue(os.listdir(path))

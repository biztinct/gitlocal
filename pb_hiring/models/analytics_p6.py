# -*- coding: utf-8 -*-
"""RECRUIT P6 — Hiring numbers, filtered and finished (G-52, G-53, G-54, G-58).

  * **Filters** — department, country (the role's market) and recruiter, on
    top of the date range. Every figure on the screen narrows with them,
    because they narrow the one cohort everything is counted over.
  * **The funnel** — offered → accepted → signed → joined, with the two ways
    out (turned down, did not join) and their reasons, per market and per
    role (`get_funnel`).
  * **Ageing** — every candidate in a working column right now, by days
    since their last move, by stage and recruiter; "stalled" from
    `pb_hiring.stalled_days` (14) (`get_ageing`). NOT bound to the date
    range: it is a picture of today.
  * **Leadership** — fill rate (filled ÷ agreed roles), open headcount by
    country, source mix, whose side moved interviews, and time to fill by
    month over the last twelve (`get_leadership`).
  * **Excel** — `export_xlsx(date_from, date_to, kind, …)`: one section, or
    the whole workbook, built from the same payload the screen draws.
"""

import base64
import io
import logging
from collections import OrderedDict
from datetime import date, timedelta

from dateutil.relativedelta import relativedelta

from odoo import _, api, fields, models
from odoo.exceptions import UserError

from .analytics import SCAN_LIMIT, _median, _span
from .hiring_common import (
    DROP_REASONS, P_STALLED_DAYS, REQUEST_STATES, REQUISITION_STATES, as_id,
    number,
)

_logger = logging.getLogger(__name__)

#: How many candidates the ageing table lists (the summaries count them all).
AGEING_ROWS = 400

EXPORT_KINDS = ('all', 'summary', 'funnel', 'ageing', 'leadership', 'stages',
                'sources', 'delays', 'agency', 'roles')


class PbHiringAnalyticsP6(models.AbstractModel):
    _inherit = 'pb.hiring.analytics'

    # =====================================================================
    #  The filters
    # =====================================================================
    @api.model
    def _filters(self, department_id=None, country_id=None, recruiter_id=None):
        return {'department_id': as_id(department_id),
                'country_id': as_id(country_id),
                'recruiter_id': as_id(recruiter_id)}

    @api.model
    def _filter_domain(self, filters=None):
        """The role clause the filters add (department, market, recruiter)."""
        f = filters if filters is not None else (self.env.context.get('pb_hn_filters') or {})
        dom = []
        if f.get('department_id'):
            dom.append(('department_id', 'child_of', f['department_id']))
        if f.get('country_id'):
            cid = f['country_id']
            dom += ['|', ('country_id', '=', cid), '&', ('country_id', '=', False),
                    ('company_id.country_id', '=', cid)]
        if f.get('recruiter_id'):
            dom.append(('recruiter_id', '=', f['recruiter_id']))
        return dom

    @api.model
    def _requests(self, start, end):
        dom = self._filter_domain()
        if not dom:
            return super()._requests(start, end)
        co_ids = self.env.companies.ids or [self.env.company.id]
        return self.env['pb.hiring.requisition'].sudo().search([
            ('company_id', 'in', co_ids), ('opened_on', '!=', False),
            ('opened_on', '>=', start), ('opened_on', '<=', end)] + dom,
            order='opened_on', limit=SCAN_LIMIT)

    @api.model
    def _options(self):
        """What the filter bar offers: only what the company's roles use."""
        co_ids = self.env.companies.ids or [self.env.company.id]
        roles = self.env['pb.hiring.requisition'].sudo().search(
            [('company_id', 'in', co_ids)], limit=SCAN_LIMIT)
        deps = roles.mapped('department_id')
        countries = roles.mapped('country_id') | roles.filtered(
            lambda r: not r.country_id).mapped('company_id.country_id')
        recruiters = roles.mapped('recruiter_id')
        return {
            'departments': sorted([{'id': d.id, 'name': d.name or ''} for d in deps],
                                  key=lambda r: r['name'].lower()),
            'countries': sorted([{'id': c.id, 'name': c.name or ''} for c in countries],
                                key=lambda r: r['name'].lower()),
            'recruiters': sorted([{'id': u.id, 'name': u.name or ''} for u in recruiters],
                                 key=lambda r: r['name'].lower()),
        }

    @api.model
    def _country_of(self, req):
        return req.country_id or req.company_id.country_id

    # =====================================================================
    #  The board, with everything
    # =====================================================================
    @api.model
    def get_board(self, date_from=None, date_to=None, department_id=None,
                  country_id=None, recruiter_id=None):
        filters = self._filters(department_id, country_id, recruiter_id)
        me = self.with_context(pb_hn_filters=filters)
        board = super(PbHiringAnalyticsP6, me).get_board(date_from, date_to)
        if not board.get('allowed'):
            return board
        board['filters'] = filters
        board['options'] = me._safe(me._options, {})
        board['stalled_days'] = number(self.env, P_STALLED_DAYS, 14)
        board['ageing'] = me._safe(lambda: me.get_ageing(**filters), {})
        if board.get('empty'):
            board['funnel'] = {}
            board['leadership'] = me._safe(
                lambda: me.get_leadership(date_from, date_to, **filters), {})
            return board
        board['funnel'] = me._safe(lambda: me.get_funnel(date_from, date_to, **filters), {})
        board['leadership'] = me._safe(
            lambda: me.get_leadership(date_from, date_to, **filters), {})
        # The funnel's last step is a tile too: what a signed offer is worth.
        fun = board['funnel'] or {}
        if fun.get('signed'):
            board['tiles'].append({
                'key': 'joined', 'label': _('Signed and joined'),
                'value': int(round(fun['joined'] * 100.0 / fun['signed'])),
                'unit': '%', 'icon': 'logIn',
                'hint': _('Of the people who signed, the share who actually started. '
                          'The rest did not join, or are still on their way.')})
        return board

    @api.model
    def _safe(self, fn, default):
        try:
            return fn()
        except Exception:               # noqa: BLE001 — one section never takes the screen down
            _logger.warning('pb_hiring: a numbers section failed', exc_info=True)
            return default

    # =====================================================================
    #  The funnel (G-58)
    # =====================================================================
    @api.model
    def get_funnel(self, date_from=None, date_to=None, department_id=None,
                   country_id=None, recruiter_id=None):
        if not self.env['pb.hiring']._can_read():
            return {}
        filters = self._filters(department_id, country_id, recruiter_id)
        me = self.with_context(pb_hn_filters=filters)
        start, end = me._window(date_from, date_to)
        requests = me._requests(start, end)
        offers = self.env['pb.hiring.offer'].sudo().search(
            [('requisition_id', 'in', requests.ids)])

        def classify(o):
            offered = bool(o.sent_on) or o.state in (
                'sent', 'accepted', 'declined', 'signed', 'joined', 'dropped')
            accepted = o.candidate_decision == 'accepted' or o.state in (
                'accepted', 'signed', 'joined', 'dropped')
            signed = bool(o.signed_on) or o.state in ('signed', 'joined', 'dropped')
            return {'offered': offered, 'accepted': accepted and offered,
                    'signed': signed and offered, 'joined': o.state == 'joined',
                    'declined': o.state == 'declined' or o.candidate_decision == 'declined',
                    'dropped': o.state == 'dropped'}

        keys = ('offered', 'accepted', 'signed', 'joined', 'declined', 'dropped')
        total = dict.fromkeys(keys, 0)
        markets = OrderedDict()
        roles = OrderedDict()
        declined_notes, drop_rows = [], {}
        for o in offers:
            c = classify(o)
            if not c['offered']:
                continue
            req = o.requisition_id
            country = self._country_of(req)
            m = markets.setdefault(country.id or 0, dict(
                {'country': country.name or _('No country'), 'country_id': country.id or 0},
                **dict.fromkeys(keys, 0)))
            r = roles.setdefault(req.id, dict(
                {'role': req.title or '', 'requisition_id': req.id,
                 'country': country.name or ''}, **dict.fromkeys(keys, 0)))
            for k in keys:
                if c[k]:
                    total[k] += 1
                    m[k] += 1
                    r[k] += 1
            if c['declined']:
                declined_notes.append({'who': o.candidate_name or '',
                                       'text': (o.candidate_comment or '').strip()})
            if c['dropped']:
                row = drop_rows.setdefault(o.drop_reason or 'other', {
                    'key': o.drop_reason or 'other', 'kind': 'dropped',
                    'reason': dict(DROP_REASONS).get(o.drop_reason or 'other', ''),
                    'n': 0, 'examples': []})
                row['n'] += 1
                if o.drop_note and len(row['examples']) < 3:
                    row['examples'].append(o.drop_note.strip()[:160])
        reasons = []
        if total['declined']:
            reasons.append({'key': 'declined', 'kind': 'declined',
                            'reason': _('Turned the offer down'), 'n': total['declined'],
                            'examples': [d['text'][:160] for d in declined_notes
                                         if d['text']][:3]})
        reasons += sorted(drop_rows.values(), key=lambda r: -r['n'])

        def pct(a, b):
            return int(round(a * 100.0 / b)) if b else None

        steps = [
            {'key': 'offered', 'label': _('Offered'), 'n': total['offered'],
             'pct': 100 if total['offered'] else None, 'lost': 0, 'lost_label': '',
             'reasons': []},
            {'key': 'accepted', 'label': _('Accepted'), 'n': total['accepted'],
             'pct': pct(total['accepted'], total['offered']), 'lost': total['declined'],
             'lost_label': _('turned it down'),
             'reasons': [r for r in reasons if r['kind'] == 'declined']},
            {'key': 'signed', 'label': _('Signed'), 'n': total['signed'],
             'pct': pct(total['signed'], total['offered']), 'lost': 0,
             'lost_label': '', 'reasons': []},
            {'key': 'joined', 'label': _('Joined'), 'n': total['joined'],
             'pct': pct(total['joined'], total['offered']), 'lost': total['dropped'],
             'lost_label': _('did not join'),
             'reasons': [r for r in reasons if r['kind'] == 'dropped']},
        ]
        waiting_answer = total['offered'] - total['accepted'] - total['declined']
        waiting_join = total['signed'] - total['joined'] - total['dropped']
        return dict(total, steps=steps, reasons=reasons,
                    by_market=list(markets.values()),
                    by_role=sorted(roles.values(), key=lambda r: -r['offered'])[:30],
                    waiting_answer=max(0, waiting_answer),
                    waiting_join=max(0, waiting_join),
                    sentence=self._funnel_sentence(total, waiting_join))

    @api.model
    def _funnel_sentence(self, t, waiting_join):
        if not t['offered']:
            return _("No offer reached a candidate on these roles yet.")
        return _("%(o)s offered, %(a)s accepted, %(s)s signed and %(j)s joined. "
                 "%(d)s turned it down and %(x)s did not join; %(w)s are still on "
                 "their way.", o=t['offered'], a=t['accepted'], s=t['signed'],
                 j=t['joined'], d=t['declined'], x=t['dropped'], w=waiting_join)

    # =====================================================================
    #  Ageing (G-54)
    # =====================================================================
    @api.model
    def get_ageing(self, department_id=None, country_id=None, recruiter_id=None):
        if not self.env['pb.hiring']._can_read():
            return {}
        filters = self._filters(department_id, country_id, recruiter_id)
        co_ids = self.env.companies.ids or [self.env.company.id]
        roles = self.env['pb.hiring.requisition'].sudo().search(
            [('company_id', 'in', co_ids), ('state', 'in', ('setup', 'open')),
             ('job_id', '!=', False)] + self._filter_domain(filters), limit=SCAN_LIMIT)
        by_job = {r.job_id.id: r for r in roles}
        threshold = max(1, number(self.env, P_STALLED_DAYS, 14))
        apps = self.env['hr.applicant'].sudo().search([
            ('job_id', 'in', list(by_job)), ('active', '=', True),
            ('stage_id.pb_family', '=', 'open')], limit=SCAN_LIMIT * 2)
        now = fields.Datetime.now()
        rows, stages, recruiters = [], OrderedDict(), OrderedDict()
        for app in apps:
            req = by_job.get(app.job_id.id)
            since = app.date_last_stage_update or app.create_date or now
            days = max(0, (now - since).days)
            stalled = days >= threshold and app.stage_id.pb_key not in ('post_offer',)
            st = stages.setdefault(app.stage_id.id, {
                'stage': app.stage_id.name or '', 'key': app.stage_id.pb_key or '',
                'sequence': app.stage_id.sequence or 0, 'n': 0, 'stalled': 0, 'days': []})
            st['n'] += 1
            st['stalled'] += 1 if stalled else 0
            st['days'].append(days)
            rec = req.recruiter_id
            rr = recruiters.setdefault(rec.id or 0, {
                'recruiter': rec.name or _('Nobody yet'), 'recruiter_id': rec.id or 0,
                'n': 0, 'stalled': 0})
            rr['n'] += 1
            rr['stalled'] += 1 if stalled else 0
            rows.append({'applicant_id': app.id, 'requisition_id': req.id,
                         'candidate': app.partner_name or app.email_from or '',
                         'role': req.title or '', 'stage': app.stage_id.name or '',
                         'stage_key': app.stage_id.pb_key or '', 'days': days,
                         'since': str(since.date()), 'recruiter': rec.name or '',
                         'stalled': stalled})
        rows.sort(key=lambda r: (-r['days'], r['candidate'].lower()))
        stage_rows = []
        for st in sorted(stages.values(), key=lambda s: s['sequence']):
            days = st.pop('days')
            st['median'] = _median(days)
            stage_rows.append(st)
        return {
            'threshold': threshold,
            'total': len(rows), 'stalled': len([r for r in rows if r['stalled']]),
            'rows': rows[:AGEING_ROWS], 'capped': len(rows) > AGEING_ROWS,
            'by_stage': stage_rows,
            'by_recruiter': sorted(recruiters.values(), key=lambda r: (-r['stalled'], -r['n'])),
        }

    # =====================================================================
    #  Leadership (G-53)
    # =====================================================================
    @api.model
    def get_leadership(self, date_from=None, date_to=None, department_id=None,
                       country_id=None, recruiter_id=None):
        if not self.env['pb.hiring']._can_read():
            return {}
        filters = self._filters(department_id, country_id, recruiter_id)
        me = self.with_context(pb_hn_filters=filters)
        start, end = me._window(date_from, date_to)
        requests = me._requests(start, end)
        agreed = requests.filtered(lambda r: r.request_state == 'agreed')
        filled = agreed.filtered(lambda r: r.state == 'filled' or r.filled_on)
        co_ids = self.env.companies.ids or [self.env.company.id]
        open_roles = self.env['pb.hiring.requisition'].sudo().search(
            [('company_id', 'in', co_ids), ('state', 'in', ('setup', 'open'))]
            + me._filter_domain(filters), limit=SCAN_LIMIT)
        heads = OrderedDict()
        for req in open_roles:
            country = me._country_of(req)
            row = heads.setdefault(country.id or 0, {
                'country': country.name or _('No country'), 'country_id': country.id or 0,
                'roles': 0, 'open': 0})
            row['roles'] += 1
            row['open'] += max(0, (req.headcount or 1) - (req.filled_count or 0))
        headcount = sorted(heads.values(), key=lambda r: -r['open'])
        sources = me._sources(requests) if requests else []
        total_apps = sum(s['applicants'] for s in sources) or 0
        for s in sources:
            s['share'] = int(round(s['applicants'] * 100.0 / total_apps)) if total_apps else 0
        delays = me._delays(requests) if requests else []
        total_moves = sum(d['count'] for d in delays) or 0
        for d in delays:
            d['share'] = int(round(d['count'] * 100.0 / total_moves)) if total_moves else 0
        # Time to fill by month: the twelve months up to the end of the range.
        month_end = date(end.year, end.month, 1)
        first = month_end - relativedelta(months=11)
        fills = self.env['pb.hiring.requisition'].sudo().search(
            [('company_id', 'in', co_ids), ('filled_on', '>=', first),
             ('filled_on', '<', month_end + relativedelta(months=1)),
             ('opened_on', '!=', False)] + me._filter_domain(filters), limit=SCAN_LIMIT)
        months = []
        for i in range(12):
            m0 = first + relativedelta(months=i)
            m1 = m0 + relativedelta(months=1)
            spans = [_span(r.filled_on, r.opened_on) for r in fills
                     if r.filled_on and m0 <= r.filled_on < m1]
            months.append({'month': str(m0), 'label': m0.strftime('%b'),
                           'year': m0.year, 'filled': len(spans),
                           'median': _median(spans)})
        all_spans = [_span(r.filled_on, r.opened_on) for r in fills
                     if r.filled_on and r.opened_on]
        rate = int(round(len(filled) * 100.0 / len(agreed))) if agreed else None
        return {
            'agreed': len(agreed), 'filled': len(filled), 'fill_rate': rate,
            'open_headcount': sum(r['open'] for r in headcount),
            'open_roles': len(open_roles),
            'headcount': headcount,
            'sources': sources[:8],
            'delays': delays,
            'months': months,
            'fill_median_12m': _median(all_spans),
            'tiles': [
                {'key': 'fill_rate', 'label': _('Fill rate'), 'value': rate, 'unit': '%',
                 'icon': 'target', 'hint': _('Filled roles out of the roles agreed in this range.')},
                {'key': 'open_headcount', 'label': _('People still to hire'),
                 'value': sum(r['open'] for r in headcount), 'unit': '', 'icon': 'users',
                 'hint': _('Across every open role right now: seats asked for, less the '
                           'people who have joined.')},
                {'key': 'open_roles', 'label': _('Roles open'), 'value': len(open_roles),
                 'unit': '', 'icon': 'briefcase', 'hint': _('Roles being set up or open for candidates now.')},
                {'key': 'fill_12m', 'label': _('Days to fill, last 12 months'),
                 'value': _median(all_spans), 'unit': _('days'), 'icon': 'clock',
                 'hint': _('The middle role filled in the twelve months shown on the chart.')},
            ],
        }

    # =====================================================================
    #  The workbook — per section, or everything (G-54)
    # =====================================================================
    @api.model
    def export_xlsx(self, date_from=None, date_to=None, kind='all', department_id=None,
                    country_id=None, recruiter_id=None):
        if not self.env['pb.hiring']._can_read():
            raise UserError(_(
                "The hiring numbers are for the hiring team. Ask them, or ask "
                "the HR team to add you."))
        kind = kind if kind in EXPORT_KINDS else 'all'
        board = self.get_board(date_from, date_to, department_id, country_id, recruiter_id)
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise UserError(_("This system cannot build spreadsheets at the moment. "
                              "Ask an administrator to look at it."))
        wb = openpyxl.Workbook()
        head = Font(bold=True, color='FFFFFF')
        fill = PatternFill('solid', fgColor='6355C7')
        stalled_fill = PatternFill('solid', fgColor='FDE2E4')
        made = []

        def sheet(title, columns, rows, highlight=None):
            ws = wb.active if not made else wb.create_sheet()
            made.append(title)
            ws.title = title[:31]
            for i, label in enumerate(columns, start=1):
                cell = ws.cell(row=1, column=i, value=label)
                cell.font, cell.fill = head, fill
                cell.alignment = Alignment(horizontal='center', wrap_text=True)
            for r, values in enumerate(rows, start=2):
                for i, value in enumerate(values, start=1):
                    cell = ws.cell(row=r, column=i, value=value)
                    if highlight and highlight(r - 2):
                        cell.fill = stalled_fill
            ws.column_dimensions['A'].width = 34
            for i in range(2, len(columns) + 1):
                ws.column_dimensions[get_column_letter(i)].width = 18
            ws.freeze_panes = 'A2'
            return ws

        def want(k):
            return kind == 'all' or kind == k

        f = board.get('filters') or {}
        scope = self._scope_words(board, f)
        if want('summary'):
            ws = sheet(_('Summary'), [_('What is measured'), _('Number'), _('Unit')],
                       [[t['label'], '' if t['value'] is None else t['value'], t['unit']]
                        for t in board.get('tiles') or []])
            n = len(board.get('tiles') or [])
            ws.cell(row=n + 3, column=1, value=board.get('headline') or '').font = Font(bold=True)
            ws.cell(row=n + 4, column=1, value=scope)
        fun = board.get('funnel') or {}
        if want('funnel'):
            sheet(_('Offer funnel'), [_('Step'), _('People'), _('Share of offered %'),
                                      _('Lost here'), _('How')],
                  [[s['label'], s['n'], s['pct'], s['lost'], s['lost_label']]
                   for s in fun.get('steps') or []])
            sheet(_('Offer drops'), [_('Reason'), _('People'), _('In their words')],
                  [[r['reason'], r['n'], ' | '.join(r['examples'])]
                   for r in fun.get('reasons') or []])
            sheet(_('Funnel by market'), [_('Market'), _('Offered'), _('Accepted'),
                                          _('Signed'), _('Joined'), _('Turned down'),
                                          _('Did not join')],
                  [[m['country'], m['offered'], m['accepted'], m['signed'], m['joined'],
                    m['declined'], m['dropped']] for m in fun.get('by_market') or []])
            sheet(_('Funnel by role'), [_('Role'), _('Market'), _('Offered'), _('Accepted'),
                                        _('Signed'), _('Joined'), _('Turned down'),
                                        _('Did not join')],
                  [[r['role'], r['country'], r['offered'], r['accepted'], r['signed'],
                    r['joined'], r['declined'], r['dropped']] for r in fun.get('by_role') or []])
        age = board.get('ageing') or {}
        if want('ageing'):
            rows = age.get('rows') or []
            sheet(_('Ageing'), [_('Candidate'), _('Role'), _('Stage'), _('Days since last move'),
                                _('Since'), _('Recruiter'), _('Stalled')],
                  [[r['candidate'], r['role'], r['stage'], r['days'], r['since'],
                    r['recruiter'], _('Yes') if r['stalled'] else '']
                   for r in rows], highlight=lambda i: rows[i]['stalled'])
            sheet(_('Ageing by stage'), [_('Stage'), _('People'), _('Stalled'),
                                         _('Typical days')],
                  [[s['stage'], s['n'], s['stalled'], s['median']] for s in age.get('by_stage') or []])
            sheet(_('Ageing by recruiter'), [_('Recruiter'), _('People'), _('Stalled')],
                  [[r['recruiter'], r['n'], r['stalled']] for r in age.get('by_recruiter') or []])
        lead = board.get('leadership') or {}
        if want('leadership'):
            sheet(_('Leadership'), [_('What is measured'), _('Number'), _('Unit')],
                  [[t['label'], '' if t['value'] is None else t['value'], t['unit']]
                   for t in lead.get('tiles') or []])
            sheet(_('Open headcount'), [_('Market'), _('Roles open'), _('People to hire')],
                  [[h['country'], h['roles'], h['open']] for h in lead.get('headcount') or []])
            sheet(_('Time to fill by month'), [_('Month'), _('Roles filled'), _('Typical days')],
                  [['%s %s' % (m['label'], m['year']), m['filled'], m['median']]
                   for m in lead.get('months') or []])
            sheet(_('Source mix'), [_('Source'), _('Candidates'), _('Share %'), _('Joined')],
                  [[s['name'], s['applicants'], s.get('share'), s['hires']]
                   for s in lead.get('sources') or []])
        if want('stages'):
            sheet(_('Time in stage'), [_('Stage'), _('Moves'), _('Average days'),
                                       _('Typical days')],
                  [[r['name'], r['moves'], r['mean'], r['median']] for r in board.get('stages') or []])
        if want('sources'):
            sheet(_('Where they came from'), [_('Source'), _('Candidates'), _('Joined'),
                                              _('Share %')],
                  [[r['name'], r['applicants'], r['hires'], r['rate']]
                   for r in board.get('sources') or []])
        if want('delays'):
            sheet(_('Interviews moved'), [_('Whose side'), _('Times'), _('At short notice')],
                  [[r['name'], r['count'], r['late']] for r in board.get('delays') or []])
            sheet(_('Nobody came'), [_('Who did not come'), _('Times')],
                  [[r['name'], r['count']] for r in board.get('no_shows') or []])
        if want('agency'):
            sheet(_('Agency or our own'), [_('Who filled it'), _('Roles'), _('Filled'),
                                           _('Typical days'), _('Average days')],
                  [[r['name'], r['roles'], r['filled'], r['median'], r['mean']]
                   for r in board.get('agency') or []])
        if want('roles'):
            sheet(_('Roles'), *self._roles_sheet(f))
        if not made:
            sheet(_('Summary'), [_('Nothing to show')], [])
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        names = {'all': _('Hiring'), 'funnel': _('Hiring offer funnel'),
                 'ageing': _('Hiring ageing'), 'leadership': _('Hiring leadership'),
                 'roles': _('Hiring roles'), 'summary': _('Hiring summary'),
                 'stages': _('Hiring time in stage'), 'sources': _('Hiring sources'),
                 'delays': _('Hiring interviews moved'), 'agency': _('Hiring agency')}
        return {
            'ok': True,
            'file_b64': base64.b64encode(out.read()).decode(),
            'filename': _('%(what)s %(from)s to %(to)s.xlsx', what=names.get(kind, _('Hiring')),
                          **{'from': board.get('from', ''), 'to': board.get('to', '')}),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'rows': board.get('requests', 0),
            'sheets': made,
        }

    @api.model
    def _scope_words(self, board, f):
        bits = [_('Roles agreed between %(from)s and %(to)s',
                  **{'from': board.get('from', ''), 'to': board.get('to', '')})]
        if f.get('department_id'):
            bits.append(self.env['hr.department'].sudo().browse(f['department_id']).name or '')
        if f.get('country_id'):
            bits.append(self.env['res.country'].sudo().browse(f['country_id']).name or '')
        if f.get('recruiter_id'):
            bits.append(_('recruited by %s',
                          self.env['res.users'].sudo().browse(f['recruiter_id']).name or ''))
        return ' · '.join(b for b in bits if b)

    @api.model
    def _roles_sheet(self, f):
        """The requisition status snapshot (G-54): every role, where it is."""
        co_ids = self.env.companies.ids or [self.env.company.id]
        roles = self.env['pb.hiring.requisition'].sudo().search(
            [('company_id', 'in', co_ids)] + self._filter_domain(f),
            order='opened_on desc, id desc', limit=SCAN_LIMIT)
        states, reqs = dict(REQUISITION_STATES), dict(REQUEST_STATES)
        today = fields.Date.context_today(self)
        rows = [[r.title or '', r.department_id.name or '', self._country_of(r).name or '',
                 states.get(r.state, ''), reqs.get(r.request_state, ''), r.headcount or 0,
                 r.filled_count or 0, str(r.opened_on or ''),
                 _span(r.filled_on or today, r.opened_on) if r.opened_on else '',
                 r.recruiter_id.name or ''] for r in roles]
        return ([_('Role'), _('Department'), _('Market'), _('Where it is'), _('Request'),
                 _('Seats'), _('Joined'), _('Opened'), _('Days open'), _('Recruiter')], rows)

# -*- coding: utf-8 -*-
"""RIZE W2 A3 — the offer, the closure and the numbers.

Every test here is written against a failure this phase could have and that
NOTHING AT RUNTIME WOULD REPORT:

  * a background check that did not actually block an offer would be a note
    dressed up as a control, and the screen would look identical;
  * an offer total that counted a one-off the way a pay package does would
    understate a first year by a sign-on bonus, on the one document a
    candidate weighs the whole decision on;
  * a closure that half ran would leave a person with an employee record, no
    joining date and no checklist — and the board would say they had joined;
  * a cover that answered "yes" without being asked WHICH ROLE would turn a
    fortnight's stand-in into the run of the place;
  * a token that reached a payload would be a candidate's offer, readable by
    anybody who could see the board.
"""

import base64
import io
import re
from datetime import date, timedelta

from odoo import fields
from odoo.exceptions import AccessError, UserError
from odoo.modules.module import get_module_path
from odoo.tests import TransactionCase, tagged

_RE_ADJACENT_STRINGS = re.compile(r"""["']\s*\n\s*["']""")

#: A one-pixel PNG. Small enough to be an argument, real enough to be a file.
_PNG = base64.b64decode(
    b'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQ'
    b'GAhKmMIQAAAABJRU5ErkJggg==')


def _src(*parts):
    path = get_module_path('pb_hiring')
    with open(path + '/' + '/'.join(parts), encoding='utf-8') as fh:
        return fh.read()


class OfferCase(TransactionCase):
    """One role, one candidate a panel picked, and the people around them."""

    def setUp(self):
        super().setUp()
        self.company = self.env.company
        Employee = self.env['hr.employee']
        self.boss = Employee.create({
            'name': 'DEMO A3 Boss', 'company_id': self.company.id})
        self.head = Employee.create({
            'name': 'DEMO A3 Head', 'company_id': self.company.id,
            'parent_id': self.boss.id})
        self.dept = self.env['hr.department'].create({
            'name': 'DEMO A3 Function', 'company_id': self.company.id,
            'manager_id': self.head.id})
        self.job = self.env['hr.job'].create({
            'name': 'DEMO A3 Role', 'company_id': self.company.id,
            'department_id': self.dept.id})
        self.applicant = self.env['hr.applicant'].create({
            'partner_name': 'DEMO A3 Candidate',
            'email_from': 'demo.a3.candidate@example.com',
            'job_id': self.job.id,
            'company_id': self.company.id,
        })
        self.req = self.env['pb.hiring.requisition'].sudo().create({
            'title': 'DEMO A3 Role',
            'department_id': self.dept.id,
            'company_id': self.company.id,
            'requested_by_id': self.head.id,
            'reporting_manager_id': self.boss.id,
            'headcount': 2,
            'budget_cost': 100.0,
            'requirements': 'Able to do the job.',
            'job_id': self.job.id,
            'opened_on': fields.Date.context_today(self.env.user)
            - timedelta(days=20),
            'state': 'open',
            'selected_applicant_id': self.applicant.id,
        })
        self.applicant.sudo().write({'pb_requisition_id': self.req.id})

    # ------------------------------------------------------------- helpers
    def _clear_bgv(self, bgv, result='ok'):
        bgv.item_ids.sudo().write({'result': result})
        bgv.invalidate_recordset(['state', 'pending_count', 'flag_count'])
        return bgv

    def _offer(self, lines=True):
        # The check is always answered first: the door it holds shut is the
        # thing `TestTheBackgroundCheck` proves, and every other test in this
        # file is about what is on the other side of it.
        self._clear_bgv(self.env['pb.hiring.bgv'].open_for(self.req.id))
        offer = self.env['pb.hiring.offer'].draft_for(self.req.id)
        if lines:
            for name, kind, amount, period in (
                ('Basic pay', 'earning', 25000000.0, 'monthly'),
                ('Allowance', 'earning', 2000000.0, 'monthly'),
                ('Sign-on bonus', 'bonus', 10000000.0, 'one_time'),
            ):
                offer.action_add_line({'name': name, 'kind': kind,
                                       'amount': amount, 'period': period})
            offer.invalidate_recordset(['monthly_total', 'annual_total'])
        return offer

    def _complete_documents(self, offer):
        docreq = self.env['pb.hiring.docreq'].open_for(offer.id)
        docreq.sudo().write({'sent_on': fields.Datetime.now(),
                             'deadline': fields.Date.context_today(
                                 self.env.user) + timedelta(days=2)})
        for item in docreq.item_ids.filtered(lambda i: i.required):
            docreq.receive(item.id, 'passport.png', _PNG,
                           mimetype='image/png')
        docreq.invalidate_recordset(['state', 'in_count'])
        return docreq


# =========================================================================
#  T2 — the background check is a DOOR
# =========================================================================
@tagged('post_install', '-at_install')
class TestTheBackgroundCheck(OfferCase):

    def test_t2_the_check_seeds_itself_from_the_shipped_lines(self):
        bgv = self.env['pb.hiring.bgv'].open_for(self.req.id)
        self.assertTrue(bgv.item_ids,
                        'a check with no lines is a check nobody fills in')
        self.assertEqual(bgv.state, 'open')
        # Idempotent: a second press finds the same row, not a second one.
        again = self.env['pb.hiring.bgv'].open_for(self.req.id)
        self.assertEqual(bgv.id, again.id)
        self.assertEqual(len(bgv.item_ids), len(again.item_ids))

    def test_t2_an_offer_is_refused_while_a_line_is_unanswered(self):
        with self.assertRaises(UserError) as caught:
            self.env['pb.hiring.offer'].draft_for(self.req.id)
        # THE SENTENCE NAMES THE LINES. "The check is not complete" is a wall;
        # naming them is an instruction.
        message = str(caught.exception)
        first = self.env['pb.hiring.bgv'].open_for(self.req.id).item_ids[0]
        self.assertIn(first.name, message)

    def test_t2_answering_every_line_opens_the_door(self):
        bgv = self._clear_bgv(self.env['pb.hiring.bgv'].open_for(self.req.id))
        self.assertEqual(bgv.state, 'complete')
        offer = self.env['pb.hiring.offer'].draft_for(self.req.id)
        self.assertTrue(offer.exists())

    def test_t2_a_flag_refuses_until_somebody_says_to_go_ahead(self):
        bgv = self._clear_bgv(self.env['pb.hiring.bgv'].open_for(self.req.id))
        bgv.item_ids[0].action_set('flag', note='They would only confirm '
                                                'the dates.')
        bgv.invalidate_recordset(['state', 'flag_count'])
        self.assertEqual(bgv.state, 'flagged')
        with self.assertRaises(UserError):
            self.env['pb.hiring.offer'].draft_for(self.req.id)
        bgv.action_override(note='Policy of that employer. Known in advance.')
        self.assertTrue(bgv.override_user_id)
        offer = self.env['pb.hiring.offer'].draft_for(self.req.id)
        self.assertTrue(offer.exists())

    def test_t2_going_ahead_needs_a_reason_in_words(self):
        bgv = self._clear_bgv(self.env['pb.hiring.bgv'].open_for(self.req.id))
        bgv.item_ids[0].action_set('flag')
        with self.assertRaises(UserError):
            bgv.action_override(note='  ')

    def test_t2_does_not_apply_is_a_real_answer_and_not_a_gap(self):
        bgv = self.env['pb.hiring.bgv'].open_for(self.req.id)
        bgv.item_ids.sudo().write({'result': 'na'})
        bgv.invalidate_recordset(['state', 'pending_count'])
        self.assertEqual(bgv.state, 'complete')

    def test_t2_who_answered_is_stamped_without_anybody_filling_a_field(self):
        bgv = self.env['pb.hiring.bgv'].open_for(self.req.id)
        item = bgv.item_ids[0]
        item.action_set('ok')
        self.assertEqual(item.checked_by_id.id, self.env.uid)
        self.assertTrue(item.checked_on)


# =========================================================================
#  T3 — the papers
# =========================================================================
@tagged('post_install', '-at_install')
class TestTheDocumentRequest(OfferCase):

    def test_t3_asking_mints_a_link_and_a_working_day_deadline(self):
        offer = self._offer()
        docreq = offer.action_request_documents()
        self.assertTrue(docreq.token and len(docreq.token) >= 20)
        self.assertTrue(docreq.deadline)
        self.assertGreaterEqual(docreq.deadline,
                                fields.Date.context_today(self.env.user))
        self.assertIn('/hiring/d/', docreq._token_url())
        self.assertTrue(docreq.item_ids)

    def test_t3_one_file_at_a_time_moves_it_through_its_states(self):
        offer = self._offer()
        docreq = offer.action_request_documents()
        required = docreq.item_ids.filtered(lambda i: i.required)
        self.assertEqual(docreq.state, 'sent')
        docreq.receive(required[0].id, 'passport.png', _PNG,
                       mimetype='image/png')
        docreq.invalidate_recordset(['state', 'in_count'])
        self.assertEqual(docreq.state, 'partial')
        for item in required[1:]:
            docreq.receive(item.id, 'thing.pdf', b'%PDF-1.4 hello',
                           mimetype='application/pdf')
        docreq.invalidate_recordset(['state', 'in_count'])
        self.assertEqual(docreq.state, 'complete')

    def test_t3_a_file_that_is_too_big_is_refused_with_a_sentence(self):
        offer = self._offer()
        docreq = offer.action_request_documents()
        with self.assertRaises(UserError):
            docreq.receive(docreq.item_ids[0].id, 'huge.pdf',
                           b'x' * (5 * 1024 * 1024 + 1),
                           mimetype='application/pdf')

    def test_t3_a_kind_of_file_we_cannot_open_is_refused(self):
        offer = self._offer()
        docreq = offer.action_request_documents()
        with self.assertRaises(UserError):
            docreq.receive(docreq.item_ids[0].id, 'thing.exe', b'MZ',
                           mimetype='application/x-msdownload')

    def test_t3_a_line_that_is_not_on_this_request_is_refused(self):
        """The line id is the ONE thing the public form sends, so it is
        checked against this request's own lines and nothing else."""
        offer = self._offer()
        docreq = offer.action_request_documents()
        with self.assertRaises(UserError):
            docreq.receive(0, 'x.pdf', b'%PDF', mimetype='application/pdf')
        with self.assertRaises(UserError):
            docreq.receive(999999, 'x.pdf', b'%PDF',
                           mimetype='application/pdf')

    def test_t3_a_reminder_goes_once_a_day_and_never_twice(self):
        offer = self._offer()
        docreq = offer.action_request_documents()
        self.assertTrue(docreq.action_remind())
        self.assertFalse(docreq.action_remind(),
                         'a candidate chased twice in a morning stops reading')

    def test_t3_the_daily_job_is_idempotent(self):
        offer = self._offer()
        docreq = offer.action_request_documents()
        docreq.sudo().write({'last_reminder_on': False})
        first = self.env['pb.hiring.automation']._chase_documents()
        second = self.env['pb.hiring.automation']._chase_documents()
        self.assertGreaterEqual(first, 1)
        self.assertEqual(second, 0)

    def test_t3_a_replaced_file_leaves_no_second_answer(self):
        offer = self._offer()
        docreq = offer.action_request_documents()
        item = docreq.item_ids[0]
        docreq.receive(item.id, 'one.png', _PNG, mimetype='image/png')
        first = item.attachment_id
        docreq.receive(item.id, 'two.png', _PNG, mimetype='image/png')
        self.assertNotEqual(item.attachment_id.id, first.id)
        self.assertFalse(first.exists(),
                         'two downloadable answers to one question')

    def test_t3_the_token_resolves_and_a_junk_key_does_not(self):
        offer = self._offer()
        docreq = offer.action_request_documents()
        Docreq = self.env['pb.hiring.docreq']
        row, status = Docreq._request_for_token(docreq.token)
        self.assertEqual((row.id, status), (docreq.id, 'ok'))
        self.assertEqual(Docreq._request_for_token('short')[1], 'invalid')
        self.assertEqual(Docreq._request_for_token('a' * 30)[1], 'invalid')

    def test_t3_the_page_says_nothing_about_the_money(self):
        offer = self._offer()
        docreq = offer.action_request_documents()
        facts = docreq.page_facts()
        blob = str(facts)
        self.assertNotIn('25000000', blob)
        self.assertNotIn(docreq.token, blob)


# =========================================================================
#  T4 — the offer, the totals and the letter
# =========================================================================
@tagged('post_install', '-at_install')
class TestTheOffer(OfferCase):

    def test_t4_the_two_totals(self):
        """A month is the monthly lines; a YEAR counts the one-off ONCE.

        The pay package scores a one-off at zero, which is right for "what is
        this person paid every year" and wrong for an offer: a sign-on bonus
        is real money in the first year and it is the number a candidate
        weighs the whole decision on.
        """
        offer = self._offer()
        self.assertAlmostEqual(offer.monthly_total, 27000000.0, places=2)
        self.assertAlmostEqual(offer.annual_total, 334000000.0, places=2)

    def test_t4_the_letter_leaves_no_placeholder_behind(self):
        offer = self._offer()
        offer.action_prepare_letter()
        self.assertTrue(offer.rendered_html)
        self.assertNotIn('${', str(offer.rendered_html),
                         'a hole the offer could not fill is on the letter')
        self.assertIn(offer.candidate_name, str(offer.rendered_html))
        self.assertTrue(offer.attachment_id,
                        'the PDF is what is attached to the email')

    def test_t4_a_letter_with_no_lines_is_refused(self):
        offer = self._offer(lines=False)
        with self.assertRaises(UserError):
            offer.action_prepare_letter()

    def test_t4_an_offer_with_no_lines_cannot_be_sent_in(self):
        offer = self._offer(lines=False)
        with self.assertRaises(UserError):
            offer.action_submit()

    def test_t4_the_numbers_are_in_the_revision_stamp(self):
        """A stamp that did not carry the money would let somebody change
        25 million into 30 million after the hiring manager agreed it."""
        offer = self._offer()
        before = offer._chain_revision_values()
        offer.line_ids[0].write({'amount': 30000000.0})
        self.assertNotEqual(before, offer._chain_revision_values())

    def test_t4_a_country_template_beats_the_generic_one(self):
        vietnam = self.env.ref('base.vn', raise_if_not_found=False)
        if not vietnam:
            self.skipTest('this database has no Vietnam')
        self.req.sudo().write({'country_id': vietnam.id})
        picked = self.env['pb.hiring.offer']._pick_template(self.req)
        self.assertTrue(picked)
        self.assertIn('Vietnam', picked.name)

    def test_t4_the_offer_line_is_editable_only_while_it_is_a_draft(self):
        """A route stamps the numbers when it is sent in and re-reads them at
        the last rung, so a figure changed mid-route does not reopen anything
        — it makes the final approval refuse to carry out. Refusing the edit
        at the door is the same rule said while somebody can still act."""
        offer = self._offer()
        offer.sudo().write({'state': 'submitted'})
        with self.assertRaises(UserError):
            self.env['pb.hiring'].act('offer_line', {
                'offer_id': offer.id, 'name': 'Late extra', 'amount': 1.0})

    def test_t4_two_offers_are_never_opened_for_one_candidate(self):
        first = self._offer()
        again = self.env['pb.hiring.offer'].draft_for(self.req.id)
        self.assertEqual(first.id, again.id)

    def test_t4_an_offer_cannot_be_drafted_before_anybody_is_picked(self):
        self.req.sudo().write({'selected_applicant_id': False})
        with self.assertRaises(UserError):
            self.env['pb.hiring.offer'].draft_for(self.req.id)


# =========================================================================
#  T5 — sending it, and the candidate's own page
# =========================================================================
@tagged('post_install', '-at_install')
class TestSendingIt(OfferCase):

    def _agreed(self):
        offer = self._offer()
        offer.sudo().write({'state': 'hr_ok'})
        return offer

    def test_t5_it_is_refused_while_the_papers_are_missing(self):
        offer = self._agreed()
        with self.assertRaises(UserError) as caught:
            offer.action_send_to_candidate()
        self.assertIn('papers', str(caught.exception).lower())

    def test_t5_the_hr_lead_can_send_it_anyway(self):
        offer = self._agreed()
        offer.action_request_documents()
        offer.sudo().action_send_to_candidate(force=True)
        self.assertEqual(offer.state, 'sent')

    def test_t5_with_the_papers_in_it_goes(self):
        offer = self._agreed()
        self._complete_documents(offer)
        Mail = self.env['mail.mail']
        before = Mail.sudo().search_count([])
        offer.action_send_to_candidate()
        self.assertEqual(offer.state, 'sent')
        self.assertTrue(offer.sent_on)
        self.assertGreater(Mail.sudo().search_count([]), before,
                           'an offer that is agreed and never sent is the '
                           'most expensive thing that can go wrong here')

    def test_t5_the_page_answers_and_a_second_visit_shows_the_decision(self):
        offer = self._agreed()
        self._complete_documents(offer)
        offer.action_send_to_candidate()
        Offer = self.env['pb.hiring.offer']
        row, status = Offer._request_for_token(offer.token)
        self.assertEqual((row.id, status), (offer.id, 'ok'))
        self.assertTrue(row.record_decision('accepted', comment='Delighted.'))
        self.assertEqual(offer.state, 'accepted')
        self.assertEqual(offer.candidate_decision, 'accepted')
        self.assertEqual(Offer._request_for_token(offer.token)[1], 'used')
        # A replay writes nothing.
        self.assertFalse(offer.record_decision('declined'))
        self.assertEqual(offer.candidate_decision, 'accepted')

    def test_t5_a_decline_leaves_the_candidate_exactly_where_they_are(self):
        offer = self._agreed()
        self._complete_documents(offer)
        offer.action_send_to_candidate()
        stage = self.applicant.stage_id
        offer.record_decision('declined', comment='Counter-offer at home.')
        self.assertEqual(offer.state, 'declined')
        self.assertEqual(self.applicant.stage_id, stage)
        self.assertTrue(self.applicant.active,
                        'a person who says no to one offer is often the '
                        'person who says yes to the next role')

    def test_t5_an_offer_nobody_has_agreed_shows_a_closed_page(self):
        offer = self._offer()
        self.assertEqual(
            self.env['pb.hiring.offer']._request_for_token(offer.token)[1],
            'closed')

    def test_t5_the_page_carries_no_token_and_no_id_of_anybody_else(self):
        offer = self._agreed()
        blob = str(offer.page_facts())
        self.assertNotIn(offer.token, blob)


# =========================================================================
#  T6 — the closure
# =========================================================================
@tagged('post_install', '-at_install')
class TestTheClosure(OfferCase):

    def _signed(self):
        offer = self._offer()
        offer.sudo().write({'state': 'hr_ok'})
        self._complete_documents(offer)
        offer.action_send_to_candidate()
        offer.record_decision('accepted')
        offer.action_record_signed(filename='signed.pdf',
                                   content=b'%PDF-1.4 signed',
                                   mimetype='application/pdf')
        return offer

    def test_t6_a_signature_with_nothing_behind_it_is_refused(self):
        offer = self._offer()
        offer.sudo().write({'state': 'accepted'})
        with self.assertRaises(UserError):
            offer.action_record_signed()

    def test_t6_the_whole_chain_runs_in_one_press(self):
        offer = self._signed()
        self.assertEqual(offer.state, 'signed')
        offer.action_close()
        self.assertEqual(offer.state, 'closed')

        # the employee
        self.assertTrue(offer.employee_id)
        self.assertEqual(offer.employee_id.sudo().work_email,
                         self.applicant.email_from,
                         'a joiner must not be given the switchboard mailbox')
        self.assertEqual(offer.employee_id.sudo().company_id, self.company)
        self.assertEqual(offer.employee_id.sudo().parent_id, self.boss)

        # the candidate is hired
        self.assertTrue(self.applicant.sudo().stage_id)

        # the contract: a joining date is a contract, not a field (R77)
        self.assertTrue(offer.contract_id)
        self.assertEqual(offer.contract_id.sudo().date_start,
                         offer.start_date)
        self.assertAlmostEqual(offer.contract_id.sudo().wage, 25000000.0,
                               places=2)

        # the pay package, in draft with every line ticked
        self.assertTrue(offer.comp_id)
        self.assertEqual(offer.comp_id.sudo().state, 'draft')
        self.assertEqual(len(offer.comp_id.sudo().line_ids), 3)
        self.assertTrue(all(offer.comp_id.sudo().line_ids.mapped('checked')))
        self.assertEqual(offer.comp_id.sudo().effective_date,
                         offer.start_date)

        # the joining checklist
        self.assertTrue(offer.case_id)

        # the signed copy, filed
        docs = self.env['pb.employee.document'].sudo().search(
            [('employee_id', '=', offer.employee_id.id)])
        self.assertTrue(docs, 'the signed offer belongs on their record')

    def test_t6_closing_twice_makes_nothing_twice(self):
        offer = self._signed()
        offer.action_close()
        employee, contract = offer.employee_id, offer.contract_id
        case, comp = offer.case_id, offer.comp_id
        offer.action_close()
        self.assertEqual(offer.employee_id, employee)
        self.assertEqual(offer.contract_id, contract)
        self.assertEqual(offer.case_id, case)
        self.assertEqual(offer.comp_id, comp)

    def test_t6_the_role_fills_only_when_the_head_count_is_reached(self):
        offer = self._signed()
        offer.action_close()
        self.req.invalidate_recordset(['filled_count'])
        self.assertEqual(self.req.filled_count, 1)
        self.assertEqual(self.req.state, 'open',
                         'a request for two is not filled by the first')

        second = self.env['hr.applicant'].create({
            'partner_name': 'DEMO A3 Second Candidate',
            'email_from': 'demo.a3.second@example.com',
            'job_id': self.job.id, 'company_id': self.company.id,
            'pb_requisition_id': self.req.id,
        })
        self.req.sudo().write({'selected_applicant_id': second.id})
        offer2 = self._offer()
        offer2.sudo().write({'state': 'hr_ok'})
        self._complete_documents(offer2)
        offer2.action_send_to_candidate()
        offer2.record_decision('accepted')
        offer2.action_record_signed(filename='signed.pdf',
                                    content=b'%PDF-1.4 signed',
                                    mimetype='application/pdf')
        offer2.action_close()
        self.req.invalidate_recordset(['filled_count'])
        self.assertEqual(self.req.filled_count, 2)
        self.assertEqual(self.req.state, 'filled')
        self.assertTrue(self.req.filled_on)
        self.assertFalse(self.req.referral_open)
        self.assertFalse(self.job.sudo().is_published,
                         'the advert comes off the careers page')

    def test_t6_the_switch_turns_the_contract_off_and_says_so(self):
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_hiring.create_contract', '0')
        offer = self._signed()
        offer.action_close()
        self.assertTrue(offer.employee_id)
        self.assertFalse(offer.contract_id)
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_hiring.create_contract', '1')

    def test_t6_an_offer_that_is_not_signed_cannot_be_closed(self):
        offer = self._offer()
        with self.assertRaises(UserError):
            offer.action_close()


# =========================================================================
#  T7 — the cover
# =========================================================================
@tagged('post_install', '-at_install')
class TestTheCover(OfferCase):

    def setUp(self):
        super().setUp()
        self.recruiter = self.env['res.users'].sudo().create({
            'name': 'DEMO A3 Recruiter', 'login': 'demo.a3.recruiter',
            'company_id': self.company.id,
            'company_ids': [(6, 0, [self.company.id])],
            'group_ids': [(4, self.env.ref('pb_hiring.group_hiring_user').id)],
        })
        self.stand_in = self.env['res.users'].sudo().create({
            'name': 'DEMO A3 Stand-in', 'login': 'demo.a3.standin',
            'company_id': self.company.id,
            'company_ids': [(6, 0, [self.company.id])],
            'group_ids': [(4, self.env.ref('base.group_user').id)],
        })
        self.req.sudo().write({'recruiter_id': self.recruiter.id})

    def _cover(self, **extra):
        # THE SERVER'S OWN DATE, because that is the clock the gate reads —
        # a window built on the reader's timezone is a window that is live
        # for one colleague and not for another (see `_window_today`).
        today = fields.Date.today()
        vals = {'recruiter_id': self.recruiter.id,
                'cover_user_id': self.stand_in.id,
                'date_from': today, 'date_to': today + timedelta(days=14),
                'reason': 'Two weeks away.',
                'company_id': self.company.id}
        vals.update(extra)
        return self.env['pb.hiring.cover'].sudo().create(vals)

    def test_t7_nobody_can_cover_for_themselves(self):
        from odoo.exceptions import ValidationError
        with self.assertRaises(ValidationError):
            self._cover(cover_user_id=self.recruiter.id)

    def test_t7_a_cover_that_ends_before_it_starts_is_refused(self):
        from odoo.exceptions import ValidationError
        today = fields.Date.today()
        with self.assertRaises(ValidationError):
            self._cover(date_to=today - timedelta(days=1))

    def test_t7_an_agreed_cover_names_exactly_the_roles_it_covers(self):
        cover = self._cover()
        cover.sudo().write({'state': 'approved'})
        Cover = self.env['pb.hiring.cover']
        covered = Cover.with_user(self.stand_in).covered_recruiter_uids()
        self.assertEqual(covered, [self.recruiter.id])
        # The answer must NOT depend on who is asking: two people in two
        # timezones reading the same cover have to get the same answer.
        self.assertEqual(Cover.covered_recruiter_uids(self.stand_in),
                         [self.recruiter.id])
        # Somebody with no cover covers nobody.
        self.assertEqual(
            Cover.with_user(self.recruiter).covered_recruiter_uids(), [])

    def test_t7_the_gate_asks_WHICH_role(self):
        cover = self._cover()
        cover.sudo().write({'state': 'active'})
        Facade = self.env['pb.hiring'].with_user(self.stand_in)
        # Their own recruiter's role: allowed.
        self.assertTrue(Facade._require_recruit(self.req))
        # Somebody else's: refused, by name.
        other = self.env['pb.hiring.requisition'].sudo().create({
            'title': 'DEMO A3 Other Role',
            'department_id': self.dept.id, 'company_id': self.company.id,
            'requested_by_id': self.head.id, 'headcount': 1,
            'recruiter_id': self.env.user.id,
        })
        with self.assertRaises(AccessError):
            Facade._require_recruit(other)
        # And a verb with no role named at all is refused too: a cover that
        # could not see which role it was about would be the run of the place.
        with self.assertRaises(AccessError):
            Facade._require_recruit()

    def test_t7_a_cover_outside_its_window_covers_nobody(self):
        today = fields.Date.today()
        cover = self._cover(date_from=today - timedelta(days=30),
                            date_to=today - timedelta(days=10))
        cover.sudo().write({'state': 'active'})
        self.assertEqual(
            self.env['pb.hiring.cover'].with_user(
                self.stand_in).covered_recruiter_uids(), [])

    def test_t7_the_daily_step_starts_and_ends_them_and_is_idempotent(self):
        today = fields.Date.today()
        due = self._cover()
        due.sudo().write({'state': 'approved'})
        over = self._cover(date_from=today - timedelta(days=30),
                           date_to=today - timedelta(days=1))
        over.sudo().write({'state': 'active'})
        counts = self.env['pb.hiring.cover'].run_window()
        self.assertGreaterEqual(counts['started'], 1)
        self.assertGreaterEqual(counts['ended'], 1)
        self.assertEqual(due.state, 'active')
        self.assertEqual(over.state, 'ended')
        again = self.env['pb.hiring.cover'].run_window()
        self.assertEqual(again, {'started': 0, 'ended': 0})

    def test_t7_who_is_asked_to_agree_it_is_said_out_loud(self):
        """"Your manager approved this" means nothing if nobody can tell
        whether "your manager" came from the hiring rule or the org chart."""
        cover = self._cover()
        self.assertIn(cover.approver_source, ('rule', 'chart', 'none'))
        title = cover._chain_title()
        self.assertIn(cover.cover_user_id.name, title)
        self.assertIn(cover.recruiter_id.name, title)
        self.assertTrue(re.search(r'\(.+\)$', title.strip()),
                        'the title must say where the approver came from')


# =========================================================================
#  T8 — the agency
# =========================================================================
@tagged('post_install', '-at_install')
class TestTheAgency(OfferCase):

    def setUp(self):
        super().setUp()
        self.vendor = self.env['pb.vendor'].sudo().create({
            'name': 'DEMO A3 Talent Partners',
            'vendor_type': 'recruitment',
            'responsible_user_id': self.env.uid,
            'company_id': self.company.id,
        })

    def test_t8_the_vendor_card_counts_what_the_agency_delivered(self):
        self.req.sudo().write({'agency_vendor_id': self.vendor.id})
        self.vendor.invalidate_recordset(['hiring_count', 'hiring_avg_days',
                                          'hiring_open_count'])
        self.assertEqual(self.vendor.hiring_count, 0)
        self.assertEqual(self.vendor.hiring_open_count, 1)
        # A role that filled is timed; a role still open is not.
        self.req.sudo().write({
            'filled_on': fields.Date.context_today(self.env.user)})
        self.vendor.invalidate_recordset(['hiring_avg_days'])
        self.assertEqual(self.vendor.hiring_avg_days, 20)

    def test_t8_the_stat_button_opens_only_that_agency_s_roles(self):
        self.req.sudo().write({'agency_vendor_id': self.vendor.id})
        action = self.vendor.action_open_hiring()
        self.assertIn('views', action, 'R125 — a hand-built action needs it')
        self.assertIn(('agency_vendor_id', '=', self.vendor.id),
                      action['domain'])

    def test_t8_setting_an_agency_is_written_in_the_chatter(self):
        before = len(self.req.message_ids)
        self.req.sudo().write({'agency_vendor_id': self.vendor.id})
        self.assertGreater(len(self.req.message_ids), before)


# =========================================================================
#  T9 — the numbers
# =========================================================================
@tagged('post_install', '-at_install')
class TestTheNumbers(OfferCase):

    def test_t9_an_empty_range_says_so_in_words(self):
        board = self.env['pb.hiring.analytics'].get_board(
            '2001-01-01', '2001-12-31')
        self.assertTrue(board['empty'])
        self.assertIn('2001', board['headline'])
        self.assertFalse(board['tiles'],
                         'eight zeros is how a working screen gets reported '
                         'as broken (R27)')

    def test_t9_the_figures_equal_a_python_recomputation(self):
        offer = self._offer()
        offer.sudo().write({'state': 'hr_ok'})
        self._complete_documents(offer)
        offer.action_send_to_candidate()
        offer.record_decision('accepted')
        offer.action_record_signed(filename='s.pdf', content=b'%PDF',
                                   mimetype='application/pdf')
        offer.action_close()

        today = fields.Date.context_today(self.env.user)
        board = self.env['pb.hiring.analytics'].get_board(
            today - timedelta(days=90), today)
        tiles = {t['key']: t['value'] for t in board['tiles']}

        Requisition = self.env['pb.hiring.requisition'].sudo()
        cohort = Requisition.search([
            ('company_id', 'in', self.env.companies.ids
             or [self.env.company.id]),
            ('opened_on', '>=', today - timedelta(days=90)),
            ('opened_on', '<=', today)])
        self.assertEqual(tiles['requests'], len(cohort))
        spans = [(r.filled_on - r.opened_on).days for r in cohort
                 if r.filled_on and r.opened_on]
        self.assertEqual(tiles['filled'], len(spans))
        sent = self.env['pb.hiring.offer'].sudo().search_count([
            ('requisition_id', 'in', cohort.ids), ('sent_on', '!=', False)])
        self.assertEqual(tiles['offers'], sent)

    def test_t9_the_spreadsheet_opens_and_says_the_same_thing(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl is not on this system')
        self._offer()
        today = fields.Date.context_today(self.env.user)
        res = self.env['pb.hiring.analytics'].export_xlsx(
            today - timedelta(days=90), today)
        self.assertTrue(res['ok'])
        self.assertTrue(res['filename'].endswith('.xlsx'))
        book = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(res['file_b64'])))
        self.assertIn('Summary', book.sheetnames)
        self.assertIn('Where they came from', book.sheetnames)
        summary = book['Summary']
        self.assertEqual(summary.cell(row=1, column=1).value,
                         'What is measured')

    def test_t9_a_role_still_open_is_never_counted_as_taking_so_far(self):
        today = fields.Date.context_today(self.env.user)
        board = self.env['pb.hiring.analytics'].get_board(
            today - timedelta(days=90), today)
        tiles = {t['key']: t['value'] for t in board['tiles']}
        if not tiles['filled']:
            self.assertIsNone(tiles['fill_median'],
                              'an unfilled role has no time-to-fill and a '
                              'zero would be a lie')


# =========================================================================
#  T1 / T12 — the gates
# =========================================================================
@tagged('post_install', '-at_install')
class TestA3SourceGates(TransactionCase):

    def test_t1_both_new_routes_are_in_the_catalogue_and_wired_up(self):
        for xmlid, model in (
            ('pb_hiring.process_hiring_offer', 'pb.hiring.offer'),
            ('pb_hiring.process_hiring_cover', 'pb.hiring.cover'),
        ):
            process = self.env.ref(xmlid)
            self.assertEqual(process.model_name, model)
            self.assertTrue(process.key)

    def test_t1_seeding_the_routes_twice_lays_one_of_everything(self):
        """`post_init_hook` runs on INSTALL only, so the migration lays the
        same routes on a database that already has the module — which means
        both run on some databases and neither may double anything."""
        from odoo.addons.pb_hiring.models.offer_approval import seed_all as so
        from odoo.addons.pb_hiring.models.cover_approval import seed_all as sc
        first = (so(self.env), sc(self.env))
        second = (so(self.env), sc(self.env))
        self.assertEqual(first, second,
                         'seeding is not idempotent across companies')

    def test_t1_the_migration_is_guarded_on_a_version(self):
        src = _src('migrations', '19.0.1.2.0', 'post-offer_approval.py')
        self.assertIn('if not version:', src,
                      'without the guard it runs on a fresh install too')

    def test_t1_the_shipped_lists_are_company_less(self):
        for model in ('pb.hiring.bgv.template', 'pb.hiring.doc.template'):
            rows = self.env[model].sudo().search([])
            self.assertTrue(rows, '%s ships nothing' % model)
            self.assertFalse(
                rows.filtered(lambda r: r.company_id),
                'a seed with a company installs onto one company and is '
                'invisible to every other (R8)')

    def test_t1_the_offer_letter_type_reaches_the_letter_for_free(self):
        """A RELATED SELECTION TAKES ITS LIST FROM THE SOURCE FIELD, so
        adding the value to the template is the whole of the change — and
        adding it to the letter as well trips a hard assertion that fails the
        entire registry load."""
        keys = dict(self.env['pb.letter.template']
                    ._fields['letter_type'].selection)
        self.assertIn('offer', keys)
        letter = self.env['pb.hr.letter']._fields['letter_type']
        self.assertTrue(letter.related,
                        'if this ever stops being related, the offer value '
                        'has to be added here too')

    def test_t1_two_offer_letters_ship_and_both_have_a_body(self):
        rows = self.env['pb.letter.template'].sudo().search(
            [('letter_type', '=', 'offer')])
        self.assertGreaterEqual(len(rows), 2)
        for row in rows:
            self.assertTrue(row.body_html)
            self.assertIn('${candidate_name}', str(row.body_html))
            self.assertIn('${lines}', str(row.body_html))

    def test_t1_every_mail_this_phase_sends_resolves(self):
        for xmlid in ('mail_template_docreq_ask',
                      'mail_template_docreq_remind',
                      'mail_template_offer_candidate',
                      'mail_template_offer_answered',
                      'mail_template_offer_closed',
                      'mail_template_requisition_filled',
                      'mail_template_requisition_agency'):
            self.assertTrue(self.env.ref('pb_hiring.%s' % xmlid),
                            'a mail nobody can resolve is a mail nobody gets')

    def test_t12_the_palette_stays_inside_the_3500_block(self):
        src = _src('static', 'src', 'js', 'hiring_palette.js')
        numbers = [int(n) for n in re.findall(r'sequence:\s*(\d{4})', src)]
        self.assertTrue(numbers)
        self.assertTrue(all(3500 <= n < 3600 for n in numbers),
                        'B1 starts at 3600: %s' % numbers)
        self.assertIn('3570', src)
        self.assertIn('3580', src)

    def test_t12_the_two_new_doors_point_at_real_actions(self):
        for xmlid in ('pb_hiring.action_pb_hiring_numbers',
                      'pb_hiring.action_pb_hiring_cover',
                      'pb_hiring.action_pb_hiring_offer',
                      'pb_hiring.action_pb_hiring_bgv_template',
                      'pb_hiring.action_pb_hiring_doc_template'):
            self.assertTrue(self.env.ref(xmlid), xmlid)

    def test_t12_the_insights_lens_is_registered_at_thirty(self):
        src = _src('static', 'src', 'js', 'hiring_palette.js')
        self.assertIn('INSIGHTS_LENSES', src)
        block = src[src.index('INSIGHTS_LENSES).add'):]
        self.assertIn('sequence: 30', block[:800])

    def test_every_icon_name_including_the_MAPS_is_in_the_registry(self):
        """THE GATE READS THE OBJECT-LITERAL MAPS TOO (R146).

        A gate that matched only `icon: "…"` never saw
        `SCREEN_ICON = { rejected: "xCircle" }` — which shipped against a
        registry that did not have it, and `ic()` falls back to a plain
        circle with no error at all. And it is checked against the INSTALLED
        copy (R147): a local check against the repo passes while the live
        screen draws blank circles.
        """
        registry_src = _src('..', 'pb_import_kit', 'static', 'src', 'js',
                            'import_icons.js')
        known = set(re.findall(r'^\s{4}([A-Za-z][A-Za-z0-9]*):',
                               registry_src, re.M))
        self.assertGreater(len(known), 60, 'the icon registry did not parse')
        used = set()
        for name in ('hiring_board.js', 'hiring_numbers.js',
                     'hiring_palette.js'):
            src = _src('static', 'src', 'js', name)
            used |= set(re.findall(r'\bic\(\s*["\']([A-Za-z0-9]+)["\']', src))
            used |= set(re.findall(r'icon:\s*["\']([A-Za-z0-9]+)["\']', src))
            for block in re.findall(r'_ICON\s*=\s*\{(.*?)\}', src, re.S):
                used |= set(re.findall(r':\s*["\']([A-Za-z0-9]+)["\']', block))
        # The analytics tiles name their icon on the SERVER, so the gate has
        # to read that file too or it checks half of them.
        server = _src('models', 'analytics.py')
        used |= set(re.findall(r"'icon':\s*'([A-Za-z0-9]+)'", server))
        missing = sorted(used - known)
        self.assertFalse(missing, 'not in the shared registry: %s' % missing)

    def test_no_python_style_implicit_string_concatenation(self):
        """Two adjacent string literals are a SyntaxError that blanks the
        ENTIRE backend asset bundle for every user (R2)."""
        for name in ('hiring_board.js', 'hiring_numbers.js',
                     'hiring_palette.js'):
            src = _src('static', 'src', 'js', name)
            body = re.sub(r'//[^\n]*', '', src)
            body = re.sub(r'/\*.*?\*/', '', body, flags=re.S)
            self.assertIsNone(_RE_ADJACENT_STRINGS.search(body), name)

    def test_the_word_odoo_appears_in_no_user_visible_string(self):
        """The white-label rule, and only where it actually binds.

        It covers user-visible STRINGS. Engineering comments MUST be able to
        say the real name (R118), so the comments are stripped and everything
        that is left is checked — which is the half a person can read.
        """
        for parts in (('views', 'offer_views.xml'),
                      ('views', 'offer_token_templates.xml'),
                      ('views', 'vendor_views.xml'),
                      ('data', 'mail_template_offer.xml'),
                      ('data', 'letter_template_offer.xml'),
                      ('data', 'hiring_bgv_templates.xml'),
                      ('data', 'hiring_doc_templates.xml'),
                      ('report', 'hiring_offer_report.xml'),
                      ('static', 'src', 'xml', 'hiring_numbers.xml')):
            src = re.sub(r'<!--.*?-->', '', _src(*parts), flags=re.S)
            self.assertNotIn('Odoo', src,
                             '%s shows the word Odoo to a user' % parts[-1])
            self.assertNotIn('odoo.com', src.lower())

    def test_no_bracketed_plurals(self):
        """"1 document(s)" is how a screen announces a programme wrote it
        rather than a person (R46)."""
        for parts in (('models', 'offer.py'), ('models', 'bgv.py'),
                      ('models', 'docreq.py'), ('models', 'cover.py'),
                      ('models', 'analytics.py'),
                      ('models', 'offer_closure.py'),
                      ('models', 'requisition_a3.py'),
                      ('views', 'offer_views.xml'),
                      ('views', 'offer_token_templates.xml'),
                      ('static', 'src', 'xml', 'hiring_numbers.xml')):
            src = _src(*parts)
            self.assertFalse(re.search(r'\w\(s\)', src),
                             '%s has a bracketed plural' % parts[-1])

    def test_every_hand_built_window_action_carries_views(self):
        """R125 — `_preprocessAction` maps `action.views` unconditionally, and
        a dict returned from a facade has `view_mode` and no `views`. The
        client then throws and the theme shows the generic "something went
        wrong" dialog with no useful console line."""
        import ast
        offenders = []
        for parts in (('models', 'offer.py'), ('models', 'bgv.py'),
                      ('models', 'docreq.py'),
                      ('models', 'offer_closure.py'),
                      ('models', 'requisition_a3.py'),
                      ('models', 'vendor_ext.py'),
                      ('models', 'pb_hiring_a3.py')):
            tree = ast.parse(_src(*parts))
            for node in ast.walk(tree):
                if not isinstance(node, ast.Dict):
                    continue
                keys = [k.value for k in node.keys
                        if isinstance(k, ast.Constant)]
                if 'type' not in keys or 'view_mode' not in keys:
                    continue
                kind = node.values[keys.index('type')]
                if not isinstance(kind, ast.Constant) \
                        or kind.value != 'ir.actions.act_window':
                    continue
                if 'views' not in keys:
                    offenders.append('%s line %s' % (parts[-1], node.lineno))
        self.assertFalse(offenders, 'act_window without views: %s' % offenders)

    def test_every_nolabel_field_in_a_group_carries_a_colspan(self):
        """R128 — an inner group is a two-column grid and a label-less field
        takes the NARROW cell, so a reason box renders about 150px wide with
        a thousand pixels of empty row beside it."""
        from lxml import etree
        offenders = []
        for parts in (('views', 'offer_views.xml'),
                      ('views', 'vendor_views.xml')):
            tree = etree.fromstring(_src(*parts).encode('utf-8'))
            for field in tree.iter('field'):
                if field.get('nolabel') != '1':
                    continue
                parent = field.getparent()
                if parent is None or parent.tag != 'group':
                    continue
                if not field.get('colspan'):
                    offenders.append('%s line %s'
                                     % (parts[-1], field.sourceline))
        self.assertFalse(offenders, 'nolabel without colspan: %s' % offenders)

    def test_no_search_group_carries_a_string_or_expand(self):
        """R129 — either one fails RNG validation and aborts the WHOLE module
        load, with the real error only in the server log."""
        from lxml import etree
        tree = etree.fromstring(_src('views', 'offer_views.xml')
                                .encode('utf-8'))
        for search in tree.iter('search'):
            for group in search.iter('group'):
                self.assertIsNone(group.get('string'))
                self.assertIsNone(group.get('expand'))

    def test_every_paperwork_leg_runs_in_a_savepoint(self):
        """R131 — a try/except is NOT enough when the thing that failed
        reached the database: Postgres aborts the whole transaction and every
        statement after it fails too, including the write the paperwork was
        about."""
        for name in ('offer.py', 'offer_closure.py', 'docreq.py',
                     'requisition_a3.py', 'hiring_automation.py'):
            src = _src('models', name)
            self.assertIn('leg(', src,
                          '%s does paperwork without a savepoint' % name)

    def test_the_token_is_in_no_view_and_no_payload(self):
        """R13 — a token is protected by the access list, the record rule and
        by never appearing in a view or a payload."""
        for name in ('offer_views.xml', 'offer_token_templates.xml'):
            # COMMENTS ARE STRIPPED FIRST (R118): the rule binds what SHIPS,
            # and the sentence that stops the next contributor putting a
            # token in a view is worth more than a gate that forbids it.
            src = re.sub(r'<!--.*?-->', '', _src('views', name), flags=re.S)
            self.assertNotIn('name="token"', src, name)
        facade = _src('models', 'pb_hiring_a3.py')
        row = facade[facade.index('def _offer_row'):
                     facade.index('#  The verbs')]
        self.assertNotIn("'token'", row)

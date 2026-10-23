# -*- coding: utf-8 -*-
"""Morning "new demo sign-ups" report for the Payobook team.

An hourly scheduled action asks: is it between 5 am and noon in Sydney, and
has today's report not run yet? If so it takes the demo-signup leads created
since the last report (cursor = the last ``crm.lead`` id reported, kept in
``pb_demo_portal.signup_report_last_lead_id``), keeps the people whose email
was never reported before, and mails them as an Excel workbook to the same
list as the instant sign-up email (``pb_demo_portal.lead_notify_to``).
No new people -> no mail; the next report still starts from the cursor.

Hourly + a Sydney-clock check (instead of a daily nextcall) keeps 5 am right
across daylight-saving changes and catches up if the server was down at 5.
"""
import base64
import io
import logging
import re
from collections import OrderedDict
from datetime import datetime

import pytz

from odoo import api, models, tools

_logger = logging.getLogger(__name__)

EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')
# Owner ruling 2026-10-06: ash + hello + lehoa. Overridable without a deploy via
# the ``pb_demo_portal.lead_notify_to`` system parameter (comma-separated).
DEFAULT_LEAD_NOTIFY_TO = 'ash@biztinct.com,hello@payobook.com,lehoa@payobook.com'

PARAM_CURSOR = 'pb_demo_portal.signup_report_last_lead_id'
PARAM_LAST_DAY = 'pb_demo_portal.signup_report_last_day'
REPORT_TZ = 'Australia/Sydney'
SEND_FROM_HOUR = 5
SEND_UNTIL_HOUR = 12   # a later catch-up waits for tomorrow; the cursor keeps the delta
SHEET_TZ = 'Asia/Ho_Chi_Minh'   # the sign-ups are Vietnamese; same clock as the first report
SOURCE_NAME = 'Demo Signup'
SIZE_ORDER = ['1–50', '51–200', '201–500', '501–1,000', '1,001–5,000', '5,000+']
BODY_ROWS = 25


class PbDemoSignupReport(models.AbstractModel):
    _name = 'pb.demo.signup.report'
    _description = 'Demo sign-up morning report'

    # ------------------------------------------------------------------ #
    #  Recipients / sender — shared with the instant sign-up email       #
    # ------------------------------------------------------------------ #
    @api.model
    def _notify_company(self):
        return (self.env['res.company'].sudo().with_context(active_test=False).search(
            [('name', '=', 'Payobook Vietnam JSC')], limit=1) or self.env.company)

    @api.model
    def _notify_recipients(self, company=None):
        """Comma-separated recipient list -> ordered, de-duplicated, valid emails.
        Falls back to the notify company's email when the parameter is blank."""
        raw = self.env['ir.config_parameter'].sudo().get_param(
            'pb_demo_portal.lead_notify_to', DEFAULT_LEAD_NOTIFY_TO) or ''
        seen, out = set(), []
        for part in re.split(r'[,;\s]+', raw):
            addr = part.strip().lower()
            if addr and EMAIL_RE.match(addr) and addr not in seen:
                seen.add(addr)
                out.append(addr)
        if not out:
            company = company or self._notify_company()
            fallback = company.email or company.partner_id.email_normalized
            if fallback:
                out.append(fallback)
        return out

    @api.model
    def _notify_from(self, company=None):
        """The system's configured default sender, else the notify company's
        email, else a no-reply on the product domain."""
        default_from = self.env['ir.mail_server'].sudo()._get_default_from_address()
        if default_from:
            return default_from
        company = company or self._notify_company()
        return company.email or company.partner_id.email_normalized or 'noreply@payobook.com'

    # ------------------------------------------------------------------ #
    #  Scheduled entry point                                             #
    # ------------------------------------------------------------------ #
    @api.model
    def _cron_send_daily_report(self):
        ICP = self.env['ir.config_parameter'].sudo()
        now_local = datetime.now(pytz.timezone(REPORT_TZ))
        today = now_local.date().isoformat()
        if not (SEND_FROM_HOUR <= now_local.hour < SEND_UNTIL_HOUR):
            return False
        if ICP.get_param(PARAM_LAST_DAY) == today:
            return False
        result = self.send_report()
        if result != 'failed':   # a failed send retries on the next hourly tick
            ICP.set_param(PARAM_LAST_DAY, today)
        return result

    @api.model
    def send_report(self):
        """Mail the people who signed up since the last report.
        Returns 'sent', 'nothing' or 'failed'."""
        ICP = self.env['ir.config_parameter'].sudo()
        cursor = int(ICP.get_param(PARAM_CURSOR) or 0)
        source = self.env['utm.source'].sudo().search([('name', '=', SOURCE_NAME)], limit=1)
        if not source:
            return 'nothing'
        Lead = self.env['crm.lead'].sudo().with_context(active_test=False)
        new_leads = Lead.search([('source_id', '=', source.id), ('id', '>', cursor)], order='id')
        if not new_leads:
            return 'nothing'
        top = max(new_leads.ids)
        reported = {(r['email_from'] or '').strip().lower() for r in Lead.search_read(
            [('source_id', '=', source.id), ('id', '<=', cursor)], ['email_from'])}
        people = OrderedDict()
        for lead in new_leads:
            email = (lead.email_from or '').strip().lower()
            if email in reported:
                continue
            people.setdefault(email, []).append(lead)
        if not people:   # only repeat sends from people already reported
            ICP.set_param(PARAM_CURSOR, str(top))
            return 'nothing'

        since = Lead.browse(cursor).exists().create_date if cursor else False
        rows = self._person_rows(people)
        now_sheet = datetime.now(pytz.timezone(SHEET_TZ)).replace(tzinfo=None)
        xlsx = self._build_workbook(rows, people, since, now_sheet)
        stamp = datetime.now(pytz.timezone(REPORT_TZ))
        n = len(rows)
        company = self._notify_company()
        mail = self.env['mail.mail'].sudo().create({
            'subject': u"Payobook demo sign-ups: %d new (%s)" % (n, stamp.strftime('%d %b %Y')),
            'email_from': self._notify_from(company),
            'email_to': ', '.join(self._notify_recipients(company)),
            'body_html': self._mail_body(rows, since),
            'auto_delete': True,
            'attachment_ids': [(0, 0, {
                'name': 'Payobook_new_signups_%s.xlsx' % stamp.strftime('%Y-%m-%d'),
                'datas': base64.b64encode(xlsx),
                'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            })],
        })
        mail.send()
        if mail.exists() and mail.state == 'exception':
            _logger.warning("Demo sign-up report not sent: %s", mail.failure_reason)
            return 'failed'
        ICP.set_param(PARAM_CURSOR, str(top))
        _logger.info("Demo sign-up report sent: %d new people (leads %d..%d)", n, cursor + 1, top)
        return 'sent'

    # ------------------------------------------------------------------ #
    #  Data                                                              #
    # ------------------------------------------------------------------ #
    @staticmethod
    def _note_value(desc, key):
        m = re.search(key + r':\s*([^\n]*)', tools.html2plaintext(desc or ''))
        return m.group(1).strip() if m else ''

    @staticmethod
    def _sheet_time(dt):
        return pytz.utc.localize(dt).astimezone(pytz.timezone(SHEET_TZ)).replace(tzinfo=None)

    def _person_rows(self, people):
        self.env.cr.execute("""
            SELECT lower(login), count(*), bool_or(password IS NOT NULL)
              FROM res_users WHERE lower(login) = ANY(%s) GROUP BY lower(login)
        """, [list(people)])
        accounts = {r[0]: (r[1], r[2]) for r in self.env.cr.fetchall()}
        rows = []
        for email, leads in people.items():
            first = leads[0]
            n_acc, has_pw = accounts.get(email, (0, False))
            rows.append({
                'when': self._sheet_time(first.create_date),
                'name': first.contact_name or '',
                'company': first.partner_name or '',
                'email': email,
                'phone': first.phone or '',
                'country': first.country_id.name or '',
                'industry': self._note_value(first.description, 'Industry'),
                'size': self._note_value(first.description, 'Company size'),
                'sent': len(leads),
                'accounts': n_acc,
                'password': 'Yes' if has_pw else 'No',
                'stage': first.stage_id.name or '',
                'lead_nos': ', '.join(str(l.id) for l in leads),
            })
        rows.sort(key=lambda r: r['when'], reverse=True)
        return rows

    # ------------------------------------------------------------------ #
    #  Excel (Summary / Leads / All submissions — same look as the       #
    #  first one-off report; values, not formulas, so mail previews      #
    #  show the numbers)                                                 #
    # ------------------------------------------------------------------ #
    def _build_workbook(self, rows, people, since, now_sheet):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo

        head_fill = PatternFill('solid', fgColor='5A4BB0')
        head_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
        body = Font(name='Arial', size=10)
        bold = Font(name='Arial', size=10, bold=True)
        title = Font(name='Arial', size=14, bold=True, color='1B1733')
        sub = Font(name='Arial', size=10, italic=True, color='6B6880')
        when_fmt = 'dd mmm yyyy hh:mm'

        def header(ws, row, cols):
            for i, c in enumerate(cols, 1):
                cell = ws.cell(row=row, column=i, value=c)
                cell.font, cell.fill = head_font, head_fill
                cell.alignment = Alignment(vertical='center', wrap_text=True)
            ws.row_dimensions[row].height = 30

        def widths(ws, ws_widths):
            for i, w in enumerate(ws_widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w

        def table(ws, name, ncols, last_row):
            t = Table(displayName=name, ref='A1:%s%d' % (get_column_letter(ncols), last_row))
            t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
            ws.add_table(t)

        wb = Workbook()

        # Summary
        ws = wb.active
        ws.title = 'Summary'
        ws['A1'] = 'New demo sign-ups since the last report'
        ws['A1'].font = title
        period = (u"Sign-ups after %s up to %s (Vietnam time)."
                  % (self._sheet_time(since).strftime('%d %b %Y %H:%M'),
                     now_sheet.strftime('%d %b %Y %H:%M'))) if since else \
            u"All sign-ups up to %s (Vietnam time)." % now_sheet.strftime('%d %b %Y %H:%M')
        ws['A2'] = period + u" Only people not in an earlier report."
        ws['A2'].font = sub
        kpis = [
            ('New people who signed up', len(rows)),
            ('Times the form was sent', sum(r['sent'] for r in rows)),
            ('People who sent the form more than once', sum(1 for r in rows if r['sent'] > 1)),
            ('Demo accounts created', sum(r['accounts'] for r in rows)),
            ('People who have set a password', sum(1 for r in rows if r['password'] == 'Yes')),
            ('People who have not set a password yet', sum(1 for r in rows if r['password'] == 'No')),
        ]
        row = 4
        for label, val in kpis:
            ws.cell(row=row, column=1, value=label).font = body
            c = ws.cell(row=row, column=2, value=val)
            c.font, c.number_format = bold, '#,##0'
            row += 1
        row += 1
        total = len(rows) or 1
        for heading, key in (('Industry', 'industry'), ('Company size', 'size'), ('Country', 'country')):
            counts = {}
            for r in rows:
                counts[r[key] or '—'] = counts.get(r[key] or '—', 0) + 1
            if key == 'size':
                order = [s for s in SIZE_ORDER if s in counts] + \
                        [s for s in counts if s not in SIZE_ORDER]
            else:
                order = sorted(counts, key=lambda k: (-counts[k], k))
            for col, text in enumerate((heading, 'People', 'Share'), 1):
                cell = ws.cell(row=row, column=col, value=text)
                cell.font, cell.fill = head_font, head_fill
            row += 1
            for k in order:
                ws.cell(row=row, column=1, value=k).font = body
                ws.cell(row=row, column=2, value=counts[k]).font = body
                c = ws.cell(row=row, column=3, value=counts[k] / total)
                c.font, c.number_format = body, '0.0%'
                row += 1
            row += 1
        widths(ws, [44, 12, 10])

        # Leads — one row per new person
        ws = wb.create_sheet('Leads')
        cols = ['First signed up (Vietnam time)', 'Full name', 'Company', 'Business email',
                'Mobile number', 'Country', 'Industry', 'Company size',
                'Times form was sent', 'Demo accounts created', 'Has set a password',
                'Lead stage', 'Lead numbers']
        header(ws, 1, cols)
        keys = ['when', 'name', 'company', 'email', 'phone', 'country', 'industry', 'size',
                'sent', 'accounts', 'password', 'stage', 'lead_nos']
        for i, r in enumerate(rows, 2):
            for c, k in enumerate(keys, 1):
                ws.cell(row=i, column=c, value=r[k]).font = body
            ws.cell(row=i, column=1).number_format = when_fmt
            ws.cell(row=i, column=5).number_format = '@'
        table(ws, 'Leads', len(cols), len(rows) + 1)
        widths(ws, [20, 26, 30, 34, 16, 14, 16, 13, 11, 11, 11, 10, 22])
        ws.freeze_panes = 'C2'

        # All submissions of those people
        ws = wb.create_sheet('All submissions')
        cols = ['Lead number', 'Sent (Vietnam time)', 'Full name', 'Company', 'Business email',
                'Mobile number', 'Country', 'Industry', 'Company size', 'Lead stage']
        header(ws, 1, cols)
        subs = sorted((l for leads in people.values() for l in leads), key=lambda l: l.id)
        for i, l in enumerate(subs, 2):
            vals = [l.id, self._sheet_time(l.create_date), l.contact_name or '',
                    l.partner_name or '', (l.email_from or '').strip().lower(), l.phone or '',
                    l.country_id.name or '', self._note_value(l.description, 'Industry'),
                    self._note_value(l.description, 'Company size'), l.stage_id.name or '']
            for c, v in enumerate(vals, 1):
                ws.cell(row=i, column=c, value=v).font = body
            ws.cell(row=i, column=2).number_format = when_fmt
            ws.cell(row=i, column=6).number_format = '@'
        table(ws, 'Submissions', len(cols), len(subs) + 1)
        widths(ws, [10, 20, 26, 30, 34, 16, 14, 16, 13, 10])
        ws.freeze_panes = 'D2'

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------ #
    #  Email body                                                        #
    # ------------------------------------------------------------------ #
    def _mail_body(self, rows, since):
        esc = tools.html_escape
        n = len(rows)
        lead_in = (u"since the last report (sign-ups after %s, Vietnam time)"
                   % self._sheet_time(since).strftime('%d %b %Y %H:%M')) if since else u"so far"
        cell = u"padding:6px 12px;font-size:13px;border-bottom:1px solid #EEF0F8;"
        head = u"padding:6px 12px;font-size:12px;color:#8A87A0;text-align:left;"
        trs = u''.join(
            u"<tr><td style='%s font-weight:600;'>%s</td><td style='%s'>%s</td>"
            u"<td style='%s'>%s</td><td style='%s'>%s</td><td style='%s'>%s</td></tr>"
            % (cell, esc(r['name']), cell, esc(r['company']), cell, esc(r['industry']),
               cell, esc(r['size']), cell, esc(r['country']))
            for r in rows[:BODY_ROWS])
        more = (u"<p style='margin:10px 0 0;font-size:13px;color:#4B5168;'>…and %d more in the "
                u"attached file.</p>" % (n - BODY_ROWS)) if n > BODY_ROWS else u''
        return (
            u"<div style='font-family:Segoe UI,Roboto,Arial,sans-serif;color:#1B1733;'>"
            u"<h2 style='margin:0 0 4px;font-size:20px;'>%d new demo sign-up%s</h2>"
            u"<p style='margin:0 0 16px;color:#8A87A0;font-size:13px;'>New people %s. "
            u"The full details are in the attached Excel file.</p>"
            u"<table style='border-collapse:collapse;'><tr><th style='%s'>Name</th>"
            u"<th style='%s'>Company</th><th style='%s'>Industry</th>"
            u"<th style='%s'>Company size</th><th style='%s'>Country</th></tr>%s</table>%s</div>"
            % (n, u'' if n == 1 else u's', lead_in, head, head, head, head, head, trs, more))

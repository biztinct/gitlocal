#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build ``tests/fixtures/small_payroll.xlsx`` — the workbook the round trip uses.

B1 could not drive the Excel starting point to the end because every payroll
workbook in this repository is a Vietnamese template whose headers the review
cannot read as a key column, and the review requires a Primary Key Column that
exists in every selected worksheet. This builds the smallest workbook that the
product's own reader accepts.

**It has to be colour-coded**, because `hr.formula.config.use_color_coded_excel_import`
defaults to True and every configuration the guided setup creates therefore
takes that path (`multisheet_import_wizard.py:2054`). The reader's rules, in the
order it applies them:

* the HEADER BLOCK is every row carrying a yellow or amber fill; a workbook with
  none is refused outright with "Could not detect header rows";
* inside it, a YELLOW cell is a component's name and an AMBER cell above it is
  that component's type;
* the row above the block is the payslip-identifier row;
* the FORMULA ROW is the first GREEN-filled row after the block, and it is also
  the first data row: a column whose cell there begins with ``=`` is a
  calculation, and everything else is a number the payroll is given.

So the layout is four fixed rows and then people:

    row 1   payslip identifiers (left empty here)
    row 2   amber   component types
    row 3   yellow  component names  ← the Primary Key Column lives here
    row 4   green   the first person, and the formulas
    row 5+  the rest of the people

Run it from the repository root:

    .venv/bin/python pb_blueprint/tools/gen_fixture_workbook.py

The workbook is committed, so this only has to be run when the fixture changes.
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

#: The three fills the reader looks for. RGB chosen to sit well inside its own
#: thresholds (`_is_yellow_fill` / `_is_amber_fill` / `_is_green_fill`).
YELLOW = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
AMBER = PatternFill(start_color="FFFFC000", end_color="FFFFC000", fill_type="solid")
GREEN = PatternFill(start_color="FF92D050", end_color="FF92D050", fill_type="solid")

#: One tuple per column: the type (amber row), the name (yellow row), and how
#: the first data row is filled — a number, a piece of text, or a formula.
COLUMNS = [
    ("Input", "Employee ID", "EMP001"),          # A — the primary key
    ("Input", "Employee Name", "Nguyen Van An"),  # B
    ("Earning", "Basic Salary", 20000000),        # C
    ("Earning", "Allowance", 1500000),            # D
    ("Input", "Days Worked", 26),                 # E
    ("Gross", "Gross Pay", "=C{row}+D{row}"),     # F — a calculation
    ("Deduction", "Insurance", "=ROUND(C{row}*0.105,0)"),   # G
    ("Net", "Net Pay", "=F{row}-G{row}"),         # H
]

#: The other two people. Same columns, same formulas.
PEOPLE = [
    ("EMP002", "Tran Thi Binh", 30000000, 2000000, 22),
    ("EMP003", "Le Minh Chau", 45000000, 0, 26),
]

TYPE_ROW = 2
HEADER_ROW = 3
FIRST_DATA_ROW = 4

OUT = Path(__file__).resolve().parents[1] / 'tests' / 'fixtures' / 'small_payroll.xlsx'


def build(path=OUT):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Payroll"

    for col, (kind, name, first) in enumerate(COLUMNS, start=1):
        type_cell = sheet.cell(row=TYPE_ROW, column=col, value=kind)
        type_cell.fill = AMBER
        head_cell = sheet.cell(row=HEADER_ROW, column=col, value=name)
        head_cell.fill = YELLOW
        value = first.format(row=FIRST_DATA_ROW) if isinstance(first, str) else first
        cell = sheet.cell(row=FIRST_DATA_ROW, column=col, value=value)
        cell.fill = GREEN                      # this row IS the formula row

    for i, (code, name, basic, allowance, days) in enumerate(PEOPLE):
        row = FIRST_DATA_ROW + 1 + i
        sheet.cell(row=row, column=1, value=code)
        sheet.cell(row=row, column=2, value=name)
        sheet.cell(row=row, column=3, value=basic)
        sheet.cell(row=row, column=4, value=allowance)
        sheet.cell(row=row, column=5, value=days)
        sheet.cell(row=row, column=6, value="=C%d+D%d" % (row, row))
        sheet.cell(row=row, column=7, value="=ROUND(C%d*0.105,0)" % row)
        sheet.cell(row=row, column=8, value="=F%d-G%d" % (row, row))

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


if __name__ == '__main__':
    print(build())

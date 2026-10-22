# -*- coding: utf-8 -*-
"""LEARNING L6 — Numbers, finished.

The Numbers lens gains five blocks and a flat sheet, each read through
`_safe` with a default and each with an empty sentence (handover §3.1):

  * HOURS — finished training hours per person, department, course and kind
    (online · in a room or on site · toolbox talk), over the range, by the
    day the course was FINISHED or the session was HELD. A course finished by
    attending a session is counted once, through the attendance (its hours
    are the same hours); a blended course adds only what its lessons count
    for above the room time. So a person's hours are their courses plus their
    sessions and never both for the same thing (T1).
  * COST — claims agreed (and of those, paid), outside courses agreed and not
    claimed yet (the money set aside), and what sessions cost, against the
    yearly allowances and the HR-operations budget line when `pb_budget` is
    installed (read through `env.get`, never written). ONE CURRENCY PER TOTAL:
    when two currencies meet, no total is drawn — one line per currency.
  * COMPLIANCE — the Required matrix's own summary for the same department,
    site and country, with splits by site, department, country and safety
    group (`required_reports._required_numbers`, so the matrix headline and
    this block are one number). `None` when the company has no live rule.
  * EFFECT — "Did it work?": course ratings, % who would recommend, and the
    managers' follow-up answers, from L5's own roll-up.
  * COUNTRY REPORTS — the four packs with on/off, for the buttons.
  * the ASSIGNMENTS sheet — one row per assignment in the range with every
    dimension resolved (the BI door; no connector).

Filters gain site (work location), country and kind. Money is never written
here: nothing in Learning writes a payslip or a budget (ledger L6).
"""

import base64
import io
import json
import logging
from datetime import datetime

from odoo import _, api, fields, models
from odoo.exceptions import UserError

from .session import SESSION_DELIVERY
from .training_common import as_id, fold, money_words, tr

_logger = logging.getLogger(__name__)

#: The kinds of training hours, in screen order (sign-off waits for a later
#: phase, ledger LD5 — kept as a zero so the shape is the handover's).
HOURS_KINDS = [
    ('online', 'Online'),
    ('session', 'In a room or on site'),
    ('toolbox', 'Toolbox talk'),
]
HOURS_KIND_LABEL = dict(HOURS_KINDS)

#: People listed on the screen's Hours block; the spreadsheet has everybody.
HOURS_PEOPLE_SHOWN = 50

#: The context key that carries a Numbers payload already built, so the
#: spreadsheet and the pack never work the whole page out twice.
CACHE_KEY = 'pb_numbers_cache'
#: The context key that asks for every person on the Hours block (the
#: spreadsheet and the pack); the screen gets the first rows only.
FULL_KEY = 'pb_numbers_full'


def one_currency(sums):
    """The one currency these amounts are in, or a UserError naming the
    rule: a total over two currencies is a number nobody can read (ledger
    "no total ever adds two currencies"). `sums` maps currency → amount;
    zero amounts do not count."""
    used = [c for c, v in sums.items() if c and v]
    if len(used) > 1:
        raise UserError(_(
            "These amounts are in more than one currency, so they are not "
            "added up."))
    return used[0] if used else None


def _headline_hours(value):
    """Hours as a sentence says them: one decimal, or whole hours from 10
    up, rounded half up ("22.05" read like a ledger — live walk, L6)."""
    from decimal import ROUND_HALF_UP, Decimal
    value = Decimal(str(float(value or 0.0)))
    step = Decimal('1') if value >= 10 else Decimal('0.1')
    value = value.quantize(step, rounding=ROUND_HALF_UP)
    return int(value) if value == value.to_integral_value() else float(value)


def _hours_num(value):
    value = round(float(value or 0.0), 2)
    return int(value) if value == int(value) else value


class PbTrainingAnalyticsL6(models.AbstractModel):
    _inherit = 'pb.training.analytics'

    # =====================================================================
    #  filters: site, country, kind
    # =====================================================================
    @api.model
    def _employee_domain(self, filters, prefix='employee_id.'):
        """Department, site and country, on whatever field leads to the
        employee. Country is the work address's, else the company's — the
        Required matrix's own rule (`_person_facts`)."""
        filters = filters or {}
        domain = []
        if filters.get('department_id'):
            domain.append((prefix + 'department_id', 'child_of',
                           as_id(filters['department_id'])))
        if filters.get('work_location_id'):
            domain.append((prefix + 'work_location_id', '=',
                           as_id(filters['work_location_id'])))
        code = (filters.get('country_code') or '').upper()
        if code:
            domain += ['|', (prefix + 'address_id.country_id.code', '=', code),
                       '&', (prefix + 'address_id.country_id', '=', False),
                       (prefix + 'company_id.partner_id.country_id.code', '=',
                        code)]
        return domain

    @api.model
    def _kind_domain(self, kind):
        """An assignment's kind: a toolbox talk's course, a course taught in
        a room, or an online course."""
        if kind == 'toolbox':
            return [('session_id.kind', '=', 'toolbox')]
        if kind == 'session':
            return [('channel_id.delivery', 'in', SESSION_DELIVERY),
                    '|', ('session_id', '=', False),
                    ('session_id.kind', '!=', 'toolbox')]
        if kind == 'online':
            return [('channel_id.delivery', 'not in', SESSION_DELIVERY),
                    ('session_id', '=', False)]
        return []

    def _domain(self, start, end, filters=None):
        domain = super()._domain(start, end, filters)
        filters = filters or {}
        extra = {k: filters.get(k) for k in ('work_location_id',
                                             'country_code')}
        domain += self._employee_domain(extra)
        if filters.get('kind') in HOURS_KIND_LABEL:
            domain += self._kind_domain(filters['kind'])
        return domain

    @api.model
    def _assignment_kind(self, row):
        session = row.session_id.sudo() if 'session_id' in row._fields \
            else False
        if session and session.kind == 'toolbox':
            return 'toolbox'
        if row.channel_id.sudo().delivery in SESSION_DELIVERY:
            return 'session'
        return 'online'

    def _lists(self):
        out = super()._lists()
        rows = self.env['pb.training.assignment'].sudo().search(
            [('company_id', 'in', self.env.companies.ids)])
        employees = rows.mapped('employee_id')
        facts = self.env['pb.training']._person_facts(employees) \
            if employees else {}
        sites, countries = {}, {}
        for f in facts.values():
            if f.get('site_id'):
                sites[f['site_id']] = f.get('site') or ''
            if f.get('country_code'):
                countries[f['country_code']] = f.get('country') or \
                    f['country_code']
        kinds = {self._assignment_kind(r) for r in rows}
        out['sites'] = sorted([{'id': k, 'name': v} for k, v in sites.items()],
                              key=lambda r: fold(r['name']))
        out['countries'] = sorted(
            [{'code': k, 'name': v} for k, v in countries.items()],
            key=lambda r: fold(r['name']))
        out['kinds'] = [{'key': k, 'label': tr(self.env, v)}
                        for k, v in HOURS_KINDS if k in kinds]
        return out

    # =====================================================================
    #  HOURS
    # =====================================================================
    @api.model
    def _hours_rows(self, start, end, filters=None):
        """(done assignments finished in the range, attended attendances of
        closed sessions held in the range), both filtered."""
        filters = filters or {}
        companies = self.env.companies.ids
        adom = [('company_id', 'in', companies), ('state', '=', 'done'),
                ('completed_on', '>=', start), ('completed_on', '<=', end)]
        adom += self._employee_domain(filters)
        if filters.get('channel_id'):
            adom.append(('channel_id', '=', as_id(filters['channel_id'])))
        if filters.get('reason'):
            adom.append(('reason', '=', filters['reason']))
        if filters.get('path_id'):
            adom.append(('path_enrolment_id.path_id', '=',
                         as_id(filters['path_id'])))
        done = self.env['pb.training.assignment'].sudo().search(adom)
        Att = self.env['pb.training.attendance'].sudo()
        tdom = [('company_id', 'in', companies), ('state', '=', 'attended'),
                ('session_id.state', '=', 'closed'),
                ('session_id.start', '>=', datetime.combine(
                    start, datetime.min.time())),
                ('session_id.start', '<=', datetime.combine(
                    end, datetime.max.time()))]
        tdom += self._employee_domain(filters)
        if filters.get('channel_id'):
            tdom.append(('channel_id', '=', as_id(filters['channel_id'])))
        if filters.get('reason'):
            tdom.append(('assignment_id.reason', '=', filters['reason']))
        if filters.get('path_id'):
            tdom.append(('assignment_id.path_enrolment_id.path_id', '=',
                         as_id(filters['path_id'])))
        atts = Att.search(tdom)
        kind = filters.get('kind')
        if kind in HOURS_KIND_LABEL:
            atts = atts.filtered(lambda a: self._attendance_kind(a) == kind)
        return done, atts

    @api.model
    def _attendance_kind(self, att):
        return 'toolbox' if att.session_id.kind == 'toolbox' else 'session'

    @api.model
    def _hours_numbers(self, start, end, filters=None):
        filters = filters or {}
        done, atts = self._hours_rows(start, end, filters)
        kind_filter = filters.get('kind')
        people, depts, courses = {}, {}, {}
        by_kind = {k: 0.0 for k, _l in HOURS_KINDS}
        by_kind['signoff'] = 0.0

        def add(emp, course, kind, hours, is_session):
            if hours <= 0 and not is_session:
                return
            p = people.setdefault(emp.id, {
                'employee_id': emp.id, 'employee': emp.name or '',
                'department': emp.department_id.name or '',
                'hours': 0.0, 'courses': 0, 'sessions': 0})
            p['hours'] += hours
            if is_session:
                p['sessions'] += 1
            else:
                p['courses'] += 1
            dname = emp.department_id.name or _('No department')
            depts[dname] = depts.get(dname, 0.0) + hours
            cname = course.name or _('A course that is gone')
            courses[cname] = courses.get(cname, 0.0) + hours
            by_kind[kind] = by_kind.get(kind, 0.0) + hours

        for row in done:
            emp = row.employee_id.sudo()
            linked = sum(a.hours or 0.0 for a in row.attendance_ids.sudo()
                         if a.state == 'attended') \
                if 'attendance_ids' in row._fields else 0.0
            part = max((row.hours or 0.0) - linked, 0.0)
            if linked and not part:
                continue            # counted once, through the attendance
            channel = row.channel_id.sudo()
            kind = self._assignment_kind(row)
            if linked or channel.delivery == 'blended':
                kind = 'online'     # the lessons' part of a blended course
            if kind_filter in HOURS_KIND_LABEL and kind != kind_filter:
                continue
            add(emp, channel, kind, part, False)
        for att in atts:
            add(att.employee_id.sudo(), att.session_id.channel_id.sudo(),
                self._attendance_kind(att), att.hours or 0.0, True)

        rows = sorted(people.values(),
                      key=lambda r: (-r['hours'], fold(r['employee'])))
        for r in rows:
            r['hours'] = _hours_num(r['hours'])
        total = sum(r['hours'] for r in rows)

        def table(src):
            out = [{'name': k, 'hours': _hours_num(v)}
                   for k, v in src.items() if v]
            out.sort(key=lambda r: (-r['hours'], fold(r['name'])))
            return out

        kinds = [{'key': k, 'label': tr(self.env, v),
                  'hours': _hours_num(by_kind.get(k, 0.0))}
                 for k, v in HOURS_KINDS]
        start_s, end_s = fields.Date.to_string(start), fields.Date.to_string(end)
        if not rows:
            headline = _("No training hours were finished between %(from)s "
                         "and %(to)s.", **{'from': start_s, 'to': end_s})
        elif len(rows) == 1:
            headline = _("%(hours)s hours of training, by 1 person.",
                         hours=_headline_hours(total))
        elif _headline_hours(total / len(rows)) == 1:
            # "1 hours each" read wrong (found in the live seed read-out).
            headline = _("%(hours)s hours of training, by %(n)s people · "
                         "1 hour each on average.",
                         hours=_headline_hours(total), n=len(rows))
        else:
            headline = _("%(hours)s hours of training, by %(n)s people · "
                         "%(avg)s hours each on average.",
                         hours=_headline_hours(total), n=len(rows),
                         avg=_headline_hours(total / len(rows)))
        return {
            'total': _hours_num(total),
            'people': len(rows),
            'empty': not rows,
            'headline': headline,
            'by_person': rows,
            'by_person_shown': rows[:HOURS_PEOPLE_SHOWN],
            'by_department': table(depts),
            'by_course': table(courses),
            'by_kind': {k: _hours_num(v) for k, v in by_kind.items()},
            'kinds': kinds,
        }

    # =====================================================================
    #  COST
    # =====================================================================
    @api.model
    def _cost_numbers(self, start, end, filters=None):
        filters = filters or {}
        env = self.env
        companies = env.companies.ids
        lo = datetime.combine(start, datetime.min.time())
        hi = datetime.combine(end, datetime.max.time())
        emp_dom = self._employee_domain(filters)
        channel = as_id(filters.get('channel_id'))
        # ---- claims, by the day the person paid ---------------------------
        cdom = [('company_id', 'in', companies), ('state', '=', 'approved'),
                ('paid_on', '>=', start), ('paid_on', '<=', end)] + emp_dom
        if channel:
            cdom.append(('assignment_id.channel_id', '=', channel))
        claims = env['pb.training.claim'].sudo().search(cdom)
        paid = claims.filtered(lambda c: c.fulfilment == 'paid')
        # ---- outside courses agreed and not claimed yet ------------------
        # By the day it was agreed (the day it was asked, for a row that
        # never went through a decision).
        rdom = [('company_id', 'in', companies),
                ('state', 'in', ('approved', 'taken')), ('amount', '>', 0),
                '|', '&', ('decided_on', '>=', lo), ('decided_on', '<=', hi),
                '&', ('decided_on', '=', False),
                '&', ('create_date', '>=', lo), ('create_date', '<=', hi),
                ] + emp_dom
        if channel:
            rdom.append(('channel_id', '=', channel))
        requests = env['pb.training.request'].sudo().search(rdom)
        # ---- sessions held (not cancelled) -------------------------------
        sdom = [('company_id', 'in', companies), ('state', '!=', 'cancelled'),
                ('cost', '>', 0), ('start', '>=', lo), ('start', '<=', hi)]
        if channel:
            sdom.append(('channel_id', '=', channel))
        sessions = env['pb.training.session'].sudo().search(sdom)

        def sums(recs, field):
            out = {}
            for rec in recs:
                cur = rec.currency_id or env.company.currency_id
                out[cur] = out.get(cur, 0.0) + (rec[field] or 0.0)
            return out

        parts = {'claims_agreed': sums(claims, 'amount'),
                 'claims_paid': sums(paid, 'amount'),
                 'requests_approved': sums(requests, 'amount'),
                 'sessions': sums(sessions, 'cost')}
        everything = {}
        for key in ('claims_agreed', 'requests_approved', 'sessions'):
            for cur, v in parts[key].items():
                everything[cur] = everything.get(cur, 0.0) + v
        counts = {'claims_count': len(claims), 'paid_count': len(paid),
                  'requests_count': len(requests),
                  'sessions_count': len(sessions)}
        empty = not any(everything.values())
        try:
            currency = one_currency(everything) or env.company.currency_id
        except UserError:
            # Never add two currencies: one phrase per currency instead.
            lines = [money_words(env, v, c) for c, v in everything.items()
                     if v]
            return dict(counts, mixed=True, empty=False, total=None,
                        total_words=' + '.join(lines),
                        headline=_("Training cost is in more than one "
                                   "currency: %s. No total is drawn.",
                                   ' + '.join(lines)),
                        currency='', allowance_total=None, budget=None)

        def amount(key):
            return round(parts[key].get(currency, 0.0), 2)

        out = dict(counts, mixed=False, empty=empty,
                   currency=currency.name or '')
        for key in ('claims_agreed', 'claims_paid', 'requests_approved',
                    'sessions'):
            out[key] = amount(key)
            out[key + '_words'] = money_words(env, out[key], currency)
        out['total'] = round(everything.get(currency, 0.0), 2)
        out['total_words'] = money_words(env, out['total'], currency)
        out.update(self._allowance_numbers(end.year, filters, currency))
        out['budget'] = self._budget_line(end.year, filters, currency)
        if empty:
            out['headline'] = _(
                "Nothing was spent on training between %(from)s and %(to)s.",
                **{'from': fields.Date.to_string(start),
                   'to': fields.Date.to_string(end)})
        elif out.get('allowance_used_pct') is not None:
            out['headline'] = _(
                "%(total)s spent on training · %(pct)s%% of this year's "
                "allowances used.", total=out['total_words'],
                pct=out['allowance_used_pct'])
        else:
            out['headline'] = _("%s spent on training.", out['total_words'])
        return out

    @api.model
    def _allowance_numbers(self, year, filters, currency):
        """Every person's allowance for the year added up (their own figure,
        else the company's), and how much of it is used: claims agreed and
        outside courses agreed for that year. In one currency or not at
        all."""
        env = self.env
        blank = {'allowance_total': None, 'allowance_used': None,
                 'allowance_used_pct': None, 'allowance_year': year}
        emps = env['hr.employee'].sudo().search(
            [('company_id', 'in', env.companies.ids), ('active', '=', True)]
            + self._employee_domain(filters, prefix=''))
        if not emps:
            return blank
        Allowance = env['pb.training.allowance'].sudo()
        own = {r.employee_id.id: r for r in Allowance.search(
            [('employee_id', 'in', emps.ids), ('year', '=', year)])}
        company_rows = {r.company_id.id: r for r in Allowance.search(
            [('employee_id', '=', False), ('year', '=', year),
             ('company_id', 'in', env.companies.ids)])}
        total = 0.0
        for emp in emps:
            row = own.get(emp.id) or company_rows.get(emp.company_id.id)
            if not row or not row.amount:
                continue
            if row.currency_id and row.currency_id != currency:
                return blank        # two currencies: no percentage
            total += row.amount
        if not total:
            return blank
        used = 0.0
        for model, states in (('pb.training.claim', ('approved',)),
                              ('pb.training.request', ('approved', 'taken'))):
            for rec in env[model].sudo().search(
                    [('employee_id', 'in', emps.ids), ('year', '=', year),
                     ('state', 'in', states)]):
                if rec.currency_id and rec.currency_id != currency:
                    return blank
                used += rec.amount or 0.0
        return {
            'allowance_year': year,
            'allowance_total': round(total, 2),
            'allowance_total_words': money_words(env, total, currency),
            'allowance_used': round(used, 2),
            'allowance_used_words': money_words(env, used, currency),
            'allowance_used_pct': int(round(100.0 * used / total)),
        }

    @api.model
    def _budget_line(self, year, filters, currency):
        """The HR-operations budget for the year, when `pb_budget` is here —
        READ ONLY (binding non-goal: no pb_budget edits). Training spend
        sits under HR operations there (`pb_budget_line.py` help)."""
        Line = self.env.get('pb.budget.line')
        if Line is None or 'pb_budget_type' not in Line._fields:
            return None
        domain = [('company_id', 'in', self.env.companies.ids),
                  ('pb_budget_type', '=', 'hr_ops'),
                  ('period_month', '>=', '%s-01-01' % year),
                  ('period_month', '<=', '%s-12-31' % year)]
        if (filters or {}).get('department_id'):
            domain.append(('department_id', 'child_of',
                           as_id(filters['department_id'])))
        lines = Line.sudo().search(domain)
        if not lines:
            return None
        cur = lines.mapped('pb_currency_id')
        if len(cur) > 1 or (cur and cur != currency):
            return None             # two currencies: not added up
        budget = sum(lines.mapped('forecast_cost'))
        spent = sum(lines.mapped('actual_cost'))
        env = self.env
        return {
            'line': _("HR operations budget %s", year),
            'budget': round(budget, 2),
            'spent': round(spent, 2),
            'left': round(budget - spent, 2),
            'budget_words': money_words(env, budget, currency),
            'spent_words': money_words(env, spent, currency),
            'left_words': money_words(env, budget - spent, currency),
        }

    # =====================================================================
    #  EFFECT and the country buttons
    # =====================================================================
    @api.model
    def _effect_numbers(self, payload):
        ratings = payload.get('ratings') or {}
        work = payload.get('did_it_work') or {}
        return {
            'ratings_avg': ratings.get('avg'),
            'ratings_count': ratings.get('count') or 0,
            'recommend_pct': ratings.get('recommend'),
            'followup_effect_avg': work.get('effect'),
            'followup_used_pct': work.get('used'),
            'followups_answered': work.get('answered') or 0,
            'followups_waiting': work.get('waiting') or 0,
            'by_course_top': (ratings.get('rows') or [])[:5],
            'empty': not (ratings.get('count') or work.get('answered')),
        }

    @api.model
    def _countries(self):
        from .required_reports import REPORT_TITLES
        from .requirement import PACKS
        Req = self.env['pb.training.requirement']
        return [{'key': p['cc'], 'label': tr(self.env, p['label']),
                 'title': tr(self.env, REPORT_TITLES[p['cc']]),
                 'on': Req._pack_is_on(p['cc'], self.env.company)}
                for p in PACKS]

    @api.model
    def _l6_tiles(self, payload):
        hours = payload.get('hours') or {}
        cost = payload.get('cost') or {}
        effect = payload.get('effect') or {}
        tiles = [
            {'key': 'hours_total', 'label': _('Hours of training'),
             'value': hours.get('total') or 0, 'icon': 'clock',
             'sub': _('finished in this period'), 'zero': not hours.get('total')},
            {'key': 'cost_total', 'label': _('Spent on training'),
             'value': (cost.get('total_words') or '—') if not cost.get('mixed')
             else '—', 'icon': 'receipt',
             'sub': _('claims, outside courses and sessions'),
             'zero': not cost.get('total')},
            {'key': 'ratings_avg', 'label': _('Course rating'),
             'value': effect.get('ratings_avg') or '—', 'icon': 'star',
             'sub': _('out of 5') if effect.get('ratings_count')
             else _('nobody has rated yet'),
             'zero': not effect.get('ratings_count')},
        ]
        return tiles

    # =====================================================================
    #  the one read
    # =====================================================================
    @api.model
    def _cache_key(self, date_from, date_to, filters):
        start, end = self._range(date_from, date_to)
        return json.dumps([fields.Date.to_string(start),
                           fields.Date.to_string(end), filters or {}],
                          sort_keys=True, default=str)

    @api.model
    def get_numbers(self, date_from=None, date_to=None, filters=None):
        cached = self.env.context.get(CACHE_KEY)
        if cached and cached.get('key') == self._cache_key(
                date_from, date_to, filters):
            return cached['payload']
        payload = super().get_numbers(date_from, date_to, filters)
        if not payload.get('allowed'):
            return payload
        start, end = self._range(date_from, date_to)
        payload['hours'] = self._safe(
            lambda: self._hours_numbers(start, end, filters),
            default={'total': 0, 'people': 0, 'empty': True, 'headline': '',
                     'by_person': [], 'by_person_shown': [],
                     'by_department': [], 'by_course': [], 'by_kind': {},
                     'kinds': []})
        if not self.env.context.get(FULL_KEY):
            # The screen gets the first rows; the spreadsheet gets everybody.
            hours = payload['hours']
            hours['by_person'] = hours.get('by_person_shown') or []
        payload['cost'] = self._safe(
            lambda: self._cost_numbers(start, end, filters),
            default={'empty': True, 'mixed': False, 'total': 0,
                     'headline': '', 'budget': None})
        req = payload.get('required') or {}
        payload['compliance'] = {
            'percent': req.get('percent'),
            'missing': req.get('missing', 0),
            'expiring': req.get('expiring', 0),
            'people': req.get('people', 0),
            'people_missing': req.get('people_missing', 0),
            'headline': req.get('headline', ''),
            'by_site': req.get('by_site') or [],
            'by_department': req.get('by_department') or [],
            'by_country': req.get('by_country') or [],
            'by_group': req.get('by_group') or [],
            'by_requirement': req.get('by_requirement') or [],
        } if req.get('has_rules') else None
        payload['effect'] = self._safe(lambda: self._effect_numbers(payload),
                                       default={'empty': True,
                                                'by_course_top': []})
        payload['countries'] = self._safe(self._countries, default=[])
        if not payload.get('empty'):
            payload['tiles'] = (payload.get('tiles') or []) + \
                self._l6_tiles(payload)
        return payload

    # =====================================================================
    #  the spreadsheet: Hours, Cost, Compliance, Effect, Assignments
    # =====================================================================
    @api.model
    def export_xlsx(self, date_from=None, date_to=None, filters=None):
        self._require()
        board = self.with_context(**{FULL_KEY: True}).get_numbers(
            date_from, date_to, filters)
        me = self.with_context(**{CACHE_KEY: {
            'key': self._cache_key(date_from, date_to, filters),
            'payload': board}})
        res = super(PbTrainingAnalyticsL6, me).export_xlsx(
            date_from, date_to, filters)
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
            wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(
                res['file_b64'])))
            me._l6_sheets(wb, board, date_from, date_to, filters,
                          Font, PatternFill)
            out = io.BytesIO()
            wb.save(out)
            res['file_b64'] = base64.b64encode(out.getvalue()).decode()
        except Exception:               # noqa: BLE001 — the sheets are extra
            _logger.warning('pb_training: the L6 sheets could not be added',
                            exc_info=True)
        return res

    @api.model
    def _l6_sheets(self, wb, board, date_from, date_to, filters, Font,
                   PatternFill):
        def sheet(title, head):
            ws = wb.create_sheet(title[:31])
            ws.append(head)
            for i in range(1, len(head) + 1):
                cell = ws.cell(row=1, column=i)
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='6355C7')
            ws.column_dimensions['A'].width = 30
            ws.freeze_panes = 'A2'
            return ws

        # Summary: the hours line is the Hours block's (one number, R80).
        hours = board.get('hours') or {}
        summary = wb[wb.sheetnames[0]]
        for row in summary.iter_rows(min_col=1, max_col=2):
            if row[0].value == _('Hours of training finished'):
                row[1].value = hours.get('total', 0)

        ws = sheet(_('Hours'), [_('Person'), _('Department'), _('Hours'),
                                _('Courses'), _('Sessions')])
        for r in hours.get('by_person') or []:
            ws.append([r['employee'], r['department'], r['hours'],
                       r['courses'], r['sessions']])

        cost = board.get('cost') or {}
        ws = sheet(_('Cost'), [_('Figure'), _('Amount'), _('Currency')])
        cur = cost.get('currency') or ''
        if cost.get('mixed'):
            ws.append([_('More than one currency'), cost.get('total_words'),
                       ''])
        else:
            for label, key in ((_('Claims agreed'), 'claims_agreed'),
                               (_('Claims paid'), 'claims_paid'),
                               (_('Outside courses agreed'),
                                'requests_approved'),
                               (_('Sessions'), 'sessions'),
                               (_('Total'), 'total'),
                               (_('Allowances for the year'),
                                'allowance_total'),
                               (_('Allowances used'), 'allowance_used')):
                ws.append([label, cost.get(key), cur])
            ws.append([_('Allowances used %'), cost.get('allowance_used_pct'),
                       ''])
            budget = cost.get('budget')
            if budget:
                ws.append([budget['line'], budget['budget'], cur])
                ws.append([_('Spent from it'), budget['spent'], cur])
                ws.append([_('Left'), budget['left'], cur])

        comp = board.get('compliance') or {}
        ws = sheet(_('Compliance'), [_('Split'), _('Name'), _('In date %'),
                                     _('Counted'), _('Missing'),
                                     _('Expiring')])
        for split, key in ((_('By rule'), 'by_requirement'),
                           (_('By department'), 'by_department'),
                           (_('By site'), 'by_site'),
                           (_('By country'), 'by_country'),
                           (_('By safety group'), 'by_group')):
            for r in comp.get(key) or []:
                ws.append([split, r['name'], r.get('percent'),
                           r.get('measured'), r.get('missing'),
                           r.get('expiring')])

        effect = board.get('effect') or {}
        ratings = (board.get('ratings') or {}).get('rows') or []
        ws = sheet(_('Effect'), [_('Course'), _('Average rating'),
                                 _('Ratings'), _('Would recommend %'),
                                 _('Work changed (out of 5)'),
                                 _('Used what they learned %')])
        ws.append([_('Every course'), effect.get('ratings_avg'),
                   effect.get('ratings_count'), effect.get('recommend_pct'),
                   effect.get('followup_effect_avg'),
                   effect.get('followup_used_pct')])
        for r in ratings:
            ws.append([r['name'], r.get('avg'), r.get('count'),
                       r.get('recommend'), r.get('effect'), r.get('used')])

        ws = sheet(_('Assignments'), self._assignment_columns())
        for line in self._assignment_lines(date_from, date_to, filters):
            ws.append(line)
        for col in 'BCDEFGHIJK':
            ws.column_dimensions[col].width = 18

    @api.model
    def _assignment_columns(self):
        return [_('Employee'), _('Department'), _('Job'), _('Site'),
                _('Country'), _('Safety group'), _('Course'), _('Kind'),
                _('Why'), _('Path'), _('Requirement'), _('Assigned on'),
                _('Due'), _('Finished on'), _('State'), _('Score'),
                _('Hours'), _('Certificate valid until'), _('Version')]

    @api.model
    def _assignment_lines(self, date_from=None, date_to=None, filters=None):
        """One row per assignment in the range (the cohort the Numbers
        count), every dimension resolved — the flat sheet for BI (T5)."""
        from .hr_employee_ext import SAFETY_GROUP_LABEL
        from .training_common import reason_word, state_word
        start, end = self._range(date_from, date_to)
        rows = self.env['pb.training.assignment'].sudo().search(
            self._domain(start, end, filters or {}),
            order='assigned_on, id')
        employees = rows.mapped('employee_id')
        facts = self.env['pb.training']._person_facts(employees) \
            if employees else {}
        certs = {}
        if rows and 'pb.employee.document' in self.env:
            for doc in self.env['pb.employee.document'].sudo().search(
                    [('employee_id', 'in', employees.ids),
                     ('channel_id', 'in', rows.mapped('channel_id').ids),
                     ('renewed_by_id', '=', False)],
                    order='expiry_date desc'):
                certs.setdefault((doc.employee_id.id, doc.channel_id.id),
                                 doc.expiry_date)
        out = []
        for row in rows:
            f = facts.get(row.employee_id.id) or {}
            group = f.get('safety_group') or ''
            until = certs.get((row.employee_id.id, row.channel_id.id))
            out.append([
                f.get('name') or row.employee_id.sudo().name or '',
                f.get('department') or '',
                f.get('job') or '',
                f.get('site') or '',
                f.get('country') or '',
                tr(self.env, SAFETY_GROUP_LABEL.get(group, '')) if group
                else '',
                row.channel_id.sudo().name or '',
                tr(self.env, HOURS_KIND_LABEL[self._assignment_kind(row)]),
                reason_word(self.env, row.reason),
                row.path_enrolment_id.sudo().path_id.name or ''
                if 'path_enrolment_id' in row._fields else '',
                row.requirement_id.sudo().name or ''
                if 'requirement_id' in row._fields else '',
                fields.Date.to_string(row.assigned_on) or '',
                fields.Date.to_string(row.due_date) or '',
                fields.Date.to_string(row.completed_on) or '',
                state_word(self.env, row.state),
                row.score or '',
                row.hours or '',
                fields.Date.to_string(until) if until else '',
                row.course_version if 'course_version' in row._fields
                else '',
            ])
        return out


class PbTrainingSessionCost(models.Model):
    """What a session cost (trainer, room, materials) — the Numbers' Cost
    block adds it up. One currency: the session's company's."""
    _inherit = 'pb.training.session'

    cost = fields.Monetary(string='What it cost',
                           currency_field='currency_id')
    currency_id = fields.Many2one(
        'res.currency', string='Currency',
        related='company_id.currency_id', store=True, readonly=True)

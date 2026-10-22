# -*- coding: utf-8 -*-
"""Learning › Courses › SESSIONS — the board, the session page, the drawers,
the attendance screen and the toolbox talk (LEARNING L4).

One read per screen, behind the board's gate (`{'allowed': False}` without a
training group, never an AccessError):

  * `sessions_board(filters)`   the week strip (the hero), the four numbers,
                                upcoming and recent sessions as cards;
  * `session_page(session_id)`  one session: header, the people on it, what
                                can be pressed next;
  * `session_lists()`           the drawers' lists (classroom courses,
                                trainers, departments, jobs, kinds);
  * `attendance_poll(id)`       the attendance screen's live count.

WHO MAY DO WHAT: a trainer runs sessions (plan, book, invite, start, tick,
close, cancel) — that is the trainer's job; only a training manager opens a
closed session again. A TOOLBOX TALK may also be logged by anybody who manages
people, for the people under them (handover §3.3) — `toolbox_lists`,
`toolbox_people` and `log_toolbox` check that themselves.

Also here: the Overview's sessions block and "Sessions to close" chip, the
Numbers' sessions section and sheet, the course drawer's sessions, and My
learning's "My sessions" (`pb.my.training.home()['sessions']`), self-booking
and the QR scan.
"""

import base64
import io
import logging
from datetime import datetime, timedelta

from odoo import _, api, fields, models
from odoo.exceptions import AccessError, UserError
from odoo.tools.mimetypes import guess_mimetype

from .pb_training_paths import PbTrainingPaths
from .session import (
    ATT_LIVE, ATT_SEATED, ATT_STATE_LABEL, PHOTO_MAX_BYTES, PHOTO_MIMES,
    SESSION_DELIVERY, SESSION_KIND_LABEL, SESSION_KINDS, SESSION_LIVE,
    SESSION_STATE_LABEL,
)
from .training_common import (
    P_BULK_CAP, P_PEOPLE_LIMIT, P_REMINDER_CAP, P_REMINDERS, as_id, flag,
    fold, number, tr,
)

_logger = logging.getLogger(__name__)

#: Recent sessions on the board (the past list).
PAST_DAYS = 30
PAST_CAP = 60
#: Upcoming sessions on the board.
UPCOMING_CAP = 120

#: Problem first on the roster.
_ATT_RANK = {'waitlisted': 1, 'booked': 0, 'no_show': 2, 'attended': 3,
             'cancelled': 4}


class PbTrainingSessions(models.AbstractModel):
    _inherit = 'pb.training'

    _VERBS = PbTrainingPaths._VERBS + (
        'save_session', 'book_session', 'remove_booking', 'mark_attendance',
        'everyone_here', 'send_invites', 'start_session', 'close_session',
        'reopen_session', 'cancel_session', 'log_toolbox', 'attendance_xlsx',
    )

    # =====================================================================
    #  helpers
    # =====================================================================
    @api.model
    def _session(self, session_id):
        rec = self.env['pb.training.session'].sudo().browse(
            as_id(session_id)).exists()
        if not rec or rec.company_id.id not in self.env.companies.ids:
            raise UserError(_("That session is not here any more."))
        return rec

    @api.model
    def _me_employee(self):
        return self.env['hr.employee'].sudo().search(
            [('user_id', '=', self.env.uid),
             ('company_id', 'in', self.env.companies.ids)], limit=1)

    @api.model
    def _my_reports(self):
        """Everybody under the caller in the org chart (and the caller): the
        people a supervisor may log a toolbox talk for."""
        me = self._me_employee()
        if not me:
            return self.env['hr.employee'].sudo().browse()
        return self.env['hr.employee'].sudo().search(
            [('id', 'child_of', me.id), ('active', '=', True)])

    @api.model
    def _can_toolbox(self):
        if self._can_read():
            return True
        return len(self._my_reports()) > 1

    @api.model
    def _local_today(self):
        tz = self.env.context.get('tz') or self.env.user.tz or 'UTC'
        return fields.Datetime.context_timestamp(
            self.with_context(tz=tz), fields.Datetime.now()).date()

    @api.model
    def _local_day(self, session):
        return session._local(session.start).date() if session.start \
            else None

    @api.model
    def _kind_word(self, key):
        return tr(self.env, SESSION_KIND_LABEL.get(key or '', ''))

    @api.model
    def _session_card(self, rec):
        day = self._local_day(rec)
        tone = ''
        if rec.state == 'cancelled':
            tone = 'rose'
        elif rec.seats and rec.waiting_count:
            tone = 'amber'
        elif rec.state in ('open', 'held') and rec.stop and \
                rec.stop < fields.Datetime.now():
            tone = 'amber'
        return {
            'id': rec.id,
            'name': rec.name or '',
            'course': rec.channel_id.sudo().name or '',
            'channel_id': rec.channel_id.id,
            'kind': rec.kind,
            'kind_word': self._kind_word(rec.kind),
            'state': rec.state,
            'state_word': tr(self.env, SESSION_STATE_LABEL.get(rec.state, '')),
            'date': fields.Date.to_string(day) if day else '',
            'when_words': rec._when_words(),
            'where_words': rec._where_words(),
            'trainer': rec._trainer_words(),
            'seats': rec.seats,
            'booked': rec.booked_count,
            'attended': rec.attended_count,
            'waiting': rec.waiting_count,
            'no_show': rec.no_show_count,
            'seats_words': rec._seats_words(),
            'tone': tone,
            'to_close': rec.state in ('open', 'held') and bool(rec.stop)
            and rec.stop < fields.Datetime.now(),
        }

    @api.model
    def _sessions_scope(self):
        return [('company_id', 'in', self.env.companies.ids)]

    # =====================================================================
    #  reading
    # =====================================================================
    @api.model
    def _week(self, sessions, today, picked=''):
        """Seven columns from today: date, labels, how many sessions."""
        counts = {}
        for rec in sessions:
            if rec.state == 'cancelled':
                continue
            day = self._local_day(rec)
            if day:
                counts[day] = counts.get(day, 0) + 1
        out = []
        for i in range(7):
            day = today + timedelta(days=i)
            key = fields.Date.to_string(day)
            out.append({
                'date': key,
                'weekday': tr(self.env, day.strftime('%a')),
                'day': day.day,
                'count': counts.get(day, 0),
                'today': i == 0,
                'on': picked == key,
            })
        return out

    @api.model
    def _sessions_headline(self, k):
        if not k['upcoming']:
            return _("No sessions coming up.")
        bits = [_("1 session coming up") if k['upcoming'] == 1
                else _("%s sessions coming up", k['upcoming'])]
        if k['seats_booked']:
            bits.append(_("1 seat booked") if k['seats_booked'] == 1
                        else _("%s seats booked", k['seats_booked']))
        if k['to_close']:
            bits.append(_("1 to close") if k['to_close'] == 1
                        else _("%s to close", k['to_close']))
        return ' · '.join(bits)

    @api.model
    def sessions_board(self, filters=None):
        if not self._can_read():
            return {'allowed': False, 'upcoming': [], 'past': []}
        filters = filters or {}
        Session = self.env['pb.training.session'].sudo()
        now = fields.Datetime.now()
        today = self._local_today()
        scope = self._sessions_scope()
        live = Session.search(scope + [('state', 'in', SESSION_LIVE)],
                              order='start asc')
        upcoming = live.filtered(lambda s: not s.stop or s.stop >= now)
        to_close = live.filtered(lambda s: s.stop and s.stop < now
                                 and s.state in ('open', 'held'))
        past = Session.search(scope + [
            ('start', '>=', now - timedelta(days=PAST_DAYS)),
            '|', ('state', 'in', ('closed', 'cancelled')),
            '&', ('state', 'in', SESSION_LIVE), ('stop', '<', now)],
            order='start desc', limit=PAST_CAP)
        week_end = now + timedelta(days=7)
        this_week = live.filtered(lambda s: s.start and s.start <= week_end)
        recent_closed = Session.search(scope + [
            ('state', '=', 'closed'),
            ('start', '>=', now - timedelta(days=30))])
        kpis = {
            'this_week': len(this_week),
            'seats_free': sum(max(s.seats - s.booked_count, 0)
                              for s in upcoming if s.seats),
            'waiting': sum(upcoming.mapped('waiting_count')),
            'no_shows_30d': sum(recent_closed.mapped('no_show_count')),
            'upcoming': len(upcoming),
            'seats_booked': sum(upcoming.mapped('booked_count')),
            'to_close': len(to_close),
        }
        day = filters.get('day') or ''
        state = filters.get('state') or ''
        needle = fold(filters.get('q') or '')
        channel = as_id(filters.get('channel_id'))

        def keep(rec):
            if day and fields.Date.to_string(self._local_day(rec)) != day:
                return False
            if channel and rec.channel_id.id != channel:
                return False
            if needle and needle not in fold(rec.channel_id.sudo().name) \
                    and needle not in fold(rec.place or '') \
                    and needle not in fold(rec._trainer_words()):
                return False
            return True

        up_rows = upcoming
        if state == 'to_close':
            up_rows = to_close
        cards_up = [self._session_card(s) for s in up_rows if keep(s)]
        cards_past = [] if state == 'to_close' else \
            [self._session_card(s) for s in past if keep(s)
             and s not in up_rows]
        return {
            'allowed': True,
            'can_write': self._can_write(),
            'can_manage': self._can_manage(),
            'today': fields.Date.to_string(today),
            'week': self._week(live, today, day),
            'kpis': kpis,
            'headline': self._sessions_headline(kpis),
            'upcoming': cards_up[:UPCOMING_CAP],
            'past': cards_past,
            'filters': {'day': day, 'state': state,
                        'q': filters.get('q') or '', 'channel_id': channel},
        }

    @api.model
    def _roster(self, rec):
        rows = rec.attendance_ids.sudo().filtered(
            lambda a: a.state != 'cancelled')
        facts = {}
        if rows:
            for row in rows.mapped('employee_id').sudo().read(
                    ['name', 'department_id', 'user_id']):
                facts[row['id']] = row
        out = []
        for att in rows:
            f = facts.get(att.employee_id.id, {})
            out.append({
                'id': att.id,
                'employee_id': att.employee_id.id,
                'name': f.get('name') or '',
                'department': f.get('department_id') and
                f['department_id'][1] or '',
                'state': att.state,
                'state_word': att._state_word(),
                'checked_in': att._checked_in_words(),
                'how': att.how or '',
                'hours': att.hours or 0.0,
                'note': att.note or '',
                'has_login': bool(f.get('user_id')),
                'invited': bool(att.invited),
            })
        out.sort(key=lambda r: (_ATT_RANK.get(r['state'], 9),
                                fold(r['name'])))
        return out

    @api.model
    def session_page(self, session_id):
        if not self._can_read():
            return {'allowed': False}
        rec = self._session(session_id)
        card = self._session_card(rec)
        editable = rec.state not in ('closed', 'cancelled')
        return dict(card, **{
            'allowed': True,
            'can_write': self._can_write(),
            'can_manage': self._can_manage(),
            'place': rec.place or '',
            'join_url': rec.join_url or '',
            'hours': rec.hours or 0.0,
            'minutes': rec.minutes or 0,
            'notes': rec.notes or '',
            'roster': self._roster(rec),
            'editable': editable,
            'can_invite': rec.state in ('planned', 'open') and editable,
            'can_start': rec.state in ('planned', 'open'),
            'can_close': rec.state in ('open', 'held'),
            'can_reopen': rec.state == 'closed' and self._can_manage(),
            'can_cancel': rec.state in SESSION_LIVE,
            'attend_url': rec._attend_url(),
            'qr_src': rec._qr_src(400),
            'pdf_url': '/report/pdf/pb_training.report_session_sheet/%s'
            % rec.id,
            'photo_url': '/web/content/%s?download=true' % rec.photo_id.id
            if rec.photo_id else '',
            'values': {
                'id': rec.id,
                'channel_id': rec.channel_id.id,
                'kind': rec.kind,
                'start': fields.Datetime.to_string(rec.start) or '',
                'minutes': rec.minutes or 0,
                'place': rec.place or '',
                'join_url': rec.join_url or '',
                'trainer_id': rec.trainer_id.id or 0,
                'trainer_name': rec.trainer_name or '',
                'seats': rec.seats or 0,
                'hours': rec.hours or 0.0,
            },
        })

    @api.model
    def attendance_poll(self, session_id):
        """The attendance screen asks every 10 seconds: who is here."""
        if not self._can_read():
            return {'allowed': False}
        rec = self._session(session_id)
        rows = rec.attendance_ids.sudo().filtered(
            lambda a: a.state in ATT_LIVE)
        here = len(rows.filtered(lambda a: a.state == 'attended'))
        expected = len(rows.filtered(lambda a: a.state in ATT_SEATED))
        return {
            'allowed': True,
            'state': rec.state,
            'here': here,
            'expected': expected,
            'here_words': _("%(here)s of %(n)s here", here=here, n=expected),
            'roster': self._roster(rec),
        }

    @api.model
    def _classroom_courses(self):
        rows = self.env['slide.channel'].sudo().search(
            [('active', '=', True), ('delivery', 'in', SESSION_DELIVERY)],
            limit=number(self.env, 'pb_training.course_limit', 200))
        return sorted([{
            'id': c.id, 'name': c.name or '',
            'hours': c._pb_session_hours(),
            'delivery': c.delivery,
            'published': bool(c.is_published),
        } for c in rows], key=lambda r: fold(r['name']))

    @api.model
    def session_lists(self):
        if not self._can_read():
            return {'allowed': False}
        comps = self.env.companies.ids
        return {
            'allowed': True,
            'courses': self._classroom_courses(),
            'kinds': [{'key': k, 'label': tr(self.env, v)}
                      for k, v in SESSION_KINDS],
            'departments': sorted(
                [{'id': d['id'], 'name': d['complete_name'] or d['name']}
                 for d in self.env['hr.department'].sudo().search_read(
                     [('company_id', 'in', comps)],
                     ['name', 'complete_name'])],
                key=lambda r: fold(r['name'])),
            'jobs': sorted([{'id': j['id'], 'name': j['name']}
                            for j in self.env['hr.job'].sudo().search_read(
                                [('company_id', 'in', comps)], ['name'])],
                           key=lambda r: fold(r['name'])),
        }

    @api.model
    def _people_rows(self, term, domain=None):
        """Accent-blind people search (F7) that ALSO offers people with no
        login: a person at a session or a toolbox talk counts in hours and on
        their record even when nothing can be put on a training page."""
        needle = fold(term or '')
        base = [('company_id', 'in', self.env.companies.ids),
                ('active', '=', True)] + (domain or [])
        rows = self.env['hr.employee'].sudo().search_read(
            base, ['name', 'user_id', 'department_id'])
        cap = number(self.env, P_PEOPLE_LIMIT)
        out = []
        for row in rows:
            if needle and needle not in fold(row['name']):
                continue
            out.append({'id': row['id'], 'name': row['name'],
                        'department': row['department_id'] and
                        row['department_id'][1] or '',
                        'has_login': bool(row['user_id'])})
            if len(out) >= cap:
                break
        return sorted(out, key=lambda r: fold(r['name']))

    @api.model
    def session_people(self, term=''):
        self._require_read()
        return self._people_rows(term)

    # =====================================================================
    #  writing a session (trainers)
    # =====================================================================
    @api.model
    def _parse_start(self, raw):
        if not raw:
            return None
        try:
            return fields.Datetime.to_datetime(raw)
        except (TypeError, ValueError):
            return None

    @api.model
    def _new_classroom_course(self, name, hours, toolbox=False):
        name = (name or '').strip()[:200]
        if not name:
            raise UserError(_("Give the new course a name."))
        channel = self.env['slide.channel'].sudo().create({
            'name': name, 'channel_type': 'training',
            'visibility': 'members', 'enroll': 'invite',
            'completed_template_id': False, 'is_published': True,
            'delivery': 'session',
            'session_hours': hours,
            'pb_hours': hours,
        })
        # The content engine puts a course's responsible (the caller) on it
        # as a member, so whoever logged a toolbox talk found the course on
        # their own page (walk finding). Take that membership off again; a
        # caller who attended is booked like everybody else.
        partner = self.env.user.partner_id
        if partner:
            channel._remove_membership(partner.ids)
        return channel

    @api.model
    def save_session(self, payload=None):
        self._require_write()
        payload = payload or {}
        start = self._parse_start(payload.get('start'))
        if not start:
            raise UserError(_("Say when the session starts."))
        try:
            mins = max(int(payload.get('minutes') or 60), 1)
            seats = max(int(payload.get('seats') or 0), 0)
            hours = max(float(payload.get('hours') or 0.0), 0.0)
        except (TypeError, ValueError):
            raise UserError(_("These numbers cannot be read."))
        kind = payload.get('kind') or 'classroom'
        if kind not in dict(SESSION_KINDS):
            kind = 'classroom'
        cid = as_id(payload.get('channel_id'))
        if not cid and (payload.get('new_course') or '').strip():
            cid = self._new_classroom_course(payload['new_course'],
                                             round(mins / 60.0, 2)).id
        if not cid:
            raise UserError(_("Pick the course this session is for."))
        trainer = as_id(payload.get('trainer_id'))
        if trainer and not self.env['hr.employee'].sudo().browse(
                trainer).exists():
            trainer = 0
        vals = {
            'channel_id': cid,
            'kind': kind,
            'start': start,
            'minutes': mins,
            'place': (payload.get('place') or '').strip()[:200],
            'join_url': (payload.get('join_url') or '').strip()[:500],
            'trainer_id': trainer or False,
            'trainer_name': (payload.get('trainer_name') or '').strip()[:200]
            if not trainer else '',
            'seats': seats,
            'hours': hours or round(mins / 60.0, 2),
        }
        sid = as_id(payload.get('id'))
        if sid:
            rec = self._session(sid)
            rec._check_editable()
            moved = rec.start != start or rec.minutes != mins or \
                (rec.place or '') != vals['place']
            rec.write(vals)
            rec._promote()
            if rec.state in ('open', 'held'):
                rec._sync_event()
                if moved:
                    # A moved session: everybody booked gets the new time.
                    rec.sudo().write(
                        {'invite_sequence': rec.invite_sequence + 1})
                    for att in rec.attendance_ids.sudo().filtered(
                            lambda a: a.state == 'booked'):
                        att._send_invite()
            message = _("The session is saved.")
        else:
            vals['company_id'] = self.env.company.id
            rec = self.env['pb.training.session'].sudo().create(vals)
            message = _("%s is planned. Book people on it, then send the "
                        "invitations.", rec.name)
        return {'id': rec.id, 'message': message}

    @api.model
    def _book_audience(self, rec, payload):
        Employee = self.env['hr.employee'].sudo()
        base = [('active', '=', True),
                ('company_id', 'in', self.env.companies.ids)]
        if as_id(payload.get('department_id')):
            return Employee.search(base + [(
                'department_id', 'child_of', as_id(payload['department_id']))])
        if as_id(payload.get('job_id')):
            return Employee.search(base + [(
                'job_id', '=', as_id(payload['job_id']))])
        if payload.get('on_course'):
            # Everybody who still owes this course.
            rows = self.env['pb.training.assignment'].sudo().search([
                ('channel_id', '=', rec.channel_id.id),
                ('state', 'in', ('assigned', 'in_progress', 'overdue',
                                 'excused')),
                ('company_id', 'in', self.env.companies.ids)])
            return rows.mapped('employee_id')
        ids = [as_id(e) for e in (payload.get('employee_ids') or [])
               if as_id(e)]
        return Employee.search(base + [('id', 'in', ids)])

    @api.model
    def book_session(self, payload=None):
        self._require_write()
        payload = payload or {}
        rec = self._session(payload.get('session_id'))
        rec._check_editable()
        people = self._book_audience(rec, payload)
        if not people:
            raise UserError(_("Pick at least one person."))
        cap = number(self.env, P_BULK_CAP, 200)
        if len(people) > cap:
            raise UserError(_(
                "That is %(n)s people at once. The most this screen will do "
                "in one press is %(cap)s — pick a smaller group.",
                n=len(people), cap=cap))
        res = rec._book(people, by_user=self.env.user)
        bits = []
        if res['booked']:
            bits.append(_("1 person is booked.") if len(res['booked']) == 1
                        else _("%s people are booked.", len(res['booked'])))
        if res['waitlisted']:
            bits.append(_("1 person is waiting for a seat.")
                        if len(res['waitlisted']) == 1
                        else _("%s people are waiting for a seat.",
                               len(res['waitlisted'])))
        if res['already']:
            bits.append(_("%s was already on it.", ', '.join(
                res['already'][:5])) if len(res['already']) == 1
                else _("%s were already on it.", ', '.join(
                    res['already'][:5])))
        if not bits:
            bits.append(_("Nobody new was booked."))
        if rec.state in ('open', 'held') and res['booked']:
            bits.append(_("Their invitations have gone out."))
        return dict(res, message=' '.join(bits),
                    type='success' if res['booked'] or res['waitlisted']
                    else 'warning')

    @api.model
    def _attendance(self, attendance_id):
        att = self.env['pb.training.attendance'].sudo().browse(
            as_id(attendance_id)).exists()
        if not att or att.session_id.company_id.id not in \
                self.env.companies.ids:
            raise UserError(_("That person is not on the session any more."))
        return att

    @api.model
    def remove_booking(self, attendance_id):
        self._require_write()
        att = self._attendance(attendance_id)
        rec = att.session_id
        rec._check_editable()
        name = att.employee_id.name or ''
        rec._unbook(att)
        return {'message': _("%s is off this session.", name)}

    @api.model
    def mark_attendance(self, payload=None):
        self._require_write()
        payload = payload or {}
        att = self._attendance(payload.get('attendance_id'))
        rec = att.session_id
        rec._check_editable()
        state = payload.get('state') or 'attended'
        if state not in ('attended', 'no_show', 'booked'):
            raise UserError(_("That is not something this screen can do."))
        if state == 'attended' and rec.state in ('planned', 'open'):
            rec.action_start()
        att._mark(state, 'trainer')
        if 'hours' in payload:
            try:
                att.write({'hours': max(float(payload['hours'] or 0), 0.0)})
            except (TypeError, ValueError):
                pass
        return {'state': att.state, 'message': _(
            "%(who)s: %(state)s.", who=att.employee_id.name or '',
            state=att._state_word())}

    @api.model
    def everyone_here(self, session_id):
        self._require_write()
        rec = self._session(session_id)
        rec._check_editable()
        if rec.state in ('planned', 'open'):
            rec.action_start()
        rows = rec.attendance_ids.sudo().filtered(
            lambda a: a.state == 'booked')
        rows._mark('attended', 'trainer')
        return {'marked': len(rows), 'message': _(
            "Everybody booked is marked here (%s).", len(rows))}

    @api.model
    def send_invites(self, session_id):
        self._require_write()
        rec = self._session(session_id)
        if not rec.attendance_ids.filtered(lambda a: a.state == 'booked'):
            raise UserError(_("Book people on the session first — there is "
                              "nobody to invite yet."))
        sent = rec.action_send_invites()
        if not sent:
            return {'message': _("Everybody booked already has their "
                                 "invitation."), 'type': 'warning'}
        return {'message': _("1 invitation sent, with a calendar file.")
                if sent == 1 else
                _("%s invitations sent, each with a calendar file.", sent)}

    @api.model
    def start_session(self, session_id):
        self._require_write()
        rec = self._session(session_id)
        rec.action_start()
        return {'message': _("The session is under way. Show the code so "
                             "people can scan it.")}

    @api.model
    def close_session(self, session_id):
        self._require_write()
        rec = self._session(session_id)
        if rec.state not in ('open', 'held', 'planned'):
            raise UserError(_("This session is already %s.", tr(
                self.env, SESSION_STATE_LABEL[rec.state]).lower()))
        rec.action_close()
        return {'message': _(
            "Closed. %(here)s attended and their course is marked finished; "
            "%(away)s did not come.", here=rec.attended_count,
            away=rec.no_show_count)}

    @api.model
    def reopen_session(self, session_id):
        self._require_manager()
        rec = self._session(session_id)
        rec.action_reopen()
        return {'message': _("The session is open again.")}

    @api.model
    def cancel_session(self, session_id):
        self._require_write()
        rec = self._session(session_id)
        told = len(rec.attendance_ids.filtered(
            lambda a: a.state in ('booked', 'waitlisted') and a.invited))
        rec.action_cancel()
        if told:
            return {'message': _("Cancelled. 1 person was told, with a "
                                 "calendar update.") if told == 1 else
                    _("Cancelled. %s people were told, with a calendar "
                      "update.", told)}
        return {'message': _("Cancelled. Nobody had an invitation yet, so "
                             "nobody was written to.")}

    # =====================================================================
    #  the toolbox talk
    # =====================================================================
    @api.model
    def toolbox_lists(self):
        if not self._can_toolbox():
            return {'allowed': False}
        return {
            'allowed': True,
            'is_team': not self._can_read(),
            'courses': self._classroom_courses(),
            'now': fields.Datetime.to_string(fields.Datetime.now()),
        }

    @api.model
    def toolbox_people(self, term=''):
        if self._can_read():
            return self._people_rows(term)
        reports = self._my_reports()
        if len(reports) <= 1:
            raise AccessError(_("A toolbox talk is logged by the training "
                                "team or by somebody who manages people."))
        return self._people_rows(term, [('id', 'in', reports.ids)])

    @api.model
    def _photo(self, payload):
        raw = payload.get('photo')
        if not raw:
            return None
        try:
            data = base64.b64decode(raw, validate=False)
        except Exception:               # noqa: BLE001 — a bad file is refused
            raise UserError(_("That photo could not be read."))
        if len(data) > PHOTO_MAX_BYTES:
            raise UserError(_("The photo is bigger than 5 MB."))
        mime = guess_mimetype(data[:4096] if len(data) > 4096 else data)
        if mime not in PHOTO_MIMES:
            raise UserError(_("Only a JPG, PNG, WEBP or PDF file."))
        return data, mime

    @api.model
    def log_toolbox(self, payload=None):
        """One drawer → a closed toolbox session with everybody attended, in
        one go. A supervisor may include only the people under them."""
        payload = payload or {}
        is_team = not self._can_read()
        if is_team and not self._can_toolbox():
            raise AccessError(_("A toolbox talk is logged by the training "
                                "team or by somebody who manages people."))
        ids = [as_id(e) for e in (payload.get('employee_ids') or [])
               if as_id(e)]
        if not ids:
            raise UserError(_("Tick who was there."))
        people = self.env['hr.employee'].sudo().search([
            ('id', 'in', ids), ('company_id', 'in', self.env.companies.ids)])
        if is_team:
            allowed = set(self._my_reports().ids)
            if any(p.id not in allowed for p in people):
                raise AccessError(_("You can log a toolbox talk only for the "
                                    "people who report to you."))
        try:
            mins = max(int(payload.get('minutes') or 15), 1)
        except (TypeError, ValueError):
            mins = 15
        start = self._parse_start(payload.get('start')) or \
            fields.Datetime.now() - timedelta(minutes=mins)
        photo = self._photo(payload)
        cid = as_id(payload.get('channel_id'))
        Channel = self.env['slide.channel'].sudo()
        if cid:
            channel = Channel.browse(cid).exists()
            if not channel or not channel._pb_is_session():
                raise UserError(_("Pick a classroom course, or name a new "
                                  "toolbox topic."))
        else:
            channel = self._new_classroom_course(
                payload.get('new_course') or '', 0.25, toolbox=True)
        rec = self.env['pb.training.session'].sudo().create({
            'channel_id': channel.id,
            'kind': 'toolbox',
            'start': start,
            'minutes': mins,
            'hours': round(mins / 60.0, 2),
            'place': (payload.get('place') or '').strip()[:200],
            'trainer_id': self._me_employee().id or False,
            'state': 'held',
            'notes': payload.get('notes') or False,
            'logged_by_id': self.env.uid,
            'company_id': (people[:1].company_id or self.env.company).id,
        })
        rec._book(people, by_user=self.env.user, send=False,
                  state='attended', how='trainer')
        if photo:
            data, mime = photo
            attachment = self.env['ir.attachment'].sudo().create({
                'name': (payload.get('photo_name') or 'sign-in sheet')[:120],
                'raw': data, 'mimetype': mime,
                'res_model': 'pb.training.session', 'res_id': rec.id,
            })
            rec.write({'photo_id': attachment.id})
        n = len(people)
        rec.action_close(note=_(
            "Toolbox talk logged by %(who)s: 1 person.",
            who=self.env.user.name or '') if n == 1 else _(
            "Toolbox talk logged by %(who)s: %(n)s people.",
            who=self.env.user.name or '', n=n))
        if n == 1:
            message = _("Toolbox talk logged: %(topic)s, 1 person.",
                        topic=channel.name)
        else:
            message = _("Toolbox talk logged: %(topic)s, %(n)s people.",
                        topic=channel.name, n=n)
        return {'id': rec.id, 'message': message}

    # =====================================================================
    #  the attendance sheet (spreadsheet)
    # =====================================================================
    @api.model
    def attendance_xlsx(self, session_id):
        self._require_read()
        rec = self._session(session_id)
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            raise UserError(_("This system cannot build spreadsheets at the "
                              "moment. Ask an administrator to look at it."))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = _('Attendance')[:31]
        head = [_('Name'), _('Department'), _('Where they stand'),
                _('Checked in'), _('Hours'), _('Note')]
        ws.append([rec.channel_id.sudo().name or ''])
        ws.append([rec._when_words(), rec._where_words(),
                   rec._trainer_words()])
        ws.append([])
        ws.append(head)
        for i in range(1, len(head) + 1):
            cell = ws.cell(row=4, column=i)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='5A4BB0')
            cell.alignment = Alignment(wrap_text=True)
        ws.cell(row=1, column=1).font = Font(bold=True, size=13)
        for row in self._roster(rec):
            ws.append([row['name'], row['department'], row['state_word'],
                       row['checked_in'], row['hours'], row['note']])
        ws.column_dimensions['A'].width = 32
        ws.column_dimensions['B'].width = 26
        ws.column_dimensions['C'].width = 20
        out = io.BytesIO()
        wb.save(out)
        return {
            'ok': True,
            'file_b64': base64.b64encode(out.getvalue()).decode(),
            'filename': '%s.xlsx' % (rec.name or _('Attendance')),
            'mimetype': ('application/vnd.openxmlformats-officedocument.'
                         'spreadsheetml.sheet'),
        }

    # =====================================================================
    #  the course drawer
    # =====================================================================
    @api.model
    def get_course(self, channel_id):
        data = super().get_course(channel_id)
        try:
            channel = self.env['slide.channel'].sudo().browse(
                as_id(channel_id)).exists()
            data['delivery'] = channel.delivery or 'online'
            data['delivery_word'] = tr(self.env, dict(
                channel._fields['delivery'].selection).get(
                    channel.delivery, ''))
            sessions = self.env['pb.training.session'].sudo().search(
                self._sessions_scope() + [('channel_id', '=', channel.id),
                                          ('state', '!=', 'cancelled')],
                order='start desc', limit=12)
            data['sessions'] = [self._session_card(s) for s in sessions]
        except Exception:               # noqa: BLE001 — never takes the drawer
            _logger.warning('pb_training: the course sessions could not be '
                            'read', exc_info=True)
            data.setdefault('delivery', 'online')
            data.setdefault('sessions', [])
        return data

    # =====================================================================
    #  the Overview's block
    # =====================================================================
    @api.model
    def _sessions_overview(self):
        Session = self.env['pb.training.session'].sudo()
        now = fields.Datetime.now()
        scope = self._sessions_scope()
        if not Session.search_count(scope):
            return {'has_sessions': False}
        week = Session.search(scope + [
            ('state', 'in', SESSION_LIVE), ('start', '>=', now - timedelta(
                hours=12)), ('start', '<=', now + timedelta(days=7))])
        to_close = Session.search_count(scope + [
            ('state', 'in', ('open', 'held')), ('stop', '<', now)])
        return {
            'has_sessions': True,
            'this_week': len(week),
            'seats_booked': sum(week.mapped('booked_count')),
            'waiting': sum(week.mapped('waiting_count')),
            'to_close': to_close,
        }

    @api.model
    def overview(self):
        data = super().overview()
        if not data.get('allowed'):
            return data
        block = self._safe(self._sessions_overview,
                           default={'has_sessions': False})
        data['sessions'] = block
        if block.get('has_sessions'):
            n = block.get('to_close', 0)
            if n:
                data['calm'] = False
            data['queue'] = data.get('queue', []) + [{
                'key': 'sessions', 'label': _("Sessions to close"),
                'count': n, 'lens': 'courses', 'tab': 'sessions',
                'session_state': 'to_close',
            }]
        return data

    # =====================================================================
    #  the dispatcher
    # =====================================================================
    @api.model
    def act(self, verb, payload=None):
        payload = payload or {}
        if verb == 'save_session':
            return self.save_session(payload)
        if verb == 'book_session':
            return self.book_session(payload)
        if verb == 'remove_booking':
            return self.remove_booking(payload.get('attendance_id'))
        if verb == 'mark_attendance':
            return self.mark_attendance(payload)
        if verb == 'everyone_here':
            return self.everyone_here(payload.get('session_id'))
        if verb == 'send_invites':
            return self.send_invites(payload.get('session_id'))
        if verb == 'start_session':
            return self.start_session(payload.get('session_id'))
        if verb == 'close_session':
            return self.close_session(payload.get('session_id'))
        if verb == 'reopen_session':
            return self.reopen_session(payload.get('session_id'))
        if verb == 'cancel_session':
            return self.cancel_session(payload.get('session_id'))
        if verb == 'log_toolbox':
            return self.log_toolbox(payload)
        if verb == 'attendance_xlsx':
            return self.attendance_xlsx(payload.get('session_id'))
        return super().act(verb, payload)


# =========================================================================
#  the Numbers
# =========================================================================
class PbTrainingAnalyticsSessions(models.AbstractModel):
    _inherit = 'pb.training.analytics'

    @api.model
    def _sessions_numbers(self, start, end, filters=None):
        """Sessions that started in the range: hours by kind, attendance and
        no-show rates, by trainer and by course."""
        filters = filters or {}
        domain = [('company_id', 'in', self.env.companies.ids),
                  ('state', '=', 'closed'),
                  ('start', '>=', datetime.combine(start, datetime.min.time())),
                  ('start', '<=', datetime.combine(end, datetime.max.time()))]
        if filters.get('channel_id'):
            domain.append(('channel_id', '=', as_id(filters['channel_id'])))
        sessions = self.env['pb.training.session'].sudo().search(
            domain, order='start asc')
        if not sessions:
            return {'has_sessions': False, 'sessions': 0, 'rows': []}
        atts = sessions.mapped('attendance_ids').filtered(
            lambda a: a.state in ('attended', 'no_show'))
        dept = as_id(filters.get('department_id'))
        if dept:
            covered = set(self.env['hr.department'].sudo().search(
                [('id', 'child_of', dept)]).ids)
            atts = atts.filtered(
                lambda a: a.employee_id.department_id.id in covered)
        attended = atts.filtered(lambda a: a.state == 'attended')
        expected = len(atts)

        def rate(part, whole):
            return int(round(100.0 * part / whole)) if whole else None

        def split(keyfn):
            buckets = {}
            for att in atts:
                key = keyfn(att) or _('Not said')
                row = buckets.setdefault(key, {'name': key, 'sessions': set(),
                                               'attended': 0, 'no_show': 0,
                                               'hours': 0.0})
                row['sessions'].add(att.session_id.id)
                if att.state == 'attended':
                    row['attended'] += 1
                    row['hours'] += att.hours or 0.0
                else:
                    row['no_show'] += 1
            out = []
            for row in buckets.values():
                whole = row['attended'] + row['no_show']
                out.append({'name': row['name'],
                            'sessions': len(row['sessions']),
                            'attended': row['attended'],
                            'no_show': row['no_show'],
                            'hours': round(row['hours'], 1),
                            'rate': rate(row['attended'], whole)})
            out.sort(key=lambda r: (-r['hours'], fold(r['name'])))
            return out

        return {
            'has_sessions': True,
            'sessions': len(sessions),
            'attended': len(attended),
            'no_show': expected - len(attended),
            'hours': round(sum(attended.mapped('hours')), 1),
            'attendance_rate': rate(len(attended), expected),
            'no_show_rate': rate(expected - len(attended), expected),
            'by_kind': split(lambda a: tr(self.env, SESSION_KIND_LABEL.get(
                a.session_id.kind, ''))),
            'by_trainer': split(lambda a: a.session_id._trainer_words()),
            'by_course': split(lambda a: a.session_id.channel_id.sudo().name),
            'rows': [{
                'name': s.channel_id.sudo().name or '',
                'kind': tr(self.env, SESSION_KIND_LABEL.get(s.kind, '')),
                'when': s._when_words(),
                'where': s._where_words(),
                'trainer': s._trainer_words(),
                'attended': s.attended_count,
                'no_show': s.no_show_count,
                'hours': s.hours or 0.0,
            } for s in sessions],
        }

    @api.model
    def get_numbers(self, date_from=None, date_to=None, filters=None):
        payload = super().get_numbers(date_from, date_to, filters)
        if not payload.get('allowed'):
            return payload
        start, end = self._range(date_from, date_to)
        payload['sessions'] = self._safe(
            lambda: self._sessions_numbers(start, end, filters),
            default={'has_sessions': False, 'rows': []})
        return payload

    @api.model
    def export_xlsx(self, date_from=None, date_to=None, filters=None):
        res = super().export_xlsx(date_from, date_to, filters)
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
            start, end = self._range(date_from, date_to)
            data = self._sessions_numbers(start, end, filters or {})
            wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(
                res['file_b64'])))
            ws = wb.create_sheet(_('Sessions')[:31])
            head = [_('Course'), _('Kind'), _('When'), _('Where'),
                    _('Trainer'), _('Attended'), _('Did not come'),
                    _('Hours')]
            ws.append(head)
            for i in range(1, len(head) + 1):
                cell = ws.cell(row=1, column=i)
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='6355C7')
            for row in data.get('rows') or []:
                ws.append([row['name'], row['kind'], row['when'],
                           row['where'], row['trainer'], row['attended'],
                           row['no_show'], row['hours']])
            ws.column_dimensions['A'].width = 30
            ws.column_dimensions['C'].width = 28
            out = io.BytesIO()
            wb.save(out)
            res['file_b64'] = base64.b64encode(out.getvalue()).decode()
        except Exception:               # noqa: BLE001 — the sheet is extra
            _logger.warning('pb_training: the Sessions sheet could not be '
                            'added', exc_info=True)
        return res


# =========================================================================
#  the learner: My sessions, booking a seat, the QR scan
# =========================================================================
class PbMyTrainingSessions(models.AbstractModel):
    _inherit = 'pb.my.training'

    @api.model
    def _session_row(self, att, today_dt):
        rec = att.session_id
        return {
            'id': rec.id,
            'attendance_id': att.id if att else 0,
            'course': rec.channel_id.sudo().name or '',
            'channel_id': rec.channel_id.id,
            'kind_word': tr(self.env, SESSION_KIND_LABEL.get(rec.kind, '')),
            'when_words': rec._when_words(),
            'where_words': rec._where_words(),
            'trainer': rec._trainer_words(),
            'join_url': rec.join_url or '',
            'state': rec.state,
            'my_state': att.state if att else '',
            'my_state_word': att._state_word() if att else '',
            'ics_url': '/my/training/sessions/%s/calendar.ics' % rec.id,
            'start': fields.Datetime.to_string(rec.start) or '',
        }

    @api.model
    def _bookable(self, emp, channel_ids=None):
        """Open sessions of courses this person owes and may book
        themselves, with a free seat, not booked already."""
        if not emp:
            return self.env['pb.training.session']
        Assignment = self.env['pb.training.assignment'].sudo()
        owed = Assignment.search([
            ('employee_id', '=', emp.id),
            ('state', 'in', ('assigned', 'in_progress', 'overdue',
                             'excused'))]).mapped('channel_id')
        owed = owed.filtered(lambda c: c.self_booking
                             and c.delivery in SESSION_DELIVERY)
        if channel_ids is not None:
            owed = owed.filtered(lambda c: c.id in channel_ids)
        if not owed:
            return self.env['pb.training.session']
        sessions = self.env['pb.training.session'].sudo().search([
            ('channel_id', 'in', owed.ids),
            ('state', 'in', ('planned', 'open')),
            ('start', '>', fields.Datetime.now()),
            ('company_id', '=', emp.company_id.id)], order='start asc')
        mine = set(self.env['pb.training.attendance'].sudo().search([
            ('employee_id', '=', emp.id),
            ('session_id', 'in', sessions.ids),
            ('state', '!=', 'cancelled')]).mapped('session_id').ids)
        return sessions.filtered(
            lambda s: s.id not in mine and s._free_seats() > 0)

    @api.model
    def _my_sessions(self):
        emp = self._employee()
        if not emp:
            return {'next': None, 'rows': [], 'bookable': []}
        now = fields.Datetime.now()
        rows = self.env['pb.training.attendance'].sudo().search([
            ('employee_id', '=', emp.id),
            ('state', 'in', ('booked', 'waitlisted', 'attended')),
            ('session_id.state', '!=', 'cancelled'),
            ('session_start', '>=', now - timedelta(days=1))],
            order='session_start asc', limit=20)
        out = [self._session_row(a, now) for a in rows]
        upcoming = [r for r, a in zip(out, rows)
                    if a.state in ('booked', 'waitlisted')
                    and (a.session_id.stop or a.session_id.start) >= now]
        bookable = []
        for rec in self._bookable(emp)[:6]:
            bookable.append({
                'id': rec.id, 'course': rec.channel_id.sudo().name or '',
                'channel_id': rec.channel_id.id,
                'when_words': rec._when_words(),
                'where_words': rec._where_words(),
                'seats_words': rec._seats_words(),
            })
        return {'next': upcoming[0] if upcoming else None, 'rows': out,
                'bookable': bookable}

    @api.model
    def home(self):
        data = super().home()
        try:
            data['sessions'] = self._my_sessions()
        except Exception:               # noqa: BLE001 — never a 500 on /my
            _logger.warning('pb_training: my sessions could not be read',
                            exc_info=True)
            data['sessions'] = {'next': None, 'rows': [], 'bookable': []}
        return data

    def _tile(self, row):
        data = super()._tile(row)
        channel = row.channel_id.sudo()
        data['delivery'] = channel.delivery or 'online'
        if channel.delivery not in SESSION_DELIVERY:
            return data
        emp = self._employee()
        att = self.env['pb.training.attendance'].sudo().search([
            ('employee_id', '=', emp.id if emp else 0),
            ('channel_id', '=', channel.id),
            ('state', 'in', ('booked', 'waitlisted', 'attended'))],
            order='session_start desc', limit=1) if emp else None
        data['session'] = {
            'when_words': att.session_id._when_words(),
            'where_words': att.session_id._where_words(),
            'my_state': att.state,
            'my_state_word': att._state_word(),
        } if att else None
        if channel.delivery == 'session' and channel.attendance_completes:
            attended = bool(att and att.state == 'attended')
            data['finished'] = attended and data.get('test') in (
                'passed', 'none')
            data['percent'] = 100 if data['finished'] else 0
        return data

    @api.model
    def course(self, channel_id):
        data = super().course(channel_id)
        if not data:
            return data
        try:
            channel = self.env['slide.channel'].sudo().browse(
                as_id(channel_id))
            data['delivery'] = channel.delivery or 'online'
            emp = self._employee()
            mine = self.env['pb.training.attendance'].sudo().search([
                ('employee_id', '=', emp.id if emp else 0),
                ('channel_id', '=', channel.id),
                ('state', '!=', 'cancelled')], order='session_start desc',
                limit=5) if emp else []
            now = fields.Datetime.now()
            data['my_sessions'] = [self._session_row(a, now) for a in mine]
            data['bookable'] = [{
                'id': s.id, 'when_words': s._when_words(),
                'where_words': s._where_words(),
                'seats_words': s._seats_words(),
            } for s in self._bookable(emp, [channel.id])] if emp else []
        except Exception:               # noqa: BLE001 — never a 500 on /my
            _logger.warning('pb_training: course sessions could not be read',
                            exc_info=True)
            data.setdefault('my_sessions', [])
            data.setdefault('bookable', [])
        return data

    @api.model
    def book_my_seat(self, session_id):
        """Self-booking: only on a course the person owes, that allows it,
        on a session not started yet with a free seat (T12)."""
        emp = self._employee()
        if not emp:
            raise UserError(_("You do not have an employee record yet, so "
                              "there is nothing to book a seat for."))
        rec = self.env['pb.training.session'].sudo().browse(
            as_id(session_id)).exists()
        if not rec or rec not in self._bookable(emp):
            raise UserError(_("That session cannot be booked from here. It "
                              "may be full, already started, or not open to "
                              "booking — ask the training team."))
        res = rec._book(emp)
        if res['booked']:
            return {'ok': True, 'code': 'booked',
                    'message': _("You have a seat. The invitation comes by "
                                 "email.")
                    if rec.state in ('open', 'held') else
                    _("You have a seat. The invitation comes by email when "
                      "the training team sends them.")}
        return {'ok': True, 'code': 'waiting',
                'message': _("The session is full, so you are waiting for a "
                             "seat.")}

    @api.model
    def attend_by_token(self, token, join=False):
        """The QR scan. Returns {'code', 'session'}: ok / already / ask /
        full / invalid / closed / not_yet / no_employee."""
        token = (token or '').strip()
        Session = self.env['pb.training.session'].sudo()
        rec = Session.search([('qr_token', '=', token)], limit=1) \
            if len(token) >= 12 else Session.browse()
        if not rec:
            return {'code': 'invalid', 'session': None}
        info = {'course': rec.channel_id.sudo().name or '',
                'when_words': rec._when_words(),
                'where_words': rec._where_words()}
        if rec.state in ('closed', 'cancelled'):
            return {'code': 'closed', 'session': info}
        emp = self.env['hr.employee'].sudo().search(
            [('user_id', '=', self.env.uid),
             ('company_id', '=', rec.company_id.id)], limit=1)
        if not emp:
            return {'code': 'no_employee', 'session': info}
        if not rec._scan_window_open():
            return {'code': 'not_yet', 'session': info}
        att = rec.attendance_ids.sudo().filtered(
            lambda a: a.employee_id == emp)[:1]
        if att and att.state == 'attended':
            return {'code': 'already', 'session': info}
        if att and att.state in ('booked', 'no_show'):
            att._mark('attended', 'qr')
            return {'code': 'ok', 'session': info}
        if rec._free_seats() <= 0:
            return {'code': 'full', 'session': info}
        if not join:
            return {'code': 'ask', 'session': info,
                    'seats_words': rec._seats_words()}
        if att:
            att.write({'state': 'booked'})
            att._mark('attended', 'qr')
        else:
            rec._book(emp, send=False, state='attended', how='qr')
        return {'code': 'ok', 'session': info}

    @api.model
    def session_ics(self, session_id):
        """"Add to my calendar": the same file the invitation carries, for a
        person on the session (or the training team)."""
        rec = self.env['pb.training.session'].sudo().browse(
            as_id(session_id)).exists()
        if not rec or rec.state == 'cancelled':
            return None
        emp = self._employee()
        mine = emp and rec.attendance_ids.filtered(
            lambda a: a.employee_id == emp and a.state != 'cancelled')
        if not mine and not self.env['pb.training']._can_read():
            return None
        return rec._ics()


# =========================================================================
#  the morning job
# =========================================================================
class PbTrainingAutomationSessions(models.AbstractModel):
    """Two legs (handover §3.5): the day-before reminder (under the shared
    send cap, behind the one reminders switch) and closing a session that
    has been "under way" for a whole day after it ended."""
    _inherit = 'pb.training.automation'

    @api.model
    def _cron_daily(self):
        counts = super()._cron_daily() or {}
        now = fields.Datetime.now()
        counts['sessions_closed'] = self._leg(
            'close sessions that ended a day ago',
            lambda: self._sessions_autoclose(now))
        if flag(self.env, P_REMINDERS):
            counts['session_reminders'] = self._leg(
                'remind people of tomorrow\'s sessions',
                lambda: self._session_reminders(now, counts))
        return counts

    @api.model
    def _session_reminders(self, now=None, counts=None):
        """Everybody booked on a session that starts tomorrow (in the
        session's own time zone) gets one reminder. Shares the send cap with
        the rest of the morning's mail."""
        now = now or fields.Datetime.now()
        counts = counts if counts is not None else {}
        cap = number(self.env, P_REMINDER_CAP, 400)
        used = sum(counts.get(k) or 0 for k in
                   ('employee', 'manager', 'hr', 'certificates'))
        sent = 0
        sessions = self.env['pb.training.session'].sudo().search([
            ('state', '=', 'open'), ('reminded', '=', False),
            ('start', '>', now), ('start', '<', now + timedelta(hours=60))])
        for rec in sessions:
            zone = rec.create_uid.tz or rec.company_id.partner_id.tz or 'UTC'
            local = rec.with_context(tz=zone)
            today = local._local(now).date()
            if local._local(rec.start).date() != today + timedelta(days=1):
                continue
            for att in rec.attendance_ids.filtered(
                    lambda a: a.state == 'booked' and not a.reminded):
                if used + sent >= cap:
                    counts['capped'] = (counts.get('capped') or 0) + 1
                    return sent
                try:
                    with self.env.cr.savepoint():
                        if att._send_mail(
                                'pb_training.mail_template_session_reminder',
                                with_ics=False):
                            sent += 1
                        att.write({'reminded': True})
                except Exception:       # noqa: BLE001 — one row, one grave
                    _logger.warning('pb_training: no reminder for attendance '
                                    '%s', att.id, exc_info=True)
            rec.write({'reminded': True})
        return sent

    @api.model
    def _sessions_autoclose(self, now=None):
        now = now or fields.Datetime.now()
        rows = self.env['pb.training.session'].sudo().search([
            ('state', '=', 'held'), ('stop', '<', now - timedelta(hours=24))])
        closed = 0
        for rec in rows:
            try:
                with self.env.cr.savepoint():
                    rec.action_close(note=_(
                        "Closed by itself a day after it ended: %(here)s "
                        "attended, %(away)s did not come.",
                        here=len(rec.attendance_ids.filtered(
                            lambda a: a.state == 'attended')),
                        away=len(rec.attendance_ids.filtered(
                            lambda a: a.state == 'booked'))))
                closed += 1
            except Exception:           # noqa: BLE001 — one row, one grave
                _logger.warning('pb_training: session %s could not be closed '
                                'by itself', rec.id, exc_info=True)
        return closed

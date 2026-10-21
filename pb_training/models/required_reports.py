# -*- coding: utf-8 -*-
"""LEARNING L2 — required training on the Numbers, the Overview and My
learning, and the four country reports.

Every figure here comes from `pb.training._matrix_data` /
`pb.training.requirement.evaluate` — the matrix's own cells — so the Overview
tile, the Numbers section and the matrix headline are the same number (R80).

THE COUNTRY REPORTS (handover §4.3), one spreadsheet each, offered only for
packs switched on for the reader's company:

  * Vietnam   "Renewals due in 30 days" — person, safety group, course or
              card, hours done, valid until, card number, trainer;
  * India     "Awareness session register" — person, type (contractors
              included), session date, evidence; per calendar year;
  * Indonesia "Licence register" — person, licence, number, issuer, valid
              until;
  * Singapore "Training hours" — person, course, hours, finished on,
              residency words; per calendar year.
"""

import base64
import io
import logging
from datetime import timedelta

from odoo import _, api, fields, models
from odoo.exceptions import UserError

from .hr_employee_ext import SAFETY_GROUP_LABEL
from .pb_training_required import EMPLOYEE_TYPES, _date_words, _num
from .requirement import PACK_BY_CC, PACKS
from .training_common import fold, tr

_logger = logging.getLogger(__name__)

#: The documented columns of each report (T12 reads these).
REPORT_COLUMNS = {
    'vn': ('Person', 'Safety group', 'Course or card', 'Hours done',
           'Valid until', 'Card number', 'Trainer'),
    'in': ('Person', 'Type', 'Session date', 'Evidence'),
    'id': ('Person', 'Licence', 'Number', 'Issuer', 'Valid until'),
    'sg': ('Person', 'Course', 'Hours', 'Finished on', 'Residency'),
}
REPORT_TITLES = {
    'vn': 'Renewals due in 30 days',
    'in': 'Awareness session register',
    'id': 'Licence register',
    'sg': 'Training hours',
}


class PbTrainingAnalyticsRequired(models.AbstractModel):
    _inherit = 'pb.training.analytics'

    # =====================================================================
    #  the Numbers section
    # =====================================================================
    @api.model
    def _required_numbers(self):
        board = self.env['pb.training']
        if not board._live_requirements():
            return {'has_rules': False}
        data = board._matrix_data({}, paged=False)
        s = data['summary']

        def split(key_of):
            groups = {}
            for row in data['rows']:
                name = key_of(row)
                g = groups.setdefault(name, {'name': name, 'in_date': 0,
                                             'measured': 0, 'missing': 0,
                                             'expiring': 0})
                for cell in row['cells'].values():
                    st = cell['state']
                    if st in ('valid', 'expiring', 'missing', 'in_progress'):
                        g['measured'] += 1
                    if st in ('valid', 'expiring'):
                        g['in_date'] += 1
                    if st == 'missing':
                        g['missing'] += 1
                    if st == 'expiring':
                        g['expiring'] += 1
            out = []
            for g in groups.values():
                g['percent'] = int(round(100.0 * g['in_date'] / g['measured'])) \
                    if g['measured'] else None
                out.append(g)
            out.sort(key=lambda g: (-g['missing'], fold(g['name'])))
            return out

        by_req = []
        for col in data['columns']:
            g = {'name': col['name'], 'in_date': 0, 'measured': 0,
                 'missing': 0, 'expiring': 0}
            for row in data['rows']:
                cell = row['cells'].get(str(col['id']))
                if not cell:
                    continue
                st = cell['state']
                if st in ('valid', 'expiring', 'missing', 'in_progress'):
                    g['measured'] += 1
                if st in ('valid', 'expiring'):
                    g['in_date'] += 1
                if st == 'missing':
                    g['missing'] += 1
                if st == 'expiring':
                    g['expiring'] += 1
            g['percent'] = int(round(100.0 * g['in_date'] / g['measured'])) \
                if g['measured'] else None
            by_req.append(g)
        by_req.sort(key=lambda g: (-g['missing'], fold(g['name'])))
        none = _('No department')
        no_site = _('No site')
        return {
            'has_rules': True,
            'percent': s['percent'],
            'missing': s['missing'],
            'expiring': s['expiring'],
            'people': s['people'],
            'people_missing': s['people_missing'],
            'headline': data['headline'],
            'by_department': split(lambda r: r['department'] or none),
            'by_requirement': by_req,
            'by_site': split(lambda r: r['site'] or no_site),
        }

    @api.model
    def _required_tiles(self, req):
        if not req.get('has_rules'):
            return []
        pct = req.get('percent')
        return [
            {'key': 'required_pct', 'label': _('Required training in date'),
             'value': '—' if pct is None else '%s%%' % pct,
             'icon': 'listChecks',
             'sub': _('across %s people', req.get('people', 0))},
            {'key': 'required_missing', 'label': _('Required, missing'),
             'value': req.get('missing', 0), 'icon': 'triangleAlert',
             'sub': _('courses or papers nobody has yet')},
            {'key': 'required_expiring', 'label': _('Required, expiring'),
             'value': req.get('expiring', 0), 'icon': 'hourglass',
             'sub': _('inside their renewal window')},
        ]

    @api.model
    def get_numbers(self, date_from=None, date_to=None, filters=None):
        payload = super().get_numbers(date_from, date_to, filters)
        if not payload.get('allowed'):
            return payload
        req = self._safe(self._required_numbers, default={'has_rules': False})
        payload['required'] = req
        payload['country_reports'] = self._safe(self._country_reports,
                                                default=[])
        if payload.get('tiles') is not None and not payload.get('empty'):
            payload['tiles'] = payload['tiles'] + self._required_tiles(req)
        return payload

    # =====================================================================
    #  the country reports
    # =====================================================================
    @api.model
    def _country_reports(self):
        Req = self.env['pb.training.requirement']
        out = []
        for pack in PACKS:
            if Req._pack_is_on(pack['cc'], self.env.company):
                out.append({'cc': pack['cc'],
                            'title': tr(self.env, REPORT_TITLES[pack['cc']]),
                            'pack': tr(self.env, pack['label'])})
        return out

    @api.model
    def _pack_rules(self, cc, active_only=True):
        pack = PACK_BY_CC[cc]
        domain = [('company_id', '=', self.env.company.id),
                  ('pack_key', '=', pack['key'])]
        if active_only:
            domain.append(('active', '=', True))
        return self.env['pb.training.requirement'].sudo().search(domain)

    @api.model
    def country_report_rows(self, cc, year=None):
        """(columns, rows) for one country report — the spreadsheet's and
        the tests' one source."""
        self._require()
        if cc not in PACK_BY_CC:
            raise UserError(_("There is no such country report."))
        today = fields.Date.today()
        year = int(year or today.year)
        columns = [tr(self.env, c) for c in REPORT_COLUMNS[cc]]
        rows = getattr(self, '_report_%s' % cc)(today, year)
        return columns, rows

    @api.model
    def _report_vn(self, today, year):
        rules = self._pack_rules('vn').filtered(
            lambda r: not r._is_headcount())
        Req = self.env['pb.training.requirement']
        cells = Req.evaluate(rules, today=today)
        horizon = today + timedelta(days=30)
        docs, rows_by_id = {}, {}
        doc_ids = [c['evidence']['id'] for c in cells.values()
                   if c.get('evidence') and c['evidence']['kind'] == 'document']
        if doc_ids:
            for d in self.env['pb.employee.document'].sudo().browse(
                    doc_ids).read(['reference', 'issuer', 'hours']):
                docs[d['id']] = d
        as_ids = [c['evidence']['id'] for c in cells.values()
                  if c.get('evidence')
                  and c['evidence']['kind'] == 'assignment']
        if as_ids:
            for a in self.env['pb.training.assignment'].sudo().browse(
                    as_ids).read(['hours']):
                rows_by_id[a['id']] = a
        emps = self.env['hr.employee'].sudo().browse(
            sorted({e for _r, e in cells})).read(['name', 'pb_safety_group'])
        emp_by = {e['id']: e for e in emps}
        out = []
        for rule in rules:
            for (rid, emp_id), cell in cells.items():
                if rid != rule.id or cell['state'] not in ('expiring', 'valid',
                                                           'missing'):
                    continue
                until = cell.get('valid_until')
                if cell['state'] == 'missing' or not until or until > horizon:
                    continue
                ev = cell.get('evidence') or {}
                doc = docs.get(ev.get('id')) if ev.get('kind') == 'document' \
                    else None
                row = rows_by_id.get(ev.get('id')) \
                    if ev.get('kind') == 'assignment' else None
                emp = emp_by.get(emp_id, {})
                hours = (doc or {}).get('hours') or (row or {}).get('hours') \
                    or 0
                out.append([
                    emp.get('name') or '',
                    tr(self.env, SAFETY_GROUP_LABEL.get(
                        emp.get('pb_safety_group') or '', '')),
                    rule.name,
                    _num(hours) if hours else '',
                    _date_words(self.env, until),
                    (doc or {}).get('reference') or '',
                    (doc or {}).get('issuer') or '',
                ])
        out.sort(key=lambda r: fold(r[0]))
        return out

    @api.model
    def _report_in(self, today, year):
        rules = self._pack_rules('in', active_only=False).filtered(
            lambda r: r.code == 'in_posh_all' and r.channel_id)
        if not rules:
            return []
        start = fields.Date.to_date('%s-01-01' % year)
        end = fields.Date.to_date('%s-12-31' % year)
        types = dict(EMPLOYEE_TYPES)
        rows = self.env['pb.training.assignment'].sudo().search([
            ('channel_id', 'in', rules.mapped('channel_id').ids),
            ('company_id', '=', self.env.company.id),
            ('state', '=', 'done'),
            ('completed_on', '>=', start), ('completed_on', '<=', end),
        ], order='completed_on, id')
        out = []
        for row in rows:
            emp = row.employee_id.sudo()
            out.append([
                emp.name or '',
                tr(self.env, types.get(emp.employee_type or '', '')
                   or 'Employee'),
                _date_words(self.env, row.completed_on),
                _('Finished %s', row.channel_id.sudo().name or ''),
            ])
        return out

    @api.model
    def _report_id(self, today, year):
        rules = self._pack_rules('id').filtered(
            lambda r: r.target == 'licence')
        if not rules:
            return []
        people = self.env['hr.employee'].sudo().search(
            [('company_id', '=', self.env.company.id), ('active', '=', True)])
        docs = self.env['pb.employee.document'].sudo().search_read(
            [('employee_id', 'in', people.ids),
             ('renewed_by_id', '=', False)],
            ['employee_id', 'name', 'category_id', 'survey_id', 'channel_id',
             'issue_date', 'expiry_date', 'training_source', 'company_id',
             'reference', 'issuer'])
        company_of = {e.id: e.company_id.id for e in people}
        out = []
        for doc in docs:
            emp_id = doc['employee_id'][0]
            if not any(r._doc_matches(doc, company_of.get(emp_id))
                       for r in rules):
                continue
            out.append([doc['employee_id'][1], doc['name'] or '',
                        doc['reference'] or '', doc['issuer'] or '',
                        _date_words(self.env, doc['expiry_date'])])
        out.sort(key=lambda r: (fold(r[0]), fold(r[1])))
        return out

    @api.model
    def _report_sg(self, today, year):
        start = fields.Date.to_date('%s-01-01' % year)
        end = fields.Date.to_date('%s-12-31' % year)
        rows = self.env['pb.training.assignment'].sudo().search([
            ('company_id', '=', self.env.company.id),
            ('state', '=', 'done'),
            ('completed_on', '>=', start), ('completed_on', '<=', end),
        ], order='completed_on, id')
        sg = self.env.ref('base.sg', raise_if_not_found=False)
        out = []
        for row in rows:
            emp = row.employee_id.sudo()
            hours = row.hours or row.channel_id._pb_hours()
            residency = _('Singapore national') \
                if sg and emp.country_id == sg else ''
            out.append([emp.name or '', row.channel_id.sudo().name or '',
                        _num(hours) if hours else '',
                        _date_words(self.env, row.completed_on), residency])
        return out

    @api.model
    def country_report_xlsx(self, cc, year=None):
        columns, rows = self.country_report_rows(cc, year)
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise UserError(_("This system cannot build spreadsheets at the "
                              "moment. Ask an administrator to look at it."))
        wb = openpyxl.Workbook()
        ws = wb.active
        title = tr(self.env, REPORT_TITLES[cc])
        ws.title = title[:31]
        for i, label in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=i, value=label)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='5A4BB0')
            cell.alignment = Alignment(wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = \
                30 if i == 1 else 20
        ws.freeze_panes = 'A2'
        for r, line in enumerate(rows, start=2):
            for c, value in enumerate(line, start=1):
                ws.cell(row=r, column=c, value=value)
        if cc == 'in':
            dates = {line[2] for line in rows}
            summary = wb.create_sheet(_('Summary')[:31])
            for r, (label, value) in enumerate((
                    (_('Year'), int(year or fields.Date.today().year)),
                    (_('Sessions held'), len(dates)),
                    (_('People who attended'), len({l[0] for l in rows})),
                    (_('Contractors among them'), len(
                        [l for l in rows if l[1] in (
                            tr(self.env, 'Contractor'),
                            tr(self.env, 'Freelancer'))]))), start=1):
                summary.cell(row=r, column=1, value=label).font = \
                    Font(bold=True)
                summary.cell(row=r, column=2, value=value)
        out = io.BytesIO()
        wb.save(out)
        return {
            'ok': True,
            'file_b64': base64.b64encode(out.getvalue()).decode(),
            'filename': '%s %s.xlsx' % (title, fields.Date.to_string(
                fields.Date.today())),
            'mimetype': ('application/vnd.openxmlformats-officedocument.'
                         'spreadsheetml.sheet'),
            'rows': len(rows),
        }


class PbTrainingOverviewRequired(models.AbstractModel):
    """The Overview's compliance tile, queue chip and headline fragment."""
    _inherit = 'pb.training'

    @api.model
    def overview(self):
        data = super().overview()
        if not data.get('allowed'):
            return data
        req = self._safe(self._required_overview,
                         default={'has_rules': False})
        data['required'] = req
        if req.get('has_rules'):
            fragment = self._required_fragment(req)
            if fragment:
                data['headline'] = '%s %s' % (data.get('headline', ''),
                                              fragment)
            data['queue'] = data.get('queue', []) + [{
                'key': 'required', 'label': _("Missing required training"),
                'count': req.get('people_missing', 0),
                'lens': 'required', 'tab': 'matrix', 'state': 'missing',
            }]
            if req.get('people_missing'):
                data['calm'] = False
        return data


class PbMyTrainingRequired(models.AbstractModel):
    """"Required for you" on My learning."""
    _inherit = 'pb.my.training'

    @api.model
    def required_for_me(self):
        emp = self.env['hr.employee'].sudo().search(
            [('user_id', '=', self.env.uid)], limit=1)
        if not emp:
            return {'has_employee': False, 'rows': []}
        Req = self.env['pb.training.requirement'].sudo()
        reqs = Req.search([('company_id', '=', emp.company_id.id),
                           ('active', '=', True)]).filtered(
            lambda r: not r._is_headcount())
        cells = Req.evaluate(reqs, employees=emp) if reqs else {}
        open_ids = [c['open_id'] for c in cells.values() if c.get('open_id')]
        channel_of = {}
        if open_ids:
            for a in self.env['pb.training.assignment'].sudo().browse(
                    open_ids).read(['channel_id']):
                channel_of[a['id']] = a['channel_id'][0] \
                    if a['channel_id'] else 0
        rank = {'missing': 0, 'expiring': 1, 'in_progress': 2, 'not_due': 3,
                'valid': 4}
        rows = []
        for req in reqs:
            cell = cells.get((req.id, emp.id))
            if not cell or cell['state'] == 'exempt':
                continue
            rows.append({
                'id': req.id,
                'name': req.name,
                'state': cell['state'],
                'words': cell['words'],
                'channel_id': channel_of.get(cell.get('open_id'), 0),
            })
        rows.sort(key=lambda r: (rank.get(r['state'], 5), fold(r['name'])))
        return {'has_employee': True, 'rows': rows}

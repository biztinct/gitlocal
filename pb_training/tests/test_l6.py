# -*- coding: utf-8 -*-
"""LEARNING L6 — Numbers and the Overview, finished. Handover tests T1–T10
(T11 is the Chrome walk).

The rails: an hour counted twice (a course finished by attending a session),
a total over two currencies, a compliance figure that is not the matrix's, a
"did it work" figure that is not the answers', a flat sheet with a missing or
shuffled column, a headline that buries the worst news, a queue chip in
colour at zero, a setting anybody can change, a bad address saved, a test
pack that reaches somebody other than the person who pressed it, an amount
on Home › TEAM, and a lens that says "no access" when the server only
failed to answer.

Fixtures are "DEMO" and roll back. The classes run in a company of their own
(ledger LN32) and pin `lang='en_US'` (LN14).
"""

import base64
import io
import json
import re
from datetime import timedelta

from odoo import fields
from odoo.exceptions import AccessError, UserError
from odoo.tests import tagged

from .test_assignments import AssignmentCase
from .test_training import _PDF, _src

NEW_FILES = (
    ('models', 'numbers_l6.py'),
    ('models', 'overview_l6.py'),
    ('views', 'learning_settings.xml'),
    ('static', 'src', 'js', 'learning_settings.js'),
    ('static', 'src', 'xml', 'learning_settings.xml'),
    ('static', 'src', 'scss', 'learning_settings.scss'),
)


class L6Case(AssignmentCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env['res.company'].create({'name': 'DEMO L6 Co'})
        cls.env = cls.env(context=dict(
            cls.env.context, lang='en_US', tz='UTC',
            allowed_company_ids=[cls.company.id]))
        (cls.boss | cls.e_one | cls.e_two | cls.e_three).write(
            {'company_id': cls.company.id})
        for user in (cls.boss_user, cls.u_one, cls.u_two, cls.u_three):
            user.write({'company_ids': [(4, cls.company.id)],
                        'company_id': cls.company.id, 'tz': 'UTC'})
        cls.dept = cls.env['hr.department'].create(
            {'name': 'DEMO L6 Plant', 'company_id': cls.company.id})
        cls.office = cls.env['hr.department'].create(
            {'name': 'DEMO L6 Office', 'company_id': cls.company.id})
        (cls.e_one | cls.e_two).write({'department_id': cls.dept.id})
        cls.e_three.department_id = cls.office
        cls.room = cls.env['slide.channel'].create({
            'name': 'DEMO L6 Site safety walk', 'channel_type': 'training',
            'visibility': 'members', 'enroll': 'invite',
            'completed_template_id': False, 'is_published': True,
            'delivery': 'session', 'session_hours': 2.0})
        cls.simple.pb_hours = 3.0
        Users = cls.env['res.users'].with_context(no_reset_password=True)

        def user(login, *groups):
            return Users.create({
                'name': login, 'login': login, 'email': login, 'tz': 'UTC',
                'company_id': cls.company.id,
                'company_ids': [(6, 0, cls.company.ids)],
                'group_ids': [(6, 0, [cls.env.ref('base.group_user').id]
                               + [cls.env.ref(g).id for g in groups])]})
        cls.trainer = user('demo.l6.trainer@example.com',
                           'pb_training.group_training_user')
        cls.tmanager = user('demo.l6.tmanager@example.com',
                            'pb_training.group_training_manager')
        cls.head = user('demo.l6.head@example.com',
                        'pb_training.group_training_admin')
        cls.plain = user('demo.l6.plain@example.com')
        cls.Analytics = cls.env['pb.training.analytics']
        cls.Board = cls.env['pb.training']
        cls.today = fields.Date.today()

    # ------------------------------------------------------------- helpers
    def _numbers(self, user=None, filters=None, date_from=None, date_to=None):
        return self.Analytics.with_user(user or self.tmanager).get_numbers(
            date_from or fields.Date.to_string(self.today - timedelta(days=30)),
            date_to or fields.Date.to_string(self.today), filters or {})

    def _finished(self, user, emp):
        row = self.env['pb.training.assignment'].sudo().create({
            'channel_id': self.simple.id, 'employee_id': emp.id,
            'company_id': self.company.id})
        self._finish(user, self.simple, [self.simple_lesson])
        row._refresh()
        self.assertEqual(row.state, 'done')
        return row

    def _session(self, days=-1, **vals):
        start = fields.Datetime.now().replace(microsecond=0) + \
            timedelta(days=days)
        base = {'channel_id': self.room.id, 'start': start, 'minutes': 120,
                'hours': 2.0, 'company_id': self.company.id,
                'trainer_id': self.boss.id, 'state': 'held'}
        base.update(vals)
        return self.env['pb.training.session'].sudo().create(base)

    def _dormant(self, keys=('training_request', 'training_claim')):
        if 'biz.approval.binding' not in self.env:
            return
        for key in keys:
            process = self.env['biz.approval.process'].sudo().search(
                [('key', '=', key)], limit=1)
            if process:
                self.env['biz.approval.binding'].sudo().search(
                    [('process_id', '=', process.id)]).write(
                        {'active': False})


# =========================================================================
#  T1 — hours add up, by kind, over the period
# =========================================================================
@tagged('post_install', '-at_install')
class TestT1Hours(L6Case):

    def _world(self):
        self._finished(self.u_one, self.e_one)              # 3 h online
        sess = self._session()
        sess._book(self.e_one, send=False, state='attended', how='trainer')
        sess._book(self.e_two, send=False)
        sess.action_close()                                 # 2 h, e_two no-show
        self.Board.with_user(self.trainer).log_toolbox({
            'new_course': 'DEMO L6 Ladders', 'minutes': 15,
            'employee_ids': [self.e_three.id]})            # 0.25 h toolbox
        return sess

    def test_t1_person_is_courses_plus_sessions_and_counted_once(self):
        self._world()
        hours = self._numbers()['hours']
        people = {r['employee_id']: r for r in hours['by_person']}
        one = people[self.e_one.id]
        self.assertEqual(one['hours'], 5, 'three online + two in the room')
        self.assertEqual((one['courses'], one['sessions']), (1, 1),
                         'the room course is counted once, through the '
                         'attendance')
        self.assertNotIn(self.e_two.id, people, 'did not come = no hours')
        self.assertEqual(people[self.e_three.id]['hours'], 0.25)
        self.assertEqual(hours['total'], 5.25)
        self.assertEqual(hours['total'],
                         sum(r['hours'] for r in hours['by_person']))
        self.assertEqual(hours['by_kind']['online'], 3)
        self.assertEqual(hours['by_kind']['session'], 2)
        self.assertEqual(hours['by_kind']['toolbox'], 0.25)
        self.assertEqual(sum(r['hours'] for r in hours['by_department']),
                         hours['total'])
        self.assertIn('5.25 hours of training', hours['headline'])

    def test_t1_period_and_kind_filters(self):
        self._world()
        old = self._numbers(
            date_from=fields.Date.to_string(self.today - timedelta(days=90)),
            date_to=fields.Date.to_string(self.today - timedelta(days=60)))
        self.assertEqual(old['hours']['total'], 0)
        self.assertTrue(old['hours']['empty'])
        self.assertIn('No training hours were finished',
                      old['hours']['headline'])
        toolbox = self._numbers(filters={'kind': 'toolbox'})['hours']
        self.assertEqual(toolbox['total'], 0.25)
        plant = self._numbers(filters={'department_id': self.dept.id})
        self.assertEqual(plant['hours']['total'], 5)

    def test_t1_tile_and_screen_rows(self):
        self._world()
        data = self._numbers()
        tiles = {t['key']: t for t in data['tiles']}
        self.assertEqual(tiles['hours_total']['value'], 5.25)
        self.assertIn('cost_total', tiles)
        self.assertIn('ratings_avg', tiles)
        self.assertLessEqual(len(data['hours']['by_person']), 50,
                             'the screen gets the first rows only')


# =========================================================================
#  T2 — cost
# =========================================================================
@tagged('post_install', '-at_install')
class TestT2Cost(L6Case):

    def _claim(self, emp, amount):
        claim = self.env['pb.training.claim'].create({
            'employee_id': emp.id, 'course_name': 'DEMO L6 Evening class',
            'provider': 'DEMO College', 'amount': amount,
            'currency_id': self.company.currency_id.id,
            'paid_on': self.today, 'company_id': self.company.id,
            'invoice_file': _PDF, 'invoice_filename': 'r.pdf',
            'certificate_file': _PDF, 'certificate_filename': 'c.pdf'})
        claim.sudo()._advance_state('submitted')
        claim.sudo()._advance_state('approved')
        return claim

    def test_t2_claims_requests_sessions_and_allowance(self):
        self._dormant()
        self.env['pb.training.allowance'].create({
            'company_id': self.company.id, 'year': self.today.year,
            'amount': 2000000.0, 'currency_id': self.company.currency_id.id})
        self._claim(self.e_one, 1000000.0)
        ext = self.env['pb.training.outside.course'].create({
            'name': 'DEMO L6 Licence', 'cost': 500000.0,
            'currency_id': self.company.currency_id.id, 'kind': 'exam',
            'company_id': self.company.id})
        res = self.env['pb.my.training'].with_user(self.u_two).ask_course({
            'kind': 'outside', 'outside_course_id': ext.id,
            'why': 'DEMO for the job'})
        req = self.env['pb.training.request'].browse(res['id']).sudo()
        req._advance_state('approved')
        self._session(cost=250000.0)
        cost = self._numbers()['cost']
        self.assertFalse(cost['mixed'])
        self.assertEqual(cost['claims_agreed'], 1000000.0)
        self.assertEqual(cost['requests_approved'], 500000.0)
        self.assertEqual(cost['sessions'], 250000.0)
        self.assertEqual(cost['total'], 1750000.0)
        # four people with the company's 2,000,000 each = 8,000,000; used
        # = the claim and the agreed request = 1,500,000 → 19%.
        people = self.env['hr.employee'].search_count(
            [('company_id', '=', self.company.id), ('active', '=', True)])
        self.assertEqual(cost['allowance_total'], 2000000.0 * people)
        self.assertEqual(cost['allowance_used'], 1500000.0)
        self.assertEqual(cost['allowance_used_pct'],
                         int(round(100.0 * 1500000.0 / (2000000.0 * people))))
        self.assertIn('spent on training', cost['headline'])

    def test_t2_budget_line_only_when_budget_is_installed(self):
        cost = self._numbers()['cost']
        if 'pb.budget.line' not in self.env:
            self.assertIsNone(cost.get('budget'))
            return
        self.assertIsNone(cost.get('budget'), 'no line, no budget')
        Line = self.env['pb.budget.line'].sudo()
        line = Line.create({
            'period_month': self.today.replace(day=1),
            'company_id': self.company.id,
            'pb_budget_type': 'hr_ops', 'forecast_cost': 900000.0,
            'actual_cost': 300000.0,
            'pb_currency_id': self.company.currency_id.id})
        budget = self._numbers()['cost']['budget']
        self.assertEqual((budget['budget'], budget['spent'], budget['left']),
                         (900000.0, 300000.0, 600000.0))
        self.assertTrue(line.exists(), 'read only: the line is untouched')

    def test_t2_two_currencies_are_never_added(self):
        from odoo.addons.pb_training.models.numbers_l6 import one_currency
        vnd = self.env.ref('base.VND')
        usd = self.env.ref('base.USD')
        self.assertEqual(one_currency({vnd: 5.0, usd: 0.0}), vnd)
        with self.assertRaises(UserError):
            one_currency({vnd: 5.0, usd: 2.0})

    def test_t2_empty_sentence(self):
        cost = self._numbers()['cost']
        self.assertTrue(cost['empty'])
        self.assertIn('Nothing was spent on training', cost['headline'])


# =========================================================================
#  T3 — compliance is the matrix's own summary
# =========================================================================
@tagged('post_install', '-at_install')
class TestT3Compliance(L6Case):

    def test_t3_no_rules_is_null(self):
        self.assertIsNone(self._numbers()['compliance'])

    def test_t3_equals_the_matrix_for_the_same_filters(self):
        self.env['pb.training.requirement'].create({
            'name': 'DEMO L6 rule', 'company_id': self.company.id,
            'kind': 'always', 'target': 'course',
            'channel_id': self.simple.id, 'auto_assign': False,
            'department_ids': [(6, 0, (self.dept | self.office).ids)]})
        self._finished(self.u_one, self.e_one)
        board = self.Board.with_user(self.tmanager)
        for flt in ({}, {'department_id': self.dept.id}):
            comp = self._numbers(filters=flt)['compliance']
            summary = board._matrix_data(flt, paged=False)['summary']
            self.assertEqual(comp['percent'], summary['percent'], flt)
            self.assertEqual(comp['missing'], summary['missing'], flt)
            self.assertEqual(comp['people'], summary['people'], flt)
        comp = self._numbers()['compliance']
        for key in ('by_site', 'by_department', 'by_country', 'by_group'):
            self.assertIn(key, comp)
        self.assertEqual(sum(r['measured'] for r in comp['by_department']),
                         sum(r['measured'] for r in comp['by_country']))


# =========================================================================
#  T4 — did it work? = the answers, recomputed
# =========================================================================
@tagged('post_install', '-at_install')
class TestT4Effect(L6Case):

    def test_t4_effect_equals_an_independent_recount(self):
        Feedback = self.env['pb.feedback.request'].sudo()
        answers = [(5, 5), (4, 3), (2, 4)]
        for useful, recommend in answers:
            Feedback.create({
                'subject_employee_id': self.e_one.id,
                'kind': 'training', 'channel_id': self.simple.id,
                'company_id': self.company.id,
            }).write({
                'state': 'submitted', 'submitted_at': fields.Datetime.now(),
                'answers_json': json.dumps({
                    'useful': {'value': useful},
                    'recommend': {'value': recommend}}),
            })
        effect = self._numbers()['effect']
        self.assertEqual(effect['ratings_count'], 3)
        self.assertEqual(effect['ratings_avg'], round((5 + 4 + 2) / 3.0, 1))
        self.assertEqual(effect['recommend_pct'],
                         int(round(100.0 * 2 / 3)), 'recommend = 4 or 5')
        self.assertEqual(effect['by_course_top'][0]['count'], 3)
        self.assertFalse(effect['empty'])

    def test_t4_empty_sentence(self):
        data = self._numbers()
        self.assertTrue(data['effect']['empty'])
        self.assertEqual(data['effect']['ratings_count'], 0)


# =========================================================================
#  T5 — the flat Assignments sheet
# =========================================================================
@tagged('post_install', '-at_install')
class TestT5Assignments(L6Case):

    def test_t5_one_row_per_assignment_columns_in_order(self):
        self._finished(self.u_one, self.e_one)
        self.env['pb.training.assignment'].sudo().create({
            'channel_id': self.room.id, 'employee_id': self.e_two.id,
            'company_id': self.company.id, 'reason': 'compliance'})
        cols = self.Analytics.with_user(self.tmanager)._assignment_columns()
        self.assertEqual(cols, [
            'Employee', 'Department', 'Job', 'Site', 'Country',
            'Safety group', 'Course', 'Kind', 'Why', 'Path', 'Requirement',
            'Assigned on', 'Due', 'Finished on', 'State', 'Score', 'Hours',
            'Certificate valid until', 'Version'])
        res = self.Analytics.with_user(self.tmanager).export_xlsx(
            fields.Date.to_string(self.today - timedelta(days=30)),
            fields.Date.to_string(self.today), {})
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(
            res['file_b64'])))
        for name in ('Hours', 'Cost', 'Compliance', 'Effect',
                     'Assignments'):
            self.assertIn(name, wb.sheetnames, name)
        ws = wb['Assignments']
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(list(rows[0]), cols)
        body = rows[1:]
        self.assertEqual(len(body), 2, 'one row per assignment')
        by_who = {r[0]: r for r in body}
        one = by_who['DEMO Bùi Hữu Dũng']
        self.assertEqual(one[1], 'DEMO L6 Plant')
        self.assertEqual(one[6], 'DEMO Walking the rows')
        self.assertEqual(one[7], 'Online')
        self.assertEqual(one[14], 'Finished')
        self.assertEqual(one[16], 3)
        two = by_who['DEMO Tran Thi Hoa']
        self.assertEqual(two[7], 'In a room or on site')
        self.assertEqual(two[8], 'Compliance')
        hours = list(wb['Hours'].iter_rows(values_only=True))
        self.assertEqual(hours[1][0], 'DEMO Bùi Hữu Dũng')


# =========================================================================
#  T6 — the Overview headline, the calm state, the chips
# =========================================================================
@tagged('post_install', '-at_install')
class TestT6Overview(L6Case):

    def test_t6_calm_when_nothing_waits(self):
        data = self.Board.with_user(self.trainer).overview()
        self.assertTrue(data['calm'])
        self.assertEqual(data['waiting'], [])
        self.assertEqual(data['headline'], 'Nobody is on a course yet.')
        self.assertTrue(data['queue'], 'every chip stays, quiet at zero')
        self.assertTrue(all(not q['count'] for q in data['queue']))

    def test_t6_all_clear_names_the_people_in_date(self):
        self._finished(self.u_one, self.e_one)
        data = self.Board.with_user(self.trainer).overview()
        self.assertTrue(data['calm'])
        self.assertEqual(data['headline'], 'All clear. 1 person is in date.')

    def test_t6_priority_and_at_most_three(self):
        late = self.env['pb.training.assignment'].sudo()
        for emp in (self.e_one, self.e_two):
            late.create({'channel_id': self.simple.id, 'employee_id': emp.id,
                         'company_id': self.company.id,
                         'due_date': self.today - timedelta(days=3)})
        self.env['pb.training.requirement'].create({
            'name': 'DEMO L6 rule', 'company_id': self.company.id,
            'kind': 'always', 'target': 'course',
            'channel_id': self.room.id, 'auto_assign': False,
            'department_ids': [(6, 0, self.office.ids)]})
        sess = self._session(days=-2)
        sess.write({'state': 'held'})
        data = self.Board.with_user(self.trainer).overview()
        self.assertFalse(data['calm'])
        self.assertEqual(data['concerns'][:3], ['late', 'required', 'sessions'])
        text = data['headline']
        self.assertTrue(text.startswith('2 people are on 1 course.'), text)
        self.assertLess(text.index('2 are past their date.'),
                        text.index('missing required training'))
        self.assertLess(text.index('missing required training'),
                        text.index('waiting to be closed'))
        self.assertLessEqual(len(re.findall(r'\.\s', text + ' ')), 4,
                             'the situation and at most three concerns')
        keys = [q['key'] for q in data['waiting']]
        self.assertIn('late', keys)
        self.assertIn('required', keys)
        self.assertTrue(all(q['count'] > 0 for q in data['waiting']))
        self.assertIn('requests', data)
        self.assertIn('followups', data)
        self.assertEqual(data['compliance']['people_missing'], 1)


# =========================================================================
#  T7 — the settings
# =========================================================================
@tagged('post_install', '-at_install')
class TestT7Settings(L6Case):

    def test_t7_head_only_writes_managers_read_plain_refused(self):
        self.assertEqual(self.Board.with_user(self.plain).settings_get(),
                         {'allowed': False})
        self.assertEqual(self.Board.with_user(self.trainer).settings_get(),
                         {'allowed': False})
        read = self.Board.with_user(self.tmanager).settings_get()
        self.assertTrue(read['allowed'])
        self.assertFalse(read['can_write'])
        with self.assertRaises(AccessError):
            self.Board.with_user(self.tmanager).settings_set(
                {'reminders': False})
        self.assertTrue(self.Board.with_user(self.head).settings_get()
                        ['can_write'])

    def test_t7_round_trip_and_words(self):
        board = self.Board.with_user(self.head)
        res = board.settings_set({
            'report_pack': True, 'report_pack_period': 'weekly',
            'report_pack_email': 'a@example.com; b@example.com\nA@example.com',
            'manager_days': 4, 'hr_days': 9, 'renew_before': 45,
            'enforce_cap': 120, 'feedback_on': False, 'followup_days': 21,
            'day1': True, 'cert_reminders': False})
        v = res['values']
        self.assertEqual(v['report_pack_email'], 'a@example.com, b@example.com')
        self.assertEqual((v['report_pack'], v['report_pack_period']),
                         (True, 'weekly'))
        self.assertEqual((v['manager_days'], v['hr_days'], v['renew_before'],
                          v['enforce_cap'], v['followup_days']),
                         (4, 9, 45, 120, 21))
        self.assertFalse(v['feedback_on'])
        self.assertTrue(v['day1'])
        self.assertIn('the manager after 4 days late; the HR lead after 9.',
                      res['words']['reminders'])
        self.assertIn('a@example.com, b@example.com',
                      res['words']['report_pack'])
        self.assertEqual(self.env['pb.training.pack']._recipients(),
                         ['a@example.com', 'b@example.com'])

    def test_t7_bad_address_and_bad_number_change_nothing(self):
        board = self.Board.with_user(self.head)
        board.settings_set({'report_pack_email': 'ok@example.com'})
        with self.assertRaises(UserError):
            board.settings_set({'report_pack_email': 'ok@example.com, nope',
                                'manager_days': 7})
        v = board.settings_get()['values']
        self.assertEqual(v['report_pack_email'], 'ok@example.com')
        self.assertNotEqual(v['manager_days'], 7, 'nothing half-written')
        with self.assertRaises(UserError):
            board.settings_set({'hr_days': 500})
        many = ', '.join('p%s@example.com' % i for i in range(25))
        with self.assertRaises(UserError):
            board.settings_set({'report_pack_email': many})

    def test_t7_the_daily_job_honours_the_switch(self):
        Auto = self.env['pb.training.automation']
        self.Board.with_user(self.head).settings_set({'reminders': False})
        counts = Auto._cron_daily()
        self.assertNotIn('employee', [k for k, v in counts.items()
                                      if k == 'employee' and v])
        self.assertEqual(counts.get('renewals', 0), 0)
        self.assertEqual(counts.get('employee', 0), 0)
        self.Board.with_user(self.head).settings_set({'manager_days': 6})
        from odoo.addons.pb_training.models.training_common import (
            P_MANAGER_DAYS, number)
        self.assertEqual(number(self.env, P_MANAGER_DAYS, 2), 6)


# =========================================================================
#  T8 — the pack: new lines, the test send, the cap
# =========================================================================
@tagged('post_install', '-at_install')
class TestT8Pack(L6Case):

    def test_t8_body_carries_the_new_lines(self):
        self._finished(self.u_one, self.e_one)
        Pack = self.env['pb.training.pack'].with_user(self.tmanager) \
            .with_context(pb_numbers_full=True)
        board = self.Analytics.with_user(self.tmanager).with_context(
            pb_numbers_full=True).get_numbers(
            fields.Date.to_string(self.today - timedelta(days=30)),
            fields.Date.to_string(self.today), {})
        html = str(Pack._body(board, 'DEMO period'))
        self.assertIn('Most hours of training', html)
        self.assertIn('DEMO Bùi Hữu Dũng', html)
        self.assertIn('3 hours of training, by 1 person.', html)
        self.assertIn('Open Learning', html)
        prev = Pack.preview()
        self.assertIn('Training', prev['subject'])
        self.assertIn('.xlsx', prev['attachment'])
        with self.assertRaises(AccessError):
            self.env['pb.training.pack'].with_user(self.plain).preview()

    def test_t8_test_send_goes_to_the_presser_only(self):
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_training.report_pack_email', 'boss@example.com, hr@example.com')
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_training.report_pack', '1')
        before = self.env['mail.mail'].sudo().search([]).ids
        stamp = self.env['ir.config_parameter'].sudo().get_param(
            'pb_training.report_pack_last')
        res = self.env['pb.training.pack'].with_user(self.tmanager) \
            .send_test()
        new = self.env['mail.mail'].sudo().search(
            [('id', 'not in', before)])
        self.assertEqual(len(new), 1)
        self.assertEqual(new.email_to, 'demo.l6.tmanager@example.com')
        self.assertTrue(new.attachment_ids, 'the full spreadsheet rides it')
        self.assertEqual(res['to'], 'demo.l6.tmanager@example.com')
        self.assertEqual(self.env['ir.config_parameter'].sudo().get_param(
            'pb_training.report_pack_last'), stamp,
            'a test never uses up the period')
        with self.assertRaises(AccessError):
            self.env['pb.training.pack'].with_user(self.plain).send_test()

    def test_t8_cap(self):
        from odoo.addons.pb_training.models.pack import MAIL_CAP
        many = ', '.join('p%s@example.com' % i for i in range(30))
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_training.report_pack_email', many)
        self.assertEqual(len(self.env['pb.training.pack']._recipients()),
                         MAIL_CAP)


# =========================================================================
#  T9 — Home › TEAM: no amounts, chips counted
# =========================================================================
@tagged('post_install', '-at_install')
class TestT9Team(L6Case):

    def _doc(self, emp, days):
        att = self.env['ir.attachment'].sudo().create({
            'name': 'DEMO.pdf', 'raw': base64.b64decode(_PDF),
            'mimetype': 'application/pdf',
            'res_model': 'hr.employee', 'res_id': emp.id})
        cat = self.env['pb.employee.document.category'].sudo().search(
            [], limit=1)
        return self.env['pb.employee.document'].sudo().create({
            'employee_id': emp.id, 'category_id': cat.id,
            'name': 'DEMO L6 First aid', 'attachment_id': att.id,
            'issue_date': self.today - timedelta(days=300),
            'expiry_date': self.today + timedelta(days=days),
            'training_source': 'upload', 'company_id': self.company.id})

    def test_t9_expiring_and_waiting_on_me(self):
        self._doc(self.e_one, 10)
        self._doc(self.e_two, 90)
        from odoo.addons.pb_training.models.request_approval import seed_all
        seed_all(self.env)
        free = self.env['pb.training.outside.course'].create({
            'name': 'DEMO L6 Free webinar', 'cost': 0.0, 'kind': 'online',
            'company_id': self.company.id})
        self.env['pb.my.training'].with_user(self.u_one).ask_course({
            'kind': 'outside', 'outside_course_id': free.id,
            'why': 'DEMO it helps'})
        team = self.env['pb.my.training'].with_user(self.boss_user).team()
        self.assertEqual(team['expiring_certs']['count'], 1)
        self.assertEqual(team['expiring_certs']['names'],
                         ['DEMO Bùi Hữu Dũng'])
        self.assertEqual(team['waiting_on_me']['count'], 1,
                         'the request waits on its manager')
        other = self.env['pb.my.training'].with_user(self.u_two).team()
        self.assertEqual(other['waiting_on_me']['count'], 0)
        if 'pb.home.team' not in self.env:
            return
        home = self.env['pb.home.team'].with_user(self.boss_user) \
            .with_company(self.company).get_team()
        training = home['training']
        self.assertEqual((training['expiring'], training['waiting']), (1, 1))
        forbidden = ('amount', 'salary', 'wage', 'budget', 'cost', 'money',
                     'currency')

        def keys(value):
            if isinstance(value, dict):
                for k, v in value.items():
                    yield str(k)
                    yield from keys(v)
            elif isinstance(value, (list, tuple)):
                for v in value:
                    yield from keys(v)
        bad = [k for k in keys(training)
               if any(w in k.lower() for w in forbidden)]
        self.assertEqual(bad, [], 'no amount on Home › TEAM')


# =========================================================================
#  T10 — white label, emoji, Vietnamese; every lens tells "failed" apart
# =========================================================================
@tagged('post_install', '-at_install')
class TestT10Words(L6Case):

    EMOJI = re.compile('[\U0001F300-\U0001FAFF☀-➿]')

    def test_t10_no_vendor_word_no_emoji(self):
        for parts in NEW_FILES:
            text = _src(*parts)
            for line in text.splitlines():
                if 'from odoo' in line or 'odoo.' in line or '@odoo' in line \
                        or 'odoo-module' in line:
                    continue
                self.assertNotIn('Odoo', line, parts)
            self.assertFalse(self.EMOJI.search(text), parts)

    def test_t10_vietnamese_present(self):
        po = _src('i18n', 'vi_VN.po')
        for word in ('Report pack', 'That did not load',
                     'Hours of training', 'Learning settings',
                     'All clear. %s people are in date.',
                     'Send me a test'):
            self.assertIn('msgid "%s"' % word, po, word)

    def test_t10_every_lens_tells_failed_from_refused(self):
        lenses = ('training_board', 'training_numbers', 'learning_overview',
                  'my_learning_lens', 'training_tests_lens',
                  'certificates_lens', 'required_lens', 'paths_lens',
                  'sessions_lens', 'outside_lens', 'requests_lens',
                  'skills_lens', 'learning_settings', 'toolbox_drawer')
        for name in lenses:
            js = _src('static', 'src', 'js', '%s.js' % name)
            xml = _src('static', 'src', 'xml', '%s.xml' % name)
            self.assertIn('failed', js, name)
            self.assertNotRegex(
                js, r'catch \(e\) \{[^}]*state\.allowed = false', name)
            self.assertIn('state.failed', xml, name)
            block = xml[xml.index('state.failed'):][:900]
            self.assertTrue('retry' in block or 'load()' in block, name)

    def test_t10_help_lists_the_eleven_tabs(self):
        import os
        here = os.path.dirname(os.path.dirname(os.path.dirname(
            os.path.abspath(__file__))))
        path = os.path.join(here, 'pb_learn', 'static', 'content',
                            'learn_content.json')
        if not os.path.exists(path):
            self.skipTest('pb_learn is not beside pb_training here')
        with open(path, encoding='utf-8') as fh:
            content = json.load(fh)
        screen = next(s for s in content['screens']
                      if s.get('key') == 'hub_learning')
        self.assertEqual(set(screen['hub']['tabs']), {
            'overview', 'mine', 'courses', 'paths', 'assigned', 'required',
            'tests', 'certificates', 'claims', 'skills', 'numbers'})

    def test_t10_settings_category_is_learning(self):
        js = _src('static', 'src', 'js', 'training_palette.js')
        self.assertIn('add("learning"', js)
        self.assertNotIn('add("required_training"', js)
        for card in ('report_pack', 'reminders', 'feedback', 'country_packs',
                     'day_one'):
            self.assertIn('id: "%s"' % card, js)

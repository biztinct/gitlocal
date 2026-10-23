# -*- coding: utf-8 -*-
"""LEARNING L1 — certificates and expiry. Handover tests T1–T13.

The register is the vault document with training facts; these are the rails
under it that nothing at runtime would report: a renewal that silently files
nothing, a back-fill that invents an expiry date, a reminder that goes twice,
a register count that disagrees with the Overview, and somebody reading or
printing a record that is not theirs.

Fixtures are named "DEMO" (ledger rule 9) and roll back with the transaction.
The sentences are asserted in English, so every class pins `lang='en_US'`
(ledger LN14); the Vietnamese words are checked against the `.po` itself.
"""

import base64
import io
from datetime import timedelta
from unittest.mock import patch

from dateutil.relativedelta import relativedelta

from odoo import fields
from odoo.exceptions import AccessError, UserError
from odoo.tests import TransactionCase, tagged

from .test_assignments import AssignmentCase
from .test_training import _PDF, _src

_PDF_BYTES = base64.b64decode(_PDF)


class CertCase(AssignmentCase):
    """E1/E2's course and people, plus a trainer and helpers to pass the test
    and to file documents by hand."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.env = cls.env(context=dict(cls.env.context, lang='en_US'))
        cls.Doc = cls.env['pb.employee.document'].sudo()
        cls.Cert = cls.env['pb.training.certificate']
        cls.cat = cls.Cert._category()
        cls.lic = cls.Cert._licence_category()
        Users = cls.env['res.users'].with_context(no_reset_password=True)
        cls.trainer = Users.create({
            'name': 'DEMO Trainer L1',
            'login': 'demo.trainer.l1@example.com',
            'email': 'demo.trainer.l1@example.com',
            'group_ids': [(6, 0, [
                cls.env.ref('base.group_user').id,
                cls.env.ref('pb_training.group_training_user').id])],
        })
        cls.plain = Users.create({
            'name': 'DEMO Plain L1',
            'login': 'demo.plain.l1@example.com',
            'email': 'demo.plain.l1@example.com',
            'group_ids': [(6, 0, [cls.env.ref('base.group_user').id])],
        })

    # ------------------------------------------------------------ helpers
    def _attachment(self, emp):
        return self.env['ir.attachment'].sudo().create({
            'name': 'DEMO certificate.pdf', 'raw': _PDF_BYTES,
            'mimetype': 'application/pdf',
            'res_model': 'hr.employee', 'res_id': emp.id})

    def _doc(self, emp, name='DEMO Forklift licence', days=None,
             source='statutory', category=None, **extra):
        vals = {
            'employee_id': emp.id,
            'category_id': (category or self.cat).id,
            'name': name,
            'attachment_id': self._attachment(emp).id,
            'issue_date': fields.Date.today() - timedelta(days=30),
            'expiry_date': fields.Date.today() + timedelta(days=days)
            if days is not None else False,
            'training_source': source,
            'company_id': emp.company_id.id,
        }
        vals.update(extra)
        return self.Doc.create(vals)

    def _answer_right(self, answer):
        for question in self.survey.question_and_page_ids.filtered(
                lambda q: not q.is_page):
            right = question.suggested_answer_ids.filtered('is_correct')[:1]
            self.env['survey.user_input.line'].sudo().create({
                'user_input_id': answer.id, 'question_id': question.id,
                'answer_type': 'suggestion',
                'suggested_answer_id': right.id})
        answer._mark_done()
        self.env.flush_all()
        answer.invalidate_recordset()
        return answer

    def _pass(self, user):
        """Finish the lessons and pass the test through the real doors."""
        self.course._action_add_members(user.partner_id)
        self._finish_every_lesson(user)
        self.env['pb.my.training'].with_user(user).start_test(self.course.id)
        answer = self.env['survey.user_input'].sudo().search(
            [('survey_id', '=', self.survey.id),
             ('partner_id', '=', user.partner_id.id),
             ('state', '!=', 'done')], limit=1)
        self.assertTrue(answer, 'the test did not open')
        return self._answer_right(answer)

    def _pass_again(self, user):
        """A second pass (a renewal): a fresh attempt, answered right."""
        answer = self.survey.sudo()._create_answer(
            partner=user.partner_id, check_attempts=False,
            slide_id=self.test_slide.id)
        return self._answer_right(answer)

    def _test_docs(self, emp):
        return self.Doc.with_context(active_test=False).search(
            [('employee_id', '=', emp.id), ('training_source', '=', 'test')],
            order='id')


# =========================================================================
#  T1/T2 — validity on a passed test, and renewals
# =========================================================================
@tagged('post_install', '-at_install')
class TestT1T2Validity(CertCase):

    def test_t1_a_pass_with_validity_files_a_dated_certificate(self):
        self.survey.sudo().certification_validity_months = 12
        attempt = self._pass(self.u_one)
        self.assertTrue(attempt.scoring_success)
        docs = self._test_docs(self.e_one)
        self.assertEqual(len(docs), 1)
        doc = docs
        self.assertEqual(doc.attempt_id, attempt)
        self.assertEqual(doc.survey_id, self.survey)
        self.assertEqual(doc.channel_id, self.course)
        issued = fields.Date.to_date(attempt.end_datetime)
        self.assertEqual(doc.issue_date, issued)
        self.assertEqual(doc.expiry_date, issued + relativedelta(months=12))
        self.assertIn('(%s)' % issued.year, doc.name)
        self.assertTrue(doc.issuer)

    def test_t1_validity_zero_means_no_expiry(self):
        self.survey.sudo().certification_validity_months = 0
        self._pass(self.u_one)
        doc = self._test_docs(self.e_one)
        self.assertEqual(len(doc), 1)
        self.assertFalse(doc.expiry_date)
        self.assertEqual(doc.training_state, 'none')

    def test_t2_passing_again_files_a_second_and_renews_the_first(self):
        self.survey.sudo().certification_validity_months = 12
        self._pass(self.u_one)
        first = self._test_docs(self.e_one)
        year_ago = fields.Date.today() - relativedelta(years=1)
        first.write({'issue_date': year_ago,
                     'expiry_date': year_ago + relativedelta(months=12)})
        second_attempt = self._pass_again(self.u_one)
        docs = self._test_docs(self.e_one)
        self.assertEqual(len(docs), 2, 'the renewal filed nothing')
        second = docs - first
        self.assertEqual(second.attempt_id, second_attempt)
        self.assertEqual(first.renewed_by_id, second)
        self.assertEqual(first.training_state, 'renewed')
        self.assertFalse(second.renewed_by_id)
        # The same pass noticed again files nothing new.
        self.Cert.file_course_certificate(self.course, self.u_one.partner_id)
        self.assertEqual(len(self._test_docs(self.e_one)), 2)
        # The résumé line (hr_skills_survey) still ran, without error.
        self.assertTrue(self.env['hr.resume.line'].sudo().search_count(
            [('employee_id', '=', self.e_one.id),
             ('survey_id', '=', self.survey.id)]))

    def test_t2_a_row_filed_before_l1_is_linked_not_duplicated(self):
        attempt = self._pass(self.u_one)
        doc = self._test_docs(self.e_one)
        # Make it look like a pre-L1 row: old name, no links.
        doc.write({'training_source': False, 'attempt_id': False,
                   'name': self.Cert._document_name(self.course.name)})
        again = self.Cert.file_course_certificate(self.course,
                                                  self.u_one.partner_id)
        self.assertEqual(again, doc)
        self.assertEqual(doc.attempt_id, attempt)
        self.assertEqual(doc.training_source, 'test')


# =========================================================================
#  T3 — the back-fill (the migration's job)
# =========================================================================
@tagged('post_install', '-at_install')
class TestT3Backfill(CertCase):

    def test_t3_duplicate_name_uses_the_employees_pass(self):
        self.course.name = 'DEMO Duplicate legacy course'
        duplicate = self.env['slide.channel'].create({
            'name': self.course.name, 'channel_type': 'training'})
        self._pass(self.u_one)
        doc = self._test_docs(self.e_one)
        doc.write({'training_source': False, 'attempt_id': False,
                   'survey_id': False, 'channel_id': False,
                   'name': self.Cert._document_name(self.course.name)})
        self.Cert._backfill_test(doc, self.course.name,
                                {'test': 0, 'linked': 0})
        self.assertEqual(doc.channel_id, self.course)
        self.assertNotEqual(doc.channel_id, duplicate)

    def test_t3_ambiguous_name_does_not_invent_a_course_or_expiry(self):
        name = 'DEMO Ambiguous legacy course'
        self.env['slide.channel'].create([
            {'name': name, 'channel_type': 'training'},
            {'name': name, 'channel_type': 'training'}])
        doc = self._doc(self.e_two,
                        name=self.Cert._document_name(name), source=False)
        self.Cert._backfill_test(doc, name, {'test': 0, 'linked': 0})
        self.assertEqual(doc.training_source, 'test')
        self.assertFalse(doc.channel_id)
        self.assertFalse(doc.survey_id)
        self.assertFalse(doc.expiry_date)

    def test_t3_old_rows_get_their_source_and_only_honest_dates(self):
        self.survey.sudo().certification_validity_months = 12
        attempt = self._pass(self.u_one)
        test_doc = self._test_docs(self.e_one)
        test_doc.write({'training_source': False, 'attempt_id': False,
                        'survey_id': False, 'channel_id': False,
                        'expiry_date': False, 'issuer': False,
                        'name': 'DEMO Agronomist basics — certificate'})
        claim_doc = self._doc(self.e_two,
                              name='DEMO Excel — certificate (DEMO School)',
                              source=False)
        other = self._doc(self.e_three, name='DEMO Unrelated contract',
                          source=False)
        # A course whose test carries no validity: no date may be invented.
        plain_course = self.simple
        nodate = self._doc(self.e_three,
                           name='%s — certificate' % plain_course.name,
                           source=False)
        counts = self.Cert._backfill()
        self.assertEqual(test_doc.training_source, 'test')
        self.assertEqual(test_doc.channel_id, self.course)
        self.assertEqual(test_doc.survey_id, self.survey)
        self.assertEqual(test_doc.attempt_id, attempt)
        self.assertEqual(test_doc.expiry_date,
                         test_doc.issue_date + relativedelta(months=12))
        self.assertEqual(claim_doc.training_source, 'claim')
        self.assertEqual(claim_doc.issuer, 'DEMO School')
        self.assertFalse(claim_doc.expiry_date)
        self.assertFalse(other.training_source, 'an ordinary file is left')
        self.assertEqual(nodate.training_source, 'test')
        self.assertFalse(nodate.expiry_date, 'no validity, no date')
        self.assertGreaterEqual(counts['test'], 2)
        self.assertGreaterEqual(counts['claim'], 1)
        # Idempotent: a second run changes nothing.
        before = (test_doc.expiry_date, claim_doc.training_source)
        again = self.Cert._backfill()
        self.assertEqual((test_doc.expiry_date, claim_doc.training_source),
                         before)
        self.assertEqual(again['test'] + again['claim'], 0)


# =========================================================================
#  T4 — the state, and searching on it
# =========================================================================
@tagged('post_install', '-at_install')
class TestT4State(CertCase):

    def test_t4_every_state_and_its_words(self):
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_training.renew_before_days', '60')
        valid = self._doc(self.e_one, 'DEMO Valid', days=200)
        soon = self._doc(self.e_one, 'DEMO Soon', days=59)
        gone = self._doc(self.e_one, 'DEMO Gone', days=-3)
        none = self._doc(self.e_one, 'DEMO None')
        old = self._doc(self.e_one, 'DEMO Old', days=10)
        old.renewed_by_id = valid
        self.assertEqual([d.training_state for d in (valid, soon, gone, none,
                                                     old)],
                         ['valid', 'soon', 'expired', 'none', 'renewed'])
        self.assertTrue(valid.valid_words.startswith('Valid until'))
        self.assertEqual(soon.valid_words, 'Expires in 59 days')
        self.assertEqual(gone.valid_words, 'Expired 3 days ago')
        self.assertEqual(none.valid_words, 'Never expires')
        self.assertEqual(old.valid_words, 'Renewed')

    def test_t4_the_state_searches(self):
        mine = [('employee_id', '=', self.e_one.id)]
        valid = self._doc(self.e_one, 'DEMO Valid', days=200)
        soon = self._doc(self.e_one, 'DEMO Soon', days=5)
        gone = self._doc(self.e_one, 'DEMO Gone', days=-1)
        none = self._doc(self.e_one, 'DEMO None')
        old = self._doc(self.e_one, 'DEMO Old', days=10, renewed_by_id=valid.id)
        for state, want in (('valid', valid), ('soon', soon),
                            ('expired', gone), ('none', none),
                            ('renewed', old)):
            got = self.Doc.search(mine + [('training_state', '=', state)])
            self.assertEqual(got, want, state)
        both = self.Doc.search(mine + [('training_state', 'in',
                                        ['soon', 'expired'])])
        self.assertEqual(both, soon | gone)
        rest = self.Doc.search(mine + [('training_state', '!=', 'renewed')])
        self.assertNotIn(old, rest)
        self.assertEqual(len(rest), 4)


# =========================================================================
#  T5/T6 — uploads, the employee's and HR's
# =========================================================================
@tagged('post_install', '-at_install')
class TestT5T6Uploads(CertCase):

    def _values(self, **kw):
        vals = {'name': 'DEMO First aid', 'kind': 'CERT',
                'issuer': 'DEMO Red Cross', 'reference': 'FA-1',
                'issue_date': fields.Date.to_string(fields.Date.today()),
                'expiry_date': fields.Date.to_string(
                    fields.Date.today() + timedelta(days=75))}
        vals.update(kw)
        return vals

    def test_t5_the_employee_upload_checks_the_paper(self):
        me = self.env['pb.my.training'].with_user(self.u_one)
        with self.assertRaises(UserError) as big:
            me.add_certificate(self._values(),
                               ('big.pdf', b'%PDF-1.4' + b'0' * (11 * 1024 * 1024)))
        self.assertIn('bigger than 10 MB', str(big.exception))
        with self.assertRaises(UserError) as exe:
            me.add_certificate(self._values(),
                               ('run.exe', b'MZ\x90\x00\x03' + b'\x00' * 200))
        self.assertIn('Only PDF, JPG or PNG files.', str(exe.exception))
        res = me.add_certificate(self._values(), ('fa.pdf', _PDF_BYTES))
        doc = self.Doc.browse(res['id'])
        self.assertEqual(doc.training_source, 'upload')
        self.assertFalse(doc.verified)
        self.assertEqual(doc.employee_id, self.e_one)
        self.assertEqual(doc.attachment_id.res_model, 'hr.employee')
        self.assertEqual(doc.attachment_id.res_id, self.e_one.id)
        self.assertEqual(doc.issuer, 'DEMO Red Cross')

    def test_t5_a_licence_needs_a_date_and_a_login_needs_an_employee(self):
        me = self.env['pb.my.training'].with_user(self.u_one)
        with self.assertRaises(UserError) as nodate:
            me.add_certificate(self._values(kind='LIC', expiry_date=False),
                               ('lic.pdf', _PDF_BYTES))
        self.assertIn('A licence needs a valid-until date.',
                      str(nodate.exception))
        with self.assertRaises(AccessError) as caught:
            self.env['pb.my.training'].with_user(self.plain).add_certificate(
                self._values(), ('fa.pdf', _PDF_BYTES))
        self.assertIn('Only employees can add certificates.',
                      str(caught.exception))

    def test_t6_hr_adds_one_and_it_is_verified_the_vaults_way(self):
        payload = dict(self._values(name='DEMO Forklift', kind='LIC'),
                       employee_id=self.e_two.id, source='statutory',
                       file_name='lic.pdf',
                       file_b64=base64.b64encode(_PDF_BYTES).decode())
        res = self.env['pb.training'].with_user(self.trainer).act(
            'add_certificate', payload)
        doc = self.Doc.browse(res['id'])
        self.assertEqual(doc.training_source, 'statutory')
        self.assertEqual(doc.category_id, self.lic)
        self.assertTrue(doc.verified)
        self.assertEqual(doc.verified_by, self.trainer)
        self.assertTrue(doc.verified_at)

    def test_t6_a_plain_user_cannot_use_the_verb(self):
        payload = dict(self._values(), employee_id=self.e_two.id,
                       file_name='fa.pdf',
                       file_b64=base64.b64encode(_PDF_BYTES).decode())
        with self.assertRaises(AccessError):
            self.env['pb.training'].with_user(self.plain).act(
                'add_certificate', payload)
        self.assertEqual(
            self.env['pb.training'].with_user(self.plain)
            .certificates_board({}), {'allowed': False, 'rows': [],
                                      'tiles': []})

    def test_t6_verify_verb(self):
        doc = self._doc(self.e_two, days=40, source='upload')
        self.assertFalse(doc.verified)
        self.env['pb.training'].with_user(self.trainer).act(
            'verify_certificate', {'document_id': doc.id})
        self.assertTrue(doc.verified)
        self.assertEqual(doc.verified_by, self.trainer)


# =========================================================================
#  T7 — renewal courses, put on by themselves
# =========================================================================
@tagged('post_install', '-at_install')
class TestT7Renewals(CertCase):

    def setUp(self):
        super().setUp()
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_training.renew_before_days', '60')

    def test_t7_inside_the_window_one_renewal_and_only_one(self):
        doc = self._doc(self.e_one, 'DEMO Agronomist cert', days=59,
                        source='test', channel_id=self.course.id)
        made = self.Auto._renewals(fields.Date.today())
        self.assertEqual(made, 1)
        row = doc.renewal_assignment_id
        self.assertTrue(row)
        self.assertEqual(row.reason, 'renewal')
        self.assertEqual(row.due_date, doc.expiry_date)
        self.assertEqual(row.employee_id, self.e_one)
        self.assertEqual(self.Auto._renewals(fields.Date.today()), 0)

    def test_t7_outside_the_window_or_already_expired_nothing(self):
        far = self._doc(self.e_one, 'DEMO Far', days=90, source='test',
                        channel_id=self.course.id)
        gone = self._doc(self.e_two, 'DEMO Gone', days=-5, source='test',
                         channel_id=self.course.id)
        self.Auto._renewals(fields.Date.today())
        self.assertFalse(far.renewal_assignment_id)
        self.assertFalse(gone.renewal_assignment_id)

    def test_t7_an_open_assignment_means_none_and_no_error(self):
        self.Assignment.create({'channel_id': self.course.id,
                                'employee_id': self.e_two.id})
        doc = self._doc(self.e_two, 'DEMO Busy', days=30, source='test',
                        channel_id=self.course.id)
        self.assertEqual(self.Auto._renewals(fields.Date.today()), 0)
        self.assertFalse(doc.renewal_assignment_id)
        with self.assertRaises(UserError):
            self.env['pb.training'].with_user(self.trainer).act(
                'assign_refresher', {'document_id': doc.id})

    def test_t7_an_archived_course_means_none(self):
        self.simple.active = False
        doc = self._doc(self.e_three, 'DEMO Archived', days=30, source='test',
                        channel_id=self.simple.id)
        self.assertEqual(self.Auto._renewals(fields.Date.today()), 0)
        self.assertFalse(doc.renewal_assignment_id)

    def test_t7_the_register_button_assigns_now(self):
        doc = self._doc(self.e_three, 'DEMO By hand', days=100,
                        source='test', channel_id=self.course.id)
        res = self.env['pb.training'].with_user(self.trainer).act(
            'assign_refresher', {'document_id': doc.id})
        self.assertIn('again', res['message'])
        self.assertEqual(doc.renewal_assignment_id.reason, 'renewal')
        self.assertEqual(doc.renewal_assignment_id.assigned_by_id,
                         self.trainer)


# =========================================================================
#  T8 — the certificate ladder
# =========================================================================
@tagged('post_install', '-at_install')
class TestT8Ladder(CertCase):

    def _chase(self, doc, today):
        counts = {'certificates': 0}
        with patch.object(type(self.Auto), '_hr_email',
                          lambda self, company: 'demo.hr.lead@example.com'):
            self.Auto._chase_cert_one(doc, today, 14, counts)
        return counts['certificates']

    def test_t8_every_key_once_at_its_day(self):
        doc = self._doc(self.e_one, 'DEMO Ladder', days=60)
        expiry = doc.expiry_date
        steps = [(60, {'x60'}), (45, set()), (30, {'x30'}), (14, {'xm14'}),
                 (7, {'x7'}), (3, set()), (0, {'x0', 'xh0'})]
        seen = set()
        for left, want in steps:
            today = expiry - timedelta(days=left)
            before = self._mails()
            self._chase(doc, today)
            sent = set(doc._pb_sent()) - seen
            self.assertEqual(sent, want, 'day %s' % left)
            self.assertEqual(self._mails() - before, len(want), left)
            seen |= sent
            # Same morning again: nothing twice.
            self.assertEqual(self._chase(doc, today), 0)
        # Nothing after the day.
        self.assertEqual(self._chase(doc, expiry + timedelta(days=1)), 0)
        self.assertEqual(seen, {'x60', 'x30', 'xm14', 'x7', 'x0', 'xh0'})

    def test_t8_a_late_starter_gets_its_band_not_every_band(self):
        doc = self._doc(self.e_one, 'DEMO Late', days=20)
        self._chase(doc, fields.Date.today())
        self.assertEqual(set(doc._pb_sent()), {'x30'})

    def test_t8_renewed_rows_are_skipped(self):
        new = self._doc(self.e_one, 'DEMO New', days=400)
        old = self._doc(self.e_one, 'DEMO Old', days=7,
                        renewed_by_id=new.id)
        self.assertEqual(self._chase(old, fields.Date.today()), 0)

    def test_t8_the_cap_and_the_switch(self):
        doc = self._doc(self.e_one, 'DEMO Capped', days=7)
        icp = self.env['ir.config_parameter'].sudo()
        icp.set_param('pb_training.reminder_cap', '5')
        counts = {'employee': 5, 'manager': 0, 'hr': 0, 'certificates': 0,
                  'capped': 0}
        self.Auto._chase_certificates(fields.Date.today(), counts)
        self.assertEqual(doc._pb_sent(), {})
        self.assertGreaterEqual(counts['capped'], 1)
        icp.set_param('pb_training.cert_reminders', '0')
        icp.set_param('pb_training.reminder_cap', '400')
        result = self.Auto._cron_daily()
        self.assertEqual(result['certificates'], 0)
        self.assertEqual(doc._pb_sent(), {})


# =========================================================================
#  T9 — the templates
# =========================================================================
@tagged('post_install', '-at_install')
class TestT9Templates(CertCase):

    def test_t9_they_render_and_say_valid_words(self):
        doc = self._doc(self.e_one, 'DEMO Render', days=30)
        for xmlid in ('pb_training.mail_template_certificate_expiring',
                      'pb_training.mail_template_certificate_expiring_manager',
                      'pb_training.mail_template_certificate_expired_hr'):
            template = self.env.ref(xmlid)
            subject = template._render_field('subject', doc.ids)[doc.id]
            body = template._render_field('body_html', doc.ids)[doc.id]
            self.assertIn(doc.name, subject)
            self.assertNotIn('odoo', (subject + str(body)).lower())
            self.assertNotIn('#4F46E5', str(body))
        subject = self.env.ref(
            'pb_training.mail_template_certificate_expiring')._render_field(
                'subject', doc.ids)[doc.id]
        self.assertIn('Expires in 30 days', subject)
        src = _src('data', 'mail_template_certificates.xml')
        self.assertNotIn('Odoo', src)
        self.assertNotIn('#4F46E5', src)
        self.assertIn('#5A4BB0', src)


# =========================================================================
#  T10/T11 — the register, the Overview and the Numbers, in their own company
# =========================================================================
class IsolatedCase(CertCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.co = cls.env['res.company'].create({'name': 'DEMO L1 Cert Co'})
        cls.boss_trainer = cls.env['res.users'].with_context(
            no_reset_password=True).create({
                'name': 'DEMO Head Trainer L1',
                'login': 'demo.headtrainer.l1@example.com',
                'email': 'demo.headtrainer.l1@example.com',
                'company_id': cls.co.id,
                'company_ids': [(6, 0, [cls.co.id])],
                'group_ids': [(6, 0, [
                    cls.env.ref('base.group_user').id,
                    cls.env.ref('pb_training.group_training_manager').id])],
            })
        Dept = cls.env['hr.department']
        cls.dept = Dept.create({'name': 'DEMO Farm', 'company_id': cls.co.id})
        cls.sub = Dept.create({'name': 'DEMO Farm East',
                               'parent_id': cls.dept.id,
                               'company_id': cls.co.id})
        Emp = cls.env['hr.employee']
        cls.p1 = Emp.create({'name': 'DEMO Bùi Hữu Dũng L1',
                             'company_id': cls.co.id,
                             'department_id': cls.sub.id})
        cls.p2 = Emp.create({'name': 'DEMO Trần Thị Hoa L1',
                             'company_id': cls.co.id,
                             'department_id': cls.dept.id})
        cls.p3 = Emp.create({'name': 'DEMO Lê Văn Nam L1',
                             'company_id': cls.co.id})

    def _board(self, **filters):
        return self.env['pb.training'].with_user(
            self.boss_trainer).certificates_board(filters)


@tagged('post_install', '-at_install')
class TestT10Register(IsolatedCase):

    def _seed(self):
        days = [5, 25, 45, 75, 200, -3, -40, None]
        docs = self.Doc.browse()
        for i, d in enumerate(days):
            emp = (self.p1, self.p2, self.p3)[i % 3]
            docs |= self._doc(emp, 'DEMO Cert %s' % i, days=d)
        return docs

    def test_t10_the_tiles_match_an_independent_count(self):
        docs = self._seed()
        today = fields.Date.today()
        live = docs.filtered(lambda d: d.expiry_date)
        want = {
            'expiring_30': len(live.filtered(
                lambda d: today <= d.expiry_date <= today + timedelta(30))),
            'expiring_60': len(live.filtered(
                lambda d: today <= d.expiry_date <= today + timedelta(60))),
            'expiring_90': len(live.filtered(
                lambda d: today <= d.expiry_date <= today + timedelta(90))),
            'expired': len(live.filtered(lambda d: d.expiry_date < today)),
            'valid': len(live.filtered(lambda d: d.expiry_date >= today)),
            'no_expiry': len(docs.filtered(lambda d: not d.expiry_date)),
        }
        data = self._board()
        self.assertTrue(data['allowed'])
        got = {t['key']: t['count'] for t in data['tiles']}
        self.assertEqual(got, want)
        self.assertEqual(want, {'expiring_30': 2, 'expiring_60': 3,
                                'expiring_90': 4, 'expired': 2, 'valid': 5,
                                'no_expiry': 1})
        self.assertEqual(data['total'], 8)
        self.assertEqual(data['headline'],
                         '8 certificates. 2 expire within 30 days, 2 have '
                         'expired.')
        # Problem first: the expired rows head the table.
        self.assertEqual([r['state'] for r in data['rows'][:2]],
                         ['expired', 'expired'])
        for key, n in want.items():
            self.assertEqual(self._board(state=key)['total'], n, key)

    def test_t10_search_ignores_accents_and_departments_include_children(self):
        self._seed()
        found = self._board(q='bui huu')
        self.assertTrue(found['rows'])
        self.assertEqual({r['employee'] for r in found['rows']},
                         {'DEMO Bùi Hữu Dũng L1'})
        by_dept = self._board(department_id=self.dept.id)
        self.assertEqual({r['employee'] for r in by_dept['rows']},
                         {'DEMO Bùi Hữu Dũng L1', 'DEMO Trần Thị Hoa L1'})

    def test_t10_paging(self):
        self._seed()
        with patch('odoo.addons.pb_training.models.pb_training_certificates.'
                   'PAGE_SIZE', 3):
            first = self._board()
            third = self._board(page=3)
        self.assertEqual(first['pages'], 3)
        self.assertEqual(len(first['rows']), 3)
        self.assertEqual(len(third['rows']), 2)
        self.assertFalse({r['id'] for r in first['rows']}
                         & {r['id'] for r in third['rows']})

    def test_t10_another_companys_certificates_never_show(self):
        self._doc(self.e_one, 'DEMO Elsewhere', days=10)
        self.assertEqual(self._board()['total'], 0)


@tagged('post_install', '-at_install')
class TestT11OverviewNumbers(IsolatedCase):

    def test_t11_overview_block_and_headline(self):
        self._doc(self.p1, 'DEMO A', days=10)
        self._doc(self.p2, 'DEMO B', days=-2)
        self._doc(self.p3, 'DEMO C')
        data = self.env['pb.training'].with_user(self.boss_trainer).overview()
        self.assertEqual(data['certificates'], {
            'expiring_30': 1, 'expired': 1, 'valid': 1, 'no_expiry': 1,
            'total': 3})
        self.assertIn('1 certificate expires within 30 days.',
                      data['headline'])
        queue = {q['key']: q for q in data['queue']}
        self.assertEqual(queue['certs']['count'], 1)
        self.assertEqual(queue['certs']['lens'], 'certificates')
        self.assertEqual(queue['certs']['filter'], 'expiring_30')

    def test_t11_zero_certificates_says_nothing(self):
        data = self.env['pb.training'].with_user(self.boss_trainer).overview()
        self.assertNotIn('certificate', data['headline'])

    def test_t11_numbers_tiles_breakdown_and_sheet(self):
        self._doc(self.p1, 'DEMO Forklift', days=10)
        self._doc(self.p2, 'DEMO Forklift', days=-2)
        self._doc(self.p3, 'DEMO First aid', days=300)
        # A cohort to measure, so the tiles are drawn.
        user = self.env['res.users'].with_context(
            no_reset_password=True).create({
                'name': 'DEMO Learner L1 Co', 'login': 'demo.lco.l1@example.com',
                'company_id': self.co.id, 'company_ids': [(6, 0, [self.co.id])],
                'group_ids': [(6, 0, [self.env.ref('base.group_user').id])]})
        self.p3.user_id = user
        self.Assignment.create({'channel_id': self.simple.id,
                                'employee_id': self.p3.id,
                                'company_id': self.co.id})
        Analytics = self.env['pb.training.analytics'].with_user(
            self.boss_trainer)
        data = Analytics.get_numbers()
        tiles = {t['key']: t for t in data['tiles']}
        for key in ('expiring_30', 'expiring_90', 'expired', 'renewals_open'):
            self.assertIn(key, tiles)
        self.assertEqual(tiles['expiring_30']['value'], 1)
        self.assertEqual(tiles['expired']['value'], 1)
        by = {r['name']: r for r in data['by_certificate']}
        self.assertEqual(by['DEMO Forklift']['holders'], 2)
        self.assertEqual(by['DEMO Forklift']['expiring'], 1)
        self.assertEqual(by['DEMO Forklift']['expired'], 1)
        self.assertEqual(sum(r['holders'] for r in data['by_certificate']), 3)
        self.assertEqual(data['cert_expiring'], 1)
        export = Analytics.export_xlsx()
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(
            export['file_b64'])))
        self.assertIn('Certificates', wb.sheetnames)
        sheet = wb['Certificates']
        self.assertEqual(sheet.cell(row=1, column=1).value, 'Certificate')


# =========================================================================
#  T12/T13 — the learning record, the wallet and the home card
# =========================================================================
@tagged('post_install', '-at_install')
class TestT12T13Learner(CertCase):

    def test_t12_the_record_renders_for_zero_one_and_many(self):
        me = self.env['pb.my.training'].with_user(self.u_one)
        pdf, name = me.record_pdf()
        self.assertTrue(pdf)
        self.assertIn('Learning record', name)
        self._doc(self.e_one, 'DEMO One', days=40)
        self.assertTrue(me.record_pdf()[0])
        for i in range(4):
            self._doc(self.e_one, 'DEMO Many %s' % i, days=10 * i)
        body = me.record_pdf()[0]
        self.assertTrue(body)
        values = self.env['report.pb_training.report_learning_record'] \
            ._get_report_values(self.e_one.ids)
        record = values['records'][self.e_one.id]
        self.assertEqual(len(record['certificates']), 5)
        self.assertTrue(record['summary'])

    def test_t12_another_persons_record_is_refused_unless_trainer(self):
        with self.assertRaises(AccessError):
            self.env['pb.my.training'].with_user(self.u_two).record_pdf(
                self.e_one.id)
        pdf, _name = self.env['pb.my.training'].with_user(
            self.trainer).record_pdf(self.e_one.id)
        self.assertTrue(pdf)

    def test_t13_the_wallet_lists_only_my_own(self):
        mine = self._doc(self.e_one, 'DEMO Mine', days=20)
        theirs = self._doc(self.e_two, 'DEMO Theirs', days=20)
        wallet = self.env['pb.my.training'].with_user(self.u_one) \
            .my_certificates()
        ids = {r['id'] for r in wallet['rows']}
        self.assertIn(mine.id, ids)
        self.assertNotIn(theirs.id, ids)
        self.assertEqual(wallet['next']['id'], mine.id)
        self.assertEqual(self.env['pb.my.training'].with_user(self.u_one)
                         .certificates_count(), 1)
        self.assertTrue(wallet['rows'][0]['download_url'].startswith(
            '/web/content/ir.attachment/'))

    def test_t13_the_home_card_and_the_routes_exist(self):
        view = self.env.ref(
            'pb_training.portal_my_home_training_certificates')
        self.assertIn('training_certificates_count', view.arch_db)
        self.assertTrue(self.env.ref(
            'pb_training.portal_training_certificates'))
        src = _src('controllers', 'portal.py')
        for route in ("'/my/training/certificates'",
                      "'/my/training/certificates/new'",
                      "'/my/training/record.pdf'"):
            self.assertIn(route, src)

    def test_t13_a_course_tile_carries_its_certificate(self):
        self.survey.sudo().certification_validity_months = 12
        self._pass(self.u_one)
        home = self.env['pb.my.training'].with_user(self.u_one).home()
        tile = [c for c in home['courses'] if c['id'] == self.course.id][0]
        self.assertTrue(tile['certificate'])
        self.assertTrue(tile['certificate']['valid_words'].startswith(
            'Valid until'))


# =========================================================================
#  Carried from L0 — the Talent gate (LN11) and the words (LN10)
# =========================================================================
@tagged('post_install', '-at_install')
class TestCarriedFromL0(TransactionCase):

    def test_ln11_talent_no_longer_gated_on_training_roles(self):
        try:
            from odoo.addons.pb_vendor_access.hooks import (
                SCREEN_GATES, SCREEN_GATES_RETIRED)
        except ImportError:
            self.skipTest('pb_vendor_access is not on this database')
        talent = SCREEN_GATES['pb_talent_hub.item_talent']
        for key in ('training-trainer', 'training-manager', 'training-head'):
            self.assertNotIn(key, talent)
            self.assertIn(key, SCREEN_GATES_RETIRED[
                'pb_talent_hub.item_talent'])
        self.assertIn('goals-head', talent)

    def test_ln10_the_server_words_have_vietnamese(self):
        from odoo.addons.pb_training.models.training_common import (
            due_words, duration_words, reason_word, state_word)
        vi = self.env(context=dict(self.env.context, lang='vi_VN'))
        if not vi['res.lang']._lang_get('vi_VN'):
            self.skipTest('Vietnamese is not installed here')
        today = fields.Date.today()
        self.assertEqual(due_words(today, today, vi), 'Đến hạn hôm nay')
        self.assertEqual(due_words(today + timedelta(days=3), today, vi),
                         'Còn 3 ngày')
        self.assertEqual(reason_word(vi, 'adhoc'), 'Một lần')
        self.assertEqual(reason_word(vi, 'renewal'), 'Gia hạn')
        self.assertEqual(state_word(vi, 'assigned'), 'Chưa bắt đầu')
        self.assertEqual(duration_words(30, vi), '30 phút')
        # No environment: the mail computes keep their English.
        self.assertEqual(due_words(today, today), 'Due today')

    def test_l1_every_new_word_has_its_vietnamese(self):
        po = _src('i18n', 'vi_VN.po')
        for word, vi in (('Certificates', 'Chứng chỉ'),
                         ('Renewal', 'Gia hạn'),
                         ('Add a certificate', 'Thêm chứng chỉ'),
                         ('My certificates', 'Chứng chỉ của tôi'),
                         ('Never expires', 'Không hết hạn'),
                         ('Licence / permit', 'Giấy phép')):
            self.assertIn('msgid "%s"\nmsgstr "%s"' % (word, vi), po, word)

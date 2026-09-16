# -*- coding: utf-8 -*-
"""RIZE W2 E3 — the rails under the allowance, the claim and the money.

Every one of these is written against a failure this phase could have and that
nothing at runtime would report:

  * a claim that reached a payslip any way other than the one money door would
    be a second way to pay people, and every screen would look right;
  * an approval that raised TWO awards pays somebody twice, and the second one
    looks exactly like the first;
  * an allowance that counted the claim being weighed against itself refuses
    every claim ever made, with a sentence that sounds correct;
  * a certificate filed on every read rather than on the crossing is one vault
    document per board refresh, for ever;
  * a report pack with no period stamp emails the whole thing again every
    morning, with a cheerful count in the log;
  * a spreadsheet built from a second query disagrees with the screen it came
    off, and nobody can tell which one is wrong.

WHY THE FIXTURES ARE NAMED "DEMO". Everything here is rolled back with the
transaction, but the names are the live convention anyway (ledger rule 9): a
fixture copied into a live script keeps the name it was written with.
"""

import base64
import io
from datetime import timedelta

from odoo import fields
from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tests import tagged

from .test_training import TrainingCase, _PDF


class ClaimCase(TrainingCase):
    """E1's course, one employee with a login, and an allowance for this year."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.Claim = cls.env['pb.training.claim']
        cls.Allowance = cls.env['pb.training.allowance']
        cls.Analytics = cls.env['pb.training.analytics']
        cls.Pack = cls.env['pb.training.pack']
        cls.Cert = cls.env['pb.training.certificate']

        cls.company = cls.env.company
        Users = cls.env['res.users'].with_context(no_reset_password=True)
        cls.claimant_user = Users.create({
            'name': 'DEMO Claimant One',
            'login': 'demo.claimant.one@example.com',
            'email': 'demo.claimant.one@example.com',
            'group_ids': [(6, 0, [cls.env.ref('base.group_user').id])],
        })
        cls.claimant = cls.env['hr.employee'].create({
            'name': 'DEMO Claimant One',
            'user_id': cls.claimant_user.id,
            'work_email': 'demo.claimant.one@example.com',
            'company_id': cls.company.id,
        })
        cls.allowance = cls.Allowance.create({
            'company_id': cls.company.id,
            'year': fields.Date.today().year,
            'amount': 5000000.0,
            'currency_id': cls.company.currency_id.id,
        })

    # ------------------------------------------------------------- helpers
    def _dormant(self, key='training_claim'):
        """Run the approval hooks without a live route.

        A published route is the real thing and it is proved LIVE, where an HR
        lead really does open an inbox. In a unit transaction the engine would
        need a resolved seat held by the acting user, which is a test of the
        engine and not of this module — so the binding is switched off for the
        length of the transaction and the record's own one-rung ladder drives
        the same two hooks.
        """
        if 'biz.approval.binding' not in self.env:
            return
        process = self.env['biz.approval.process'].sudo().search(
            [('key', '=', key)], limit=1)
        if process:
            self.env['biz.approval.binding'].sudo().search(
                [('process_id', '=', process.id)]).write({'active': False})

    def _claim(self, amount=1000000.0, files=True, employee=None, **extra):
        values = {
            'employee_id': (employee or self.claimant).id,
            'course_name': 'DEMO Advanced spreadsheets',
            'provider': 'DEMO Evening College',
            'amount': amount,
            'currency_id': self.company.currency_id.id,
            'paid_on': fields.Date.today(),
            'company_id': self.company.id,
        }
        values.update(extra)
        claim = self.Claim.create(values)
        if files:
            claim.write({
                'invoice_file': _PDF, 'invoice_filename': 'DEMO receipt.pdf',
                'certificate_file': _PDF,
                'certificate_filename': 'DEMO certificate.pdf',
            })
        return claim

    def _mails(self):
        return self.env['mail.mail'].sudo().search_count([])


# =========================================================================
#  T1 — the shape, the gates and the catalogue row
# =========================================================================
@tagged('post_install', '-at_install')
class TestShapeAndGates(ClaimCase):

    def test_t1_the_catalogue_row_exists_once(self):
        """A route can only be laid against a catalogue row that names a
        STORED model that can hold the request (AM45)."""
        rows = self.env['biz.approval.process'].sudo().search(
            [('key', '=', 'training_claim')])
        self.assertEqual(len(rows), 1,
                         "the training-claim process must be laid exactly once")
        self.assertEqual(rows.model_name, 'pb.training.claim')

    def test_t1_the_claim_is_driven_by_the_chain(self):
        """A raw write to `state` must be refused, or the chain is decorative
        and the money path pays on `state` alone."""
        claim = self._claim()
        with self.assertRaises(AccessError):
            claim.with_user(self.claimant_user).write({'state': 'approved'})

    def test_t1_create_cannot_mint_an_approved_claim(self):
        made = self.Claim.with_user(self.claimant_user).create({
            'employee_id': self.claimant.id,
            'course_name': 'DEMO Sneaky',
            'amount': 1.0,
            'paid_on': fields.Date.today(),
            'state': 'approved',
        })
        self.assertEqual(made.state, 'draft')

    def test_t1_the_training_kind_is_on_the_award(self):
        kinds = dict(self.env['pb.incentive']._fields['kind'].selection)
        self.assertIn('training', kinds)

    def test_t1_an_amount_of_nothing_is_refused(self):
        with self.assertRaises(ValidationError):
            self._claim(amount=0.0, files=False)

    def test_t1_a_receipt_from_the_future_is_refused(self):
        with self.assertRaises(ValidationError):
            self._claim(files=False,
                        paid_on=fields.Date.today() + timedelta(days=3))

    def test_t1_a_seat_is_a_read_and_a_write(self):
        """THE DEFECT THIS TEST EXISTS FOR, found live.

        The engine writes the record's own status AS THE PERSON WHO DECIDED.
        A seat rule that grants only read therefore produces the worst outcome
        there is: the HR lead presses Agree, the ROUTE records the approval,
        and the claim stays on "Waiting on the HR lead" with nothing on any
        screen to say why — the reason is buried in the request's own
        `block_reason`. Both seat rules, because the delay's has the same
        shape and only escapes it by accident.
        """
        for xmlid in ('pb_training.rule_claim_seat',
                      'pb_training.rule_delay_seat'):
            rule = self.env.ref(xmlid)
            self.assertTrue(rule.perm_read, '%s must be readable' % xmlid)
            self.assertTrue(
                rule.perm_write,
                '%s: a decider who cannot write approves the request and '
                'leaves the record where it was' % xmlid)
            self.assertFalse(rule.perm_create, '%s may not create' % xmlid)
            self.assertFalse(rule.perm_unlink, '%s may not delete' % xmlid)
            self.assertIn('seat_user_ids', rule.domain_force,
                          '%s must be narrowed by the seat and nothing '
                          'else' % xmlid)

    def test_t1_money_is_not_a_team_sport(self):
        """A manager can see their team's DUE DATES and must not see what
        somebody paid for a course out of their own pocket. The assignment
        rule reaches a team; the claim rule stops at "my own"."""
        claim_rule = self.env.ref('pb_training.rule_claim_mine')
        self.assertNotIn('parent_id', claim_rule.domain_force)
        assignment_rule = self.env.ref('pb_training.rule_assignment_team')
        self.assertIn('parent_id', assignment_rule.domain_force)

    def test_t1_the_analytics_gate_is_its_own(self):
        """R97: a trainer holds no analytics group and the other way round."""
        outsider = self.env['res.users'].with_context(
            no_reset_password=True).create({
                'name': 'DEMO Outsider',
                'login': 'demo.outsider.claims@example.com',
                'email': 'demo.outsider.claims@example.com',
                'group_ids': [(6, 0, [self.env.ref('base.group_user').id])],
            })
        self.assertFalse(
            self.Analytics.with_user(outsider)._can_read())
        with self.assertRaises(AccessError):
            self.Analytics.with_user(outsider).export_xlsx()


# =========================================================================
#  T2 — the allowance
# =========================================================================
@tagged('post_install', '-at_install')
class TestAllowance(ClaimCase):

    def test_t2_the_company_figure_answers_for_everybody(self):
        amount, currency = self.Allowance.amount_for(self.claimant)
        self.assertEqual(amount, 5000000.0)
        self.assertEqual(currency, self.company.currency_id)

    def test_t2_a_persons_own_figure_wins(self):
        self.Allowance.create({
            'company_id': self.company.id,
            'year': fields.Date.today().year,
            'employee_id': self.claimant.id,
            'amount': 8000000.0,
            'currency_id': self.company.currency_id.id,
        })
        amount, _cur = self.Allowance.amount_for(self.claimant)
        self.assertEqual(amount, 8000000.0)

    def test_t2_an_id_works_as_well_as_a_record(self):
        """R43: a record argument arrives over the wire as a plain integer."""
        by_record = self.Allowance.allowance_for(self.claimant)
        by_id = self.Allowance.allowance_for(self.claimant.id)
        self.assertEqual(by_record, by_id)

    def test_t2_no_allowance_answers_zero_and_not_a_traceback(self):
        self.allowance.unlink()
        amount, currency = self.Allowance.amount_for(self.claimant)
        self.assertEqual(amount, 0.0)
        self.assertTrue(currency)

    def test_t2_two_company_figures_for_one_year_are_refused(self):
        with self.assertRaises(ValidationError):
            self.Allowance.create({
                'company_id': self.company.id,
                'year': fields.Date.today().year,
                'amount': 1.0,
                'currency_id': self.company.currency_id.id,
            })

    def test_t2_the_claim_being_weighed_is_not_counted_against_itself(self):
        """The question is "would this fit". A claim that counted itself as
        already agreed would always be over."""
        claim = self._claim(amount=5000000.0)
        facts = claim._allowance_facts()
        self.assertEqual(facts['agreed'], 0.0)
        self.assertEqual(facts['left'], 5000000.0)


# =========================================================================
#  T3 — the claim, from asking to agreed
# =========================================================================
@tagged('post_install', '-at_install')
class TestClaimJourney(ClaimCase):

    def setUp(self):
        super().setUp()
        self._dormant()

    def test_t3_a_claim_with_no_certificate_cannot_go_in(self):
        claim = self._claim(files=False)
        claim.write({'invoice_file': _PDF,
                     'invoice_filename': 'DEMO receipt.pdf'})
        with self.assertRaises(UserError) as caught:
            claim.action_submit()
        self.assertIn('certificate', str(caught.exception).lower())
        self.assertEqual(claim.state, 'draft')

    def test_t3_a_claim_with_no_receipt_cannot_go_in(self):
        claim = self._claim(files=False)
        claim.write({'certificate_file': _PDF,
                     'certificate_filename': 'DEMO certificate.pdf'})
        with self.assertRaises(UserError) as caught:
            claim.action_submit()
        self.assertIn('receipt', str(caught.exception).lower())

    def test_t3_with_both_files_it_goes_in(self):
        claim = self._claim()
        claim.action_submit()
        self.assertEqual(claim.state, 'submitted')
        self.assertTrue(claim.submitted_on)
        self.assertTrue(claim.invoice_attachment_id)
        self.assertTrue(claim.certificate_attachment_id)

    def test_t3_agreeing_raises_exactly_one_award_and_mails_them(self):
        claim = self._claim(amount=2000000.0)
        claim.action_submit()
        before = self._mails()
        claim.action_approve()
        self.assertEqual(claim.state, 'approved')
        award = claim.incentive_id
        self.assertTrue(award, "an agreed claim must become an award")
        self.assertEqual(award.kind, 'training')
        self.assertEqual(award.state, 'approved')
        self.assertEqual(award.amount, 2000000.0)
        self.assertEqual(award.employee_id, self.claimant)
        # R81: the month it was AGREED in, which is the month the awards
        # dialog picks by — not the month the course was paid for.
        self.assertEqual(award.period_month,
                         fields.Date.today().replace(day=1))
        self.assertGreater(self._mails(), before,
                           "the person has to be told")

    def test_t3_agreeing_twice_never_raises_a_second_award(self):
        claim = self._claim()
        claim.action_submit()
        claim.action_approve()
        first = claim.incentive_id
        claim._make_award()
        claim._make_award()
        self.assertEqual(claim.incentive_id, first)
        self.assertEqual(
            self.env['pb.incentive'].sudo().search_count(
                [('employee_id', '=', self.claimant.id),
                 ('kind', '=', 'training')]), 1)

    def test_t3_the_remaining_allowance_drops_when_one_is_agreed(self):
        first = self._claim(amount=4000000.0)
        first.action_submit()
        first.action_approve()
        second = self._claim(amount=2000000.0)
        facts = second._allowance_facts()
        self.assertEqual(facts['agreed'], 4000000.0)
        self.assertEqual(facts['left'], 1000000.0)

    def test_t3_a_claim_over_the_remainder_is_refused_with_the_arithmetic(self):
        first = self._claim(amount=4000000.0)
        first.action_submit()
        first.action_approve()
        second = self._claim(amount=2000000.0)
        with self.assertRaises(UserError) as caught:
            second.action_submit()
        message = str(caught.exception)
        # THE SENTENCE HAS THE NUMBERS IN IT. "Over your allowance" is a
        # refusal somebody has to email about.
        self.assertIn('5,000,000', message.replace(' ', ','))
        self.assertIn('4,000,000', message.replace(' ', ','))
        self.assertEqual(second.state, 'draft')

    def test_t3_the_switch_lets_one_through_and_says_so(self):
        first = self._claim(amount=4000000.0)
        first.action_submit()
        first.action_approve()
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_training.claim_over_allowance', '1')
        second = self._claim(amount=2000000.0)
        second.action_submit()
        second.action_approve()
        self.assertEqual(second.state, 'approved')
        self.assertTrue(second.over_allowance,
                        "a rule set aside has to be visible on the record")
        bodies = ' '.join(second.message_ids.mapped('body'))
        self.assertIn('over the allowance', bodies)

    def test_t3_a_refusal_needs_a_reason_and_carries_it(self):
        claim = self._claim()
        claim.action_submit()
        with self.assertRaises(UserError):
            self.Board.refuse_claim(claim.id, '')
        before = self._mails()
        self.Board.refuse_claim(claim.id, 'DEMO: that is a conference.')
        self.assertEqual(claim.state, 'refused')
        self.assertIn('conference', claim.refuse_note)
        self.assertGreater(self._mails(), before)

    def test_t3_the_employees_own_door_makes_and_sends_in_one_press(self):
        me = self.env['pb.my.training'].with_user(self.claimant_user)
        answer = me.raise_claim(
            {'course_name': 'DEMO Evening statistics', 'amount': 750000.0,
             'paid_on': fields.Date.to_string(fields.Date.today())},
            invoice=('DEMO receipt.pdf', base64.b64decode(_PDF)),
            certificate=('DEMO certificate.pdf', base64.b64decode(_PDF)))
        claim = self.Claim.sudo().browse(answer['id'])
        self.assertEqual(claim.state, 'submitted')
        self.assertEqual(claim.employee_id, self.claimant)

    def test_t3_a_refused_self_claim_leaves_nothing_behind(self):
        """A half-made claim sitting on somebody's page waiting for an answer
        nobody was asked for is worse than a refusal."""
        me = self.env['pb.my.training'].with_user(self.claimant_user)
        before = self.Claim.sudo().search_count([])
        with self.assertRaises(UserError):
            me.raise_claim(
                {'course_name': 'DEMO No paperwork', 'amount': 100.0,
                 'paid_on': fields.Date.to_string(fields.Date.today())})
        self.assertEqual(self.Claim.sudo().search_count([]), before)

    def test_t3_the_form_cannot_say_who_it_is_for(self):
        me = self.env['pb.my.training'].with_user(self.claimant_user)
        answer = me.raise_claim(
            {'course_name': 'DEMO Forged', 'amount': 100.0,
             'paid_on': fields.Date.to_string(fields.Date.today()),
             'employee_id': self.env['hr.employee'].sudo().search(
                 [('id', '!=', self.claimant.id)], limit=1).id,
             'state': 'approved'},
            invoice=('a.pdf', base64.b64decode(_PDF)),
            certificate=('b.pdf', base64.b64decode(_PDF)))
        claim = self.Claim.sudo().browse(answer['id'])
        self.assertEqual(claim.employee_id, self.claimant)
        self.assertEqual(claim.state, 'submitted')


# =========================================================================
#  T4 — the one money door
# =========================================================================
@tagged('post_install', '-at_install')
class TestTheMoneyDoor(ClaimCase):

    def setUp(self):
        super().setUp()
        self._dormant()

    def test_t4_the_award_is_what_the_feed_picks_up(self):
        """`approved_for_month` is the lane's own picker. The award a claim
        raises has to satisfy it or the money simply never moves, with nothing
        anywhere to say why."""
        claim = self._claim(amount=1500000.0)
        claim.action_submit()
        claim.action_approve()
        picked = self.env['pb.incentive'].approved_for_month(
            fields.Date.today(), company_id=self.company.id)
        self.assertIn(claim.incentive_id, picked)

    def test_t4_nothing_here_writes_a_payslip_line(self):
        """THE RAIL THIS WHOLE PHASE HANGS OFF. If any of these appeared in
        this module, there would be a second way for money to reach a payslip
        and every screen would look exactly the same."""
        import os
        from odoo.modules.module import get_module_path
        root = get_module_path('pb_training')
        banned = ('formula_input_values', 'hr.payslip.line',
                  'hr.payroll.import.line', 'hr.payroll.import.batch')
        hits = []
        for folder, _dirs, files in os.walk(root):
            if '__pycache__' in folder or '/tests' in folder:
                continue
            for name in files:
                if not name.endswith('.py'):
                    continue
                body = open(os.path.join(folder, name), encoding='utf-8').read()
                for word in banned:
                    if word in body:
                        hits.append('%s: %s' % (name, word))
        self.assertFalse(hits, 'a payslip may only be reached through '
                               'pb.oneoff.feed: %s' % hits)

    def test_t4_the_claim_follows_the_award_and_never_leads_it(self):
        claim = self._claim()
        claim.action_submit()
        claim.action_approve()
        self.assertEqual(claim.fulfilment, 'pending')
        claim.incentive_id.sudo().write({'fulfilment': 'queued'})
        self.assertEqual(claim.fulfilment, 'queued')
        claim.incentive_id.sudo().write({'fulfilment': 'paid'})
        self.assertEqual(claim.fulfilment, 'paid')


# =========================================================================
#  T5 — certificates in the vault
# =========================================================================
@tagged('post_install', '-at_install')
class TestCertificates(ClaimCase):

    def setUp(self):
        super().setUp()
        self._dormant()

    def _documents(self, employee, name=None):
        domain = [('employee_id', '=', employee.id)]
        if name:
            domain.append(('name', '=', name))
        return self.env['pb.employee.document'].sudo().with_context(
            active_test=False).search(domain)

    def test_t5_the_vault_has_the_category_the_filing_needs(self):
        self.assertTrue(self.Cert._category(),
                        'the "CERT" vault category has to exist')

    def test_t5_an_uploaded_certificate_is_filed_once(self):
        claim = self._claim()
        claim.action_submit()
        claim.action_approve()
        name = self.Cert._document_name(claim.course_name, claim.provider)
        docs = self._documents(self.claimant, name)
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs.category_id.code, 'CERT')
        self.assertFalse(docs.expiry_date,
                         'a training certificate does not run out')
        # A COPY AND NOT THE CLAIM'S OWN FILE: the vault deletes its own
        # attachment when the document goes.
        self.assertNotEqual(docs.attachment_id,
                            claim.certificate_attachment_id)
        # Idempotent: running the leg again files nothing new.
        claim._file_certificate()
        self.assertEqual(len(self._documents(self.claimant, name)), 1)

    def test_t5_the_filing_never_takes_the_approval_down(self):
        """A vault that is missing, full or unhappy must not undo an approval
        somebody has made (R131).

        The failure is staged by taking the CATEGORY away, which is the real
        shape of it on a tenant that has the vault module without this
        module's category — and not by deleting the attachment, which the
        claim refuses outright (`ondelete='restrict'`: the evidence an
        approver agreed to must not be able to vanish from underneath them).
        """
        claim = self._claim()
        claim.action_submit()
        category = self.Cert._category()
        code = category.code
        category.sudo().write({'code': 'DEMO-NOT-CERT'})
        try:
            claim.action_approve()
        finally:
            category.sudo().write({'code': code})
        self.assertEqual(claim.state, 'approved')
        self.assertTrue(claim.incentive_id,
                        'the award must survive a vault that would not take '
                        'the certificate')

    def test_t5_the_evidence_cannot_vanish_from_under_an_approver(self):
        """`ondelete='restrict'` on both files, and it is deliberate: the
        revision stamp an approver agreed to names the attachment id (AM32),
        so a file that could be deleted would leave an approval standing over
        evidence nobody can look at."""
        from psycopg2 import IntegrityError
        claim = self._claim()
        with self.assertRaises(IntegrityError):
            with self.env.cr.savepoint():
                claim.certificate_attachment_id.sudo().unlink()

    def test_t5_the_document_name_is_built_in_one_place(self):
        """The write and the "already filed?" search both come here, so an
        idempotency check can never look for a name nothing writes."""
        name = self.Cert._document_name('DEMO Course', 'DEMO Provider')
        self.assertIn('DEMO Course', name)
        self.assertIn('DEMO Provider', name)
        self.assertEqual(self.Cert._document_name('DEMO Course', 'DEMO Provider'),
                         name)


# =========================================================================
#  T6 — the numbers
# =========================================================================
@tagged('post_install', '-at_install')
class TestAnalytics(ClaimCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.dept_a = cls.env['hr.department'].create(
            {'name': 'DEMO Field team', 'company_id': cls.company.id})
        cls.dept_b = cls.env['hr.department'].create(
            {'name': 'DEMO Office team', 'company_id': cls.company.id})
        Users = cls.env['res.users'].with_context(no_reset_password=True)
        cls.people = []
        for idx in range(6):
            user = Users.create({
                'name': 'DEMO Analytics %s' % idx,
                'login': 'demo.analytics.%s@example.com' % idx,
                'email': 'demo.analytics.%s@example.com' % idx,
                'group_ids': [(6, 0, [cls.env.ref('base.group_user').id])],
            })
            emp = cls.env['hr.employee'].create({
                'name': 'DEMO Analytics %s' % idx,
                'user_id': user.id,
                'work_email': 'demo.analytics.%s@example.com' % idx,
                'company_id': cls.company.id,
                'department_id': (cls.dept_a if idx < 3 else cls.dept_b).id,
            })
            cls.people.append((user, emp))

    def _seed(self):
        """Six assignments, two departments, mixed states."""
        today = fields.Date.today()
        rows = self.env['pb.training.assignment']
        for idx, (user, emp) in enumerate(self.people):
            rows |= self.env['pb.training.assignment'].create({
                'channel_id': self.course.id,
                'employee_id': emp.id,
                'reason': 'compliance' if idx % 2 else 'adhoc',
                'due_date': today + timedelta(days=5 - idx * 4),
                'company_id': self.company.id,
            })
        return rows

    def test_t6_the_ratios_equal_a_python_recomputation(self):
        rows = self._seed()
        rows._refresh()
        start = fields.Date.today() - timedelta(days=1)
        end = fields.Date.today()
        board = self.Analytics.get_numbers(
            fields.Date.to_string(start), fields.Date.to_string(end))
        core = board['core']

        # RECOMPUTED OVER THE SAME DOMAIN AND NOT OVER MY OWN SIX ROWS. This
        # runs against a live database that already has assignments in the
        # range (the demo data does), so a recomputation over the fixture
        # alone compares six against everything and fails over a facade that
        # is perfectly right. The domain is the facade's own.
        everything = self.env['pb.training.assignment'].sudo().search([
            ('company_id', 'in', self.env.companies.ids),
            ('assigned_on', '>=', start), ('assigned_on', '<=', end),
        ])
        self.assertTrue(rows <= everything,
                        'the fixture has to be inside the cohort it is '
                        'measured in')
        everything._refresh()
        counted = everything.filtered(lambda r: r.state != 'excused')
        done = counted.filtered(lambda r: r.state == 'done')
        self.assertEqual(core['total'], len(everything))
        self.assertEqual(core['counted'], len(counted))
        self.assertEqual(core['done'], len(done))
        self.assertEqual(
            core['completion'],
            int(round(100.0 * len(done) / len(counted))) if counted else None)

    def test_t6_the_splits_sum_to_the_total(self):
        self._seed()
        board = self.Analytics.get_numbers(
            fields.Date.to_string(fields.Date.today() - timedelta(days=1)),
            fields.Date.to_string(fields.Date.today()))
        for key in ('by_department', 'by_course', 'by_reason'):
            self.assertEqual(
                sum(row['total'] for row in board[key]), board['total'],
                'the %s split has to add up to the headline' % key)

    def test_t6_an_empty_range_says_so_in_words(self):
        board = self.Analytics.get_numbers('2001-01-01', '2001-01-31')
        self.assertTrue(board['empty'])
        self.assertIn('2001-01-01', board['headline'])
        self.assertFalse(board['tiles'],
                         'a screen of zeros is how a working feature gets '
                         'reported as broken')

    def test_t6_the_spreadsheet_matches_the_screen(self):
        self._seed()
        start = fields.Date.to_string(fields.Date.today() - timedelta(days=1))
        end = fields.Date.to_string(fields.Date.today())
        board = self.Analytics.get_numbers(start, end)
        out = self.Analytics.export_xlsx(start, end)
        self.assertTrue(out['ok'])
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl is not installed on this machine')
        book = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(out['file_b64'])))
        summary = book[book.sheetnames[0]]
        values = {}
        for row in summary.iter_rows(min_row=2, max_col=2, values_only=True):
            if row[0]:
                values[str(row[0])] = row[1]
        self.assertEqual(values.get('Courses set'), board['core']['total'])
        self.assertEqual(values.get('Past their date'),
                         board['core']['overdue'])
        self.assertIn('Week by week', book.sheetnames)

    def test_t6_the_trend_never_shows_two_bars_with_one_label(self):
        """Bucketed by the MONDAY and not by the ISO week number: a range that
        crosses a new year has two week 1s."""
        self._seed()
        board = self.Analytics.get_numbers(
            fields.Date.to_string(fields.Date.today() - timedelta(days=200)),
            fields.Date.to_string(fields.Date.today()))
        keys = [week['key'] for week in board['trend']]
        self.assertEqual(len(keys), len(set(keys)))

    def test_t6_the_filter_lists_only_offer_what_the_data_uses(self):
        """R27: a filter that matches nothing is a broken promise."""
        self._seed()
        board = self.Analytics.get_numbers()
        names = [row['name'] for row in board['lists']['courses']]
        self.assertIn(self.course.name, names)


# =========================================================================
#  T7 — the report pack
# =========================================================================
@tagged('post_install', '-at_install')
class TestReportPack(ClaimCase):

    def setUp(self):
        super().setUp()
        icp = self.env['ir.config_parameter'].sudo()
        icp.set_param('pb_training.report_pack_email',
                      'training.pack@example.com')
        icp.set_param('pb_training.report_pack_period', 'monthly')
        icp.set_param('pb_training.report_pack_last', 'never')

    def test_t7_switched_off_it_sends_nothing_and_says_why(self):
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_training.report_pack', '0')
        before = self._mails()
        answer = self.Pack.send_pack()
        self.assertEqual(answer['mode'], 'off')
        self.assertEqual(answer['sent'], 0)
        self.assertEqual(self._mails(), before)
        # R54: a thing that is off and does not SAY so is reported as broken.
        self.assertIn('switched off', answer['msg'])
        self.assertIn('training.pack@example.com', answer['msg'])

    def test_t7_switched_on_it_sends_once_a_period(self):
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_training.report_pack', '1')
        before = self._mails()
        first = self.Pack.send_pack()
        self.assertEqual(first['mode'], 'sent')
        self.assertEqual(first['sent'], 1)
        self.assertEqual(self._mails(), before + 1)
        again = self.Pack.send_pack()
        self.assertEqual(again['mode'], 'already')
        self.assertEqual(self._mails(), before + 1)

    def test_t7_the_spreadsheet_rides_with_it(self):
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_training.report_pack', '1')
        self.Pack.send_pack()
        mail = self.env['mail.mail'].sudo().search(
            [('email_to', '=', 'training.pack@example.com')],
            order='id desc', limit=1)
        self.assertTrue(mail)
        self.assertTrue(mail.attachment_ids,
                        'the pack is the numbers AND the spreadsheet')
        self.assertTrue(mail.attachment_ids[0].name.endswith('.xlsx'))

    def test_t7_with_nobody_to_write_to_it_says_that_instead(self):
        icp = self.env['ir.config_parameter'].sudo()
        icp.set_param('pb_training.report_pack', '1')
        icp.set_param('pb_training.report_pack_email', ' ')
        answer = self.Pack.send_pack()
        self.assertEqual(answer['mode'], 'nobody')

    def test_t7_the_window_is_the_period_that_has_ended(self):
        start, end, stamp, label = self.Pack._window(fields.Date.today())
        self.assertLess(end, fields.Date.today().replace(day=1))
        self.assertEqual(start.day, 1)
        self.assertEqual(stamp, start.strftime('%Y-%m'))
        self.assertTrue(label)

    def test_t7_weekly_is_a_different_window_and_a_different_stamp(self):
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_training.report_pack_period', 'weekly')
        start, end, stamp, _label = self.Pack._window(fields.Date.today())
        self.assertEqual((end - start).days, 6)
        self.assertIn('-W', stamp)

    def test_t7_only_a_manager_may_send_it_by_hand(self):
        with self.assertRaises(AccessError):
            self.Pack.with_user(self.claimant_user).run_now()


# =========================================================================
#  T9 — the doors, the numbers and the switches
# =========================================================================
@tagged('post_install', '-at_install')
class TestDoorsAndDials(ClaimCase):

    def test_t9_every_new_palette_row_points_at_something_that_exists(self):
        from .test_training import _src
        import re
        body = _src('static', 'src', 'js', 'training_palette.js')
        for xmlid in re.findall(r'xmlid:\s*"([^"]+)"', body):
            self.assertTrue(
                self.env.ref(xmlid, raise_if_not_found=False),
                '%s is offered in the command bar and does not exist' % xmlid)

    def test_t9_the_new_command_bar_rows_are_in_this_modules_block(self):
        from .test_training import _src
        import re
        body = _src('static', 'src', 'js', 'training_palette.js')
        sequences = [int(n) for n in re.findall(r'sequence:\s*(\d+)\s*\}',
                                                body)]
        block = [n for n in sequences if n >= 3900]
        self.assertIn(3960, block)
        self.assertIn(3970, block)
        self.assertTrue(all(3900 <= n < 4000 for n in block),
                        'E must stay inside its own 3900 block')

    def test_t9_the_insights_lens_is_registered_at_forty(self):
        from .test_training import _src
        body = _src('static', 'src', 'js', 'training_palette.js')
        self.assertIn('INSIGHTS_LENSES', body)
        self.assertIn('{ sequence: 40 }', body)

    def test_t9_every_facade_door_carries_views(self):
        """R125: a hand-built act_window dict MUST carry `views`, or the client
        throws before the screen opens and the user sees the generic error
        dialog with nothing in the console."""
        for verb, payload in (
                ('open_claims', {}), ('new_claim', {}),
                ('open_allowances', {}), ('new_allowance', {}),
                ('open_claim', {'claim_id': self._claim(files=False).id})):
            action = self.Board.act(verb, payload)
            self.assertIn('views', action, '%s has no views' % verb)
            self.assertTrue(action['views'])

    def test_t9_an_unknown_verb_is_refused(self):
        with self.assertRaises(UserError):
            self.Board.act('__import__', {})

    def test_t9_the_e3_switches_all_have_a_default_in_code(self):
        from ..models.training_common import (
            DEFAULTS, P_CLAIM_OVER, P_CLAIM_LIMIT, P_PACK, P_PACK_PERIOD,
        )
        for key in (P_CLAIM_OVER, P_CLAIM_LIMIT, P_PACK, P_PACK_PERIOD):
            self.assertIn(key, DEFAULTS)

    def test_t9_the_two_text_switches_are_not_in_defaults(self):
        """`set_param(key, '')` DELETES the row, so an empty-string default
        cannot be materialised and a hook that tried would write it, find it
        missing on the next read and write it again for ever."""
        from ..models.training_common import (
            DEFAULTS, P_PACK_EMAIL, P_PACK_STAMP,
        )
        self.assertNotIn(P_PACK_EMAIL, DEFAULTS)
        self.assertNotIn(P_PACK_STAMP, DEFAULTS)

    def test_t9_the_cron_row_exists_and_is_live(self):
        cron = self.env.ref('pb_training.cron_training_pack')
        self.assertTrue(cron.active)
        self.assertEqual(cron.interval_type, 'days')

    def test_t9_the_portal_helpers_are_all_module_prefixed(self):
        """R186: ALL `CustomerPortal` subclasses merge into ONE class, so a
        private helper without this module's prefix silently collides with
        another module's and whichever loads last wins."""
        import ast
        from .test_training import _src
        tree = ast.parse(_src('controllers', 'portal.py'))
        bad = []
        for node in ast.walk(tree):
            if not isinstance(node, ast.ClassDef):
                continue
            for item in node.body:
                if not isinstance(item, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    continue
                name = item.name
                if not name.startswith('_') or name.startswith('__'):
                    continue
                # Framework hooks are overridden on purpose.
                if name.startswith('_prepare_'):
                    continue
                if not name.startswith('_tr_'):
                    bad.append(name)
        self.assertFalse(bad, 'unprefixed portal helpers: %s' % bad)


# =========================================================================
#  T10 — the words on the screen
# =========================================================================
@tagged('post_install', '-at_install')
class TestTheWords(ClaimCase):
    """The white-label rule and the demo-data rule, checked where they break.

    THE VENDOR NAME is already swept module-wide by
    `test_training.TestSourceGates.test_the_word_odoo_appears_in_no_user_visible_string`,
    which reads every file this module ships and knows the difference between
    a string and an identifier. A second sweep here would be the same test
    written worse — and written worse it FAILED on engineering comments,
    which the rule explicitly exempts. So this class checks the one thing
    that sweep does not: the CUSTOMER's name.
    """

    #: What a viewer can read. An `id=`/`ref=` is an xmlid and an xmlid keeps
    #: the programme's name on purpose (ledger rule 9: module names, xmlids,
    #: docs and commit messages are engineering-facing).
    def _visible_text(self, body):
        import re
        body = re.sub(r'<!--.*?-->', ' ', body, flags=re.S)
        body = re.sub(r'\b(?:id|ref|name|inherit_id)="[^"]*"', ' ', body)
        return body

    def test_t10_no_customer_name_where_a_viewer_can_read_it(self):
        """Ledger rule 9 / D18: the owner shows this database to the
        customer's competitors, so no shipped string may carry the customer's
        name. Identifiers are exempt and stay as they are."""
        import os
        from odoo.modules.module import get_module_path
        root = get_module_path('pb_training')
        hits = []
        for folder, _dirs, files in os.walk(root):
            if '__pycache__' in folder or 'tests' in folder:
                continue
            for name in files:
                if not name.endswith('.xml'):
                    continue
                body = self._visible_text(
                    open(os.path.join(folder, name), encoding='utf-8').read())
                if 'rize' in body.lower():
                    hits.append(name)
        self.assertFalse(hits, 'customer name where a viewer can read it: %s'
                         % hits)

    def test_t10_every_sentence_this_phase_adds_is_a_whole_sentence(self):
        """R117: a frame with a number in it ("%s days left") is ungrammatical
        for exactly one of its values, and a translator handed the frame and
        the number separately cannot fix a plural they were never given."""
        import os
        import re
        from odoo.modules.module import get_module_path
        root = get_module_path('pb_training')
        hits = []
        for name in ('claim.py', 'allowance.py', 'analytics.py', 'pack.py',
                     'pb_training_claims.py', 'pb_my_training_claims.py',
                     'certificate.py'):
            body = open(os.path.join(root, 'models', name),
                        encoding='utf-8').read()
            for line in body.splitlines():
                if re.search(r'\(s\)', line) and '_(' in line:
                    hits.append('%s: %s' % (name, line.strip()[:80]))
        self.assertFalse(hits, 'bracketed plural in a user-visible string: %s'
                         % hits)

    def test_t10_the_refusal_sentences_carry_their_numbers(self):
        """A refusal a person cannot act on is a refusal they email about."""
        self._dormant()
        first = self._claim(amount=4500000.0)
        first.action_submit()
        first.action_approve()
        second = self._claim(amount=1000000.0)
        problem = second._allowance_problem()
        self.assertTrue(problem)
        for number in ('5,000,000', '4,500,000', '500,000'):
            self.assertIn(number, problem.replace('\u202f', ','),
                          '%s is missing from the refusal' % number)

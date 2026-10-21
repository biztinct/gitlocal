# -*- coding: utf-8 -*-
"""LEARNING L2 — required training. Handover tests T1–T9, T13–T15.

The rails under "who must hold what": an audience that silently includes the
wrong people (nationality instead of work country, a contractor a rule does
not cover), a cell that says "valid" over a certificate the rule's own validity
has already run out, a matrix that asks the database once per cell, a morning
job that assigns the same course twice (or alongside the renewal leg), and a
trainer who can rewrite the rules.

Fixtures are "DEMO" (ledger rule 9) and roll back. Sentences are asserted in
English, so the class pins `lang='en_US'` (ledger LN14).
"""

import base64
import io
import re
from datetime import timedelta

from unittest.mock import patch

from dateutil.relativedelta import relativedelta

from odoo import fields
from odoo.exceptions import AccessError, UserError
from odoo.tests import tagged

from .test_assignments import AssignmentCase
from .test_training import _PDF, _src

_PDF_BYTES = base64.b64decode(_PDF)


class RequirementCase(AssignmentCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        # A COMPANY OF ITS OWN. Every read here is `company_id in
        # env.companies`, and the database may already hold rules, packs,
        # sites and certificates (a walked local copy, p0clone before a
        # deploy): none of them may change a count asserted below.
        cls.company = cls.env['res.company'].create({'name': 'DEMO L2 Co'})
        cls.env = cls.env(context=dict(
            cls.env.context, lang='en_US',
            allowed_company_ids=[cls.company.id]))
        (cls.boss | cls.e_one | cls.e_two | cls.e_three).write(
            {'company_id': cls.company.id})
        for user in (cls.boss_user, cls.u_one, cls.u_two, cls.u_three):
            user.write({'company_ids': [(4, cls.company.id)],
                        'company_id': cls.company.id})
        cls.Req = cls.env['pb.training.requirement'].sudo()
        cls.Board = cls.env['pb.training']
        cls.Doc = cls.env['pb.employee.document'].sudo()
        cls.Cert = cls.env['pb.training.certificate']
        cls.cat = cls.Cert._category()
        cls.lic = cls.Cert._licence_category()
        cls.dept = cls.env['hr.department'].create({'name': 'DEMO L2 Plant'})
        cls.sub = cls.env['hr.department'].create(
            {'name': 'DEMO L2 Line A', 'parent_id': cls.dept.id})
        cls.other = cls.env['hr.department'].create({'name': 'DEMO L2 Office'})
        cls.e_one.department_id = cls.sub.id
        cls.e_two.department_id = cls.dept.id
        cls.e_three.department_id = cls.other.id
        Users = cls.env['res.users'].with_context(no_reset_password=True)
        cls.trainer = Users.create({
            'name': 'DEMO Trainer L2', 'login': 'demo.trainer.l2@example.com',
            'email': 'demo.trainer.l2@example.com',
            'company_id': cls.company.id,
            'company_ids': [(6, 0, cls.company.ids)],
            'group_ids': [(6, 0, [cls.env.ref('base.group_user').id,
                                  cls.env.ref('pb_training.group_training_user').id])]})
        cls.manager = Users.create({
            'name': 'DEMO Manager L2', 'login': 'demo.manager.l2@example.com',
            'email': 'demo.manager.l2@example.com',
            'company_id': cls.company.id,
            'company_ids': [(6, 0, cls.company.ids)],
            'group_ids': [(6, 0, [cls.env.ref('base.group_user').id,
                                  cls.env.ref('pb_training.group_training_manager').id])]})

    # ------------------------------------------------------------ helpers
    def _req(self, **vals):
        base = {'name': 'DEMO rule', 'company_id': self.company.id,
                'kind': 'always', 'target': 'course',
                'channel_id': self.simple.id, 'auto_assign': False}
        base.update(vals)
        return self.Req.create(base)

    def _doc(self, emp, name, days=None, category=None, issued_days=30,
             **extra):
        att = self.env['ir.attachment'].sudo().create({
            'name': 'DEMO.pdf', 'raw': _PDF_BYTES, 'mimetype': 'application/pdf',
            'res_model': 'hr.employee', 'res_id': emp.id})
        vals = {
            'employee_id': emp.id, 'category_id': (category or self.cat).id,
            'name': name, 'attachment_id': att.id,
            'issue_date': fields.Date.today() - timedelta(days=issued_days),
            'expiry_date': fields.Date.today() + timedelta(days=days)
            if days is not None else False,
            'training_source': 'upload', 'company_id': emp.company_id.id,
        }
        vals.update(extra)
        return self.Doc.create(vals)

    def _cells(self, req, employees=None):
        return self.Req.evaluate(req, employees=employees)

    def _state(self, req, emp):
        cell = self._cells(req, emp).get((req.id, emp.id))
        return cell and cell['state']


# =========================================================================
#  T1 — the audience
# =========================================================================
@tagged('post_install', '-at_install')
class TestT1Audience(RequirementCase):

    def test_t1_department_reaches_the_departments_under_it(self):
        req = self._req(department_ids=[(6, 0, self.dept.ids)])
        people = req._people()
        self.assertIn(self.e_one, people, 'Line A is under the plant')
        self.assertIn(self.e_two, people)
        self.assertNotIn(self.e_three, people)

    def test_t1_job_and_contractors(self):
        job = self.env['hr.job'].create({'name': 'DEMO L2 Forklift driver'})
        self.e_two.job_id = job.id
        self.assertEqual(self._req(job_ids=[(6, 0, job.ids)])._people(),
                         self.e_two)
        self.e_three.employee_type = 'contractor'
        everyone = self._req(department_ids=[(6, 0, self.other.ids)])
        self.assertIn(self.e_three, everyone._people())
        staff = self._req(department_ids=[(6, 0, self.other.ids)],
                          include_contractors=False)
        self.assertNotIn(self.e_three, staff._people())

    def test_t1_site_and_safety_group(self):
        site = self.env['hr.work.location'].create({
            'name': 'DEMO L2 Site North', 'location_type': 'office',
            'address_id': self.company.partner_id.id})
        self.e_one.work_location_id = site.id
        self.assertEqual(
            self._req(work_location_ids=[(6, 0, site.ids)])._people(),
            self.e_one)
        self.e_two.pb_safety_group = '3'
        self.assertEqual(self._req(safety_groups='3')._people(), self.e_two)

    def test_t1_country_is_where_they_work_never_nationality(self):
        vn = self.env.ref('base.vn')
        sg = self.env.ref('base.sg')
        self.company.country_id = vn.id
        office_sg = self.env['res.partner'].create(
            {'name': 'DEMO L2 Singapore office', 'country_id': sg.id})
        no_country = self.env['res.partner'].create(
            {'name': 'DEMO L2 Office with no country'})
        # e_one: Vietnamese national working at the Singapore office.
        self.e_one.write({'address_id': office_sg.id, 'country_id': vn.id})
        # e_two: an address with no country -> the company's (Vietnam).
        self.e_two.write({'address_id': no_country.id, 'country_id': sg.id})
        vn_rule = self._req(country_ids=[(6, 0, vn.ids)])
        sg_rule = self._req(country_ids=[(6, 0, sg.ids)])
        self.assertNotIn(self.e_one, vn_rule._people())
        self.assertIn(self.e_one, sg_rule._people())
        self.assertIn(self.e_two, vn_rule._people())
        self.assertNotIn(self.e_two, sg_rule._people())

    def test_t1_a_rule_that_needs_an_audience_is_about_nobody_until_given_one(self):
        req = self._req(needs_audience=True)
        self.assertFalse(req._people())
        req.department_ids = [(6, 0, self.dept.ids)]
        self.assertTrue(req._people())


# =========================================================================
#  T2 — the six states
# =========================================================================
@tagged('post_install', '-at_install')
class TestT2Status(RequirementCase):

    def test_t2_document_evidence_valid_expiring_missing(self):
        req = self._req(target='licence', channel_id=False,
                        document_category_id=self.lic.id,
                        name_pattern='Forklift', grace_days=30,
                        department_ids=[(6, 0, self.dept.ids)])
        self._doc(self.e_one, 'DEMO Forklift operator licence', days=200,
                  category=self.lic)
        self._doc(self.e_two, 'DEMO Forklift operator licence', days=10,
                  category=self.lic)
        cells = self._cells(req)
        self.assertEqual(cells[(req.id, self.e_one.id)]['state'], 'valid')
        self.assertEqual(cells[(req.id, self.e_two.id)]['state'], 'expiring')
        self.assertEqual(cells[(req.id, self.e_two.id)]['words'],
                         'Expires in 10 days')
        # A pension letter is not a forklift licence ("SIO" is a word).
        req2 = self._req(target='licence', channel_id=False,
                         name_pattern='SIO', department_ids=[(6, 0, self.dept.ids)])
        self._doc(self.e_one, 'DEMO Pension letter', days=300)
        self.assertEqual(self._state(req2, self.e_one), 'missing')

    def test_t2_a_course_rule_with_no_course_is_quiet_not_missing(self):
        req = self._req(channel_id=False,
                        department_ids=[(6, 0, self.dept.ids)])
        cell = self._cells(req, self.e_one)[(req.id, self.e_one.id)]
        self.assertEqual((cell['state'], cell['words']),
                         ('unlinked', 'No course linked yet'))
        s = self.Board.matrix({'requirement_id': req.id})['summary']
        self.assertEqual((s['missing'], s['cells'], s['people_missing']),
                         (0, 0, 0))
        self.assertEqual(
            self.Auto._enforce_requirements(fields.Date.today(), req)['made'],
            0)

    def test_t2_the_rules_validity_beats_the_tests(self):
        req = self._req(target='course', validity_months=12,
                        department_ids=[(6, 0, self.dept.ids)])
        # A test certificate issued 13 months ago, its own test saying 24.
        self._doc(self.e_one, 'DEMO Walking the rows — certificate (2025)',
                  days=330, issued_days=400, training_source='test',
                  channel_id=self.simple.id)
        self.assertEqual(self._state(req, self.e_one), 'missing')
        req.validity_months = 24
        self.assertEqual(self._state(req, self.e_one), 'valid')

    def test_t2_finished_course_open_course_not_due_and_exempt(self):
        req = self._req(validity_months=12,
                        department_ids=[(6, 0, self.dept.ids)])
        row = self.Assignment.create({'channel_id': self.simple.id,
                                      'employee_id': self.e_one.id})
        self.assertEqual(self._state(req, self.e_one), 'in_progress')
        self._finish(self.u_one, self.simple, [self.simple_lesson])
        row._refresh()
        self.assertEqual(row.state, 'done')
        self.assertEqual(self._state(req, self.e_one), 'valid')
        # Out of the audience -> exempt, only when asked about by name.
        cell = self._cells(req, self.e_three)[(req.id, self.e_three.id)]
        self.assertEqual(cell['state'], 'exempt')
        day1 = self._req(kind='day_one', due_days=14,
                         department_ids=[(6, 0, self.dept.ids)])
        self.e_two.version_id.sudo().write(
            {'contract_date_start': fields.Date.today() - timedelta(days=3)})
        self.assertEqual(self._state(day1, self.e_two), 'not_due')


# =========================================================================
#  T3 — the matrix payload, and the query rail
# =========================================================================
@tagged('post_install', '-at_install')
class TestT3Matrix(RequirementCase):

    def test_t3_counts_add_up_and_filters_work(self):
        req = self._req(department_ids=[(6, 0, self.dept.ids)])
        self.Assignment.create({'channel_id': self.simple.id,
                                'employee_id': self.e_one.id})
        data = self.Board.matrix({'requirement_id': req.id})
        self.assertTrue(data['allowed'])
        s = data['summary']
        self.assertEqual(s['people'], 2)
        self.assertEqual(s['valid'] + s['expiring'] + s['missing']
                         + s['in_progress'] + s['not_due'], s['cells'])
        self.assertEqual((s['in_progress'], s['missing']), (1, 1))
        # Missing first.
        self.assertEqual(data['rows'][0]['employee_id'], self.e_two.id)
        self.assertIn('in date across 2 people', data['headline'])
        only = self.Board.matrix({'requirement_id': req.id, 'state': 'missing'})
        self.assertEqual([r['employee_id'] for r in only['rows']],
                         [self.e_two.id])
        named = self.Board.matrix({'requirement_id': req.id, 'q': 'bui huu'})
        self.assertEqual([r['employee_id'] for r in named['rows']],
                         [self.e_one.id])

    def test_t3_two_hundred_and_one_people_page(self):
        dept = self.env['hr.department'].create({'name': 'DEMO L2 Big'})
        self.env['hr.employee'].create([
            {'name': 'DEMO L2 Person %03d' % i, 'department_id': dept.id}
            for i in range(201)])
        req = self._req(department_ids=[(6, 0, dept.ids)])
        data = self.Board.matrix({'requirement_id': req.id})
        self.assertTrue(data['truncated'])
        self.assertEqual((len(data['rows']), data['total'], data['pages']),
                         (200, 201, 2))
        page2 = self.Board.matrix({'requirement_id': req.id, 'page': 2})
        self.assertEqual(len(page2['rows']), 1)
        self.assertFalse(self.Board.matrix(
            {'requirement_id': req.id, 'q': 'Person 007'})['truncated'])

    def test_t3_fifty_people_five_rules_under_thirty_queries(self):
        dept = self.env['hr.department'].create({'name': 'DEMO L2 Fifty'})
        people = self.env['hr.employee'].create([
            {'name': 'DEMO L2 Q %02d' % i, 'department_id': dept.id}
            for i in range(50)])
        for emp in people[:10]:
            self._doc(emp, 'DEMO Forklift licence', days=100,
                      category=self.lic)
        audience = [(6, 0, dept.ids)]
        reqs = self._req(department_ids=audience) \
            | self._req(target='licence', channel_id=False,
                        name_pattern='Forklift', department_ids=audience) \
            | self._req(target='certificate', channel_id=False,
                        name_pattern='First aid', department_ids=audience) \
            | self._req(kind='day_one', department_ids=audience) \
            | self._req(target='hours', hours_required=8,
                        department_ids=audience)
        self.env.flush_all()
        self.env.invalidate_all()
        before = self.env.cr.sql_log_count
        cells = self.Req.evaluate(reqs)
        used = self.env.cr.sql_log_count - before
        self.assertEqual(len(cells), 250)
        self.assertLess(used, 30, 'the matrix asked %s times' % used)


# =========================================================================
#  T4 — the legacy rows
# =========================================================================
@tagged('post_install', '-at_install')
class TestT4Migration(RequirementCase):

    def test_t4_rules_and_schedules_convert_once(self):
        rule = self.env['pb.training.rule'].create({
            'name': 'DEMO Induction', 'due_days': 10,
            'channel_ids': [(6, 0, [self.simple.id, self.course.id])],
            'department_ids': [(6, 0, self.dept.ids)]})
        sch = self.env['pb.training.schedule'].create({
            'name': 'DEMO Fire safety yearly', 'channel_id': self.simple.id,
            'audience': 'department', 'department_id': self.dept.id,
            'every_months': 12, 'due_days': 30,
            'next_run': fields.Date.today() + timedelta(days=5)})
        first = self.Req._convert_legacy()
        self.assertEqual(first, {'rules': 2, 'schedules': 1})
        made = self.Req.search([('rule_id', '=', rule.id)])
        self.assertEqual(set(made.mapped('kind')), {'day_one'})
        self.assertEqual(made.mapped('department_ids'), self.dept)
        rec = self.Req.search([('schedule_id', '=', sch.id)])
        self.assertEqual((rec.kind, rec.validity_months, rec.due_days),
                         ('recurring', 12, 30))
        self.assertEqual(self.Req._convert_legacy(),
                         {'rules': 0, 'schedules': 0})

    def test_t4_day_one_rule_department_is_child_of(self):
        rule = self.env['pb.training.rule'].create({
            'name': 'DEMO Plant induction',
            'channel_ids': [(6, 0, self.simple.ids)],
            'department_ids': [(6, 0, self.dept.ids)]})
        self.assertTrue(rule._matches(self.e_one.sudo()),
                        'Line A sits under the plant')
        self.assertFalse(rule._matches(self.e_three.sudo()))


# =========================================================================
#  T5/T6/T7 — the morning leg
# =========================================================================
@tagged('post_install', '-at_install')
class TestT5T7Enforcement(RequirementCase):

    def _open(self, emp, channel=None):
        return self.Assignment.sudo().search([
            ('employee_id', '=', emp.id),
            ('channel_id', '=', (channel or self.simple).id),
            ('state', 'in', ('assigned', 'in_progress', 'overdue', 'excused'))])

    def test_t5_missing_people_are_put_on_once(self):
        no_login = self.env['hr.employee'].create(
            {'name': 'DEMO L2 No login', 'department_id': self.dept.id})
        req = self._req(auto_assign=True, due_days=20,
                        department_ids=[(6, 0, self.dept.ids)])
        counts = self.Auto._enforce_requirements(fields.Date.today(), req)
        self.assertEqual(counts['made'], 2)
        self.assertEqual(counts['no_login'], 1)
        rows = self.Assignment.sudo().search([('requirement_id', '=', req.id)])
        self.assertEqual(set(rows.mapped('reason')), {'required'})
        self.assertEqual(set(rows.mapped('employee_id').ids),
                         {self.e_one.id, self.e_two.id})
        self.assertEqual(rows[0].due_date,
                         fields.Date.today() + timedelta(days=20))
        self.assertNotIn(no_login, rows.mapped('employee_id'))
        again = self.Auto._enforce_requirements(fields.Date.today(), req)
        self.assertEqual(again['made'], 0, 'never a second open one')

    def test_t5_cap_auto_off_and_legacy_rows(self):
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_training.enforce_cap', '1')
        req = self._req(auto_assign=True,
                        department_ids=[(6, 0, self.dept.ids)])
        counts = self.Auto._enforce_requirements(fields.Date.today(), req)
        self.assertEqual((counts['made'], counts['capped']), (1, 1))
        self.env['ir.config_parameter'].sudo().set_param(
            'pb_training.enforce_cap', '500')
        req.auto_assign = False
        self.assertEqual(
            self.Auto._enforce_requirements(fields.Date.today(), req)['made'],
            0)
        rule = self.env['pb.training.rule'].create(
            {'name': 'DEMO legacy', 'channel_ids': [(6, 0, self.simple.ids)]})
        req.write({'auto_assign': True, 'rule_id': rule.id})
        self.assertEqual(
            self.Auto._enforce_requirements(fields.Date.today(), req)['made'],
            0,
            'a converted day-one rule assigns through its own mechanism')

    def test_t6_renewal_and_enforcement_one_assignment(self):
        rule = self._req(auto_assign=True,
                         department_ids=[(6, 0, self.dept.ids)],
                         grace_days=60)
        doc = self._doc(self.e_one, 'DEMO Walking the rows — certificate '
                        '(2025)', days=20, training_source='test',
                        channel_id=self.simple.id)
        today = fields.Date.today()
        self.assertNotIn(doc, self.Auto._renewal_candidates(today),
                         'the rule owns this course')
        self.Auto._renewals(today)
        self.assertFalse(doc.renewal_assignment_id,
                         'the renewal leg left it to the rule')
        self.Auto._enforce_requirements(today, rule)
        self.assertEqual(len(self._open(self.e_one)), 1)
        self.Auto._renewals(today)
        self.Auto._enforce_requirements(today, rule)
        self.assertEqual(len(self._open(self.e_one)), 1)
        # The renewal is due the day the old one runs out.
        self.assertEqual(self._open(self.e_one).due_date,
                         doc.expiry_date)

    def test_t7_day_one_due_date_is_joining_plus_days(self):
        req = self._req(kind='day_one', due_days=14, auto_assign=True,
                        department_ids=[(6, 0, self.dept.ids)])
        today = fields.Date.today()
        joined = today - timedelta(days=3)
        self.e_one.version_id.sudo().write({'contract_date_start': joined})
        self.e_two.version_id.sudo().write(
            {'contract_date_start': today - timedelta(days=400)})
        self.assertEqual(self._state(req, self.e_one), 'not_due')
        self.assertEqual(self._state(req, self.e_two), 'missing')
        self.Auto._enforce_requirements(today, req)
        self.assertFalse(self._open(self.e_one),
                         'not due yet: nothing is assigned')
        self.assertEqual(self._open(self.e_two).due_date,
                         today + timedelta(days=14))
        res = self.Board.assign_from_cell(req.id, self.e_one.id)
        self.assertTrue(res['assignment_id'])
        self.assertEqual(self._open(self.e_one).due_date,
                         joined + timedelta(days=14))


# =========================================================================
#  T8 — hours
# =========================================================================
@tagged('post_install', '-at_install')
class TestT8Hours(RequirementCase):

    def test_t8_course_hours_override_and_stored_at_done(self):
        self.simple.sudo().write({'pb_hours': 0})
        self.simple_lesson.sudo().completion_time = 0.5
        self.simple.invalidate_recordset()
        self.assertAlmostEqual(self.simple._pb_hours(),
                               self.simple.sudo().total_time)
        self.simple.sudo().pb_hours = 2.5
        self.assertEqual(self.simple._pb_hours(), 2.5)
        row = self.Assignment.create({'channel_id': self.simple.id,
                                      'employee_id': self.e_one.id})
        self.assertFalse(row.hours)
        self._finish(self.u_one, self.simple, [self.simple_lesson])
        row._refresh()
        self.assertEqual(row.state, 'done')
        self.assertEqual(row.hours, 2.5)
        self.simple.sudo().pb_hours = 9
        row._refresh()
        self.assertEqual(row.hours, 2.5, 'what they did is not rewritten')

    def test_t8_an_hours_rule_adds_up_finished_hours(self):
        self.simple.sudo().pb_hours = 6
        req = self._req(target='hours', hours_required=8,
                        department_ids=[(6, 0, self.dept.ids)])
        row = self.Assignment.create({'channel_id': self.simple.id,
                                      'employee_id': self.e_one.id})
        self._finish(self.u_one, self.simple, [self.simple_lesson])
        row._refresh()
        cell = self._cells(req, self.e_one)[(req.id, self.e_one.id)]
        self.assertEqual((cell['state'], cell['words']),
                         ('in_progress', '6 of 8 hours'))
        req.hours_required = 6
        self.assertEqual(self._state(req, self.e_one), 'valid')


# =========================================================================
#  T9 — who may do what
# =========================================================================
@tagged('post_install', '-at_install')
class TestT9Access(RequirementCase):

    def test_t9_trainer_reads_manager_writes_employee_cannot(self):
        Model = self.env['pb.training.requirement']
        vals = {'name': 'DEMO by hand', 'company_id': self.company.id,
                'channel_id': self.simple.id}
        with self.assertRaises(AccessError):
            Model.with_user(self.trainer).create(vals)
        made = Model.with_user(self.manager).create(vals)
        self.assertTrue(made)
        self.assertEqual(Model.with_user(self.trainer).browse(made.id).name,
                         'DEMO by hand')
        with self.assertRaises(AccessError):
            Model.with_user(self.u_one).browse(made.id).read(['name'])

    def test_t9_the_verbs_follow_the_same_line(self):
        req = self._req(department_ids=[(6, 0, self.dept.ids)])
        as_trainer = self.Board.with_user(self.trainer)
        with self.assertRaises(AccessError):
            as_trainer.act('save_requirement', {'name': 'DEMO no',
                                                'channel_id': self.simple.id})
        with self.assertRaises(AccessError):
            as_trainer.new_rule()
        with self.assertRaises(AccessError):
            as_trainer.act('set_safety_group',
                           {'employee_ids': [self.e_one.id], 'group': '3'})
        res = as_trainer.act('assign_from_cell', {
            'requirement_id': req.id, 'employee_id': self.e_two.id})
        self.assertIn('is on DEMO Walking the rows', res['message'])
        self.assertFalse(self.Board.with_user(self.u_one).matrix()['allowed'])
        as_manager = self.Board.with_user(self.manager)
        res = as_manager.act('save_requirement', {
            'name': 'DEMO Fire safety', 'kind': 'recurring',
            'channel_id': self.simple.id, 'every_months': 12,
            'department_ids': [self.dept.id], 'safety_groups': ['3', '9']})
        made = self.Req.browse(res['id'])
        self.assertEqual((made.validity_months, made.safety_groups),
                         (12, '3'))
        as_manager.act('set_safety_group', {'employee_ids': [self.e_one.id],
                                            'group': '2'})
        self.assertEqual(self.e_one.sudo().pb_safety_group, '2')


# =========================================================================
#  My learning — one line per course (found live)
# =========================================================================
@tagged('post_install', '-at_install')
class TestRequiredForMe(RequirementCase):

    def test_two_rules_on_one_course_are_one_line_for_the_learner(self):
        first = self._req(name='DEMO Induction', kind='day_one',
                          department_ids=[(6, 0, self.dept.ids)])
        second = self._req(name='DEMO Induction yearly', kind='recurring',
                           every_months=12, validity_months=12,
                           department_ids=[(6, 0, self.dept.ids)])
        self.Assignment.create({
            'channel_id': self.simple.id, 'employee_id': self.e_one.id,
            'due_date': fields.Date.today() + timedelta(days=5)})
        mine = self.env['pb.my.training'].with_user(self.u_one) \
            .required_for_me()
        lines = [r for r in mine['rows'] if r['id'] in (first.id, second.id)]
        self.assertEqual(len(lines), 1, 'one course, one line')
        self.assertEqual(lines[0]['channel_id'], self.simple.id)
        self.assertEqual(lines[0]['state'], 'in_progress')
        # The matrix keeps one column per rule.
        cols = {c['id'] for c in self.Board.matrix({})['columns']}
        self.assertTrue({first.id, second.id} <= cols)


# =========================================================================
#  T13 — Overview and Numbers agree with the matrix
# =========================================================================
@tagged('post_install', '-at_install')
class TestT13Tiles(RequirementCase):

    def test_t13_overview_and_numbers_equal_a_recount(self):
        req = self._req(department_ids=[(6, 0, self.dept.ids)])
        self.Assignment.create({'channel_id': self.simple.id,
                                'employee_id': self.e_one.id})
        cells = self.Req.evaluate(self.Board._live_requirements())
        missing_people = {e for (_r, e), c in cells.items()
                          if c['state'] == 'missing'}
        states = [c['state'] for c in cells.values()
                  if c['state'] != 'exempt']
        measured = len([s for s in states if s in ('valid', 'expiring',
                                                   'missing', 'in_progress')])
        in_date = len([s for s in states if s in ('valid', 'expiring')])
        pct = int(round(100.0 * in_date / measured)) if measured else None
        over = self.Board.overview()
        self.assertEqual(over['required']['people_missing'],
                         len(missing_people))
        self.assertEqual(over['required']['percent'], pct)
        chip = {q['key']: q for q in over['queue']}['required']
        self.assertEqual(chip['count'], len(missing_people))
        self.assertEqual((chip['lens'], chip['state']), ('required', 'missing'))
        self.assertIn('missing required training', over['headline'])
        nums = self.env['pb.training.analytics'].get_numbers()
        self.assertEqual(nums['required']['percent'], pct)
        self.assertEqual(nums['required']['missing'],
                         states.count('missing'))
        self.assertTrue(req)


# =========================================================================
#  T14 — the exports
# =========================================================================
@tagged('post_install', '-at_install')
class TestT14Exports(RequirementCase):

    def _sheets(self, res):
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(base64.b64decode(
            res['file_b64'])))

    def test_t14_spreadsheet_and_pdf_for_zero_one_and_many(self):
        nobody = self.Board.matrix_xlsx({'requirement_id': -1})
        self.assertEqual(nobody['rows'], 0)
        self.assertEqual(self._sheets(nobody).sheetnames,
                         ['Required training', 'Gaps'])
        self.assertTrue(self.Board.matrix_pdf({'requirement_id': -1})
                        ['file_b64'])
        req = self._req(department_ids=[(6, 0, self.sub.ids)])
        one = self.Board.matrix_xlsx({'requirement_id': req.id})
        self.assertEqual(one['rows'], 1)
        wb = self._sheets(one)
        self.assertEqual(wb['Gaps'].cell(row=2, column=3).value, req.name)
        dept = self.env['hr.department'].create({'name': 'DEMO L2 Many'})
        self.env['hr.employee'].create([
            {'name': 'DEMO L2 M %03d' % i, 'department_id': dept.id}
            for i in range(250)])
        many = self._req(department_ids=[(6, 0, dept.ids)])
        res = self.Board.matrix_xlsx({'requirement_id': many.id})
        self.assertEqual(res['rows'], 250)
        pdf = self.Board.matrix_pdf({'requirement_id': many.id})
        self.assertTrue(pdf['file_b64'])
        export = self.env['pb.training.matrix.export'].create(
            {'filters_json': '{"requirement_id": %s}' % many.id})
        self.assertEqual(len(export.data()['pages_rows']), 7, '40 a page')

    def test_t14_the_pdf_is_one_document_paginated_by_css(self):
        """FOUND LIVE: the server's wkhtmltopdf is the unpatched-QT build
        and refuses several documents in one PDF. One record id, one
        `article`, and a CSS page break after every block but the last."""
        dept = self.env['hr.department'].create({'name': 'DEMO L2 Pages'})
        self.env['hr.employee'].create([
            {'name': 'DEMO L2 P %03d' % i, 'department_id': dept.id}
            for i in range(85)])
        req = self._req(department_ids=[(6, 0, dept.ids)])
        seen = []

        def fake(report_self, report_ref, res_ids=None, data=None):
            seen.append(list(res_ids or []))
            return b'%PDF-1.4 DEMO', 'pdf'

        with patch.object(type(self.env['ir.actions.report']),
                          '_render_qweb_pdf', fake):
            res = self.Board.matrix_pdf({'requirement_id': req.id})
        self.assertEqual([len(ids) for ids in seen], [1])
        self.assertTrue(base64.b64decode(res['file_b64']).startswith(b'%PDF'))
        export = self.env['pb.training.matrix.export'].create(
            {'filters_json': '{"requirement_id": %s}' % req.id})
        html, _kind = self.env['ir.actions.report']._render_qweb_html(
            'pb_training.report_required_matrix', export.ids)
        html = html if isinstance(html, bytes) else str(html).encode()
        self.assertEqual(html.count(b'class="article'), 1,
                         'one document, never one per page')
        self.assertEqual(html.count(b'page-break-after: always'), 2,
                         '85 people = 3 blocks of up to 40 = 2 breaks')

    def test_t14_the_pdf_reads_as_the_reader_no_company_mixing(self):
        other = self.env['res.company'].create({'name': 'DEMO L2 Other Co'})
        stranger = self.env['hr.employee'].create(
            {'name': 'DEMO L2 Stranger', 'company_id': other.id})
        self.Req.create({'name': 'DEMO other co rule',
                         'company_id': other.id,
                         'channel_id': self.simple.id})
        mine = self._req(department_ids=[(6, 0, self.dept.ids)])
        export = self.env['pb.training.matrix.export'].with_context(
            allowed_company_ids=[self.company.id]).create({})
        data = export.data()
        ids = {r['employee_id'] for r in data['rows']}
        self.assertNotIn(stranger.id, ids)
        self.assertIn(mine.id, [c['id'] for c in data['columns']])
        self.assertNotIn('sudo()', re.sub(
            r'#.*', '', _src('models', 'pb_training_required.py').split(
                'def _pdf(self)')[1].split('def _num')[0]))


# =========================================================================
#  T15 — house rules
# =========================================================================
@tagged('post_install', '-at_install')
class TestT15HouseRules(RequirementCase):

    FILES = [('static', 'src', 'xml', 'required_lens.xml'),
             ('static', 'src', 'xml', 'required_settings.xml'),
             ('static', 'src', 'xml', 'safety_group_chip.xml'),
             ('static', 'src', 'js', 'required_lens.js'),
             ('static', 'src', 'js', 'required_settings.js'),
             ('static', 'src', 'js', 'safety_group_chip.js'),
             ('data', 'pack_vn.xml'), ('data', 'pack_in.xml'),
             ('data', 'pack_id.xml'), ('data', 'pack_sg.xml'),
             ('report', 'required_matrix.xml'),
             ('views', 'requirement_views.xml')]

    def test_t15_no_vendor_name_and_no_emoji(self):
        emoji = re.compile('[\U0001F300-\U0001FAFF☀-➿]')
        for parts in self.FILES:
            text = _src(*parts)
            self.assertNotIn('Odoo', text, '/'.join(parts))
            self.assertFalse(emoji.search(text), '/'.join(parts))

    def test_t15_vietnamese_words_and_the_rail_label(self):
        po = _src('i18n', 'vi_VN.po')
        self.assertIn('msgid "Required"\nmsgstr "Bắt buộc"', po)
        self.assertLessEqual(len('Bắt buộc'.split()), 2)
        for word in ('Matrix', 'Comes round again', 'Missing required training',
                     'Safety training group', 'Print for an auditor',
                     'Download the spreadsheet', 'Country reports'):
            match = re.search(r'msgid "%s"\nmsgstr "([^"]*)"' % re.escape(word),
                              po)
            self.assertTrue(match and match.group(1), word)

    def test_t15_reason_required_is_said_in_the_readers_language(self):
        from ..models.training_common import ASSIGN_REASON_LABEL
        self.assertEqual(ASSIGN_REASON_LABEL['required'], 'Required')
        vi = self.env(context=dict(self.env.context, lang='vi_VN'))
        if self.env['res.lang']._lang_get('vi_VN'):
            from ..models.training_common import reason_word
            self.assertEqual(reason_word(vi, 'required'), 'Bắt buộc')

"""Configuration-owned workbooks; imports are server-side reviewed patches.

The three storage models deliberately have no user ACL. All access goes through
the scoped facade below. Workbook metadata is a reference, never write authority.
"""
import base64
import calendar
import hashlib
import io
import json
import math
import re
import textwrap
import uuid
import zipfile
from datetime import date, datetime, timedelta

from odoo import api, fields, models, _
from odoo.exceptions import UserError, AccessError
from odoo.addons.pb_hr_payroll_formula.models.pay_period import default_standard_work_days

LIMIT = 10000
TABS = {'time': 'Time', 'adjustments': 'Monthly adjustments'}
RECORD_TAB = 'Employee records'
LAYOUT_LANES = ('time', 'adjustments', 'records', 'available')
IDENTITY = {'barcode', 'employee_id', 'pb_source_ref', 'identification_id',
            'work_email', 'private_email'}
TIME = re.compile(r'days?|hours?|\bhrs?\b|attendance|timesheet|ngày|giờ|công', re.I)
VARIABLE = re.compile(r'bonus|reimburs|deduct|arrear|adjust|severance|overtime|cpf|self.help|shg|tax|thuế|khấu trừ|thưởng', re.I)
RECURRING = re.compile(r'salary|basic|wage|fixed|recurring|allowance|phụ cấp|lương', re.I)
SALARY = re.compile(r'\bbasic\b|base.*salary|contract.*salary|monthly.*salary|^salary$|\bwage\b|lương hợp đồng|lương cơ bản', re.I)


def salary_inputs(config):
    """Find salary in the scheme currency, never an unused foreign reference."""
    candidates = config.rule_ids.filtered(lambda r: r.column_type == 'input'
        and SALARY.search(r.name or '') and not VARIABLE.search(r.name or '')
        and not re.search(r'\bannual|\byearly|\bhourly|\bdaily|\bnet\b', r.name or '', re.I))
    currency = config.currency_id.name
    explicit = candidates.filtered(lambda r: re.search(r'\b%s\b' % re.escape(currency), r.name or '', re.I)) if currency else candidates.browse()
    if explicit:
        return explicit
    currencies = set(config.env['res.currency'].sudo().with_context(active_test=False).search([]).mapped('name'))
    return candidates.filtered(lambda r: not (set(re.findall(r'\b[A-Z]{3}\b', r.name or '')) & currencies))


def plain(value):
    return json.loads(json.dumps(value, default=str))


class TemplateLayout(models.Model):
    _name = 'pb.payroll.template.layout'
    _description = 'Payroll workbook layout'
    config_id = fields.Many2one('hr.formula.config', required=True, ondelete='cascade')
    company_id = fields.Many2one('res.company', required=True, ondelete='cascade')
    payload = fields.Json(default=dict)
    _unique = models.Constraint('UNIQUE(config_id, company_id)', 'One layout per payroll configuration.')


class MonthlyInputs(models.Model):
    _name = 'pb.payroll.month.inputs'
    _description = 'Selected period payroll inputs'
    config_id = fields.Many2one('hr.formula.config', required=True, ondelete='cascade')
    company_id = fields.Many2one('res.company', required=True, ondelete='cascade')
    employee_id = fields.Many2one('hr.employee', required=True, ondelete='cascade')
    period = fields.Char(required=True, index=True)
    payload = fields.Json(default=dict)
    record_overrides = fields.Json(default=dict)
    origins = fields.Json(default=dict)
    review_snapshot = fields.Json(default=dict)
    _unique = models.Constraint('UNIQUE(config_id, company_id, employee_id, period)', 'One input record per person and pay period.')


class WorkbookExchange(models.Model):
    _name = 'pb.payroll.workbook.exchange'
    _description = 'Payroll workbook download and import receipt'
    _order = 'id desc'
    token = fields.Char(required=True, index=True, default=lambda s: uuid.uuid4().hex)
    company_id = fields.Many2one('res.company', required=True, ondelete='cascade')
    config_id = fields.Many2one('hr.formula.config', required=True, ondelete='cascade')
    owner_id = fields.Many2one('res.users', required=True, ondelete='cascade')
    kind = fields.Selection([('export', 'Download'), ('review', 'Import review'), ('applied', 'Confirmed import')], required=True)
    payload = fields.Json(default=dict)
    _unique = models.Constraint('UNIQUE(token)', 'Workbook reference must be unique.')


class PayrollTemplates(models.AbstractModel):
    _name = 'pb.payroll.templates'
    _description = 'Payroll templates workspace'

    def _config(self, config_id, write=False):
        desk = self.env['pb.records.desk']
        desk._check_read()
        cfg = desk._configs(int(config_id)).filtered(
            lambda c: not c.company_id or c.company_id == self.env.company)
        if not cfg:
            raise AccessError(_('Choose a payroll configuration in the current company.'))
        if write:
            self.env['pb.payrun.wizard']._require_runner()
        return cfg

    def _period(self, period):
        try:
            start = datetime.strptime(str(period), '%Y-%m').date()
            if start.strftime('%Y-%m') != period:
                raise ValueError()
        except (ValueError, TypeError):
            raise UserError(_('Choose a valid pay month.'))
        return start, start.replace(day=calendar.monthrange(start.year, start.month)[1])

    def _lock(self, cfg):
        self.env.cr.execute('SELECT pg_advisory_xact_lock(%s, %s)', (self.env.company.id, cfg.id))

    def _roster(self, cfg, period):
        start, end = self._period(period)
        ids = self.env['pb.payrun.wizard']._eligible_employees(
            formula_config_id=cfg.id, date_start=start, date_end=end)
        if not ids and cfg.state == 'draft':
            preview = self.env['pb.payroll.template.layout'].sudo().search([
                ('config_id', '=', cfg.id), ('company_id', '=', self.env.company.id)], limit=1)
            ids = (preview.payload or {}).get('preview_employee_ids', []) if preview else []
        people = self.env['hr.employee'].sudo().search([
            ('id', 'in', ids), ('company_id', '=', self.env.company.id)], order='name, id')
        if len(people) > LIMIT:
            raise UserError(_('This workbook supports up to 10,000 employees. Narrow the payroll population first.'))
        return people

    def _code(self, employee):
        return str(employee.barcode or employee.employee_id or '').strip()

    def _cards(self, cfg):
        return {k: dict(v, label=(v.get('component') or {}).get('name') or v['label'])
                for k, v in self.env['pb.records.desk']._cards(cfg.id).items()
                if v.get('field') not in IDENTITY}

    def _identity_rules(self, cfg):
        found = {}
        for m in self.env['pb.records.desk']._mappings(cfg.id):
            name = m.target_field_id.name
            if m.target_model_id.model == 'hr.employee' and name in IDENTITY | {'name'}:
                found[m.component_id.id] = 'name' if name == 'name' else 'email' if 'email' in name else 'code'
        for r in cfg.rule_ids.filtered(lambda r: r.column_type == 'input'):
            text = (r.name or '').strip().lower()
            if text in ('employee id', 'employee code', 'mã nhân viên', 'code'):
                found.setdefault(r.id, 'code')
            elif text in ('name', 'employee name', 'full name', 'họ và tên', 'họ và tên (full name)'):
                found.setdefault(r.id, 'name')
        return found

    def _identity_columns(self, cfg, mode='monthly'):
        """Visible identity headings belong to the scheme, not the DB adapter.

        The field is retained separately so translated/renamed headings never
        change where a cell's value comes from. Unmapped schemes retain the
        legacy fallback; no synthetic component is created for that fallback.
        """
        roles = self._identity_rules(cfg)
        mappings = self.env['pb.records.desk']._mappings(cfg.id)
        result = []
        for kind, fallback in [('code', _('Employee ID')),
                               ('email', _('Email')) if mode == 'records'
                               else ('name', _('Employee name'))]:
            rules = cfg.rule_ids.filtered(lambda r: r.column_type == 'input'
                and roles.get(r.id) == kind).sorted(key=lambda r: (r.sequence, r.id))
            rule = rules[:1]
            mapping = mappings.filtered(lambda m: m.component_id == rule
                and m.destination_type == 'field'
                and m.target_model_id.model == 'hr.employee'
                and m.target_field_id.name in (IDENTITY | {'name'}))[:1] if rule else mappings.browse()
            result.append({'kind': kind, 'label': rule.name or rule.code if rule else fallback,
                           'rule_id': rule.id if rule else False,
                           'field': mapping.target_field_id.name if mapping else False})
        return result

    def _identity_value(self, employee, column):
        field = column.get('field')
        if field:
            value = employee[field]
        else:
            value = {'code': self._code(employee), 'name': employee.name,
                     'email': employee.work_email or ''}[column['kind']]
        text = str(value or '')
        return text.strip() if column['kind'] == 'code' else text

    @staticmethod
    def _workbook_identity(snapshot, mode='monthly'):
        """Use the export receipt, so old downloads and later renames work."""
        return snapshot.get('identity_columns') or [
            {'kind': 'code', 'label': 'Employee ID'},
            {'kind': 'email', 'label': 'Email'} if mode == 'records'
            else {'kind': 'name', 'label': 'Employee name'}]

    def _validate_identity_headers(self, headers, columns):
        required = [c['label'] for c in columns] + ['_row']
        return all(headers.count(label) == 1 for label in required)

    def _workbook_code_unchanged(self, employee, snapshot_person, column, value):
        # Older export receipts used the engine identifier for both values.
        old = snapshot_person.get('identity_values', {}).get('code', snapshot_person['code'])
        current = self._identity_value(employee, column) if column.get('field') else self._code(employee)
        return str(value or '').strip() == old and current == old

    def _catalog(self, cfg):
        cards = self._cards(cfg)
        by_rule = {}
        for fid, card in cards.items():
            if card.get('component'):
                by_rule.setdefault(card['component']['id'], fid)
        # Several payroll inputs may read the same saved field. The Records
        # Desk presents one destination card; retain each input's own label and
        # classify every mapped alias as record context, never monthly input.
        for mapping in self.env['pb.records.desk']._mappings(cfg.id):
            fid = ('b:' + mapping.bank_role) if mapping.destination_type == 'bank_account' else (
                'f:%s:%s' % (mapping.target_model_id.model, mapping.target_field_id.name))
            if fid in cards and mapping.component_id:
                by_rule[mapping.component_id.id] = fid
        result = []
        identity_rules = self._identity_rules(cfg)
        for rule in cfg.rule_ids.filtered(lambda r: r.column_type == 'input'):
            sources = rule.source_ids.filtered(lambda s: s.kind == 'excel') if 'source_ids' in rule._fields else self.env['hr.formula.rule.source'].sudo().search([('rule_id', '=', rule.id), ('kind', '=', 'excel')])
            if rule.id in identity_rules:
                continue
            text = '%s %s' % (rule.name, rule.code)
            card_id = by_rule.get(rule.id)
            if card_id and cards[card_id].get('field') in IDENTITY:
                continue
            quantity = rule.value_kind == 'quantity' or bool(TIME.search(text) and rule.value_kind in ('integer', 'decimal', 'money'))
            recurring = bool(card_id)
            # Mapped dates, names, status and recurring salary are read from records.
            default = 'records' if recurring else 'available' if rule.value_kind in ('date', 'text', 'boolean') else ('time' if quantity else 'adjustments')
            result.append({'id': 'rule:%s' % rule.id, 'label': rule.name or rule.code,
                           'code': rule.code, 'kind': rule.value_kind,
                           'required': bool(rule.is_required),
                           'source_keys': [s.key for s in sources if s.key] or self.env['hr.payroll.import.batch']._rule_header_candidates(rule, allow_column_letter=False),
                           'record_id': card_id, 'readonly': recurring,
                           'record_eligible': recurring,
                           'record_group': cards[card_id]['group'] if card_id else '',
                           'default': default,
                           'reason': _('Saved employee / contract context') if recurring else (_('Time quantity') if quantity else _('Monthly input — review the suggested placement'))})
        represented = {c['record_id'] for c in result if c.get('record_id')}
        kinds = {'date': 'date', 'datetime': 'date', 'integer': 'integer',
                 'float': 'decimal', 'monetary': 'money', 'amount': 'money', 'boolean': 'boolean'}
        for fid, card in cards.items():
            if fid in represented or fid == 'f:hr.employee:name':
                continue
            result.append({'id': 'record:' + fid, 'label': card['label'], 'code': '',
                           'kind': kinds.get(card.get('ttype'), 'text'), 'source_keys': [],
                           'record_id': fid, 'readonly': True, 'record_eligible': True,
                           'record_group': card['group'], 'default': 'records',
                           'reason': _('Saved employee / contract context')})
        return result

    def _layout(self, cfg, catalog):
        row = self.env['pb.payroll.template.layout'].sudo().search([
            ('config_id', '=', cfg.id), ('company_id', '=', self.env.company.id)], limit=1)
        saved = row.payload or {} if row else {}
        known = {c['id']: c for c in catalog}
        result = {k: [] for k in LAYOUT_LANES}
        # Migrate old context placements without changing any rule or source.
        # Old layouts had no records lane; their unused master fields start in
        # the new snapshot. Once a records lane is saved, Available stays opt-out.
        for lane in LAYOUT_LANES:
            for fid in saved.get(lane, []):
                if fid not in known or any(fid in ids for ids in result.values()):
                    continue
                record = known[fid]['record_eligible']
                target = ('records' if record and (lane != 'available' or 'records' not in saved)
                          else 'available' if not record and lane == 'records' else lane)
                if record and target in TABS:
                    target = 'records'
                result[target].append(fid)
        used = {i for a in result.values() for i in a}
        for c in catalog:
            if c['id'] not in used:
                result[c['default']].append(c['id'])
        return result

    def _fingerprint(self, layout, catalog):
        return hashlib.sha256(json.dumps([layout, catalog], sort_keys=True).encode()).hexdigest()

    @api.model
    def workspace(self, config_id=0, period=''):
        desk = self.env['pb.records.desk']
        desk._check_read()
        configs = desk._configs().filtered(lambda c: not c.company_id or c.company_id == self.env.company)
        if not configs:
            return {'configs': [], 'fields': [], 'layout': {}, 'people': 0}
        cfg = self._config(config_id or configs[0].id)
        period = period or fields.Date.context_today(self).strftime('%Y-%m')
        people = self._roster(cfg, period)
        catalog = self._catalog(cfg)
        cards = self._cards(cfg)
        preview = self.env['pb.payroll.template.layout'].sudo().search([
            ('config_id', '=', cfg.id), ('company_id', '=', self.env.company.id)], limit=1)
        expected = (preview.payload or {}).get('preview_expected_count') if preview and cfg.state == 'draft' else 0
        return {'configs': [{'id': c.id, 'name': c.name, 'country': c.country_code, 'state': c.state} for c in configs],
                'config_id': cfg.id, 'name': cfg.name, 'country': cfg.country_code,
                'state': cfg.state,
                'currency': cfg.currency_id.name, 'period': period, 'people': len(people),
                'roster_note': (_('Draft review roster: %s existing Employee records out of %s source people. Payroll assignments are unchanged.') % (len(people), expected)) if expected else '',
                'record_columns': len(self._layout(cfg, catalog)['records']) + int('f:hr.employee:name' in cards), 'fields': catalog,
                'identity_labels': {c['kind']: c['label'] for c in self._identity_columns(cfg)},
                'record_identity_labels': {c['kind']: c['label'] for c in self._identity_columns(cfg, 'records')},
                'calculated_fields': [{'id': r.id, 'label': r.name or r.code,
                                       'code': r.code, 'column': r.column_letter,
                                       'formula': r.excel_formula or ''}
                                      for r in cfg.rule_ids if r.column_type != 'input'],
                'layout': self._layout(cfg, catalog),
                'can_layout': self.env.user.has_group('base.group_system') or self.env.user.has_group('pb_hr_payroll_base.group_payroll_base_manager'),
                'can_export': desk._may_export(),
                'history': self.history(cfg.id)}

    @api.model
    def save_layout(self, config_id, layout):
        cfg = self._config(config_id)
        if not (self.env.user.has_group('base.group_system') or self.env.user.has_group('pb_hr_payroll_base.group_payroll_base_manager')):
            raise AccessError(_('A Payroll Manager can change workbook layouts.'))
        self._lock(cfg)
        catalog = self._catalog(cfg)
        ids = [i for k in LAYOUT_LANES for i in layout.get(k, [])]
        if len(ids) != len(set(ids)) or set(ids) != {c['id'] for c in catalog}:
            raise UserError(_('The field list changed. Reload the layout and try again.'))
        by_id = {c['id']: c for c in catalog}
        for lane in ('time', 'adjustments', 'records'):
            for fid in layout.get(lane, []):
                if bool(by_id[fid]['record_eligible']) != (lane == 'records'):
                    raise UserError(_('Employee records accepts only saved employee, contract, contract component and bank fields. Monthly inputs belong on Time or Monthly adjustments.'))
        row = self.env['pb.payroll.template.layout'].sudo().search([
            ('config_id', '=', cfg.id), ('company_id', '=', self.env.company.id)], limit=1)
        vals = {'payload': dict(row.payload or {} if row else {}, **{k: layout[k] for k in LAYOUT_LANES})}
        if row:
            row.write(vals)
        else:
            self.env['pb.payroll.template.layout'].sudo().create(dict(vals, config_id=cfg.id, company_id=self.env.company.id))
        return {'ok': True}

    def _exchange(self, cfg, kind, payload):
        return self.env['pb.payroll.workbook.exchange'].sudo().create({
            'config_id': cfg.id, 'company_id': self.env.company.id,
            'owner_id': self.env.uid, 'kind': kind, 'payload': plain(payload)})

    def _get_exchange(self, cfg, token, kind, shared_export=False):
        domain = [
            ('token', '=', str(token)), ('config_id', '=', cfg.id),
            ('company_id', '=', self.env.company.id), ('kind', '=', kind)]
        if shared_export and kind == 'export':
            # Colleagues may fill the same monthly workbook. The download is
            # only a reference: review rechecks the current roster and values,
            # then creates a fresh receipt owned by the uploading payroll user.
            self._config(cfg.id, write=True)
            if not self.env['pb.records.desk']._may_export():
                raise AccessError(_('You need the Records export ability to use shared payroll workbooks.'))
        else:
            domain.append(('owner_id', '=', self.env.uid))
        row = self.env['pb.payroll.workbook.exchange'].sudo().search(domain, limit=1)
        if not row:
            raise UserError(_('This workbook reference is unavailable. Download a fresh copy using your account.'))
        return row

    def _current(self, cfg, period, employees):
        result = {e.id: {} for e in employees}
        # Earlier uploads from the existing import flow are considered ONLY for
        # this exact period. Never borrow a last-run value from another month.
        start, end = self._period(period)
        batches = self.env['hr.payroll.import.batch'].sudo().search([
            ('formula_config_id', '=', cfg.id), ('date_from', '=', start), ('date_to', '=', end),
            ('company_id', '=', self.env.company.id), ('state', '=', 'done')], order='id')
        catalog = self._catalog(cfg)
        probe = self.env['hr.payroll.import.batch'].new({'formula_config_id': cfg.id})
        for line in batches.mapped('import_line_ids').filtered(lambda l: l.employee_id.id in result):
            raw = line.get_raw_data() or {}
            for c in catalog:
                value, matched = probe._lookup_in_blob(raw, c['source_keys'])
                if matched and value not in (None, ''):
                    result[line.employee_id.id][c['id']] = plain(value)
        rows = self.env['pb.payroll.month.inputs'].sudo().search([
            ('config_id', '=', cfg.id), ('company_id', '=', self.env.company.id),
            ('period', '=', period), ('employee_id', 'in', employees.ids)])
        for row in rows:
            result[row.employee_id.id].update(row.payload or {})
        return result

    def _record_values(self, cfg, people, cards):
        desk = self.env['pb.records.desk']
        contracts = desk._io_contracts(people)
        probe = desk._probe(cfg.id)
        result = {}
        for employee in people:
            contract = contracts.get(employee.id) or self.env['hr.contract']
            values = {}
            for fid, card in cards.items():
                if fid.startswith('c:') and card.get('ttype') == 'amount':
                    # Records exports leave zero component amounts visually
                    # blank. Payroll must retain the numeric zero: coverage and
                    # reimbursement flags can require an explicit 0. The same
                    # reader returns None when there is no contract.
                    value = desk._read_cell(probe, card, employee, contract)['v']
                else:
                    value = desk._io_export_value(probe, card, employee, contract)[0]
                values[fid] = plain(value)
            result[employee.id] = values
        return result

    @api.model
    def download(self, config_id, period='', mode='monthly'):
        cfg = self._config(config_id)
        desk = self.env['pb.records.desk']
        if not desk._may_export():
            raise AccessError(_('You need the Records export ability to download payroll workbooks.'))
        if mode not in ('records', 'monthly', 'template'):
            raise UserError(_('Choose an employee or monthly workbook.'))
        reusable = mode == 'template'
        period = period or fields.Date.context_today(self).strftime('%Y-%m')
        people = self._roster(cfg, period)
        identity_columns = self._identity_columns(cfg, mode)
        codes = [self._identity_value(e, identity_columns[0]) for e in people]
        if any(not c for c in codes) or len(set(codes)) != len(codes):
            raise UserError(_('Every employee needs a unique value for %s before downloading. Update missing or duplicate values in the mapped employee field.') % identity_columns[0]['label'])
        cards = self._cards(cfg)
        records = self._record_values(cfg, people, cards)
        catalog = self._catalog(cfg)
        layout = self._layout(cfg, catalog)
        monthly = {e.id: {} for e in people} if reusable else self._current(cfg, period, people)
        by_id = {c['id']: c for c in catalog}
        selected_cards = [(by_id[fid]['record_id'], dict(cards[by_id[fid]['record_id']], label=by_id[fid]['label'])) for fid in layout['records']]
        if 'f:hr.employee:name' in cards:
            selected_cards.insert(0, ('f:hr.employee:name', cards['f:hr.employee:name']))
        if mode == 'records':
            record_columns, seen_destinations = [], set()
            for index, (fid, card) in enumerate(selected_cards):
                # Preserve stable native-field IDs for ordinary columns. An
                # alias gets its own workbook ID and the same write destination.
                # This keeps two payroll labels for one saved field reviewable.
                column_id = fid if fid not in seen_destinations else 'alias:%s:%s' % (index, fid)
                seen_destinations.add(fid)
                record_columns.append({'id': column_id, 'record_id': fid,
                    'label': card['label'], 'readonly': not card.get('editable'),
                    'record_group': card['group']})
            columns = {RECORD_TAB: record_columns}
        else:
            columns = {**{TABS[k]: [by_id[i] for i in layout[k]] for k in TABS},
                       RECORD_TAB: [by_id[i] for i in layout['records']]}
        snapshot = {'mode': 'monthly' if reusable else mode, 'period': '' if reusable else period, 'reusable': reusable, 'layout': self._fingerprint(layout, catalog), 'columns': columns, 'identity_columns': identity_columns, 'employees': {}}
        snapshot['readonly_sheets'] = [] if mode == 'records' else [RECORD_TAB]
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Protection, Alignment
        from openpyxl.comments import Comment
        from openpyxl.utils import get_column_letter
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        record_colors = {'employee': ('325F96', 'EAF2FF'), 'contract': ('2F7257', 'EAF7EF'),
                         'component': ('70538A', 'F3ECFA'), 'bank': ('916326', 'FFF3DD')}
        for sheet_name, cols in columns.items():
            ws = wb.create_sheet(sheet_name)
            master_snapshot = mode != 'records' and sheet_name == RECORD_TAB
            identity = [c['label'] for c in identity_columns]
            headers = identity + [c['label'] for c in cols] + ['_row']
            # Record columns are keyed by their saved field IDs in comments.
            # Two destinations may legitimately share a component name (e.g.
            # Wage and its contract component). Only identity headings must be
            # unique, because those are matched by heading on import.
            if not self._validate_identity_headers(headers, identity_columns):
                raise UserError(_('Two workbook columns have the same component name. Give them distinct names in the payroll scheme before downloading.'))
            ws.append(headers)
            for ci, cell in enumerate(ws[1], 1):
                cell.fill = PatternFill('solid', fgColor='5A4BB0')
                cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                ws.column_dimensions[get_column_letter(ci)].width = 26
            # Bilingual component names often need more than two lines. Size
            # the header to the exported width so the monthly workbook keeps
            # the complete field name readable without manual row resizing.
            lines = max((sum(max(1, len(textwrap.wrap(part, width=24)))
                             for part in str(label).split('\n')) for label in headers), default=1)
            ws.row_dimensions[1].height = max(32, lines * 15 + 8)
            for ci, c in enumerate(cols, len(identity) + 1):
                ws.cell(1, ci).comment = Comment('field: %s\nBlank = keep current value; 0 = update. %s' % (c['id'], 'Read-only context.' if c.get('readonly') else ''), 'Payobook')
                if master_snapshot:
                    color = record_colors.get(c.get('record_group'), ('5A4BB0', 'F0EEF8'))[0]
                    ws.cell(1, ci).fill = PatternFill('solid', fgColor=color)
            for e in people:
                key = str(e.id)
                if key not in snapshot['employees']:
                    snapshot['employees'][key] = {'code': self._code(e), 'email': e.work_email or '', 'name': e.name,
                                                 'identity_values': {c['kind']: self._identity_value(e, c) for c in identity_columns},
                                                 'records': records[e.id], 'monthly': monthly[e.id], 'row': uuid.uuid4().hex}
                person = snapshot['employees'][key]
                values = records[e.id] if mode == 'records' else monthly[e.id]
                vals = []
                for c in cols:
                    value = records[e.id].get(c.get('record_id')) if c.get('readonly') and mode != 'records' else values.get(c.get('record_id') or c['id'])
                    if mode == 'monthly' and value is None and not c.get('readonly') and re.search(r'standard.*days|standard working days', c['label'], re.I):
                        start, end = self._period(period)
                        value = default_standard_work_days(start, end)
                    vals.append(value)
                ws.append([person['identity_values'][c['kind']] for c in identity_columns] + vals + [person['row']])
                for ci, cell in enumerate(ws[ws.max_row], 1):
                    if isinstance(cell.value, str):
                        cell.data_type = 's'  # Excel must never execute imported text.
                    if ci == 1:
                        cell.number_format = '@'
                    editable = not master_snapshot and len(identity) < ci <= len(identity) + len(cols) and not cols[ci - len(identity) - 1].get('readonly')
                    cell.protection = Protection(locked=not editable)
                    cell.font = Font(name='Calibri', color='241F52' if editable else '63718A')
                    if master_snapshot and len(identity) < ci <= len(identity) + len(cols):
                        c = cols[ci - len(identity) - 1]
                        cell.fill = PatternFill('solid', fgColor=record_colors.get(c.get('record_group'), ('5A4BB0', 'F0EEF8'))[1])
                    elif editable:
                        cell.fill = PatternFill('solid', fgColor='FFF7DD')
            ws.column_dimensions[get_column_letter(len(headers))].hidden = True
            ws.freeze_panes = 'C2'
            ws.auto_filter.ref = 'A1:%s%s' % (get_column_letter(len(headers) - 1), max(1, ws.max_row))
            # Monthly workbooks are working sheets: Excel must allow normal
            # editing/pasting, including empty exports and additional rows.
            # Identity, read-only context and roster membership are enforced
            # by review() against the server-owned download snapshot.
            # Employee-record workbooks retain their protected ID/email cells.
            ws.protection.sheet = mode == 'records' or master_snapshot
            if master_snapshot:
                ws.sheet_properties.tabColor = '2F7257'
                ws.protection.set_password(uuid.uuid4().hex)
                ws['A1'].comment = Comment(
                    'Read-only master records. This entire worksheet is ignored during monthly upload. '
                    'Colours: blue = Employee; green = Contract; purple = Contract component; gold = Bank. '
                    'Use Download records for a separate reviewed master-data update.', 'Payobook')
            ws.protection.autoFilter = False
            ws.protection.sort = False
        exchange = self._exchange(cfg, 'export', snapshot)
        meta = wb.create_sheet('_Payobook')
        meta.append(['reference', exchange.token])
        meta.append(['configuration', cfg.name])
        meta.append(['pay month', '' if reusable else period])
        meta.sheet_state = 'veryHidden'
        out = io.BytesIO()
        wb.save(out)
        filename = 'Payobook_%s_monthly_template.xlsx' % cfg.country_code if reusable else 'Payobook_%s_%s_%s.xlsx' % (cfg.country_code, mode, period)
        return {'ok': True, 'filename': filename, 'file_b64': base64.b64encode(out.getvalue()).decode(), 'rows': len(people)}

    def _read(self, file_b64):
        try:
            if len(file_b64 or '') > 14 * 1024 * 1024:
                raise ValueError()
            raw = base64.b64decode(file_b64, validate=True)
            if len(raw) > 10 * 1024 * 1024:
                raise ValueError()
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(f.file_size for f in z.infolist()) > 80 * 1024 * 1024:
                    raise ValueError()
            import openpyxl
            return openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
        except Exception as err:
            raise UserError(_('Choose a valid .xlsx workbook under 10 MB.')) from err

    @api.model
    def is_template(self, file_b64):
        self.env['pb.records.desk']._check_read()
        return '_Payobook' in self._read(file_b64).sheetnames

    def _numeric(self, value, kind):
        if kind not in ('money', 'decimal', 'integer', 'quantity', 'rate'):
            return str(value) if isinstance(value, (date, datetime)) else value
        if isinstance(value, bool):
            raise ValueError('Enter a number.')
        n = float(value)
        if not math.isfinite(n) or (kind == 'integer' and n != int(n)):
            raise ValueError('Enter a valid number.')
        return n

    @api.model
    def review(self, config_id, period, mode, file_b64, filename=''):
        cfg = self._config(config_id, write=mode == 'monthly')
        self._period(period)
        if mode == 'monthly':
            return self._review_monthly(cfg, period, file_b64, filename)
        wb = self._read(file_b64)
        if '_Payobook' not in wb:
            raise UserError(_('Use a workbook downloaded from Payroll templates. Other records files can be imported in Records Desk.'))
        snapshot = self._get_exchange(cfg, wb['_Payobook']['B1'].value, 'export').payload
        identity_columns = self._workbook_identity(snapshot, mode)
        identity_labels = {c['kind']: c['label'] for c in identity_columns}
        if mode != snapshot['mode'] or (mode == 'monthly' and not snapshot.get('reusable') and period != snapshot['period']):
            raise UserError(_('This workbook belongs to a different workflow or pay month. Download the selected month and try again.'))
        catalog = self._catalog(cfg)
        layout = self._layout(cfg, catalog)
        if mode == 'monthly' and snapshot['layout'] != self._fingerprint(layout, catalog):
            raise UserError(_('The monthly layout changed after this download. Download a fresh workbook.'))
        people = self._roster(cfg, period)
        allowed = set(people.ids)
        by_row = {v['row']: (int(k), v) for k, v in snapshot['employees'].items()}
        cards = self._cards(cfg)
        records = self._record_values(cfg, people, cards)
        monthly = self._current(cfg, period, people)
        changes, issues, blanks, unchanged = [], [], 0, 0
        change_keys = set()
        seen_tabs = {}
        for sheet_name, cols in snapshot['columns'].items():
            if sheet_name not in wb:
                issues.append(_('Missing worksheet: %s') % sheet_name)
                continue
            ws = wb[sheet_name]
            if ws.max_row > LIMIT + 1 or ws.max_column > 1000:
                raise UserError(_('The worksheet exceeds the supported size.'))
            header = [c.value for c in ws[1]]
            ids = {}
            for ci, cell in enumerate(ws[1]):
                if cell.comment and cell.comment.text.startswith('field: '):
                    fid = cell.comment.text.splitlines()[0][7:]
                    if fid in ids:
                        issues.append(_('Duplicate field column: %s') % fid)
                    ids[fid] = ci
            expected = {c['id'] for c in cols}
            if set(ids) != expected or not self._validate_identity_headers(header, identity_columns):
                issues.append(_('The columns in %s changed. Start with a fresh download.') % sheet_name)
                continue
            seen = set()
            for rn, row in enumerate(ws.iter_rows(min_row=2, values_only=False), 2):
                vals = [c.value for c in row]
                if not any(v is not None and v != '' for v in vals):
                    continue
                key = vals[header.index('_row')]
                entry = by_row.get(key)
                if not entry or entry[0] not in allowed:
                    issues.append(_('%s row %s: employee is unknown or outside this payroll population.') % (sheet_name, rn))
                    continue
                eid, old = entry
                if eid in seen:
                    issues.append(_('%s row %s: duplicate Employee ID.') % (sheet_name, rn))
                    continue
                seen.add(eid)
                emp = people.browse(eid)
                if not self._workbook_code_unchanged(emp, old, identity_columns[0], vals[header.index(identity_labels['code'])]):
                    issues.append(_('%s row %s: Employee ID cannot be changed.') % (sheet_name, rn))
                old_email = old.get('identity_values', {}).get('email', old['email'])
                if mode == 'records' and ((vals[header.index(identity_labels['email'])] or '') != old_email
                        or self._identity_value(emp, identity_columns[1]) != old_email):
                    issues.append(_('%s row %s: email cannot be changed.') % (sheet_name, rn))
                provided_records = {}
                for c in cols:
                    destination = c.get('record_id') or c['id']
                    cell = row[ids[c['id']]]
                    value = cell.value
                    if value is None or value == '':
                        blanks += 1
                        continue
                    if cell.data_type == 'f':
                        issues.append(_('%s row %s: enter a value, not an Excel formula.') % (sheet_name, rn))
                        continue
                    current = records[eid].get(destination) if mode == 'records' else monthly[eid].get(c['id'])
                    previous = old['records'].get(destination) if mode == 'records' else old['monthly'].get(c['id'])
                    if mode == 'monthly' and snapshot.get('reusable'):
                        # A reusable blank template has no period snapshot.
                        # Review against the chosen period; confirm() still
                        # refuses concurrent changes after this review.
                        previous = current
                    if c.get('readonly') and mode == 'monthly':
                        previous = old['records'].get(c.get('record_id'))
                        current = records[eid].get(c.get('record_id'))
                    if c.get('readonly'):
                        # Record context can be text even when a legacy payroll
                        # column still has its default numeric value kind.
                        # It is never an input to coerce or write back.
                        if plain(value) == current:
                            unchanged += 1
                        else:
                            issues.append(_('%s row %s: %s is read-only context.') % (sheet_name, rn, c['label']))
                        continue
                    try:
                        value = self._numeric(value, c.get('kind')) if mode == 'monthly' else plain(value)
                    except (ValueError, TypeError):
                        issues.append(_('%s row %s: %s needs a valid number.') % (sheet_name, rn, c['label']))
                        continue
                    if mode == 'records':
                        if destination in provided_records and provided_records[destination] != plain(value):
                            issues.append(_('%s row %s: columns for the same saved field contain conflicting values. Use the same value, or leave one blank.') % (sheet_name, rn))
                            continue
                        provided_records[destination] = plain(value)
                    if plain(value) == current:
                        unchanged += 1
                        continue
                    if current != previous:
                        issues.append(_('%s · %s changed since download. Download again to review the latest value.') % (emp.name, c['label']))
                        continue
                    key = '%s:%s' % (eid, destination)
                    if key not in change_keys:
                        change_keys.add(key)
                        changes.append({'key': key, 'emp_id': eid, 'emp_name': emp.name,
                                        'field_id': destination, 'field_label': c['label'], 'tab': sheet_name,
                                        'old': current, 'value': plain(value)})
            seen_tabs[sheet_name] = seen
        if mode == 'monthly' and not seen_tabs.get('Monthly adjustments', set()).issubset(seen_tabs.get('Time', set())):
            issues.append(_('Every adjustment employee must also appear on the Time tab.'))
        if mode == 'records' and changes:
            preview = self.env['pb.records.desk'].preview_changes(cfg.id, changes)
            lookup = {(i['emp_id'], i['field_id']): i for i in preview['items']}
            validated = []
            for c in changes:
                item = lookup[(c['emp_id'], c['field_id'])]
                if item['status'] == 'refused':
                    issues.append('%s · %s: %s' % (c['emp_name'], c['field_label'], item['why']))
                elif item['status'] == 'same':
                    unchanged += 1
                else:
                    c.update(old_label=item['old_label'], new_label=item['new_label'])
                    validated.append(c)
            changes = validated
        payload = {'mode': mode, 'period': period, 'changes': changes, 'issues': issues,
                   'blanks': blanks, 'unchanged': unchanged, 'filename': filename,
                   'layout': self._fingerprint(layout, catalog)}
        receipt = self._exchange(cfg, 'review', payload)
        return dict(payload, token=receipt.token)

    @api.model
    def confirm(self, config_id, token, selected=None):
        cfg = self._config(config_id)
        self._lock(cfg)
        receipt = self._get_exchange(cfg, token, 'review')
        data = receipt.payload
        if data.get('grid'):
            return self._confirm_monthly(cfg, receipt, [])
        if data['issues']:
            raise UserError(_('Resolve all workbook issues before confirming changes.'))
        if data['mode'] == 'monthly':
            self.env['pb.payrun.wizard']._require_runner()
        keys = set(selected or [])
        all_keys = {c['key'] for c in data['changes']}
        if not keys or not keys.issubset(all_keys):
            raise UserError(_('Select the changes you want to apply.'))
        changes = [c for c in data['changes'] if c['key'] in keys]
        people = self._roster(cfg, data['period'])
        if not {c['emp_id'] for c in changes}.issubset(set(people.ids)):
            raise UserError(_('The payroll population changed. Review this import again.'))
        if data['mode'] == 'monthly':
            catalog = self._catalog(cfg)
            if data['layout'] != self._fingerprint(self._layout(cfg, catalog), catalog):
                raise UserError(_('The layout changed. Review a fresh workbook.'))
            current = self._current(cfg, data['period'], people)
        else:
            current = self._record_values(cfg, people, self._cards(cfg))
        for c in changes:
            if current[c['emp_id']].get(c['field_id']) != c['old']:
                raise UserError(_('%s · %s changed after review. Review the file again; nothing was applied.') % (c['emp_name'], c['field_label']))
        if data['mode'] == 'records':
            result = self.env['pb.records.desk'].apply_changes(cfg.id, changes, note='Payroll templates: %s' % data['filename'], source='import')
        else:
            Model = self.env['pb.payroll.month.inputs'].sudo()
            for eid in {c['emp_id'] for c in changes}:
                vals = dict(current[eid])
                vals.update({c['field_id']: c['value'] for c in changes if c['emp_id'] == eid})
                row = Model.search([('config_id', '=', cfg.id), ('company_id', '=', self.env.company.id), ('period', '=', data['period']), ('employee_id', '=', eid)], limit=1)
                if row:
                    row.write({'payload': vals})
                else:
                    Model.create({'config_id': cfg.id, 'company_id': self.env.company.id, 'employee_id': eid, 'period': data['period'], 'payload': vals})
            result = {'ok': True, 'written': len(changes), 'pending': False}
        receipt.write({'kind': 'applied', 'payload': dict(data, selected=sorted(keys), result=plain(result))})
        return result

    @api.model
    def history(self, config_id):
        cfg = self._config(config_id)
        rows = self.env['pb.payroll.workbook.exchange'].sudo().search([
            ('config_id', '=', cfg.id), ('company_id', '=', self.env.company.id),
            ('kind', '=', 'applied')], limit=30)
        return [{'id': r.id, 'when': str(r.write_date), 'by': r.owner_id.name,
                 'filename': r.payload.get('filename'), 'mode': r.payload.get('mode'),
                 'period': r.payload.get('period'), 'count': len(r.payload.get('selected', [])),
                 'pending': r.payload.get('result', {}).get('pending', False)} for r in rows]

    def _engine_file(self, cfg, period):
        """Use component headings; adapt engine bindings only while resolving.

        Layout edits never change formula bindings. Recurring fields come
        freshly from records.
        """
        import openpyxl
        from openpyxl.comments import Comment
        if cfg.state != 'active':
            raise UserError(_('Activate this payroll configuration before using its monthly inputs in Pay Run.'))
        people = self._roster(cfg, period)
        cards = self._cards(cfg)
        records = self._record_values(cfg, people, cards)
        current = self._current(cfg, period, people)
        catalog = [c for c in self._catalog(cfg) if c['id'].startswith('rule:')]
        resolved = self._resolved_month(cfg, period, people, catalog)
        salary_keys = {'rule:%s' % r.id for r in salary_inputs(cfg)}
        mandatory = [c for c in catalog if not c.get('readonly')
                     and re.search(r'paid.?days|working.?days|standard.*days|period.?days|days.?worked|ngày công', c['label'], re.I)]
        missing = []
        for employee in people:
            for c in mandatory:
                value = resolved[employee.id][c['id']]['value']
                if value is None or value == '' or (re.search(r'standard|period|working', c['label'], re.I) and float(value) <= 0):
                    missing.append('%s · %s' % (employee.name, c['label']))
        if missing:
            raise UserError(_('Complete the required Time inputs before computing payroll:\n%s') % '\n'.join(missing[:20]))
        missing_required = ['%s · %s' % (employee.name, c['label'])
            for employee in people for c in catalog
            if c.get('required') and not c.get('readonly')
            and resolved[employee.id][c['id']]['value'] in (None, '')]
        if missing_required:
            raise UserError(_('Complete the required monthly inputs before computing payroll:\n%s') % '\n'.join(missing_required[:20]))
        missing_salary = []
        for employee in people:
            for c in catalog:
                if c['id'] not in salary_keys:
                    continue
                value = resolved[employee.id][c['id']]['value']
                if value is None or value == '':
                    missing_salary.append('%s · %s' % (employee.name, c['label']))
        if missing_salary:
            raise UserError(_('Recurring salary is missing. Update employee / contract records or provide the required salary input before computing payroll:\n%s') % '\n'.join(missing_salary[:20]))
        identity_cols = self._identity_columns(cfg)
        roles = self._identity_rules(cfg)
        # Additional identity components (e.g. email) only exist if the schema
        # defines them. Never emit the adapter's duplicate ID/name aliases.
        for rule in cfg.rule_ids.sorted(key=lambda r: (r.sequence, r.id)):
            if rule.id in roles and rule.id not in {c['rule_id'] for c in identity_cols}:
                mapping = self.env['pb.records.desk']._mappings(cfg.id).filtered(lambda m:
                    m.component_id == rule and m.target_model_id.model == 'hr.employee')[:1]
                identity_cols.append({'kind': roles[rule.id], 'label': rule.name or rule.code,
                    'rule_id': rule.id, 'field': mapping.target_field_id.name if mapping else False})
        headers = [c['label'] for c in identity_cols]
        codes = [self._identity_value(e, identity_cols[0]) for e in people]
        if any(not c for c in codes) or len(set(codes)) != len(codes):
            raise UserError(_('Every employee needs a unique value for %s in the mapped employee field before preparing pay data.') % identity_cols[0]['label'])
        cols = []
        for c in catalog:
            cols.append(c)
            headers.append(c['label'])
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Payroll'
        ws.append(headers)
        # Names can repeat across worksheets in the original scheme. Keep the
        # familiar headings and bind each physical column by its stable code.
        component_codes = [cfg.rule_ids.browse(c['rule_id']).code if c['rule_id'] else ''
                           for c in identity_cols] + [c['code'] for c in cols]
        for cell, code in zip(ws[1], component_codes):
            if code:
                cell.comment = Comment('component: ' + code, 'Payobook')
        for e in people:
            vals = [self._identity_value(e, c) for c in identity_cols]
            for c in cols:
                # Saved recurring values are current record context. Earlier
                # uploads/monthly-layout values (including stale zeros) cannot
                # override them. Variable monthly zeros still remain updates.
                value = resolved[e.id][c['id']]['value']
                if value is None and not c.get('readonly') and c['kind'] in ('money', 'decimal', 'integer', 'quantity', 'rate'):
                    value = 0
                # No previous month's amounts enter this path.
                vals.append(value)
            ws.append(vals)
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue(), people.ids


class TemplatePayrun(models.AbstractModel):
    _inherit = 'pb.payrun.wizard'

    @api.model
    def prepare_run(self, vals):
        if vals.get('templates_monthly'):
            # A division left over from the demo defaults must not replace
            # the explicitly selected configuration with Retail's scheme.
            vals = dict(vals, division=False)
        return super().prepare_run(vals)

    @api.model
    def spreadsheet_gate(self, vals=None):
        if (vals or {}).get('templates_monthly'):
            self._require_runner()
            facade = self.env['pb.payroll.templates']
            cfg = facade._config(vals.get('formula_config_id'), write=True)
            all_fields = facade._catalog(cfg)
            layout = facade._layout(cfg, all_fields)
            selected = set(layout['time'] + layout['adjustments'])
            catalog = [c for c in all_fields if c['id'] in selected and not c.get('readonly')]
            return {'wanted': bool(catalog), 'config_id': cfg.id, 'config_name': cfg.name,
                    'components': [{'code': c['code'], 'name': c['label'], 'key': c['source_keys'][0]} for c in catalog if c['source_keys']],
                    'choices': [], 'lanes': self._payrun_lanes(vals)}
        return super().spreadsheet_gate(vals)

    @api.model
    def template_month_file(self, config_id, period):
        self._require_runner()
        facade = self.env['pb.payroll.templates']
        cfg = facade._config(config_id, write=True)
        raw, ids = facade._engine_file(cfg, period)
        # Used only by the pay-run server adapter, not exposed as an upload.
        return {'file_b64': base64.b64encode(raw).decode(), 'filename': 'Payobook-month-%s.xlsx' % period, 'people': len(ids)}

    @api.model
    def attach_spreadsheet(self, run_id, config_id, file_b64, filename,
                           date_start, date_end, one_time=False):
        if (filename or '').startswith('Payobook-month-'):
            self._require_runner()
            period = fields.Date.to_date(date_start).strftime('%Y-%m')
            facade = self.env['pb.payroll.templates']
            cfg = facade._config(config_id, write=True)
            start, end = facade._period(period)
            if fields.Date.to_date(date_start) != start or fields.Date.to_date(date_end) != end:
                raise UserError(_('Monthly templates require a complete calendar month.'))
            raw, _ids = facade._engine_file(cfg, period)
            # Always regenerate on the server. The browser's file is never the
            # authority for confirmed monthly values, or for recurring salary.
            monthly_codes = [c['code'] for c in facade._catalog(cfg) if not c.get('readonly')]
            return super(TemplatePayrun, self.with_context(pb_template_engine=True, pb_template_monthly_codes=monthly_codes)).attach_spreadsheet(
                run_id, config_id, base64.b64encode(raw).decode(), filename,
                date_start, date_end, one_time=True)
        return super().attach_spreadsheet(run_id, config_id, file_b64, filename,
                                          date_start, date_end, one_time=one_time)


class TemplateBatch(models.Model):
    _inherit = 'hr.payroll.import.batch'

    def _raw_data_from_row(self, headers, row):
        if self.env.context.get('pb_template_engine'):
            # Workbook positions belong to the export layout. Original Excel
            # letter bindings are restored from the component, not that layout.
            return dict(row) if isinstance(row, dict) else dict(zip(headers, row))
        return super()._raw_data_from_row(headers, row)

    def _template_alias_row(self, raw_data):
        """Adapt schema headings to existing bindings in memory, never in Excel."""
        if not (self.env.context.get('pb_template_engine') or self._template_period()):
            return raw_data
        data = dict(raw_data or {})
        original = dict(data)
        for rule in self.formula_config_id.rule_ids.filtered(lambda r: r.column_type == 'input'):
            label = rule.name or rule.code
            aliases = [s.key for s in rule.source_ids if s.kind == 'excel' and s.key]
            aliases += self._rule_header_candidates(rule, allow_column_letter=False)
            source = next((key for key in aliases if key in original), None)
            if source is None:
                continue
            value = original[source]
            data.setdefault(label, value)
            for alias in aliases:
                data.setdefault(alias, value)
        return data

    def _identity_from_file_row(self, raw_data):
        return super()._identity_from_file_row(self._template_alias_row(raw_data))

    def _template_period(self):
        if len(self) != 1 or not self.one_time or not self.formula_config_id or not self.date_from:
            return None
        period = fields.Date.to_date(self.date_from).strftime('%Y-%m')
        if self.import_filename != 'Payobook-month-%s.xlsx' % period:
            return None
        start, end = self.env['pb.payroll.templates']._period(period)
        return period if self.date_from == start and self.date_to == end else None

    def _transform_data_to_formula_inputs(self, raw_data, contract=None, employee=None,
                                         provenance=None, topup_data=None):
        # The persisted batch outlives the initial upload's context. Recompute
        # must join current recurring records again, rather than reusing the
        # original generated workbook's salary/allowance snapshot.
        raw_data = self._template_alias_row(raw_data)
        period = self._template_period()
        if not period or not employee:
            return super()._transform_data_to_formula_inputs(raw_data, contract=contract,
                employee=employee, provenance=provenance, topup_data=topup_data)
        cfg = self.formula_config_id
        if employee.company_id != self.company_id or cfg.company_id != self.company_id:
            raise UserError(_('The monthly payroll records belong to a different company.'))
        facade = self.env['pb.payroll.templates'].with_company(self.company_id)
        catalog = [c for c in facade._catalog(cfg) if c['id'].startswith('rule:')]
        cards = facade._cards(cfg)
        saved = facade._record_values(cfg, employee, cards)[employee.id]
        raw_data = dict(raw_data or {})
        override_row = self.env['pb.payroll.month.inputs'].sudo().search([('config_id', '=', cfg.id), ('company_id', '=', self.company_id.id), ('employee_id', '=', employee.id), ('period', '=', period)], limit=1)
        overrides = override_row.record_overrides or {}
        joined = []
        salary_keys = {'rule:%s' % r.id for r in salary_inputs(cfg)}
        for c in catalog:
            if not (c.get('readonly') and c.get('record_id')):
                continue
            value = overrides[c['id']] if c['id'] in overrides else saved.get(c['record_id'])
            if c['id'] in salary_keys and value in (None, ''):
                raise UserError(_('Recurring salary is missing for %s. Update the employee / contract records before recomputing payroll.') % employee.name)
            for key in c['source_keys']:
                raw_data[key] = value
            joined.append(c)
        monthly_codes = [c['code'] for c in catalog if not c.get('readonly') or c['id'] in overrides]
        contextual = self.with_context(pb_template_engine=True, pb_template_monthly_codes=monthly_codes)
        result = super(TemplateBatch, contextual)._transform_data_to_formula_inputs(
            raw_data, contract=contract, employee=employee, provenance=provenance, topup_data=topup_data)
        if provenance is not None:
            for c in joined:
                fid = c['record_id']
                kind = 'contract_component' if fid.startswith('c:') else 'contract_field' if fid.startswith('f:hr.contract:') else 'employee_field' if fid.startswith('f:hr.employee:') else 'bank_account'
                provenance[c['code']] = {'src': 'manual' if c['id'] in overrides else kind, 'key': fid, 'via': 'monthly_review' if c['id'] in overrides else 'monthly_records'}
        return result

    def _mapped_record_value(self, mapping, contract=None, employee=None):
        if (self.env.context.get('pb_template_engine') and mapping
                and mapping.component_id.code in self.env.context.get('pb_template_monthly_codes', [])):
            return None
        return super()._mapped_record_value(mapping, contract=contract, employee=employee)

    def _contract_component_amounts(self, contract):
        amounts = super()._contract_component_amounts(contract)
        if self.env.context.get('pb_template_engine'):
            monthly = {self._normalize_header_key(c) for c in self.env.context.get('pb_template_monthly_codes', [])}
            return {k: v for k, v in amounts.items() if k not in monthly}
        return amounts

    def _parse_source_file(self, file_content, filename, connector=None,
                           single_sheet_fallback=False):
        if self.env.context.get('pb_template_engine'):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
            ws = wb['Payroll']
            headers, used = [], set()
            by_code = {r.code: r for r in self.formula_config_id.rule_ids}
            for cell in ws[1]:
                key = cell.value
                tagged = False
                if cell.comment and cell.comment.text.startswith('component: '):
                    rule = by_code.get(cell.comment.text[11:])
                    if rule:
                        tagged = True
                        keys = [s.key for s in rule.source_ids if s.kind == 'excel' and s.key]
                        key = next((k for k in keys if k not in used), rule.code)
                if tagged and key in used:
                    raise UserError(_('The generated pay data contains duplicate input bindings. Review the component sources.'))
                headers.append(key)
                used.add(key)
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            return {'headers': headers, 'rows': [list(r) for r in rows], 'multisheet': False}
        return super()._parse_source_file(file_content, filename, connector=connector,
                                          single_sheet_fallback=single_sheet_fallback)


class TemplateImportLine(models.Model):
    _inherit = 'hr.payroll.import.line'

    def _get_value_for_rule(self, raw_data, rule):
        # Generated files use each component's declared source key, which may
        # be qualified by a different original worksheet. Validate the same
        # keys the monthly resolver reads, including after the upload context
        # is gone. The source row itself remains unchanged for auditing.
        batch = self.batch_id
        if batch._template_period() or self.env.context.get('pb_template_engine'):
            keys = [s.key for s in rule.source_ids if s.kind == 'excel' and s.key]
            keys += batch._rule_header_candidates(rule, allow_column_letter=False)
            value, matched = batch._lookup_in_blob(raw_data, keys)
            if matched:
                return value
        return super()._get_value_for_rule(raw_data, rule)


class TemplateRecordFields(models.AbstractModel):
    _inherit = 'pb.records.desk'

    def _cards(self, config_id):
        cards = super()._cards(config_id)
        if not config_id:
            return cards
        configs = self._configs(config_id)
        if not configs:
            return cards
        # Standard employee/contract details share the existing review, approval
        # and Undo machinery. This does not create or change formula bindings.
        standard = {'hr.employee': ('name', 'job_title', 'department_id', 'birthday', 'work_phone', 'mobile_phone'),
                    'hr.contract': ('wage', 'date_start', 'date_end')}
        for model, names in standard.items():
            for name in names:
                fid = 'f:%s:%s' % (model, name)
                field = self.env[model]._fields.get(name)
                if fid in cards or not field or field.readonly:
                    continue
                metadata = self.env['ir.model.fields'].sudo().search([('model', '=', model), ('name', '=', name)], limit=1)
                # Only the in-memory destination metadata needs elevation.
                # A normal payroll user cannot read ir.model.fields, and a
                # warm admin cache previously hid that during transaction tests.
                mapping = self.env['hr.payslip.import.mapping'].sudo().new({
                    'target_model_id': self.env['ir.model']._get(model).id,
                    'target_field_id': metadata.id, 'destination_type': 'field'})
                candidates = salary_inputs(configs[:1]) if name == 'wage' else self.env['hr.formula.rule']
                rule = candidates if len(candidates) == 1 else self.env['hr.formula.rule']
                cards[fid] = {'id': fid, 'group': 'employee' if model == 'hr.employee' else 'contract',
                              'label': field.string, 'sub': _('Standard employee / contract record'),
                              'ttype': field.type, 'model': model, 'field': name,
                              'editable': True, 'selection': self._selection_pairs(field) if field.type == 'selection' else [],
                              'm2o': {'comodel': field.comodel_name, 'creates_missing': False, 'key': 'name'} if field.type == 'many2one' else None,
                              'component': {'id': rule.id, 'code': rule.code, 'name': rule.name} if rule else None,
                              '_mapping': mapping, '_field': field}
        return cards


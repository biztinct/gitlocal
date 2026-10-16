"""Durable, server-owned choices about payroll inputs and future records.

No ACL is granted on the snapshot model. The facade checks the actual run,
company and Records Desk permissions. Decisions use the existing records
approval adapter; they never recompute, approve or pay the captured payroll.
"""
import base64
import copy
import io
import json

from odoo import api, fields, models, _
from odoo.exceptions import AccessError, UserError

from .templates import IDENTITY, plain


class RunInputSnapshot(models.Model):
    _name = 'pb.payroll.run.input.review'
    _description = 'Payroll input changes and future-record decisions'
    batch_id = fields.Many2one('hr.payroll.import.batch', required=True, ondelete='cascade', index=True)
    run_id = fields.Many2one('hr.payslip.run', required=True, ondelete='cascade', index=True)
    company_id = fields.Many2one('res.company', required=True, ondelete='cascade', index=True)
    payload = fields.Json(default=dict)
    _batch_unique = models.Constraint('UNIQUE(batch_id)', 'One input review per payroll import.')


class RunInputReview(models.AbstractModel):
    _inherit = 'pb.payroll.templates'

    def _review_run(self, run_id):
        run = self.env['hr.payslip.run'].browse(int(run_id or 0)).exists()
        if not run:
            raise UserError(_('This payroll no longer exists.'))
        run.check_access('read')
        company = run.company_id if 'company_id' in run._fields else run.slip_ids[:1].company_id
        if company and company != self.env.company:
            raise AccessError(_('Switch to this payroll’s company to review its input changes.'))
        self.env['pb.records.desk']._check_read()
        return run

    def _input_destinations(self, cfg):
        """Only explicit mappings / declared contract components confer authority."""
        desk = self.env['pb.records.desk']
        cards = desk._cards(cfg.id)
        mapped = {}
        for mapping in desk._mappings(cfg.id):
            if mapping.destination_type == 'bank_account':
                fid = 'b:%s' % mapping.bank_role
            else:
                fid = 'f:%s:%s' % (mapping.target_model_id.model, mapping.target_field_id.name)
            if fid in cards:
                mapped.setdefault(mapping.component_id.id, set()).add(fid)
        for rule in cfg.rule_ids.filtered('is_contract_component'):
            fid = 'c:%s' % rule.code
            if fid in cards and not mapped.get(rule.id):
                mapped.setdefault(rule.id, set()).add(fid)
        return cards, mapped

    def _capture_run_inputs(self, batch, historical=False):
        """Capture BEFORE any writeback, or explicitly label a legacy comparison."""
        Model = self.env['pb.payroll.run.input.review'].sudo()
        if Model.search_count([('batch_id', '=', batch.id)]) or not batch.payslip_run_id:
            return
        cfg = batch.formula_config_id
        desk = self.env['pb.records.desk']
        cards, mapped = self._input_destinations(cfg)
        lines = batch.import_line_ids.filtered(lambda l: l.employee_id and l.state != 'error')
        contracts = desk._io_contracts(lines.mapped('employee_id'))
        rules = cfg.rule_ids.filtered(lambda r: r.column_type == 'input' and not batch._is_employee_code_rule(r))
        columns, rows = {}, []
        for line in lines:
            emp = line.employee_id
            if emp.company_id != batch.company_id:
                continue
            contract = contracts.get(emp.id) or self.env['hr.contract']
            raw = line.get_raw_data() or {}
            if historical:
                if not line.payslip_id:
                    continue
                used = json.loads(line.payslip_id.formula_input_values or '{}')
            else:
                used = batch._transform_data_to_formula_inputs(raw, contract=contract, employee=emp)
            cells = {}
            for rule in rules:
                supplied, present = batch._get_rule_raw_value(raw, rule, allow_column_letter=False)
                if not present or rule.code not in used:
                    continue  # Blanks never become instructions to erase a record.
                destinations = mapped.get(rule.id, set())
                fid = next(iter(destinations)) if len(destinations) == 1 else None
                card = cards.get(fid)
                # Stable identity columns identify the person; they are not payroll edits.
                if card and card.get('model') == 'hr.employee' and card.get('field') in IDENTITY | {'name'}:
                    continue
                cid = str(rule.id)
                columns[cid] = {'id': cid, 'label': rule.name or rule.code, 'code': rule.code,
                    'kind': rule.value_kind, 'mapped': bool(destinations), 'field_id': fid,
                    'destination': card['label'] if card else '',
                    'mapping': sorted(destinations)}
                value = plain(used[rule.code])
                cell = {'value': value, 'supplied': plain(supplied), 'old': None, 'old_label': '',
                    'new_label': str(value) if value is not None else '', 'changed': False,
                    'decision': 'excel' if not destinations else 'same', 'reason': '',
                    'contract_id': contract.id or False}
                cell['ignored'] = plain(batch._normalize_rule_input_value(rule, supplied)) != value
                if card:
                    current = desk._read_cell(batch, card, emp, contract)
                    coerced, label, why = desk._coerce(batch, card, value, emp, contract)
                    cell.update(old=plain(current['v']), old_label=current['label'], new_label=label,
                                reason=why or (_('No contract to update.') if current.get('missing') else ''))
                    cell['changed'] = not why and not current.get('missing') and not desk._same(card, current['v'], coerced)
                    cell['decision'] = 'pending' if cell['changed'] and card.get('editable') else 'same'
                    if not card.get('editable'):
                        cell['reason'] = _('This mapped record field is read-only.')
                    if cell['reason']:
                        cell['decision'] = 'blocked'
                elif len(destinations) > 1:
                    cell.update(decision='blocked', reason=_('This input has multiple record destinations. Review its mapping before updating records.'))
                cells[cid] = cell
            if cells:
                rows.append({'emp_id': emp.id, 'name': emp.name, 'code': self._code(emp), 'line_id': line.id, 'cells': cells})
        Model.create({'batch_id': batch.id, 'run_id': batch.payslip_run_id.id,
            'company_id': batch.company_id.id, 'payload': {'columns': list(columns.values()), 'rows': rows,
            'revision': 0, 'historical': historical, 'captured_at': str(fields.Datetime.now()), 'events': []}})

    def _review_snapshots(self, run):
        return self.env['pb.payroll.run.input.review'].sudo().search([
            ('run_id', '=', run.id), ('company_id', '=', self.env.company.id)], order='id')

    def _review_payload(self, snapshot, check_inputs=True):
        data = copy.deepcopy(snapshot.payload or {})
        data.update(id=snapshot.id, filename=snapshot.batch_id.import_filename or '', revision=data.get('revision', 0))
        proposals = {c.get('apply_id') for r in data.get('rows', []) for c in r['cells'].values() if c.get('apply_id')}
        applies = {p.id: p for p in self.env['pb.records.apply'].sudo().browse(list(proposals)).exists()}
        lines = self.env['hr.payroll.import.line'].sudo().browse([r['line_id'] for r in data.get('rows', [])]).exists() if check_inputs else self.env['hr.payroll.import.line']
        lines.mapped('payslip_id.formula_input_values')
        by_line = {line.id: line for line in lines}
        col_by_id = {c['id']: c for c in data.get('columns', [])}
        for row in data.get('rows', []):
            line = by_line.get(row['line_id'])
            actual = json.loads(line.payslip_id.formula_input_values or '{}') if line and line.payslip_id else None
            for cid, cell in row['cells'].items():
                cell['not_computed'] = bool(check_inputs and (not line or not line.payslip_id))
                cell['payroll_changed'] = actual is not None and actual.get(col_by_id[cid]['code']) != cell['value']
                proposal = applies.get(cell.get('apply_id'))
                if proposal:
                    refused = any(r.get('emp_id') == row['emp_id'] and r.get('field_id') == cell.get('proposal_field') for r in proposal.late_refusals())
                    cell['decision'] = 'blocked' if refused else 'saved' if proposal.applied else 'pending' if proposal.state in ('returned', 'rejected') else 'approval'
                    cell['reason'] = proposal.block_note or (_('This record update was refused. Review it in Records Desk.') if refused else '')
                    cell['reference'] = proposal.name
                    cell['with_whom'] = proposal._waiting_for() if not proposal.applied else ''
        return data

    @api.model
    def run_input_changes(self, run_id):
        run = self._review_run(run_id)
        # Legacy batches have no before snapshot. Their comparison is labelled
        # "On record at first review", never presented as an original audit.
        batches = self.env['hr.payroll.import.batch'].sudo().search([
            ('payslip_run_id', '=', run.id), ('company_id', '=', self.env.company.id), ('source_type', '=', 'excel')])
        for batch in batches:
            self._capture_run_inputs(batch, historical=True)
        data = [self._review_payload(s) for s in self._review_snapshots(run)]
        can_decide = run.state != 'cancel' and self.env['pb.payrun.wizard']._may_run_payroll()
        return {'run_id': run.id, 'name': run.name, 'cancelled': run.state == 'cancel',
                'can_decide': can_decide, 'sheets': data}

    @api.model
    def decide_run_inputs(self, run_id, snapshot_id, revision, keys, decision, note=''):
        run = self._review_run(run_id)
        self.env['pb.payrun.wizard']._require_runner()
        if run.state == 'cancel':
            raise UserError(_('This payroll was rejected. Its input review remains available for reference.'))
        if decision not in ('future', 'once') or not isinstance(keys, list) or not keys or len(keys) != len(set(keys)):
            raise UserError(_('Select the changed mapped cells you want to decide.'))
        snapshot = self._review_snapshots(run).filtered(lambda s: s.id == int(snapshot_id))
        if not snapshot:
            raise AccessError(_('This input review does not belong to this payroll.'))
        self.env.cr.execute('SELECT id FROM pb_payroll_run_input_review WHERE id=%s FOR UPDATE', (snapshot.id,))
        snapshot.invalidate_recordset()
        data = self._review_payload(snapshot)
        if data['revision'] != revision:
            raise UserError(_('Someone else updated this review. Refresh it to see their decisions before continuing.'))
        columns = {c['id']: c for c in data['columns']}
        cells = {'%s:%s' % (r['line_id'], cid): (r, columns[cid], cell)
                 for r in data['rows'] for cid, cell in r['cells'].items()}
        if any(k not in cells or cells[k][2]['decision'] != 'pending' or cells[k][2].get('payroll_changed') or cells[k][2].get('not_computed') for k in keys):
            raise UserError(_('A selected cell has already been decided or cannot update a record. Refresh the review.'))
        desk = self.env['pb.records.desk']
        cfg = snapshot.batch_id.formula_config_id
        self._config(cfg.id, write=True)
        self._lock(cfg)
        result = {}
        if decision == 'future':
            cards, destinations = self._input_destinations(cfg)
            selected = [cells[k] for k in keys]
            employees = self.env['hr.employee'].sudo().browse(sorted({r['emp_id'] for r, _c, _v in selected})).exists()
            if not employees:
                raise UserError(_('The selected employees no longer exist. Refresh this review.'))
            self.env.cr.execute('SELECT id FROM hr_employee WHERE id IN %s ORDER BY id FOR UPDATE', (tuple(employees.ids),))
            contracts = desk._io_contracts(employees)
            contract_ids = sorted({c.id for c in contracts.values() if c})
            if contract_ids:
                self.env.cr.execute('SELECT id FROM hr_contract WHERE id IN %s ORDER BY id FOR UPDATE', (tuple(contract_ids),))
            changes, targets = [], set()
            for row, col, cell in selected:
                fid = col['field_id']
                card = cards.get(fid)
                if not card or sorted(destinations.get(int(col['id']), set())) != col['mapping']:
                    raise UserError(_('A mapping changed after this payroll. Review the field in Records Desk.'))
                emp = employees.filtered(lambda e: e.id == row['emp_id'])
                if not emp or emp.company_id != self.env.company:
                    raise AccessError(_('A selected employee is no longer in this company.'))
                contract = contracts.get(emp.id) or self.env['hr.contract']
                if (card.get('model') == 'hr.contract' or fid.startswith('c:')) and contract.id != cell['contract_id']:
                    raise UserError(_('%s now has a different contract. Review the change in Records Desk.') % row['name'])
                current = desk._read_cell(snapshot.batch_id, card, emp, contract)
                if not desk._same(card, current['v'], cell['old']):
                    raise UserError(_('%s · %s changed on record after this review was captured. Nothing was saved. Review the current record in Records Desk.') % (row['name'], col['label']))
                target = (emp.id, fid)
                if target in targets:
                    raise UserError(_('Two selected inputs update the same record field. Select only one.'))
                targets.add(target)
                changes.append({'emp_id': emp.id, 'field_id': fid, 'value': cell['value']})
            _items, _plan, counts = desk._evaluate(cfg.id, changes)
            if counts['refused']:
                raise UserError(_('Some selected values cannot be saved:\n%s') % '\n'.join(i['why'] for i in _items if i['status'] == 'refused'))
            # Use one transaction: even late refusals roll back all decisions.
            result = desk.apply_changes(cfg.id, changes, note=_('Payroll %s: %s') % (run.name, note or _('Keep selected inputs for future payrolls')), source='import')
            if result.get('refused'):
                raise UserError(_('The record update could not be completed. No review decisions were saved.'))
        event = {'by': self.env.user.name, 'at': str(fields.Datetime.now()), 'decision': decision,
                 'count': len(keys), 'note': str(note or '')[:1000], 'reference': result.get('reference', '')}
        data['events'].append(event)
        for key in keys:
            _row, col, cell = cells[key]
            cell.update(decision='approval' if result.get('pending') else 'saved' if decision == 'future' else 'once',
                        by=event['by'], at=event['at'], apply_id=result.get('apply_id'), proposal_field=col['field_id'])
        data['revision'] += 1
        snapshot.write({'payload': data})
        return self.run_input_changes(run.id)

    @api.model
    def export_run_input_changes(self, run_id, snapshot_id):
        import openpyxl
        from openpyxl.comments import Comment
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        run = self._review_run(run_id)
        if not self.env['pb.records.desk']._may_export():
            raise AccessError(_('Exporting this review requires permission to export employee records.'))
        snapshot = self._review_snapshots(run).filtered(lambda s: s.id == int(snapshot_id))
        if not snapshot:
            raise AccessError(_('This input review does not belong to this payroll.'))
        data = self._review_payload(snapshot)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Input changes'
        ws.append([_('Employee ID'), _('Employee')] + [c['label'] for c in data['columns']])
        colors = {'pending': 'F3EDFF', 'excel': 'FFF8E8', 'saved': 'EAF8F1', 'approval': 'ECF4FF', 'blocked': 'FFF0EF', 'once': 'F4F5F8'}
        for row in data['rows']:
            ws.append([row['code'], row['name']] + [row['cells'].get(c['id'], {}).get('value') for c in data['columns']])
            for i, col in enumerate(data['columns'], 3):
                cell = row['cells'].get(col['id'])
                if not cell:
                    continue
                out = ws.cell(ws.max_row, i)
                out.fill = PatternFill('solid', fgColor=colors.get(cell['decision'], 'FFFFFF'))
                out.comment = Comment('On record: %s\nStatus: %s\n%s' % (cell['old_label'], cell['decision'], cell.get('reason', '')), 'Payroll review')
            # Workbook values are data even when text begins with =, + or @.
            for out in ws[ws.max_row]:
                if isinstance(out.value, str):
                    out.data_type = 's'
        for out in ws[1]:
            out.fill = PatternFill('solid', fgColor='55428C')
            out.font = Font(color='FFFFFF', bold=True)
            out.alignment = Alignment(wrap_text=True)
            ws.column_dimensions[get_column_letter(out.column)].width = 28
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        guide = wb.create_sheet('Read me')
        for text in [run.name, 'Review captured: %s UTC' % data['captured_at'],
                     'Historical comparison: on record at first review.' if data['historical'] else 'On-record values were captured before this payroll was computed.',
                     'Violet: mapped change to decide. Amber: Excel only, enter again next payroll.',
                     'Green: saved for future payrolls. Blue: record update awaiting approval. Grey: this payroll only.',
                     'Read each cell comment for its on-record value and decision. This file is a read-only review; make decisions in the payroll batch.']:
            guide.append([text])
        guide.column_dimensions['A'].width = 110
        out = io.BytesIO()
        wb.save(out)
        return {'filename': 'Payroll-%s-input-changes.xlsx' % run.id, 'file_b64': base64.b64encode(out.getvalue()).decode()}


class CaptureBatchInputReview(models.Model):
    _inherit = 'hr.payroll.import.batch'

    def action_validate(self):
        result = super().action_validate()
        if self.env.context.get('pb_capture_run_inputs') and self.one_time and self.payslip_run_id:
            self.env['pb.payroll.templates']._capture_run_inputs(self)
        return result


class ReviewPayrun(models.AbstractModel):
    _inherit = 'pb.payrun.wizard'

    @api.model
    def attach_spreadsheet(self, run_id, config_id, file_b64, filename, date_start, date_end, one_time=False):
        self._require_runner()
        self.env['pb.payroll.templates']._review_run(run_id)
        # Wizard uploads compute first; lasting writes are a separate, durable
        # decision. Ordinary Import and connected-system batches are unchanged.
        return super(ReviewPayrun, self.with_context(pb_capture_run_inputs=True)).attach_spreadsheet(
            run_id, config_id, file_b64, filename, date_start, date_end, one_time=True)


class ReviewRunAction(models.Model):
    _inherit = 'hr.payslip.run'

    def action_input_changes(self):
        self.ensure_one()
        self.env['pb.payroll.templates']._review_run(self.id)
        return {'type': 'ir.actions.client', 'tag': 'pb_payroll_input_changes', 'params': {'run_id': self.id}}


class ReviewBoard(models.AbstractModel):
    _inherit = 'pb.payruns'

    @api.model
    def get_board_data(self):
        data = super().get_board_data()
        batches = {b['id']: b for b in data.get('batches', [])}
        snapshots = self.env['pb.payroll.run.input.review'].sudo().search([
            ('run_id', 'in', list(batches)), ('company_id', '=', self.env.company.id)])
        for snapshot in snapshots:
            info = batches[snapshot.run_id.id].setdefault('input_review', {'pending': 0, 'approval': 0})
            for row in self.env['pb.payroll.templates']._review_payload(snapshot, check_inputs=False)['rows']:
                for cell in row['cells'].values():
                    if cell['decision'] in info:
                        info[cell['decision']] += 1
        return data

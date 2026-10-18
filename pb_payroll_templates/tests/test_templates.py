import base64
import io
import json
from unittest.mock import patch

import openpyxl
from odoo.exceptions import UserError, AccessError
from odoo.tests import TransactionCase, tagged


@tagged('post_install', '-at_install')
class TestPayrollTemplates(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.api = cls.env['pb.payroll.templates']
        cls.env['biz.approval.seed'].sudo().set_no_approval_needed(cls.env.company, 'records')
        cls.cfg = cls.env['hr.formula.config'].create({'name': 'Workbook tests', 'code': 'WORKBOOKTEST', 'country_code': 'SG', 'state': 'active', 'company_id': cls.env.company.id})
        cls.salary = cls.rule('Salary', 'WBBASE', 'money')
        cls.paid = cls.rule('Paid days', 'WBDAYS', 'quantity')
        cls.bonus = cls.rule('Monthly bonus', 'WBBONUS', 'money')
        cls.email = cls.rule('Employee email', 'WBEMAIL', 'text')
        cls.map_field(cls.salary, 'hr.contract', 'wage')
        cls.map_field(cls.email, 'hr.employee', 'work_email')
        cls.emp = cls.env['hr.employee'].create({'name':'Workbook One', 'barcode':'000001', 'work_email':'workbook1@example.test', 'company_id':cls.env.company.id})
        cls.contract = cls.env['hr.contract'].create({'name':'Workbook contract', 'employee_id':cls.emp.id, 'wage':1000, 'state':'open', 'date_start':'2026-01-01'})

    @classmethod
    def rule(cls, name, code, kind):
        r=cls.env['hr.formula.rule'].create({'config_id':cls.cfg.id,'name':name,'code':code,'column_type':'input','value_kind':kind})
        r.set_source_binding('excel','Payroll|'+name)
        return r

    @classmethod
    def map_field(cls, rule, model, field):
        cls.env['hr.payslip.import.mapping'].create({'salary_structure_id':cls.cfg.id,'component_id':rule.id,'destination_type':'field','target_model_id':cls.env['ir.model']._get(model).id,'target_field_id':cls.env['ir.model.fields'].search([('model','=',model),('name','=',field)],limit=1).id})

    def setUp(self):
        super().setUp()
        self.roster_patch=patch.object(type(self.api),'_roster',return_value=self.emp)
        self.roster_patch.start()
        self.addCleanup(self.roster_patch.stop)

    def export(self, mode='monthly', period='2026-10'):
        d=self.api.download(self.cfg.id,period,mode)
        return openpyxl.load_workbook(io.BytesIO(base64.b64decode(d['file_b64'])))

    def review(self, wb, mode='monthly', period='2026-10'):
        b=io.BytesIO();wb.save(b)
        return self.api.review(self.cfg.id,period,mode,base64.b64encode(b.getvalue()).decode(),'workbook.xlsx')

    def set_value(self, wb, sheet, label, value):
        ws=wb[sheet];headers=[c.value for c in ws[1]];ws.cell(2,headers.index(label)+1,value)
        if value is None: ws.cell(2,headers.index(label)+1).value=None

    def apply(self, review):
        return self.api.confirm(self.cfg.id,review['token'],[c['key'] for c in review['changes']])

    def test_monthly_three_tabs_and_identity(self):
        wb=self.export()
        self.assertEqual([s.title for s in wb if s.sheet_state=='visible'],['Time','Monthly adjustments','Employee records'])
        for s in ('Time','Monthly adjustments'):
            self.assertEqual(wb[s]['A2'].value,'000001')
            self.assertFalse(wb[s].protection.sheet)
            self.assertEqual(wb[s].freeze_panes,'C2')
        self.assertEqual(self.review(wb)['changes'],[])

    def test_empty_monthly_tabs_are_unprotected(self):
        with patch.object(type(self.api),'_roster',return_value=self.emp.browse()):
            wb=self.export()
        for sheet in ('Time','Monthly adjustments'):
            self.assertEqual(wb[sheet].max_row,1)
            self.assertFalse(wb[sheet].protection.sheet)

    def test_unprotected_monthly_identity_still_validated(self):
        for sheet in ('Time','Monthly adjustments'):
            wb=self.export()
            self.assertFalse(wb[sheet].protection.sheet)
            self.set_value(wb,sheet,'Employee ID','000002')
            self.set_value(wb,'Monthly adjustments','Monthly bonus',50)
            review=self.review(wb)
            self.assertTrue(review['issues'])
            with self.assertRaises(UserError):self.apply(review)
        self.assertFalse(self.env['pb.payroll.month.inputs'].search([
            ('config_id','=',self.cfg.id)]))

    def test_zero_blank_and_selected_period(self):
        wb=self.export();self.set_value(wb,'Monthly adjustments','Monthly bonus',70)
        self.apply(self.review(wb))
        wb=self.export();self.set_value(wb,'Monthly adjustments','Monthly bonus',None)
        self.assertEqual(self.review(wb)['changes'],[])
        self.assertEqual(self.api._current(self.cfg,'2026-10',self.emp)[self.emp.id]['rule:%s'%self.bonus.id],70)
        self.set_value(wb,'Monthly adjustments','Monthly bonus',0)
        self.apply(self.review(wb))
        self.assertEqual(self.api._current(self.cfg,'2026-10',self.emp)[self.emp.id]['rule:%s'%self.bonus.id],0)
        next_month=self.export(period='2026-11')
        col=[c.value for c in next_month['Monthly adjustments'][1]].index('Monthly bonus')+1
        self.assertIsNone(next_month['Monthly adjustments'].cell(2,col).value)

    def test_records_diff_apply_and_blank_ignore(self):
        wb=self.export('records');self.set_value(wb,'Employee records','Salary',1200)
        review=self.review(wb,'records')
        self.assertFalse(review['issues']);self.assertEqual(len(review['changes']),1)
        self.assertEqual(self.contract.wage,1000)
        res=self.apply(review);self.assertFalse(res['pending']);self.assertEqual(self.contract.wage,1200)

    def test_employee_name_editable_identity_locked(self):
        wb=self.export('records');ws=wb['Employee records']
        self.assertTrue(ws.protection.sheet)
        self.assertEqual(ws['A1'].value,'Employee ID');self.assertEqual(ws['B1'].value,'Employee email')
        self.assertTrue(ws['A2'].protection.locked);self.assertTrue(ws['B2'].protection.locked)
        self.assertEqual(ws['C2'].value,self.emp.name);self.assertFalse(ws['C2'].protection.locked)
        ws['C2']='Workbook renamed'
        r=self.review(wb,'records');self.assertFalse(r['issues']);self.assertEqual(len(r['changes']),1)

    def test_id_and_email_cannot_change(self):
        for label,value in [('Employee ID','000002'),('Employee email','other@example.test')]:
            wb=self.export('records');self.set_value(wb,'Employee records',label,value)
            self.set_value(wb,'Employee records','Salary',1200)
            review=self.review(wb,'records');self.assertTrue(review['issues'])
            with self.assertRaises(UserError):self.apply(review)
        self.assertEqual(self.contract.wage,1000)

    def test_stale_review_is_atomic_and_cannot_replay(self):
        wb=self.export();self.set_value(wb,'Monthly adjustments','Monthly bonus',15)
        review=self.review(wb);self.apply(review)
        with self.assertRaises(UserError):self.apply(review)
        self.set_value(wb,'Monthly adjustments','Monthly bonus',20)
        stale=self.review(wb)
        self.assertFalse(stale['issues'])
        self.assertTrue(stale['warnings'])
        self.assertEqual(self.cell(stale,self.bonus)['value'],20)
        row=self.env['pb.payroll.month.inputs'].search([('config_id','=',self.cfg.id),('period','=','2026-10')])
        row.payload=dict(row.payload, **{'rule:%s'%self.bonus.id:30})
        with self.assertRaisesRegex(UserError,'changed after review'):self.apply(stale)

    def test_concurrent_change_after_review(self):
        wb=self.export('records');self.set_value(wb,'Employee records','Salary',1200)
        review=self.review(wb,'records');self.contract.wage=1100
        with self.assertRaises(UserError):self.apply(review)
        self.assertEqual(self.contract.wage,1100)

    def test_duplicate_and_cross_tab_missing_id(self):
        wb=self.export();ws=wb['Time'];ws.append([c.value for c in ws[2]])
        self.assertTrue(self.review(wb)['issues'])
        wb=self.export();wb['Time'].delete_rows(2)
        self.assertTrue(self.review(wb)['issues'])

    def test_wrong_month_refused_but_moved_columns_preserve_uploaded_values(self):
        wb=self.export()
        with self.assertRaises(UserError):self.review(wb,period='2026-11')
        self.set_value(wb,'Monthly adjustments','Monthly bonus',55)
        d=self.api.workspace(self.cfg.id,'2026-10');layout=d['layout'];fid='rule:%s'%self.bonus.id
        layout['adjustments'].remove(fid);layout['time'].append(fid)
        self.api.save_layout(self.cfg.id,layout)
        review=self.review(wb)
        self.assertFalse(review['issues'])
        self.assertTrue(review['warnings'])
        self.assertEqual(self.cell(review,self.bonus)['value'],55)
        self.assertEqual(self.cell(review,self.bonus)['source'],'excel')
        self.assertEqual(next(c for c in review['grid']['columns'] if c['id']==fid)['tab'],'Time')
        self.apply(review)
        self.assertEqual(self.api._current(self.cfg,'2026-10',self.emp)[self.emp.id][fid],55)
        fresh=self.export();self.assertIn('Monthly bonus',[c.value for c in fresh['Time'][1]])

    def test_readonly_salary_context(self):
        wb=self.export()
        self.assertEqual(wb['Employee records'].cell(2,[c.value for c in wb['Employee records'][1]].index('Salary')+1).value,1000)
        self.assertTrue(wb['Employee records'].protection.sheet)
        self.set_value(wb,'Employee records','Salary',1500)
        review=self.review(wb)
        self.assertFalse(review['issues'])
        fid='rule:%s'%self.salary.id
        self.assertEqual(review['grid']['rows'][0]['cells'][fid]['value'],1000)
        self.assertEqual(self.contract.wage,1000)

    def test_engine_aliases_and_recurring_records(self):
        wb=self.export();self.set_value(wb,'Time','Paid days',20);self.set_value(wb,'Monthly adjustments','Monthly bonus',0)
        self.apply(self.review(wb));raw,ids=self.api._engine_file(self.cfg,'2026-10')
        probe=self.env['hr.payroll.import.batch'].with_context(pb_template_engine=True).new({'formula_config_id':self.cfg.id})
        parsed=probe._parse_source_file(raw,'file.xlsx');blob=probe._raw_data_from_row(parsed['headers'],parsed['rows'][0])
        headers = list(openpyxl.load_workbook(io.BytesIO(raw)).active.values)[0]
        self.assertIn('Paid days', headers)
        self.assertNotIn('Payroll|Paid days', headers)
        aliases=probe._template_alias_row(blob)
        self.assertEqual(aliases['Payroll|Paid days'],20)
        self.assertEqual(aliases['Payroll|Monthly bonus'],0)
        self.assertEqual(aliases['Payroll|Salary'],1000)
        self.assertEqual(self.contract.wage,1000)

    def test_current_recurring_salary_wins_over_saved_monthly_zero(self):
        self.env['pb.payroll.month.inputs'].create({'config_id':self.cfg.id,
            'company_id':self.env.company.id,'employee_id':self.emp.id,'period':'2026-10',
            'payload':{'rule:%s'%self.salary.id:0,'rule:%s'%self.paid.id:20,
                       'rule:%s'%self.bonus.id:0}})
        for salary in (1000,1200,0):
            self.contract.wage=salary
            raw,_ids=self.api._engine_file(self.cfg,'2026-10')
            wb=openpyxl.load_workbook(io.BytesIO(raw))
            row=dict(zip([c.value for c in wb.active[1]],[c.value for c in wb.active[2]]))
            self.assertEqual(row['Salary'],salary)
            self.assertEqual(row['Monthly bonus'],0)

    def test_salary_candidate_uses_scheme_currency(self):
        cfg=self.env['hr.formula.config'].create({'name':'VND salary source',
            'code':'WBVNDSALARY','country_code':'VN','company_id':self.env.company.id})
        rules=self.env['hr.formula.rule'].create([
            {'config_id':cfg.id,'name':'Gross contract salary USD','code':'WBUSD',
             'column_type':'input','value_kind':'money','sequence':1},
            {'config_id':cfg.id,'name':'Gross contract salary VND','code':'WBVND',
             'column_type':'input','value_kind':'money','sequence':2}])
        cards=self.env['pb.records.desk']._cards(cfg.id)
        self.assertEqual(cards['f:hr.contract:wage']['component']['id'],rules[1].id)
        vnd=next(c for c in self.api._catalog(cfg) if c['code']=='WBVND')
        self.assertTrue(vnd['readonly'])
        self.assertEqual(vnd['record_id'],'f:hr.contract:wage')

    def test_ambiguous_salary_does_not_choose_first_column(self):
        cfg=self.env['hr.formula.config'].create({'name':'Ambiguous salary source',
            'code':'WBAMBSALARY','country_code':'SG','company_id':self.env.company.id})
        for name,code in [('Basic salary','WBASALA'),('Contract salary','WBASALB')]:
            self.env['hr.formula.rule'].create({'config_id':cfg.id,'name':name,'code':code,
                'column_type':'input','value_kind':'money'})
        cards=self.env['pb.records.desk']._cards(cfg.id)
        self.assertIsNone(cards['f:hr.contract:wage']['component'])

    def test_missing_recurring_salary_is_not_silently_zero(self):
        wb=self.export();self.set_value(wb,'Time','Paid days',20)
        self.apply(self.review(wb))
        values=self.api._record_values(self.cfg,self.emp,self.api._cards(self.cfg))
        values[self.emp.id]['f:hr.contract:wage']=None
        with patch.object(type(self.api),'_record_values',return_value=values):
            with self.assertRaisesRegex(UserError,'Recurring salary is missing'):
                self.api._engine_file(self.cfg,'2026-10')

    def test_saved_monthly_batch_refreshes_recurring_salary(self):
        batch=self.env['hr.payroll.import.batch'].create({'name':'Monthly recompute',
            'formula_config_id':self.cfg.id,'company_id':self.env.company.id,
            'one_time':True,'date_from':'2026-10-01','date_to':'2026-10-31',
            'import_filename':'Payobook-month-2026-10.xlsx'})
        raw={'Payroll|Salary':0,'Payroll|Paid days':20,'Payroll|Monthly bonus':0}
        provenance={}
        values=batch._transform_data_to_formula_inputs(raw,contract=self.contract,
            employee=self.emp,provenance=provenance)
        self.assertEqual(values[self.salary.code],1000)
        self.assertEqual(values[self.bonus.code],0)
        self.assertEqual(raw['Payroll|Salary'],0)
        self.assertEqual(provenance[self.salary.code]['src'],'contract_field')
        batch.import_filename='ordinary-upload.xlsx'
        ordinary=batch._transform_data_to_formula_inputs(raw,contract=self.contract,employee=self.emp)
        self.assertEqual(ordinary[self.salary.code],0)

    def test_missing_time_blocks_computation(self):
        with self.assertRaises(UserError):self.api._engine_file(self.cfg,'2026-10')

    def test_manual_time_inputs_are_available(self):
        self.paid.clear_source_binding('excel');self.paid.data_source='manual'
        d=self.api.workspace(self.cfg.id,'2026-10')
        self.assertIn('rule:%s'%self.paid.id,d['layout']['time'])

    def test_reusable_template_is_blank_and_imports_into_chosen_month(self):
        self.env['pb.payroll.month.inputs'].create({'config_id':self.cfg.id,
            'company_id':self.env.company.id,'employee_id':self.emp.id,'period':'2026-10',
            'payload':{'rule:%s'%self.bonus.id:99}})
        wb=self.export('template')
        self.assertIsNone(wb['_Payobook']['B3'].value)
        bonus_col=[c.value for c in wb['Monthly adjustments'][1]].index('Monthly bonus')+1
        self.assertIsNone(wb['Monthly adjustments'].cell(2,bonus_col).value)
        self.set_value(wb,'Monthly adjustments','Monthly bonus',25)
        for period in ('2026-10','2026-11'):
            review=self.review(wb,period=period)
            self.assertFalse(review['issues'])
            self.apply(review)
            self.assertEqual(self.api._current(self.cfg,period,self.emp)[self.emp.id]['rule:%s'%self.bonus.id],25)

    def test_reusable_template_still_guards_changes_after_review(self):
        wb=self.export('template');self.set_value(wb,'Monthly adjustments','Monthly bonus',25)
        review=self.review(wb,period='2026-11')
        self.env['pb.payroll.month.inputs'].create({'config_id':self.cfg.id,
            'company_id':self.env.company.id,'employee_id':self.emp.id,'period':'2026-11',
            'payload':{'rule:%s'%self.bonus.id:30}})
        with self.assertRaisesRegex(UserError,'changed after review'):
            self.apply(review)

    def test_calculated_fields_are_reference_only_and_cannot_enter_layout(self):
        computed=self.env['hr.formula.rule'].create({'config_id':self.cfg.id,
            'name':'Chênh lệch thuế TNCN 2 (PIT gap 2)','code':'WBPITGAP',
            'column_type':'formula','excel_formula':'=1+1','value_kind':'money'})
        d=self.api.workspace(self.cfg.id,'2026-10')
        self.assertIn(computed.id,[f['id'] for f in d['calculated_fields']])
        self.assertNotIn('rule:%s'%computed.id,[f['id'] for f in d['fields']])
        for mode in ('monthly','template'):
            wb=self.export(mode)
            for sheet in ('Time','Monthly adjustments'):
                self.assertNotIn(computed.name,[c.value for c in wb[sheet][1]])
        d['layout']['adjustments'].append('rule:%s'%computed.id)
        with self.assertRaises(UserError):self.api.save_layout(self.cfg.id,d['layout'])

    def test_text_record_context_with_legacy_numeric_kind_round_trips(self):
        role=self.rule('Position','WBJOBTITLE','money')
        self.map_field(role,'hr.employee','job_title')
        self.emp.job_title='Payroll specialist'
        d=self.api.workspace(self.cfg.id,'2026-10')
        fid='rule:%s'%role.id
        self.assertIn(fid,d['layout']['records'])
        self.api.save_layout(self.cfg.id,d['layout'])
        for mode in ('monthly','template'):
            wb=self.export(mode)
            review=self.review(wb)
            self.assertFalse(review['issues'])
            self.assertFalse(review['changes'])
            self.set_value(wb,'Employee records','Position','Changed context')
            changed=self.review(wb)
            self.assertFalse(changed['issues'])
            self.assertFalse(changed['changes'])
            self.assertEqual(changed['grid']['rows'][0]['cells'][fid]['value'],'Payroll specialist')

    def test_monthly_zero_cannot_fall_back_to_old_record(self):
        self.map_field(self.bonus,'hr.contract','wage')
        self.cfg.source_priority='records,excel,api'
        probe=self.env['hr.payroll.import.batch'].with_context(pb_template_engine=True,pb_template_monthly_codes=[self.bonus.code]).new({'formula_config_id':self.cfg.id})
        mapping=self.env['hr.payslip.import.mapping'].search([('component_id','=',self.bonus.id)],limit=1)
        self.assertIsNone(probe._mapped_record_value(mapping,contract=self.contract,employee=self.emp))
        normal=self.env['hr.payroll.import.batch'].new({'formula_config_id':self.cfg.id})
        self.assertEqual(normal._mapped_record_value(mapping,contract=self.contract,employee=self.emp),1000)

    def test_selected_scheme_wins_over_stale_demo_division(self):
        result=self.env['pb.payrun.wizard'].prepare_run({'name':'Workbook selected scheme check','formula_config_id':self.cfg.id,'templates_monthly':True,'division':'retail','date_start':'2089-12-01','date_end':'2089-12-31'})
        self.assertEqual(result['formula_config_id'],self.cfg.id)
        run=self.env['hr.payslip.run'].browse(result['run_id'])
        self.assertEqual(run.pb_formula_config_id,self.cfg)

    def test_draft_configuration_cannot_compute(self):
        self.cfg.state='draft'
        with self.assertRaises(UserError):self.api._engine_file(self.cfg,'2026-10')

    def test_formula_and_invalid_number_are_refused(self):
        for value in ('=1+1','abc'):
            wb=self.export();self.set_value(wb,'Monthly adjustments','Monthly bonus',value)
            self.assertTrue(self.review(wb)['issues'])

    def test_company_scope_and_token_owner(self):
        wb=self.export();token=wb['_Payobook']['B1'].value
        other=self.env['res.users'].create({'name':'Workbook reader','login':'workbook.reader','group_ids':[(6,0,[self.env.ref('pb_records.group_records_viewer').id])]})
        with self.assertRaises(UserError):self.api.with_user(other)._get_exchange(self.cfg,token,'export')
        company=self.env['res.company'].create({'name':'Other workbook company'})
        with self.assertRaises(AccessError):self.api.with_company(company)._config(self.cfg.id)

    def test_payroll_staff_can_open_templates_without_settings_access(self):
        officer = self.env['res.users'].with_context(no_reset_password=True).create({
            'name': 'Template payroll officer', 'login': 'template.officer@test.invalid',
            'company_id': self.env.company.id,
            'company_ids': [(6, 0, self.env.company.ids)],
            'group_ids': [(6, 0, [self.env.ref('base.group_user').id,
                self.env.ref('pb_hr_payroll_base.group_payroll_base_officer').id])],
        })
        self.assertFalse(officer.has_group('base.group_system'))
        workspace = self.api.with_user(officer).workspace(self.cfg.id, '2026-10')
        self.assertEqual(workspace['config_id'], self.cfg.id)
        self.assertTrue(workspace['can_export'])
        exported = self.api.with_user(officer).download(self.cfg.id, '', 'template')
        self.assertTrue(exported['file_b64'])

    def test_template_gate_checks_integration_on_selected_scheme(self):
        wizard = self.env['pb.payrun.wizard']
        vals = {'formula_config_id': self.cfg.id, 'templates_monthly': True}
        self.assertFalse(wizard.spreadsheet_gate(vals)['lanes']['integration_configured'])
        connector = self.env['hr.integration.connector'].create({
            'name': 'Workbook integration', 'connector_type': 'demo'})
        self.cfg.connector_id = connector
        self.assertTrue(wizard.spreadsheet_gate(vals)['lanes']['integration_configured'])

    def cell(self, review, rule):
        return review['grid']['rows'][0]['cells']['rule:%s' % rule.id]

    def test_review_includes_all_inputs_with_sources_and_no_calculateds(self):
        wb=self.export();self.set_value(wb,'Time','Paid days',21)
        review=self.review(wb)
        self.assertEqual(len(review['grid']['rows']),1)
        self.assertEqual({c['id'] for c in review['grid']['columns']},
                         {'rule:%s'%r.id for r in (self.salary,self.paid,self.bonus)})
        self.assertEqual(self.cell(review,self.salary)['source'],'system')
        self.assertIn('Contract record',self.cell(review,self.salary)['detail'])
        self.assertEqual(self.cell(review,self.salary)['value'],1000)
        self.assertEqual(self.cell(review,self.paid)['source'],'excel')
        self.assertEqual(self.cell(review,self.bonus)['source'],'default')
        self.assertFalse(self.cell(review,self.salary)['has_prior'])
        self.assertEqual(review['previous_period'],'2026-09')

    def test_rerun_compares_previous_calendar_month_not_same_month(self):
        fid='rule:%s'%self.bonus.id
        for period,value in [('2026-09',10),('2026-10',20),('2026-08',99)]:
            self.env['pb.payroll.month.inputs'].create({'config_id':self.cfg.id,
                'company_id':self.env.company.id,'employee_id':self.emp.id,
                'period':period,'payload':{fid:value}})
        wb=self.export();review=self.review(wb)
        cell=self.cell(review,self.bonus)
        self.assertEqual(cell['value'],20)
        self.assertEqual(cell['prior'],10)
        self.assertTrue(cell['changed'])
        self.assertEqual(review['changes'],[])
        self.apply(review)
        again=self.review(self.export())
        self.assertEqual(self.cell(again,self.bonus)['prior'],10)
        self.assertTrue(self.cell(again,self.bonus)['changed'])

    def test_no_history_does_not_invent_previous_record_salary(self):
        review=self.review(self.export())
        self.assertFalse(self.cell(review,self.salary)['has_prior'])
        self.assertEqual(review['stats']['changed'],0)
        _period,prior=self.api._previous_inputs(self.cfg,'2026-01',self.emp,self.api._review_catalog(self.cfg))
        self.assertEqual(_period,'2025-12')
        self.assertEqual(prior[self.emp.id],{})

    def test_finalized_prior_payroll_wins_over_draft_and_saved_inputs(self):
        import json
        self.env['pb.payroll.month.inputs'].create({'config_id':self.cfg.id,
            'company_id':self.env.company.id,'employee_id':self.emp.id,
            'period':'2026-09','payload':{'rule:%s'%self.bonus.id:999}})
        def slip(state,month,amount,config=None,credit=False):
            return self.env['hr.payslip'].create({'name':'Prior input review test',
                'employee_id':self.emp.id,'contract_id':self.contract.id,
                'company_id':self.env.company.id,'date_from':month+'-01',
                'date_to':month+'-28','state':state,'credit_note':credit,
                'formula_config_id':(config or self.cfg).id,
                'formula_input_values':json.dumps({self.salary.code:900,self.bonus.code:amount})})
        slip('done','2026-09',10)
        slip('draft','2026-09',15)
        slip('done','2026-10',20)
        slip('done','2026-09',22,credit=True)
        review=self.review(self.export())
        self.assertEqual(self.cell(review,self.bonus)['prior'],10)
        self.assertEqual(self.cell(review,self.bonus)['baseline_kind'],'finalized_payroll')
        self.assertEqual(self.cell(review,self.salary)['prior'],900)
        self.assertTrue(self.cell(review,self.salary)['changed'])

    def test_inline_edit_validates_and_round_trips_run_only_salary_override(self):
        wb=self.export();self.set_value(wb,'Time','Paid days',21)
        review=self.review(wb)
        result=self.api.confirm_monthly_review(self.cfg.id,review['token'],[
            {'emp_id':self.emp.id,'field_id':'rule:%s'%self.salary.id,'value':'1200'},
            {'emp_id':self.emp.id,'field_id':'rule:%s'%self.bonus.id,'value':'0'}])
        self.assertEqual(self.contract.wage,1000)
        self.assertEqual(result['reviewed'],3)
        raw,_ids=self.api._engine_file(self.cfg,'2026-10')
        row=dict(zip(*list(openpyxl.load_workbook(io.BytesIO(raw)).active.values)))
        self.assertEqual(row['Salary'],1200)
        batch=self.env['hr.payroll.import.batch'].create({'name':'Inline edit recompute',
            'formula_config_id':self.cfg.id,'company_id':self.env.company.id,
            'one_time':True,'date_from':'2026-10-01','date_to':'2026-10-31',
            'import_filename':'Payobook-month-2026-10.xlsx'})
        provenance={}
        values=batch._transform_data_to_formula_inputs({'Payroll|Salary':0,'Payroll|Paid days':21},
            contract=self.contract,employee=self.emp,provenance=provenance)
        self.assertEqual(values[self.salary.code],1200)
        self.assertEqual(provenance[self.salary.code]['via'],'monthly_review')
        self.assertEqual(self.contract.wage,1000)
        next_values=self.api._resolved_month(self.cfg,'2026-11',self.emp,self.api._review_catalog(self.cfg))
        self.assertEqual(next_values[self.emp.id]['rule:%s'%self.salary.id]['value'],1000)
        reopened=self.api.reopen_monthly_review(self.cfg.id,result['review_token'])
        self.assertEqual(self.cell(reopened,self.salary)['value'],1200)
        self.api.confirm_monthly_review(self.cfg.id,reopened['token'],[])
        with self.assertRaises(UserError):self.api.confirm_monthly_review(self.cfg.id,review['token'],[])

    def test_inline_edit_cannot_inject_field_or_employee_and_invalid_is_atomic(self):
        wb=self.export();review=self.review(wb)
        for edit in [
            {'emp_id':self.emp.id,'field_id':'f:hr.contract:wage','value':9999},
            {'emp_id':0,'field_id':'rule:%s'%self.bonus.id,'value':9999},
            {'emp_id':self.emp.id,'field_id':'rule:%s'%self.bonus.id,'value':'=1+1'},
            {'emp_id':self.emp.id,'field_id':'rule:%s'%self.bonus.id,'value':'NaN'},
        ]:
            with self.assertRaises(UserError):self.api.confirm_monthly_review(self.cfg.id,review['token'],[edit])
        self.assertFalse(self.env['pb.payroll.month.inputs'].search([('config_id','=',self.cfg.id)]))
        self.assertEqual(self.contract.wage,1000)

    def test_invalid_uploaded_cell_can_be_corrected_without_reupload(self):
        wb=self.export();self.set_value(wb,'Monthly adjustments','Monthly bonus','invalid')
        review=self.review(wb)
        self.assertEqual(self.cell(review,self.bonus)['issue'],'Enter a valid number.')
        self.api.confirm_monthly_review(self.cfg.id,review['token'],[
            {'emp_id':self.emp.id,'field_id':'rule:%s'%self.bonus.id,'value':30}])
        self.assertEqual(self.api._current(self.cfg,'2026-10',self.emp)[self.emp.id]['rule:%s'%self.bonus.id],30)
        reopened=self.api.reopen_monthly_review(self.cfg.id,review['token'])
        self.api.confirm_monthly_review(self.cfg.id,reopened['token'],[
            {'emp_id':self.emp.id,'field_id':'rule:%s'%self.bonus.id,'value':None}])
        self.assertEqual(self.api._current(self.cfg,'2026-10',self.emp)[self.emp.id]['rule:%s'%self.bonus.id],30)

    def test_database_changes_after_review_block_even_if_no_upload_delta(self):
        review=self.review(self.export())
        self.contract.wage=1100
        with self.assertRaisesRegex(UserError,'changed after review'):
            self.api.confirm_monthly_review(self.cfg.id,review['token'],[])
        self.assertFalse(self.env['pb.payroll.month.inputs'].search([('config_id','=',self.cfg.id)]))

    def test_previous_month_excludes_other_schemes_and_older_months(self):
        other=self.env['hr.formula.config'].create({'name':'Other review scheme',
            'code':'OTHERREVIEW','country_code':'SG','company_id':self.env.company.id})
        for config,period in [(other,'2026-09'),(self.cfg,'2026-08')]:
            self.env['pb.payroll.month.inputs'].create({'config_id':config.id,
                'company_id':self.env.company.id,'employee_id':self.emp.id,
                'period':period,'payload':{'rule:%s'%self.bonus.id:50}})
        review=self.review(self.export())
        self.assertFalse(self.cell(review,self.bonus)['has_prior'])

    def test_review_snapshot_preserves_prior_system_record_value(self):
        self.apply(self.review(self.export(period='2026-09'),period='2026-09'))
        self.contract.wage=1200
        review=self.review(self.export())
        self.assertEqual(self.cell(review,self.salary)['prior'],1000)
        self.assertEqual(self.cell(review,self.salary)['value'],1200)
        self.assertEqual(self.cell(review,self.salary)['baseline_kind'],'reviewed_inputs')
        self.assertTrue(self.cell(review,self.salary)['changed'])

    def test_renamed_field_and_legacy_fingerprint_keep_workbook_entries(self):
        wb=self.export()
        self.set_value(wb,'Monthly adjustments','Monthly bonus',65)
        self.bonus.name='Updated monthly bonus label'
        review=self.review(wb)
        self.assertFalse(review['issues'])
        self.assertEqual(self.cell(review,self.bonus)['value'],65)
        self.assertTrue(review['warnings'])
        self.assertIn('Updated monthly bonus label',[c['label'] for c in review['grid']['columns']])

    def test_removed_input_cannot_silently_drop_filled_workbook_values(self):
        wb=self.export();self.set_value(wb,'Monthly adjustments','Monthly bonus',80)
        self.cfg.state='draft'
        self.bonus.unlink()
        with self.assertRaisesRegex(UserError,'Monthly bonus'):
            self.review(wb)
        self.assertFalse(self.env['pb.payroll.month.inputs'].search([('config_id','=',self.cfg.id)]))

    def test_changed_record_source_requires_new_workbook(self):
        wb=self.export()
        self.set_value(wb,'Monthly adjustments','Monthly bonus',100)
        self.map_field(self.bonus,'hr.contract','meal_allowance')
        with self.assertRaisesRegex(UserError,'saved-record source changed'):
            self.review(wb)

    def test_record_lane_rejects_monthly_inputs_and_matches_records_download(self):
        d=self.api.workspace(self.cfg.id,'2026-10');layout=d['layout']
        paid='rule:%s'%self.paid.id
        layout['time'].remove(paid);layout['records'].append(paid)
        with self.assertRaisesRegex(UserError,'Employee records accepts only'):
            self.api.save_layout(self.cfg.id,layout)
        layout=self.api.workspace(self.cfg.id,'2026-10')['layout']
        salary='rule:%s'%self.salary.id
        layout['records'].remove(salary);layout['available'].append(salary)
        self.api.save_layout(self.cfg.id,layout)
        for mode in ('monthly','records'):
            self.assertNotIn('Salary',[c.value for c in self.export(mode)['Employee records'][1]])

    def test_master_snapshot_cannot_write_records_or_monthly_values(self):
        wb=self.export()
        self.set_value(wb,'Employee records','Salary','=999999*2')
        self.set_value(wb,'Monthly adjustments','Monthly bonus',25)
        r=self.review(wb)
        self.assertFalse(r['issues'])
        self.assertEqual(self.cell(r,self.salary)['value'],1000)
        self.apply(r)
        self.assertEqual(self.contract.wage,1000)
        self.assertEqual(self.api._current(self.cfg,'2026-10',self.emp)[self.emp.id]['rule:%s'%self.bonus.id],25)
        wb=self.export();del wb['Employee records']
        self.assertFalse(self.review(wb)['issues'])

    def schema_identity(self):
        code = self.rule('Mã nhân viên\n(Code)', 'WBEMPCODE', 'identifier')
        name = self.rule('Họ và tên (Full name)', 'WBEMPNAME', 'text')
        self.map_field(code, 'hr.employee', 'employee_id')
        self.map_field(name, 'hr.employee', 'name')
        self.emp.employee_id = 'SCHEMA-001'
        return code, name

    def test_schema_identity_headers_and_mapped_values_round_trip(self):
        code, name = self.schema_identity()
        for mode in ('monthly', 'template'):
            wb = self.export(mode)
            for sheet in ('Time', 'Monthly adjustments'):
                self.assertEqual([c.value for c in wb[sheet][1]][:2], [code.name, name.name])
                self.assertEqual(wb[sheet]['A2'].value, 'SCHEMA-001')
                self.assertEqual(wb[sheet]['B2'].value, self.emp.name)
                self.assertNotIn('employee_code', [c.value for c in wb[sheet][1]])
            self.assertFalse(self.review(wb)['issues'])
        data = self.api.workspace(self.cfg.id, '2026-10')
        self.assertEqual(data['identity_labels'], {'code': code.name, 'name': name.name})

    def test_schema_records_headers_and_values_round_trip(self):
        code, name = self.schema_identity()
        wb = self.export('records')
        headers = [c.value for c in wb['Employee records'][1]]
        self.assertEqual(headers[:2], [code.name, self.email.name])
        self.assertIn(name.name, headers)
        self.assertIn(self.salary.name, headers)
        self.set_value(wb, 'Employee records', self.salary.name, 1250)
        reviewed = self.review(wb, 'records')
        self.assertFalse(reviewed['issues'])
        self.assertEqual(reviewed['changes'][0]['value'], 1250)

    def test_records_shared_component_names_keep_stable_field_columns(self):
        cards = self.api._cards(self.cfg)
        salary = dict(cards['f:hr.contract:wage'])
        cards['f:hr.contract:name'] = dict(salary, id='f:hr.contract:name',
            field='name', ttype='char', editable=True)
        with patch.object(type(self.api), '_cards', return_value=cards):
            wb = self.export('records')
            ws = wb['Employee records']
            self.assertEqual([c.value for c in ws[1]].count(self.salary.name), 2)
            columns = {c.comment.text.splitlines()[0][7:]: c.column for c in ws[1]
                if c.comment and c.comment.text.startswith('field: ')}
            ws.cell(2, columns['f:hr.contract:wage'], 1250)
            review = self.review(wb, 'records')
            self.assertFalse(review['issues'])
            self.assertEqual(len(review['changes']), 1)
            self.assertEqual(review['changes'][0]['field_id'], 'f:hr.contract:wage')

    def test_repeated_component_names_keep_distinct_monthly_and_engine_values(self):
        other = self.rule(self.bonus.name, 'WBOTHERBONUS', 'money')
        other.set_source_binding('excel', 'Other|' + other.name)
        self.env['pb.payroll.month.inputs'].create({'config_id': self.cfg.id,
            'company_id': self.env.company.id, 'employee_id': self.emp.id, 'period': '2026-10',
            'payload': {'rule:%s' % self.paid.id: 20, 'rule:%s' % self.bonus.id: 45,
                        'rule:%s' % other.id: 90}})
        wb = self.export()
        self.assertEqual([c.value for c in wb['Monthly adjustments'][1]].count(self.bonus.name), 2)
        reviewed = self.review(wb)
        self.assertFalse(reviewed['issues'])
        self.assertEqual(self.cell(reviewed, self.bonus)['value'], 45)
        self.assertEqual(self.cell(reviewed, other)['value'], 90)
        raw, ids = self.api._engine_file(self.cfg, '2026-10')
        headers = list(openpyxl.load_workbook(io.BytesIO(raw)).active.values)[0]
        self.assertEqual(headers.count(self.bonus.name), 2)
        probe = self.env['hr.payroll.import.batch'].with_context(pb_template_engine=True,
            pb_template_monthly_codes=[self.paid.code, self.bonus.code, other.code]).new({'formula_config_id': self.cfg.id})
        parsed = probe._parse_source_file(raw, 'generated.xlsx')
        blob = probe._raw_data_from_row(parsed['headers'], parsed['rows'][0])
        values = probe._transform_data_to_formula_inputs(blob, contract=self.contract, employee=self.emp)
        self.assertEqual(values[self.bonus.code], 45)
        self.assertEqual(values[other.code], 90)

    def test_schema_identity_rename_uses_download_receipt(self):
        code, name = self.schema_identity()
        wb = self.export()
        code.name = 'Updated employee code'
        name.name = 'Updated employee name'
        self.set_value(wb, 'Monthly adjustments', self.bonus.name, 25)
        reviewed = self.review(wb)
        self.assertFalse(reviewed['issues'])
        self.assertEqual(self.cell(reviewed, self.bonus)['value'], 25)

    def test_schema_identity_edits_and_duplicate_headers_refused(self):
        code, name = self.schema_identity()
        wb = self.export()
        self.set_value(wb, 'Time', code.name, 'SOMEONE-ELSE')
        self.assertTrue(self.review(wb)['issues'])
        wb = self.export()
        wb['Time']['C1'] = code.name
        self.assertTrue(self.review(wb)['issues'])
        self.emp.employee_id = ''
        with self.assertRaisesRegex(UserError, 'unique value'):
            self.export()

    def test_legacy_generic_identity_download_still_imports(self):
        wb = self.export()
        receipt = self.env['pb.payroll.workbook.exchange'].search([('token', '=', wb['_Payobook']['B1'].value)])
        payload = dict(receipt.payload)
        payload.pop('identity_columns')
        for person in payload['employees'].values():
            person.pop('identity_values')
        receipt.payload = payload
        self.schema_identity()
        self.set_value(wb, 'Monthly adjustments', self.bonus.name, 30)
        reviewed = self.review(wb)
        self.assertFalse(reviewed['issues'])
        self.assertEqual(self.cell(reviewed, self.bonus)['value'], 30)

    def test_engine_uses_only_component_headers_and_resolves_old_bindings(self):
        code, name = self.schema_identity()
        self.bonus.set_source_binding('excel', 'C')
        self.env['pb.payroll.month.inputs'].create({'config_id': self.cfg.id,
            'company_id': self.env.company.id, 'employee_id': self.emp.id, 'period': '2026-10',
            'payload': {'rule:%s' % self.paid.id: 20, 'rule:%s' % self.bonus.id: 45}})
        raw, ids = self.api._engine_file(self.cfg, '2026-10')
        rows = list(openpyxl.load_workbook(io.BytesIO(raw)).active.values)
        self.assertEqual(rows[0][:2], (code.name, name.name))
        self.assertEqual(rows[1][:2], ('SCHEMA-001', self.emp.name))
        self.assertNotIn('employee_code', rows[0])
        self.assertNotIn('Employee ID', rows[0])
        self.assertNotIn('email', rows[0])
        self.assertTrue(all(h in self.cfg.rule_ids.mapped('name') for h in rows[0]))
        probe = self.env['hr.payroll.import.batch'].with_context(pb_template_engine=True,
            pb_template_monthly_codes=[self.paid.code, self.bonus.code]).new({'formula_config_id': self.cfg.id})
        blob = probe._raw_data_from_row(rows[0], rows[1])
        self.assertNotIn('C', blob)
        self.assertEqual(probe._identity_from_file_row(blob)[0], 'SCHEMA-001')
        values = probe._transform_data_to_formula_inputs(blob, contract=self.contract, employee=self.emp)
        self.assertEqual(values[self.paid.code], 20)
        self.assertEqual(values[self.bonus.code], 45)
        self.assertEqual(values[self.salary.code], 1000)
        self.assertNotIn('Payroll|Paid days', blob)

    def test_generated_mapping_folds_legacy_helpers_only(self):
        code, name = self.schema_identity()
        Batch = self.env['hr.payroll.import.batch']
        raw = Batch._raw_data_from_row(['Employee ID', 'employee_code', 'Employee name',
            'employee_name', 'email', 'Payroll|' + code.name, 'Payroll|' + name.name,
            'Payroll|Paid days'], ['SCHEMA-001', 'SCHEMA-001', self.emp.name,
            self.emp.name, self.emp.work_email, 'SCHEMA-001', self.emp.name, 20])
        batch = Batch.create({'name': 'Schema headings mapping', 'formula_config_id': self.cfg.id,
            'one_time': True, 'date_from': '2026-10-01', 'date_to': '2026-10-31',
            'import_filename': 'Payobook-month-2026-10.xlsx'})
        self.env['hr.payroll.import.line'].create({'batch_id': batch.id, 'raw_data_json': json.dumps(raw)})
        studio = self.env['pb.formula.studio']
        data = studio.import_mapping_data(self.cfg.id, batch.id)
        labels = [c['label'] for c in data['left'] if c['id'].startswith('c:')]
        self.assertEqual(labels.count(code.name), 1)
        self.assertEqual(labels.count(name.name), 1)
        self.assertFalse({'Employee ID', 'employee_code', 'Employee name', 'employee_name', 'email'} & set(labels))
        self.assertTrue(all(w['leftId'] in {c['id'] for c in data['left']} for w in data['wires']))
        self.assertEqual(batch.import_line_ids.get_raw_data(), raw)
        batch.import_filename = 'User-supplied.xlsx'
        data = studio.import_mapping_data(self.cfg.id, batch.id)
        labels = [c['label'] for c in data['left']]
        self.assertIn('employee_code', labels)
        self.assertIn('email', labels)

    def test_all_aliases_of_one_saved_field_stay_in_records(self):
        alias = self.rule('Other salary label', 'WBALIAS', 'money')
        self.map_field(alias, 'hr.contract', 'wage')
        catalog = self.api._catalog(self.cfg)
        for rule in (self.salary, alias):
            item = next(c for c in catalog if c['id'] == 'rule:%s' % rule.id)
            self.assertTrue(item['record_eligible'])
            self.assertEqual(item['default'], 'records')
        workbook = self.export('records')
        headers = [c.value for c in workbook['Employee records'][1]]
        self.assertIn(self.salary.name, headers)
        self.assertIn(alias.name, headers)

        self.assertFalse(self.review(workbook, 'records')['issues'])
        self.set_value(workbook, 'Employee records', self.salary.name, 1250)
        conflicting = self.review(workbook, 'records')
        self.assertTrue(any('conflicting values' in i for i in conflicting['issues']))
        self.set_value(workbook, 'Employee records', alias.name, 1250)
        reviewed = self.review(workbook, 'records')
        self.assertFalse(reviewed['issues'])
        self.assertEqual(len(reviewed['changes']), 1)
        self.apply(reviewed)
        self.assertEqual(self.contract.wage, 1250)

    def test_readonly_alias_columns_are_distinct_and_checked_separately(self):
        for label, code in [('Saved display name', 'WBDISPLAY1'), ('Other display name', 'WBDISPLAY2')]:
            self.map_field(self.rule(label, code, 'text'), 'hr.employee', 'display_name')
        workbook = self.export('records')
        self.assertFalse(self.review(workbook, 'records')['issues'])
        comments = [c.comment.text.splitlines()[0] for c in workbook['Employee records'][1] if c.comment]
        self.assertEqual(len(comments), len(set(comments)))
        self.set_value(workbook, 'Employee records', 'Saved display name', 'Altered name')
        reviewed = self.review(workbook, 'records')
        self.assertTrue(any('read-only context' in i for i in reviewed['issues']))
        self.assertEqual(self.emp.name, 'Workbook One')

#!/usr/bin/env python3
"""Extract the approved Vietnam calculation master and private review inputs.

Reads the supplied workbooks without changing them. The committed definition
contains no employee data. --private-dir receives reconciliation fixtures and
the October record proposal; keep that directory out of Git.
"""
import argparse
import copy
import datetime as dt
import hashlib
import json
from pathlib import Path
import re

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'RIZE/VIETNAM'
MASTER = 'Rize_VN_Payroll_Report_Aug_2026_FINAL v3.xlsx'
HC = 'VN_HC_TemplateVietnam_updated_07Oct2026.xlsx'
DATA = 'Rize_Vietnam_Payroll_Template_DATA_v3.xlsx'
CODE = 'RIZE_VIETNAM_MASTER_V3'


def prepare(private_dir, net_correction):
    base = json.loads((ROOT / 'docs/configuration/rize-vietnam-hc-202610.json').read_text())
    wf = openpyxl.load_workbook(SOURCE / MASTER, data_only=False)
    wv = openpyxl.load_workbook(SOURCE / MASTER, data_only=True)
    hc = openpyxl.load_workbook(SOURCE / HC, data_only=True)['VN']
    data = openpyxl.load_workbook(SOURCE / DATA, data_only=True)['Salary']
    salary = wv['Salary']
    rules = [copy.deepcopy(r) for r in base['rules'] if r.get('source_column')
             or r['group'] == 'Recurring payroll settings' or r['type'] == 'constant'
             or r['code'] in ('VNSTDDAYS', 'VNPAIDDAYS', 'VNUNPAID', 'VNWEEKHRS',
                              'VNACTBONUS', 'VNBENEFIT', 'VNADVANCE', 'VNDEDUCT', 'VNPITREF')]
    by = {r['code']: r for r in rules}

    def add(code, name, kind='money', group='Monthly adjustments', formula='',
            destination='', source_key='', tab='Monthly adjustments', role='info',
            detail=False, value=None, note=''):
        rule = dict(code=code, name=name, kind=kind, group=group, formula=formula,
                    destination=destination, source_key=source_key, source_column='',
                    tab=tab, role=role, detail=detail, value=value,
                    type='formula' if formula else 'constant' if value is not None else 'input',
                    note=note)
        if code in by:
            by[code].update(rule)
        else:
            rules.append(rule)
            by[code] = rule
        return rule

    def monthly(code, name, col=None, kind='money', tab='Monthly adjustments', note=''):
        key = ''
        if col:
            # The supplied DATA file repeats Others. Positional keys prevent
            # the non-taxable S and taxable AB amounts overwriting one another.
            header = str(data[f'{col}1'].value or '').strip()
            key = 'Salary|' + (col if col in ('S', 'AB') else header)
        return add(code, name, kind=kind, source_key=key, tab=tab,
                   note=note or 'Approved value for this month. Blank starts at zero in a new month.')

    for code, name, kind, note in [
        ('VNPTYPE', 'Payroll contract type', 'text', 'Salary E: Permanent or Probation; distinct from HC Full time. Review October status.'),
        ('VNPAYNAT', 'Payroll nationality class', 'text', 'Salary G: VN or Expat. Expat excludes UI; this does not establish tax residency.'),
        ('VNTUFLAG', 'TU Joining', 'text', 'Salary AR: x enables employee dues, together with Union Count.'),
        ('VNUMEMBER', 'Union Count', 'integer', 'Count matching employee codes in Union list. Not a dependent count.'),
        ('VNINSMODE', 'Insurance salary basis', 'text', 'Contract salary, Salary + private insurance, or Fixed base. Makes the Salary M13 exception explicit.'),
        ('VNUSDSHOW', 'USD payment reference', 'text', 'Yes only when the payroll master requests USD reference figures; payroll remains VND.'),
        ('VNPRIVATE', 'Private Insurance allowance', 'money', 'Salary Z / DATA V. Saved recurring cash allowance; included in insurance basis only under the explicit salary-basis setting.'),
        ('VNGROSSUSD', 'Gross contract salary USD', 'money', 'Salary K: reference only, never added to VND earnings.'),
    ]:
        add(code, name, kind=kind, group='Recurring payroll settings',
            destination='c:text' if kind == 'text' else 'c:amount',
            tab='', source_key='Recurring settings|' + name, note=note)
    by['VNTAXMODE']['note'] = 'Workbook follows Permanent/probation tax rules in Salary BA. Manual uses Reviewed manual PIT. This setting does not certify tax residency.'
    by['VNREGISTER']['note'] = 'Registered tax dependents from payroll Salary AX, kept separate from imported dependent details and HC children count. Review current registration.'
    for code, name, col, kind, tab in [
        ('VNSTDDAYS', 'Standard working days', 'M', 'quantity', 'Time'),
        ('VNPAIDDAYS', 'Paid days', 'N', 'quantity', 'Time'),
        ('VNWEEKHRS', 'Weekend Hours', 'O', 'quantity', 'Time'),
        ('VNUNIFORM', 'Uniform', 'P', 'money', 'Monthly adjustments'),
        ('VNFREETEL', 'Phone (non-taxable)', 'Q', 'money', 'Monthly adjustments'),
        ('VNSEVER', 'Severance', 'R', 'money', 'Monthly adjustments'),
        ('VNFREOTH', 'Others (non-taxable)', 'S', 'money', 'Monthly adjustments'),
        ('VNLEAVEWD', 'Annual leave encashment days', None, 'quantity', 'Time'),
        ('VNLEAVESTD', 'Annual leave salary divisor', None, 'quantity', 'Time'),
        ('VNEXTRANS', 'Additional Transportation', 'U', 'money', 'Monthly adjustments'),
        ('VNY2S2', 'Y2S2 Ags incentive', 'X', 'money', 'Monthly adjustments'),
        ('VNPRODUCT', 'New product launching incentive', 'Y', 'money', 'Monthly adjustments'),
        ('VNACTBONUS', 'Variable Bonus', 'Z', 'money', 'Monthly adjustments'),
        ('VNREFERRAL', 'Refferal incentive', 'AA', 'money', 'Monthly adjustments'),
        ('VNOTHTAX', 'Others (taxable)', 'AB', 'money', 'Monthly adjustments'),
        ('VNPREADD', 'Additional Adjustment before tax', 'AC', 'money', 'Monthly adjustments'),
        ('VNPREDED', 'Deduction Adjustment before tax (signed)', 'AD', 'money', 'Monthly adjustments'),
        ('VNBENEFIT', 'Non-cash benefit', 'AE', 'money', 'Monthly adjustments'),
        ('VNHOLPAY', 'Holiday Overtime', 'AF', 'money', 'Monthly adjustments'),
        ('VNADVANCE', 'Advance', 'AK', 'money', 'Monthly adjustments'),
        ('VNDEDUCT', 'Deduction', 'AL', 'money', 'Monthly adjustments'),
        ('VNNETADD', 'Addition', 'AM', 'money', 'Monthly adjustments'),
        ('VNPITREF', 'PIT refund', 'AN', 'money', 'Monthly adjustments'),
        ('VNTRIP', 'Incentive trip', 'AQ', 'money', 'Monthly adjustments'),
        ('VNPREMIUM', 'Premium insurance for dependent', None, 'money', 'Monthly adjustments'),
        ('VNSUPPORT', 'Support for incident', 'AS', 'money', 'Monthly adjustments'),
        ('VNPITMAN', 'Reviewed manual PIT', None, 'money', 'Monthly adjustments'),
    ]:
        monthly(code, name, col, kind, tab)
    for i, col in enumerate(('F', 'I', 'L', 'O'), 1):
        monthly('VNTECH' + str(i), 'Technical advisory approved payment ' + str(i),
                note='Data ' + col + '136:' + col + '165 approved paid amounts. Four payment buckets sum to Technical advisory allowance; unpaid outstanding amounts and payout percentages are not automatically paid.')
    by['VNPREDED']['note'] = 'Salary AH is added by SUM(Q:AK). Preserve the workbook convention: enter a negative amount for a before-tax reduction.'
    by['VNHOLPAY']['note'] = 'Approved holiday overtime payment. The master has no formula/rate for this column; do not invent one.'
    by['VNUNPAID']['tab'] = ''
    by['VNUNPAID']['note'] = 'Reference only; Paid days is authoritative. Not subtracted again.'
    add('VNBASE', 'Base salary for union cap VND', value=2530000, group='Statutory parameters', tab='', note='Salary AO1; workbook parameter, review effective period.')
    add('VNFX', 'USD reference exchange rate', kind='decimal', value=26071, group='Workbook parameters', tab='', note='Salary BC13/BJ13; not the HC CTC exchange rate. Review when changing month.')

    def formula(code, name, expr, group='Payroll calculation', role='info', detail=False, kind='money', note=''):
        return add(code, name, kind, group, formula=expr, tab='', role=role,
                   detail=detail, note=note or 'Calculation from the approved August FINAL v3 master. No spreadsheet input mapping.')

    formula('VNREADY', 'Payroll input validation',
            'IF(OR({VNSTDDAYS}<=0,{VNPAIDDAYS}<0,{VNPAIDDAYS}>{VNSTDDAYS}),"Review paid days",'
            'IF(OR({VNHCBK}<=0,{VNHCBH}<>"VND",{VNHCBI}<>"Gross"),"Review salary",'
            'IF(AND({VNPTYPE}<>"Permanent",{VNPTYPE}<>"Probation"),"Review payroll contract type",'
            'IF(AND({VNPAYNAT}<>"VN",{VNPAYNAT}<>"Expat"),"Review payroll nationality",'
            'IF(OR(AND({VNSIJOIN}<>"Yes",{VNSIJOIN}<>"No"),AND({VNHIJOIN}<>"Yes",{VNHIJOIN}<>"No"),AND({VNUIJOIN}<>"Yes",{VNUIJOIN}<>"No"),AND({VNTUJOIN}<>"Yes",{VNTUJOIN}<>"No")),"Review participation",'
            'IF(AND({VNINSMODE}<>"Contract salary",{VNINSMODE}<>"Salary + private insurance",{VNINSMODE}<>"Fixed base"),"Review insurance basis",'
            'IF(OR(AND({VNINSMODE}="Fixed base",{VNINSBASE}<0),{VNREGISTER}<0,MOD({VNREGISTER},1)<>0,{VNUMEMBER}<0,MOD({VNUMEMBER},1)<>0,{VNWEEKHRS}<0,{VNLEAVEWD}<0,AND({VNLEAVEWD}>0,{VNLEAVESTD}<=0)),"Review calculation inputs",'
            'IF(AND({VNTAXMODE}<>"Workbook",{VNTAXMODE}<>"Manual"),"Review tax mode","Ready"))))))))', kind='text')
    expressions = [
        ('VNEFFINS', 'Insurance salary used VND', 'IF({VNINSMODE}="Contract salary",{VNHCBK},IF({VNINSMODE}="Salary + private insurance",{VNHCBK}+{VNPRIVATE},{VNINSBASE}))', 'Insurance calculation', 'info', False),
        ('VNINSDELTA', 'Gross salary less insurance salary VND', '{VNHCBK}-{VNEFFINS}', 'Insurance calculation', 'info', False),
        ('VNBASIC', 'Actual gross salary', 'ROUND({VNHCBK}/{VNSTDDAYS}*{VNPAIDDAYS},0)', 'Earnings', 'earning', True),
        ('VNLEAVEPAY', 'Annual leave', 'IF({VNLEAVEWD}=0,0,ROUND({VNHCBK}/{VNLEAVESTD}*{VNLEAVEWD},0))', 'Earnings', 'earning', True),
        ('VNTRANSPAY', 'Transportation', 'ROUND({VNTRANSPORT}/{VNSTDDAYS}*{VNPAIDDAYS},0)', 'Earnings', 'earning', True),
        ('VNPHONEPAY', 'Phone Allowance', 'ROUND({VNPHONE}/{VNSTDDAYS}*{VNPAIDDAYS},0)', 'Earnings', 'earning', True),
        ('VNTECH', 'Technical advisory allowance', 'SUM({VNTECH1},{VNTECH2},{VNTECH3},{VNTECH4})', 'Earnings', 'earning', True),
        ('VNOVERTIME', 'Weekend Overtime', 'ROUND({VNWEEKHRS}*{VNHCBK}/{VNSTDDAYS}/8*200%,0)', 'Earnings', 'earning', True),
        ('VNTOTALINC', 'Total monthly income', 'SUM({VNBASIC},{VNUNIFORM},{VNFREETEL},{VNSEVER},{VNFREOTH},{VNLEAVEPAY},{VNEXTRANS},{VNTRANSPAY},{VNPHONEPAY},{VNPRIVATE},{VNTECH},{VNY2S2},{VNPRODUCT},{VNACTBONUS},{VNREFERRAL},{VNOTHTAX},{VNOTHERFIX},{VNPREADD},{VNPREDED},{VNBENEFIT},{VNOVERTIME},{VNHOLPAY})', 'Totals', 'info', False),
        ('VNCASHGROSS', 'Cash gross pay VND', '{VNTOTALINC}-{VNBENEFIT}', 'Totals', 'earning', False),
        ('VNTAXEXEM', 'Non-taxable income VND', 'SUM({VNUNIFORM},{VNFREETEL},{VNSEVER},{VNFREOTH},{VNOVERTIME},{VNHOLPAY})', 'Tax calculation', 'info', False),
        ('VNTAXABLE', 'Taxable income', '{VNTOTALINC}-{VNTAXEXEM}', 'Tax calculation', 'info', False),
        ('VNSIEMP', 'Employee SI 8%', 'IF({VNSIJOIN}="Yes",ROUND(MIN({VNEFFINS},{VNSICAP})*8%,0),0)', 'Employee deductions', 'deduction', True),
        ('VNHIEMP', 'Employee HI 1.5%', 'IF({VNHIJOIN}="Yes",ROUND(MIN({VNEFFINS},{VNSICAP})*1.5%,0),0)', 'Employee deductions', 'deduction', True),
        ('VNUIEMP', 'Employee UI 1%', 'IF({VNPAYNAT}="Expat",0,IF({VNUIJOIN}="Yes",ROUND(MIN({VNEFFINS},{VNUIREGCAP})*1%,0),0))', 'Employee deductions', 'deduction', True),
        ('VNUNIONEE', 'Employee union dues 0.5%', 'IF({VNTUFLAG}="x",MIN({VNUMEMBER}*0.5%*{VNEFFINS},10%*{VNBASE}),0)', 'Employee deductions', 'deduction', True),
        ('VNINSDED', 'Employee insurance total VND', 'SUM({VNSIEMP},{VNHIEMP},{VNUIEMP})', 'Tax calculation', 'info', False),
        ('VNPERSDED', 'Personal deduction', 'IF({VNPTYPE}="Permanent",{VNPERSONAL},0)', 'Tax calculation', 'info', False),
        ('VNDEPDED', 'Registered dependent deduction', '{VNREGISTER}*{VNDEPAMT}', 'Tax calculation', 'info', False),
        ('VNASSESS', 'Assessable income', 'MAX(0,{VNTAXABLE}-{VNINSDED}-{VNPERSDED}-{VNDEPDED})', 'Tax calculation', 'info', False),
        ('VNNETINC', 'Net income VND', '{VNCASHGROSS}-{VNINSDED}-{VNUNIONEE}-{VNPITDUE}', 'Totals', 'info', False),
        ('VNDEDTOTAL', 'Total employee deductions VND', 'SUM({VNINSDED},{VNUNIONEE},{VNPITDUE},{VNADVANCE},{VNDEDUCT})' if net_correction else 'SUM({VNINSDED},{VNUNIONEE},{VNPITDUE})', 'Totals', 'deduction', False),
        ('VNNETPAY', 'Net payment VND', '{VNNETINC}-{VNADVANCE}-{VNDEDUCT}+{VNNETADD}+{VNPITREF}' if net_correction else '{VNNETINC}+{VNPITREF}', 'Totals', 'net', False),
        ('VNSIER', 'Employer SI 17.5%', 'IF({VNSIJOIN}="Yes",ROUND(MIN({VNEFFINS},{VNSICAP})*17.5%,0),0)', 'Employer contributions', 'employer_cost', True),
        ('VNHIERR', 'Employer HI 3%', 'IF({VNHIJOIN}="Yes",ROUND(MIN({VNEFFINS},{VNSICAP})*3%,0),0)', 'Employer contributions', 'employer_cost', True),
        ('VNUIERR', 'Employer UI 1%', 'IF({VNPAYNAT}="Expat",0,IF({VNUIJOIN}="Yes",ROUND(MIN({VNEFFINS},{VNUIREGCAP})*1%,0),0))', 'Employer contributions', 'employer_cost', True),
        ('VNTUER', 'Employer union contribution 2%', 'IF({VNTUJOIN}="Yes",ROUND(MIN({VNEFFINS},{VNSICAP})*2%,0),0)', 'Employer contributions', 'employer_cost', True),
        ('VNERCONTR', 'Total employer contributions VND', 'SUM({VNSIER},{VNHIERR},{VNUIERR},{VNTUER})', 'Totals', 'employer_cost', False),
        ('VNEMPCTC', 'Total salary cost', '{VNTOTALINC}+{VNERCONTR}', 'Totals', 'info', False),
        ('VNREVTX', 'Revised taxable income — premium and trip', 'SUM({VNTAXABLE},{VNTRIP},{VNPREMIUM})', 'Employer PIT review', 'info', False),
        ('VNREVAI', 'Revised assessable income — premium and trip', 'MAX(0,{VNREVTX}-{VNINSDED}-{VNPERSDED}-{VNDEPDED})', 'Employer PIT review', 'info', False),
        ('VNGAP', 'PIT gap — premium and trip', '{VNREVPIT}-{VNPITDUE}', 'Employer PIT review', 'info', False),
        ('VNSUPTX', 'Revised taxable income — incident support', '{VNTAXABLE}+{VNSUPPORT}', 'Employer PIT review', 'info', False),
        ('VNSUPAI', 'Revised assessable income — incident support', 'MAX(0,{VNSUPTX}-{VNINSDED}-{VNPERSDED}-{VNDEPDED})', 'Employer PIT review', 'info', False),
        ('VNSUPGAP', 'PIT gap — incident support', '{VNSUPPIT}-{VNPITDUE}', 'Employer PIT review', 'info', False),
        ('VNUSDNET', 'Net income USD reference', 'IF({VNUSDSHOW}="Yes",{VNNETINC}/{VNFX},0)', 'Currency reference', 'info', False),
        ('VNUSDPAY', 'Net payment USD reference', 'IF({VNUSDSHOW}="Yes",{VNNETPAY}/{VNFX},0)', 'Currency reference', 'info', False),
    ]
    for code, name, expr, group, role, detail in expressions:
        formula(code, name, 'IF({VNREADY}<>"Ready","Review payroll inputs",' + expr + ')', group, role, detail)
    def tax(assess):
        a = '{' + assess + '}'
        return 'IF({VNPTYPE}="Permanent",ROUND(IF('+a+'<=10000000,'+a+'*5%,IF('+a+'<=30000000,'+a+'*10%-500000,IF('+a+'<=60000000,'+a+'*20%-3500000,IF('+a+'<=100000000,'+a+'*30%-9500000,'+a+'*35%-14500000)))),0),IF('+a+'<2000000,0,ROUND('+a+'*10%,0)))'
    for code, name, assess in [('VNPITRES', 'Workbook PIT', 'VNASSESS'), ('VNREVPIT', 'Revised PIT — premium and trip', 'VNREVAI'), ('VNSUPPIT', 'Revised PIT — incident support', 'VNSUPAI')]:
        formula(code, name, 'IF({VNREADY}<>"Ready","Review payroll inputs",'+tax(assess)+')', 'Tax calculation')
    formula('VNPITDUE', 'PIT withheld VND', 'IF({VNREADY}<>"Ready","Review payroll inputs",IF({VNTAXMODE}="Manual",{VNPITMAN},{VNPITRES}))', 'Employee deductions', 'deduction', True)
    for code in ['VNUNIFORM', 'VNFREETEL', 'VNSEVER', 'VNFREOTH', 'VNEXTRANS', 'VNPRIVATE', 'VNY2S2', 'VNPRODUCT', 'VNACTBONUS', 'VNREFERRAL', 'VNOTHTAX', 'VNPREADD', 'VNPREDED']:
        by[code].update(role='earning', detail=True, group='Earnings')
    for code in ['VNPITREF', 'VNNETADD']:
        by[code].update(role='earning' if net_correction or code == 'VNPITREF' else 'info', detail=False, group='Net payment additions')
    for code in ['VNADVANCE', 'VNDEDUCT']:
        by[code].update(role='deduction' if net_correction else 'info', detail=True, group='Employee deductions')
    by['VNOTHERFIX'].update(role='earning', detail=True)
    # Source type/count transformations are copied from the live configuration,
    # including the owner's separate dependent-details / tax-registration policy.
    required = ['VNHCBK', 'VNHCBH', 'VNHCBI', 'VNSTDDAYS', 'VNPAIDDAYS', 'VNPTYPE',
                'VNPAYNAT', 'VNINSMODE', 'VNTAXMODE', 'VNSIJOIN', 'VNHIJOIN', 'VNUIJOIN', 'VNTUJOIN']
    for r in rules:
        r['required'] = r['code'] in required
        if r['type'] == 'input' and not r.get('source_key') and r['code'] != 'VNHCQ':
            r['source_key'] = ('Recurring settings' if r['destination'] else r['tab'] or 'Reference') + '|' + r['name']
    result = dict(code=CODE, name='Rize Vietnam Payroll Master v3', country='VN',
                  source_config_id=4639, rules=rules,
                  sources=[dict(file=n, sha256=hashlib.sha256((SOURCE/n).read_bytes()).hexdigest(),
                                purpose=p) for n,p in [(MASTER,'Calculation master'),(HC,'October employee and contract master'),(DATA,'Supplemental payroll input layout')]],
                  net_payment_correction=net_correction,
                  policies=['No payroll activation or August monthly-input carryover.',
                            'Imported dependent details remain separate from Registered tax dependents.',
                            'Employee records is the protected third monthly worksheet.',
                            'All 11 recurring placeholders retained.',
                            'Salary controls calculations; SHUI is an independent reconciliation.',
                            'Employer PIT side calculations are information only, as in the master.',
                            'Sheet4 ignored. Exchange Rate tab empty. Funding request is an output.'])
    refs = {ref for r in rules for ref in re.findall(r'\{([^}]+)\}',r['formula'])}
    assert refs <= set(by), refs-set(by)
    assert len(by) == len(rules)

    # Map workbook outputs to independent cached controls. Every August person
    # is tested, including probation, the expatriate, leave and allowance cases.
    output_cols = {'VNBASIC':'Q','VNLEAVEPAY':'V','VNTRANSPAY':'X','VNPHONEPAY':'Y',
                   'VNTECH':'AA','VNOVERTIME':'AJ','VNTOTALINC':'AL','VNTAXABLE':'AM',
                   'VNSIEMP':'AN','VNHIEMP':'AO','VNUIEMP':'AP','VNUNIONEE':'AQ',
                   'VNSIER':'AS','VNHIERR':'AT','VNUIERR':'AU','VNTUER':'AV',
                   'VNPERSDED':'AW','VNDEPDED':'AY','VNASSESS':'AZ','VNPITDUE':'BA',
                   'VNNETINC':'BB','VNUSDNET':'BC','VNEMPCTC':'BD','VNNETPAY':'BI',
                   'VNUSDPAY':'BJ','VNREVTX':'BN','VNREVAI':'BO','VNREVPIT':'BP','VNGAP':'BQ',
                   'VNSUPTX':'BT','VNSUPAI':'BU','VNSUPPIT':'BV','VNSUPGAP':'BW'}
    input_cols = {'VNHCE':'B','VNHCF':'C','VNHCU':'D','VNPTYPE':'E','VNPAYNAT':'G',
                  'VNHCBK':'L','VNINSBASE':'M','VNSTDDAYS':'O','VNPAIDDAYS':'P',
                  'VNUNIFORM':'R','VNFREETEL':'S','VNSEVER':'T','VNFREOTH':'U','VNEXTRANS':'W',
                  'VNPRIVATE':'Z','VNY2S2':'AB','VNPRODUCT':'AC','VNACTBONUS':'AD',
                  'VNREFERRAL':'AE','VNOTHTAX':'AF','VNPREADD':'AG','VNPREDED':'AH',
                  'VNBENEFIT':'AI','VNHOLPAY':'AK','VNTUFLAG':'AR','VNREGISTER':'AX',
                  'VNADVANCE':'BE','VNDEDUCT':'BF','VNNETADD':'BG','VNPITREF':'BH',
                  'VNTRIP':'BL','VNPREMIUM':'BM','VNSUPPORT':'BS'}
    fixtures=[]; payroll={}; union=wv['Union list']
    union_counts={}
    for row in union.iter_rows(min_row=2, values_only=True):
        if row[1]: union_counts[str(row[1]).strip()]=union_counts.get(str(row[1]).strip(),0)+1
    advisory={str(wv['Data'].cell(r,2).value).strip():r for r in range(136,166)}
    shui={str(wv['SHUI'].cell(r,2).value).strip():r for r in range(11,81)}
    assert len(shui)==70
    transport_jobs={'Junior Agronomist','Research Agronomist','Senior Agronomist','Supply Chain Specialist','Compliance Officer – Export & Sustainability'}
    for row in range(9,80):
        if not salary[f'B{row}'].value: continue
        key=str(salary[f'B{row}'].value).strip()
        values={r['code']:'' if r['kind'] in ('text','identifier','date') else 0 for r in rules if r['type']=='input'}
        for code,col in input_cols.items():
            values[code]=salary[f'{col}{row}'].value or ('' if by[code]['kind'] in ('text','identifier','date') else 0)
        join=salary[f'F{row}'].value
        values.update(VNHCBH='VND',VNHCBI='Gross',VNTAXMODE='Workbook',VNSIJOIN=join,
                      VNHIJOIN=join,VNUIJOIN=join,VNTUJOIN=join,VNUIREGCAP=106200000,
                      VNUMEMBER=union_counts.get(key,0),VNINSMODE='Salary + private insurance' if row==13 else 'Contract salary',
                      VNUSDSHOW='Yes' if wf['Salary'][f'BC{row}'].data_type=='f' else 'No')
        position=salary[f'D{row}'].value
        values['VNTRANSPORT']=800000 if position in transport_jobs else 960000 if position=='Territory Manager' else 0
        values['VNPHONE']=300000 if position in transport_jobs else 360000 if position=='Territory Manager' else 0
        if row==76: values.update(VNTRANSPORT=1100000,VNPHONE=0)
        a=advisory.get(key)
        if a:
            for i,col in enumerate(('F','I','L','O'),1): values['VNTECH'+str(i)]=wv['Data'][f'{col}{a}'].value or 0
        if salary[f'AJ{row}'].value:
            # Recover the independent hour quantity from the keyed Data table,
            # not by solving backwards from an overtime result.
            hours=[wv['Data'][f'E{r}'].value for r in range(180,195) if wv['Data'][f'B{r}'].value==key]
            assert len(hours)==1,(key,hours)
            values['VNWEEKHRS']=hours[0]
        if row==50: values.update(VNLEAVEWD=18,VNLEAVESTD=23)
        expected={code:salary[f'{col}{row}'].value or 0 for code,col in output_cols.items()
                  if salary[f'{col}{row}'].value is not None}
        sr=shui[key]
        insurance={code:wv['SHUI'][f'{col}{sr}'].value for code,col in {
            'VNSIEMP':'F','VNHIEMP':'G','VNUIEMP':'H','VNSIER':'I',
            'VNHIERR':'J','VNUIERR':'K','VNTUER':'M'}.items()}
        insurance['SHUI_TOTAL']=wv['SHUI'][f'L{sr}'].value
        fixtures.append(dict(name='August master row '+str(row), source_row=row,
                             inputs=values, expected=expected, insurance=insurance))
        payroll[key]=dict(row=row, values=values, commencing=salary[f'H{row}'].value,
                          contract_date=salary[f'I{row}'].value, name=salary[f'C{row}'].value)
    records=[]
    for row in range(3,130):
        if hc[f'AK{row}'].value not in ('Active','Notice Period'): continue
        key=str(hc[f'E{row}'].value).strip()
        cells={openpyxl.utils.get_column_letter(c):hc.cell(row,c).value for c in range(1,105)}
        records.append(dict(row=row,key=key,name=cells['F'],cells=cells,payroll=payroll.get(key)))
    private=dict(fixtures=fixtures,records=records,
                 union_counts=union_counts,master_output_columns=output_cols,
                 master_input_columns=input_cols)
    private_dir.mkdir(parents=True,exist_ok=True,mode=0o700)
    private_dir.chmod(0o700)
    out=private_dir/'source-data.json'
    out.write_text(json.dumps(private,default=lambda v:v.isoformat() if isinstance(v,(dt.date,dt.datetime)) else str(v),ensure_ascii=False,indent=2));out.chmod(0o600)
    definition=ROOT/'docs/configuration/rize-vietnam-payroll-master-v3.json'
    definition.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(definition=str(definition),rules=len(rules),
                          types={t:sum(r['type']==t for r in rules) for t in ('input','formula','constant')},
                          fixtures=len(fixtures),roster=len(records),
                          missing_hc_salary=sum(not r['cells']['BK'] for r in records),
                          missing_supplemental_salary=sum(not r['cells']['BK'] and not r['payroll'] for r in records))))


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--private-dir',type=Path,required=True)
    parser.add_argument('--net-correction',action='store_true',help='Owner-approved correction: subtract Advance/Deduction and add Addition.')
    args=parser.parse_args()
    prepare(args.private_dir,args.net_correction)

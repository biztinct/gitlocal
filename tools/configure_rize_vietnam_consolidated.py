"""Rize-only consolidated source configuration, with rollback as the default.

Run in Odoo shell with VN_CONSOLIDATED_DIR containing the approved workbook and
rize-vietnam-consolidated-202610.json. VN_CONSOLIDATED_APPLY=1 commits only after
mapping, calculation, template and business-record invariants pass. VERIFY=1
checks the installed setup without reapplying it. No payroll or employee import.
"""
import base64
from collections import Counter
import hashlib
import io
import json
import os
from pathlib import Path

import openpyxl

assert env.cr.dbname == 'rize', 'Vietnam client configuration belongs only to Rize.'
root = Path(os.environ['VN_CONSOLIDATED_DIR'])
plan = json.loads((root / 'rize-vietnam-consolidated-202610.json').read_text())
content = (root / plan['filename']).read_bytes()
assert hashlib.sha256(content).hexdigest() == plan['sha256']
cfg = env['hr.formula.config'].browse(plan['config_id']).exists()
assert cfg and cfg.code == plan['code'] and cfg.company_id.name == 'Rize Farms'
e = env(context=dict(env.context, allowed_company_ids=cfg.company_id.ids))
studio = e['pb.formula.studio'].with_company(cfg.company_id)
facade = e['pb.payroll.templates'].with_company(cfg.company_id)
desk = e['pb.records.desk'].with_company(cfg.company_id)
facade._lock(cfg)
by = {r.code: r for r in cfg.rule_ids}
assert len(by) == 185
verify = os.getenv('VN_CONSOLIDATED_VERIFY') == '1'
apply = os.getenv('VN_CONSOLIDATED_APPLY') == '1'
assert not (verify and apply)


def digest(table):
    e.cr.execute('SELECT row_to_json(t)::text FROM ' + table + ' t ORDER BY id')
    return hashlib.sha256('\n'.join(row[0] for row in e.cr.fetchall()).encode()).hexdigest()


def expressions():
    return {r.code: (r.column_letter, r.column_type, r.excel_formula, r.constant_value,
                     r.net_role, r.net_role_detail, r.is_contract_component, r.is_text_component)
            for r in cfg.rule_ids}


business_tables = ('hr_employee', 'hr_contract', 'hr_contract_advantage',
                   'vietnam_employee_dependent', 'hr_payslip', 'hr_payslip_run',
                   'pb_payroll_month_inputs')
baseline = {table: digest(table) for table in business_tables}
original_expressions = expressions()
Mapping = e['hr.payslip.import.mapping']
original_mappings = Mapping.search([('salary_structure_id', '=', cfg.id)])
original_destinations = {m.component_id.code: (
    'b:' + m.bank_role if m.destination_type == 'bank_account'
    else 'f:%s:%s' % (m.target_model_id.model, m.target_field_id.name))
    for m in original_mappings}
original_additional = {spec['code']: list(by[spec['code']].source_ids.filtered(
    lambda s: s.kind == 'excel').mapped('key')) for spec in plan['preserve_absent_inputs']}
layout_row = e['pb.payroll.template.layout'].search([
    ('config_id', '=', cfg.id), ('company_id', '=', cfg.company_id.id)], limit=1)
roster = (layout_row.payload or {}).get('preview_employee_ids', [])
assert len(set(roster)) == len(roster) == plan['review_roster']['count']

changed_bindings = []
if not verify:
    assert hashlib.sha256(base64.b64decode(
        cfg.with_context(bin_size=False).import_sample_file)).hexdigest() == plan['sha256'], \
        'The selected source changed after review; do not replace concurrent work.'
    assert original_destinations.get('VNFREETEL') == 'f:hr.employee:work_phone'
    assert 'VNHCT' not in original_destinations
    read = studio.import_mapping_read_headers(cfg.id, base64.b64encode(content).decode(), plan['filename'])
    assert read.get('ok'), read.get('msg')
    assert read['formula_columns_hidden'] == 0
    cards = {c['label']: c for c in read['left'] if c['id'].startswith('c:')}
    for spec in plan['bindings']:
        rule = by[spec['code']]
        assert rule.column_type == 'input' and rule.name == spec['name']
        assert cards[spec['name']]['id'] == 'c:' + spec['key']
        existing = rule.source_ids.filtered(lambda s: s.kind == 'excel').mapped('key')
        if existing != [spec['key']]:
            rule.set_source_binding('excel', spec['key'], origin='user')
            changed_bindings.append(rule.code)
        rule.write({'source_sheet_name': plan['sheet'], 'original_column_letter': spec['column']})
    wrong = original_mappings.filtered(lambda m: m.component_id == by['VNFREETEL'])
    assert len(wrong) == 1
    assert studio.employee_mapping_delete(wrong.id)['ok']
    result = studio.employee_mapping_create(cfg.id, False, by['VNHCT'].id, 'f:hr.employee:work_phone')
    assert result.get('ok'), result
    for code, updates in plan['rule_updates'].items():
        vals = dict(updates)
        if 'column_role' in vals:
            vals['column_role_source'] = 'user'
        if 'value_kind' in vals:
            vals['value_kind_source'] = 'user'
        if 'net_role' in vals:
            vals.update(net_role_source='user', net_role_confidence='certain')
        by[code].write(vals)
    cfg.write(dict(plan['configuration'], description=(
        '<p>Input source: Rize_Vietnam_Consolidated_Payroll_and_Master.xlsx, Consolidated sheet. '
        '129 unique RIZ employee codes; October HC master values override matching August inputs.</p>'
        '<p>Calculations: approved August 2026 FINAL v3 master, including the approved net-payment correction. '
        'No historical time or adjustment amounts are saved as new monthly defaults.</p>'
        '<p>Saved-dependent transformation supplies Dependent/Children; registered tax dependents remain '
        'separate. Spreadsheet and employee/contract/bank mappings are configured. '
        'The 11 recurring-setting placeholders and 9 additional monthly inputs remain visible.</p>'
        '<p>Draft review population: 70 Active/Notice Period employees. Review salary gaps, tax/insurance/union '
        'settings and the exchange rate before activation. Other source rows remain available for mapping '
        'without live payroll assignments.</p>')))
    catalog = facade._catalog(cfg)
    layout = {lane: [] for lane in ('time', 'adjustments', 'records', 'available')}
    for col in catalog:
        lane = ('records' if col['record_eligible'] else
                'time' if col['code'] in plan['monthly_layout']['time_codes'] else
                'available' if col['code'] in plan['monthly_layout']['available_codes'] else 'adjustments')
        layout[lane].append(col['id'])
    assert facade.save_layout(cfg.id, layout)['ok']

assert all(cfg[field] == value for field, value in plan['configuration'].items())
assert expressions() == original_expressions, 'Calculation, pay role, storage type or fixed values changed.'
assert sum(r.column_type == 'formula' for r in cfg.rule_ids) == plan['preserve_formulas']
assert sum(r.column_type == 'constant' for r in cfg.rule_ids) == plan['preserve_constants']
assert sum(r.is_contract_component for r in cfg.rule_ids) == plan['preserve_contract_components']
assert original_additional == {spec['code']: list(by[spec['code']].source_ids.filtered(
    lambda s: s.kind == 'excel').mapped('key')) for spec in plan['preserve_absent_inputs']}
for code, updates in plan['rule_updates'].items():
    assert all(by[code][field] == value for field, value in updates.items()), code
for spec in plan['bindings']:
    assert by[spec['code']].source_ids.filtered(lambda s: s.kind == 'excel').mapped('key') == [spec['key']]
assert by['VNHCQ'].source_ids.filtered(lambda s: s.kind == 'rule').mapped('key') == ['VNDEPCOUNT']
assert not by['VNREGISTER'].source_ids.filtered(lambda s: s.kind == 'rule')

mappings = Mapping.search([('salary_structure_id', '=', cfg.id)])
destinations = {m.component_id.code: (
    'b:' + m.bank_role if m.destination_type == 'bank_account'
    else 'f:%s:%s' % (m.target_model_id.model, m.target_field_id.name)) for m in mappings}
assert 'VNFREETEL' not in destinations and destinations['VNHCT'] == 'f:hr.employee:work_phone'
if not verify:
    assert {k: v for k, v in destinations.items() if k != 'VNHCT'} == {
        k: v for k, v in original_destinations.items() if k != 'VNFREETEL'}
assert len(destinations) == len(mappings) == 27
board = studio.import_mapping_data(cfg.id, 'sample')
assert board['wired'] == 118, board['wired']
assert board['formula_columns_hidden'] == 0
assert len([c for c in board['left'] if c['id'].startswith('c:')]) == 122
assert all(any(w['leftId'] == 'c:' + spec['key'] and w['rightId'] == by[spec['code']].id
               and w['kind'] == 'mapping' for w in board['wires']) for spec in plan['bindings'])
ec = studio.employee_mapping_data(cfg.id, include_payroll=True)
assert ec['unresolved'] == 0, ec['unresolved']
assert len(ec['wires']) == 106
assert not {r.id for r in cfg.rule_ids if r.column_type != 'input'} & {c['id'] for c in ec['left']}
treatment = cfg.value_kind_board()
assert treatment['review_count'] == treatment['role_conflict_count'] == treatment['drift_count'] == 0
scenarios = cfg.run_sample_tests()
assert scenarios['total'] == scenarios['passed'] == 96

Batch = e['hr.payroll.import.batch']
batch = Batch.new(dict(formula_config_id=cfg.id, company_id=cfg.company_id.id,
                       source_type='excel', date_from='2026-10-01', date_to='2026-10-31'))
parsed = batch._parse_source_file(content, plan['filename'])
rows = parsed['rows']
assert len(rows) == plan['source_rows']
keys = [str(row[plan['primary_key']]).strip() for row in rows]
assert len(set(keys)) == len(keys) == 129
people = e['hr.employee'].with_context(active_test=False).search([('company_id', '=', cfg.company_id.id)])
contracts = desk._io_contracts(people)
source_status = Counter()
source_validation = Counter()
roster_validation = Counter()
checks = []
roster_keys = set()
for key, raw in zip(keys, rows):
    person = people.filtered(lambda p: key in (p.barcode, p.employee_id))
    assert len(person) <= 1, 'Duplicate saved RIZ identity.'
    person = person[:1]
    contract = contracts.get(person.id) if person else None
    values = batch._transform_data_to_formula_inputs(raw, contract=contract, employee=person)
    errors = {}
    actual = e['hr.formula.sample.data'].new(dict(config_id=cfg.id))._evaluate_rules_with_dependencies(
        values, readonly=True, errors=errors)
    assert not errors, (key, errors)
    source_status[str(raw.get('Status(Active/Inactive/Notice Period)') or 'Not supplied')] += 1
    source_validation[str(actual['VNREADY'])] += 1
    if person.id in roster:
        roster_keys.add(key)
        roster_validation[str(actual['VNREADY'])] += 1
        checks.append(dict(code=key, name=raw.get('Họ và tên (Full name)'),
                           ready=actual['VNREADY'], salary=values.get('VNHCBK'), notes=raw.get('Review notes')))
        assert values['VNHCQ'] == person.vn_dependent_record_count
        if raw.get('Người phụ thuộc (Dependants) — Count') not in (None, ''):
            assert values['VNREGISTER'] == raw['Người phụ thuộc (Dependants) — Count']
assert len(roster_keys) == 70
assert sum(source_status[status] for status in plan['review_roster']['statuses']) == 70
# Prove the transformation wins over a conflicting HC count, without sharing
# that number with the separately supplied registered tax-dependent count.
person = people.filtered(lambda p: p.id in roster and p.vn_dependent_record_count)[:1]
assert person
spec = next(s for s in plan['bindings'] if s['code'] == 'VNHCQ')
conflict = {spec['key']: 999, 'Dependent/Children': 999,
            'Consolidated|Người phụ thuộc (Dependants) — Count': 2}
used = batch._transform_data_to_formula_inputs(conflict, contract=contracts.get(person.id), employee=person)
assert used['VNHCQ'] == person.vn_dependent_record_count and used['VNREGISTER'] == 2

# Exercise a new month's blank adjustments with representative full paid days.
# These values are test inputs only; never persist August payments as defaults.
monthly = Batch.new(dict(formula_config_id=cfg.id, company_id=cfg.company_id.id,
    source_type='excel', one_time=True, import_filename='Payobook-month-2026-10.xlsx',
    date_from='2026-10-01', date_to='2026-10-31'))
monthly_catalog = facade._catalog(cfg)
monthly_raw = {key: 0 for col in monthly_catalog if not col['record_eligible']
               for key in col['source_keys']}
for code in ('VNSTDDAYS', 'VNPAIDDAYS'):
    for key in next(c for c in monthly_catalog if c['code'] == code)['source_keys']:
        monthly_raw[key] = 22
fresh_validation = Counter()
for person in people.filtered(lambda p: p.id in roster):
    values = monthly._transform_data_to_formula_inputs(
        monthly_raw, contract=contracts.get(person.id), employee=person)
    assert values['VNPAIDDAYS'] == values['VNSTDDAYS'] == 22
    assert values.get('VNACTBONUS', 0) == values.get('VNFREETEL', 0) == values.get('VNADVANCE', 0) == 0
    errors = {}
    actual = e['hr.formula.sample.data'].new(dict(config_id=cfg.id))._evaluate_rules_with_dependencies(
        values, readonly=True, errors=errors)
    assert not errors
    fresh_validation[str(actual['VNREADY'])] += 1

workspace = facade.workspace(cfg.id, '2026-10')
assert workspace['people'] == 70
layout = workspace['layout']
catalog = {c['id']: c for c in workspace['fields']}
assert all(catalog[fid]['record_eligible'] for fid in layout['records'])
assert not any(catalog[fid]['record_eligible'] for lane in ('time', 'adjustments') for fid in layout[lane])
assert {catalog[fid]['code'] for fid in layout['time']} == set(plan['monthly_layout']['time_codes'])
assert 'VNFREETEL' in {catalog[fid]['code'] for fid in layout['adjustments']}
workbooks = []
for mode in ('template', 'records'):
    download = facade.download(cfg.id, '2026-10', mode)
    book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(download['file_b64'])))
    assert book['Employee records'].max_row == 71
    if mode == 'template':
        assert [s.title for s in book if s.sheet_state == 'visible'] == ['Time', 'Monthly adjustments', 'Employee records']
        assert book['Employee records'].protection.sheet
        assert all(cell.protection.locked for row in book['Employee records'] for cell in row)
        assert all(cell.value is None for title in ('Time', 'Monthly adjustments')
                   for row in book[title].iter_rows(min_row=2, min_col=3,
                                                   max_col=book[title].max_column - 1)
                   for cell in row), 'Reusable monthly inputs must start blank.'
    else:
        review = facade.review(cfg.id, '2026-10', 'records', download['file_b64'], download['filename'])
        assert not review['issues'] and not review['changes'], (review['issues'], len(review['changes']))
    workbooks.append(dict(mode=mode, people=download['rows'], columns=book['Employee records'].max_column))
assert baseline == {table: digest(table) for table in business_tables}, 'Business data or live assignments changed.'
result = dict(database=e.cr.dbname, config_id=cfg.id, name=cfg.name, state=cfg.state,
              source_sha256=plan['sha256'], source_rows=len(rows), source_columns=122,
              mapped_business_fields=board['wired'], additional_inputs=20,
              recurring_placeholders=11, native_mappings=len(mappings),
              contract_text=sum(r.is_text_component for r in cfg.rule_ids),
              contract_amount=sum(r.is_contract_component and not r.is_text_component for r in cfg.rule_ids),
              employee_contract_connections=len(ec['wires']), unresolved=ec['unresolved'],
              treatment_review=0, scenarios=scenarios, source_status=dict(source_status),
              source_validation=dict(source_validation), review_roster_validation=dict(roster_validation),
              fresh_month_validation=dict(fresh_validation),
              review_roster=checks, roster_keys=sorted(roster_keys),
              layout={k: len(v) for k, v in layout.items()}, workbooks=workbooks,
              changed_bindings=changed_bindings, business_unchanged=True,
              definitions_preserved=True, dependent_priority_verified=True)
(root / ('verified.json' if verify else 'applied.json' if apply else 'dry-run.json')).write_text(
    json.dumps(result, ensure_ascii=False, indent=2, default=str))
if apply:
    e.cr.commit()
else:
    e.cr.rollback()
print('VN_CONSOLIDATED_RESULT=' + json.dumps({k: v for k, v in result.items()
    if k not in ('review_roster', 'roster_keys', 'changed_bindings')}, ensure_ascii=False))

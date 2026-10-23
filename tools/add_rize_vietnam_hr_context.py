"""Add the owner's 22 information-only Contract Text fields to the Rize Draft.

Run in the rize Odoo shell with VN_HR_CONTEXT_DEFINITION pointing at the reviewed
non-personal definition JSON. Dry-run rolls back; VN_HR_CONTEXT_APPLY=1 commits.
This adds configuration and layout rows, never employee/contract values.
"""
import base64
import hashlib
import io
import json
import os
from pathlib import Path

import openpyxl


assert env.cr.dbname == 'rize', 'Rize client configuration belongs only in rize.'
columns = 'N O AS AT AU AV AX AY AZ BA BB BC CI CJ CK CL CM CN CO CT CU CV'.split()
definition = json.loads(Path(os.environ['VN_HR_CONTEXT_DEFINITION']).read_text())
assert definition['code'] == 'RIZE_VIETNAM_HC_202610'
specs = {r['source_column']: r for r in definition['rules'] if r['source_column'] in columns}
assert set(specs) == set(columns) and len(specs) == 22
assert all(r['kind'] == 'text' and r['destination'] == 'c:text'
           and r['type'] == 'input' and r['role'] == 'info'
           and not r['detail'] and not r['formula'] for r in specs.values())
cfg = env['hr.formula.config'].browse(4639).exists()
assert cfg and cfg.code == definition['code'] and cfg.state == 'draft'
studio = env['pb.formula.studio'].with_company(cfg.company_id)
facade = env['pb.payroll.templates'].with_company(cfg.company_id)
facade._lock(cfg)


def snapshot(records):
    names = [name for name, field in records._fields.items()
             if field.store and not field.compute and field.type != 'binary'
             and name not in ('write_date', 'write_uid', 'create_date', 'create_uid')]
    return records.read(names)


def business_snapshot():
    result = {}
    for table in ('hr_employee', 'hr_contract', 'hr_contract_advantage', 'hr_payslip', 'hr_payslip_run'):
        env.cr.execute('SELECT row_to_json(t)::text FROM ' + table + ' t ORDER BY id')
        rows = [r[0] for r in env.cr.fetchall()]
        result[table] = {'count': len(rows), 'sha256': hashlib.sha256('\n'.join(rows).encode()).hexdigest()}
    return result


new_codes = {r['code'] for r in specs.values()}
old_rules = cfg.rule_ids.filtered(lambda r: r.code not in new_codes)
assert len(old_rules) == 116, 'Unexpected baseline; review instead of overwriting changes.'
before_rules = snapshot(old_rules)
before_letters = {r.code: r.column_letter for r in old_rules}
before_sources = snapshot(old_rules.source_ids)
mapping = env['hr.payslip.import.mapping'].search([('salary_structure_id', '=', cfg.id)])
before_mappings = snapshot(mapping)
before_business = business_snapshot()
before_config = cfg.read(['name', 'code', 'state', 'source_priority', 'connector_id',
                          'import_sample_filename', 'import_sample_columns_json'])[0]
before_sample = cfg.import_sample_file
before_layout = facade.workspace(cfg.id, '2026-10')
layout_row = env['pb.payroll.template.layout'].search([('config_id', '=', cfg.id),
                                                    ('company_id', '=', cfg.company_id.id)])
assert len(layout_row) == 1
before_metadata = {k: v for k, v in layout_row.payload.items()
                   if k not in ('time', 'adjustments', 'records', 'available')}
created = []
try:
    for col in columns:
        spec = specs[col]
        rule = cfg.rule_ids.filtered(lambda r: r.code == spec['code'])
        assert len(rule) <= 1
        if not rule:
            rule = env['hr.formula.rule'].create({
                'config_id': cfg.id, 'name': spec['name'], 'code': spec['code'],
                'sequence': max(cfg.rule_ids.mapped('sequence')) + 10,
                'column_type': 'input', 'value_kind': 'text', 'value_kind_source': 'user',
                'value_kind_reason': 'Owner approved information-only Contract Text.',
                'component_type': 'Employee information',
                'net_role': 'info', 'net_role_detail': False, 'net_role_source': 'user',
                'net_role_confidence': 'certain', 'net_role_reason': spec['note'],
                'appears_on_payslip': False, 'visibility_rule': 'never',
                'report_visible': False, 'is_required': False,
                'column_role': 'contract', 'column_role_source': 'user',
                'source_sheet_name': 'VN', 'original_column_letter': col,
            })
            rule.set_source_binding('excel', spec['source_key'], origin='user')
            result = studio.employee_mapping_create(cfg.id, '', rule.id, 'c:text')
            assert result.get('ok'), (spec['code'], result)
            created.append(rule.code)
        # A rerun verifies rather than overwrites any subsequent owner changes.
        assert rule.name == spec['name'] and rule.value_kind == 'text'
        assert rule.column_type == 'input' and rule.column_role == 'contract'
        assert rule.is_contract_component and rule.is_text_component
        assert rule.net_role == 'info' and not rule.net_role_detail
        assert not rule.appears_on_payslip and rule.visibility_rule == 'never'
        assert not rule.report_visible and not rule.is_required
        assert [(s.kind, s.key) for s in rule.source_ids] == [('excel', spec['source_key'])]

    added = cfg.rule_ids.filtered(lambda r: r.code in new_codes)
    assert len(added) == 22 and len(cfg.rule_ids) == 138
    catalog = facade._catalog(cfg)
    by_rule = {c['code']: c for c in catalog if c['code']}
    assert all(by_rule[r.code]['record_eligible'] and by_rule[r.code]['readonly']
               and by_rule[r.code]['record_id'] == 'c:' + r.code for r in added)
    layout = facade._layout(cfg, catalog)
    added_ids = {by_rule[r.code]['id'] for r in added}
    assert added_ids <= set(layout['records'])
    for lane in ('time', 'adjustments', 'available'):
        assert layout[lane] == before_layout['layout'][lane]
    assert [i for i in layout['records'] if i not in added_ids] == [
        i for i in before_layout['layout']['records'] if i not in added_ids]
    assert facade.save_layout(cfg.id, layout)['ok']
    assert before_metadata == {k: v for k, v in layout_row.payload.items()
                               if k not in ('time', 'adjustments', 'records', 'available')}

    source_board = studio.import_mapping_data(cfg.id, 'sample')
    people_board = studio.employee_mapping_data(cfg.id)
    assert source_board['ok'] and people_board['ok']
    assert all(any(w['rightId'] == r.id and w['leftId'] == 'c:' + specs[r.original_column_letter]['source_key']
                   for w in source_board['wires']) for r in added)
    assert all(any(w.get('componentId') == r.id and w['rightId'] == 'c:text'
                   for w in people_board['wires']) for r in added)
    scenarios = cfg.run_sample_tests()
    assert scenarios['passed'] == 30 and scenarios['failed'] == 0, scenarios
    health = studio.get_problems(cfg.id)['counts']
    assert health['error'] == 0, health

    # Verify both actual workbook generators, rolling back their download receipts.
    env.cr.execute('SAVEPOINT vn_context_export_check')
    workbook_checks = []
    for mode in ('template', 'records'):
        download = facade.download(cfg.id, '2026-10', mode)
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(download['file_b64'])))
        ws = wb['Employee records']
        headers = [c.value for c in ws[1]]
        assert all(r.name in headers for r in added)
        assert ws.max_row == 64 and download['rows'] == 63
        if mode == 'template':
            assert [s.title for s in wb if s.sheet_state == 'visible'] == ['Time', 'Monthly adjustments', 'Employee records']
            assert ws.protection.sheet and all(c.protection.locked for row in ws for c in row)
            assert not any(r.name in [c.value for c in wb[tab][1]]
                           for r in added for tab in ('Time', 'Monthly adjustments'))
        workbook_checks.append({'mode': mode, 'rows': download['rows'], 'new_columns': 22,
                                'protected': bool(ws.protection.sheet)})
    env.cr.execute('ROLLBACK TO SAVEPOINT vn_context_export_check')
    env.cr.execute('RELEASE SAVEPOINT vn_context_export_check')
    env.invalidate_all()

    assert snapshot(old_rules) == before_rules
    assert {r.code: r.column_letter for r in old_rules} == before_letters
    assert snapshot(old_rules.source_ids) == before_sources
    assert snapshot(mapping) == before_mappings
    assert business_snapshot() == before_business, 'Employee/contract/payroll values changed.'
    assert cfg.read(list(before_config.keys() - {'id'}))[0] == before_config
    assert cfg.import_sample_file == before_sample
    workspace = facade.workspace(cfg.id, '2026-10')
    assert workspace['people'] == before_layout['people'] == 63
    receipt = {'database': env.cr.dbname, 'configuration': cfg.id, 'state': cfg.state,
               'created': created, 'verified_components': 22, 'rules': len(cfg.rule_ids),
               'inputs': len(cfg.rule_ids.filtered(lambda r: r.column_type == 'input')),
               'source_mappings': 22, 'contract_text_destinations': 22,
               'layout': {k: len(v) for k, v in workspace['layout'].items()},
               'record_columns': workspace['record_columns'], 'roster': workspace['people'],
               'scenarios': scenarios, 'health': health, 'workbooks': workbook_checks,
               'existing_rules_sources_and_values_unchanged': True,
               'business': before_business}
    if os.getenv('VN_HR_CONTEXT_RECEIPT'):
        path = Path(os.environ['VN_HR_CONTEXT_RECEIPT'])
        path.write_text(json.dumps(receipt, indent=2, default=str))
        path.chmod(0o600)
    if os.getenv('VN_HR_CONTEXT_APPLY') == '1':
        env.cr.commit()
        print('VN_HR_CONTEXT_APPLIED', json.dumps(receipt, default=str))
    else:
        env.cr.rollback()
        print('VN_HR_CONTEXT_DRY_RUN', json.dumps(receipt, default=str))
except Exception:
    env.cr.rollback()
    raise

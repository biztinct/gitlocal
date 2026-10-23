"""Prepare SG/ID configuration metadata and private client import plans.

Input is a private, read-only extraction of the named client workbooks and the
existing Rize definitions. Never commit the generated employee plans/fixtures.
The original source files are never modified. Author XLSX through the separate
Artifact Tool builder, not this extractor.
"""
import argparse
from collections import Counter
from datetime import datetime
import json
from pathlib import Path
import re

from openpyxl.utils import get_column_letter, column_index_from_string


def clean(value):
    if value is None:
        return None
    if isinstance(value,str):
        value=value.strip()
        if value.upper() in {'', 'NA', 'N/A', '#N/A', '#VALUE!', 'NONE'}:
            return None
    return value


def text(value):
    value=clean(value)
    if isinstance(value,float) and value.is_integer():value=int(value)
    return str(value) if value is not None else None


def money(value):
    value=clean(value)
    if value is None:return None
    if isinstance(value,(int,float)):return float(value)
    return float(re.sub(r'[^0-9.\-]','',value))


def date(value):
    value=clean(value)
    return str(value)[:10] if value is not None else None


def norm(value):
    return re.sub(r'[^a-z0-9]','',str(value or '').lower())


def prepare(root, repo):
    data=json.loads((root/'sources.json').read_text())
    live=json.loads((root/'live.json').read_text())
    cross=json.loads((root/'id-crosswalk.json').read_text())
    results=json.loads((root/'existing-results.json').read_text())
    for country,key,cfgid in [('SG','sg',3948),('ID','id',3943)]:
        master_data=data['master_'+key]['sheets']['Singapore' if key=='sg' else 'ID']
        masters={r['row']:r for r in master_data['rows'][2:] if clean(r['values'][5])}
        headers={get_column_letter(i+1):str(v).strip() for i,v in enumerate(master_data['rows'][1]['values']) if v}
        payname='Payroll Input - August26' if key=='sg' else '0826 Payroll Report'
        pays={r['row']:r for r in data['payroll_'+key]['sheets'][payname]['rows'] if r['row'] in ([7,8] if key=='sg' else range(9,131))}
        old=live['rules'][str(cfgid)]
        specs=[]
        oldcodes={r['code'] for r in old}
        prefix=country
        keep_standard={'PIT','NET','GROSS','EEDED','ERCOST'}
        mapping={}
        for original in old:
            col=original['column_letter']
            if (key=='sg' and col in {'R','S'}) or (key=='id' and col=='C'):continue
            code=original['code'] if original['code'] in keep_standard else prefix+original['code']
            mapping[original['code']]=code
            spec=dict(code=code,name=original['name'],column=col,type=original['column_type'],
                formula=original['excel_formula'] or '',value=original['constant_value'],
                kind=original['value_kind'],role=original['net_role'],detail=original['net_role_detail'],
                group=original['component_type'] or 'Reference',destination='',tab='',required=False,
                note='Preserved client calculation.' if original['column_type']=='formula' else '',
                original_code=original['code'])
            if spec['type']=='input':spec['source_key']='Consolidated|'+spec['name']
            specs.append(spec)
        bycol={s['column']:s for s in specs}
        primary=bycol['A' if key=='sg' else 'D']
        primary['name']='Employee Code'
        primary['source_key']='Consolidated|Employee Code'
        if key=='id':
            bycol['H']['name']='Contract end date'
            bycol['H']['source_key']='Consolidated|Contract end date'
        bycode={s['code']:s for s in specs}
        top=max(column_index_from_string(s['column']) for s in specs if re.match(r'^[A-Z]+$',s['column']))
        def add(code,name,kind='text',destination='',tab='',required=False,formula='',role='info',note='',group='Employee master'):
            nonlocal top
            top+=1
            s=dict(code=code,name=name,column=get_column_letter(top),type='formula' if formula else 'input',formula=formula,
                value=0,kind=kind,role=role,detail=False,group=group,destination=destination,tab=tab,required=required,note=note)
            if not formula:s['source_key']='Consolidated|'+name
            specs.append(s);bycode[code]=s
            return s
        if key=='sg':
            destinations={'A':'f:hr.employee:barcode','B':'f:hr.employee:payroll_country','C':'f:hr.employee:name',
                'D':'f:hr.employee:job_title','E':'f:hr.employee:birthday','F':'f:hr.contract:date_start',
                'G':'c:text','H':'f:hr.contract:wage','J':'c:text','AF':'c:text',
                'AB':'b:bank_name','AC':'b:acc_number','AD':'b:acc_holder_name','AE':'b:bank_bic'}
            time={'T','U','V'}
            monthly={'I','K','L','Y','Z','AG'}
            bycol['H']['note']='Saved full-month contracted salary in SGD; no FX conversion. August payroll supplies the two active salary baselines.'
            for c in {'I','K','L'}:bycol[c]['note']='Approved contribution input. Reassess CPF residency/age/PR status and self-help eligibility for the selected month. No automatic rate is inferred.'
            for c in {'AB','AC','AD','AE'}:bycol[c]['note']='Employee-master banking where provided; otherwise the matched August payroll supplies this bank field. No payment is released by importing it.'
            # Native month totals and remittance rounding are separate. The
            # company SDL remittance is FLOOR(sum(employee SDL),1), not the sum
            # of rounded employee amounts.
            bycol['M']['formula']='=IF(T2>0,IF(AJ2>0,MIN(MAX(AJ2*AM2,AN2),AO2),0),"Working days required")'
            add('SGREADY','Payroll input validation',formula='=IF(OR(T2<=0,U2<0,U2>T2,V2<0),"Review payable days",IF(OR(I2<0,K2<0,L2<0),"Review contributions","Ready"))',
                note='Visible review check. CPF and self-help remain approved monthly amounts; their values do not carry into another period.',group='Validation')
            aliases={'D':'A','F':'C','K':'E','U':'D','X':'F','BC':'G','BF':'H','AP':'AF','CK':'AB','CL':'AC','CM':'AD','CN':'AE'}
            native={'J':'f:hr.employee:country_id','M':'f:hr.employee:place_of_birth','P':'f:hr.employee:marital',
                'R':'f:hr.employee:private_email','S':'f:hr.employee:private_phone','T':'f:hr.employee:work_phone',
                'Z':'f:hr.contract:trial_date_end','AC':'f:hr.contract:date_end','AH':'f:hr.employee:department_id',
                'AM':'f:hr.employee:identification_id','AN':'f:hr.employee:permit_no','AO':'f:hr.employee:passport_id',
                'AQ':'f:hr.contract:tax_identification_number','AX':'f:hr.employee:private_street',
                'AZ':'f:hr.employee:emergency_phone','BA':'f:hr.employee:emergency_contact','CC':'f:hr.employee:work_email',
                'CO':'f:hr.employee:study_school','CP':'f:hr.employee:study_field'}
            excluded=set('A D G H L AL BG BH BI BJ BK BO BP BQ BR BS'.split())
            datecols=set('K X Y Z AB AC BT BV BW CG CI'.split())
            moneycols=set('BE BF BL BM'.split())
        else:
            destinations={'A':'f:hr.employee:department_id','B':'c:text','D':'f:hr.employee:barcode',
                'E':'f:hr.employee:name','F':'f:hr.employee:job_title','G':'f:hr.contract:date_start',
                'H':'f:hr.contract:date_end','I':'c:text','J':'c:text','K':'c:text',
                'AS':'f:hr.contract:wage','AW':'c:amount','AX':'c:amount','AY':'c:amount','AZ':'c:amount','BA':'c:amount'}
            time={'AT','AU','BU'}
            monthly={'P','Q','R','S','X','Y','AJ','AK','AL','AM','AR','AV','BB','BC','BD','BE','BF','BV','BW'}
            rounder=add('IDTAXROUND','Regular tax allowance rounding adjustment',kind='money',tab='Monthly adjustments',required=True,
                group='Monthly adjustments',note='Signed, approved rounding adjustment inside the regular tax allowance. August row 105 uses -1 IDR; all other August rows use zero. Do not carry this historical correction into later months.')
            final=add('IDFINALALLOW','Approved final-period tax allowance',kind='money',tab='Monthly adjustments',required=True,
                group='Year to date inputs',note='Explicit allowance for the final annual reconciliation when allowance mode is 1 or 2. Enter zero in an ordinary month. A full net contract needs an independently reviewed annual gross-up amount.')
            reviewed=add('IDYTDREVIEW','Final-period YTD reviewed',kind='integer',tab='Monthly adjustments',required=True,
                group='Year to date inputs',note='1 only after prior-year-to-date gross, deductible JHT/JP, job-expense months, prior tax and final allowance have been checked. 0 in an ordinary month. December or leaving employment requires final-period reconciliation.')
            eligibility=add('IDSEVELIG','Severance eligibility confirmed',kind='integer',tab='Monthly adjustments',required=True,
                group='Severance',note='1 only for an approved qualifying severance payment. Normal PKWT compensation is not automatically qualifying severance. Enter 0 when no severance is paid.')
            fc=final['column']+'2';rc=rounder['column']+'2';yc=reviewed['column']+'2';sc=eligibility['column']+'2'
            bycol['M']['formula']=f'=IF(BB2=1,IF(BA2=0,0,{fc}),IF(BA2=0,0,MAX(0,BQ2)))+{rc}'
            bycol['BS']['formula']='=IF(BB2=1,IF(BA2=2,M2-AN2,0),'+bycol['BS']['formula'].lstrip('=')+')'
            bycol['BS']['note']='Ordinary months check the TER gross-up. Final full net contracts compare the approved allowance with the annual current-period PPh21, so a valid annual reconciliation does not display a spurious monthly residual.'
            # Retain every existing TER, annual-tax and severance band. The
            # source's unsupported final-period text is replaced by a visible
            # reviewed input, so annual net-contract reconciliation is usable.
            add('IDREADY','Payroll input validation',formula=(
                '=IF(OR(AU2<=0,AT2<0,AT2>AU2,BU2<0),"Review payable days",'
                'IF(AND(J2<>"S/0",J2<>"S/1",J2<>"S/2",J2<>"S/3",J2<>"M/0",J2<>"M/1",J2<>"M/2",J2<>"M/3"),"Review PTKP status",'
                'IF(OR(AX2<0,AX2>1,AY2<0,AY2>1,AZ2<0,AZ2>1,BA2<0,BA2>2,BB2<0,BB2>1),"Review coverage and tax flags",'
                f'IF(AND(BB2=1,{yc}<>1),"Review final-period YTD",'
                f'IF(AND(Y2<>0,{sc}<>1),"Confirm severance eligibility",'
                f'IF(AND(BB2=1,BA2=2,ABS(M2-AN2)>1),"Review annual net gross-up","Ready"))))))'),
                group='Validation',note='Visible guard for days, PTKP, coverage, final-period review and qualifying severance. Final full net contracts must reconcile approved allowance to annual current tax.')
            aliases={'F':'E','U':'F','X':'G','AC':'H','AH':'A','BK':'AS','BI':'BA'}
            native={'J':'f:hr.employee:country_id','K':'f:hr.employee:birthday','M':'f:hr.employee:place_of_birth',
                'P':'f:hr.employee:marital','R':'f:hr.employee:private_email','S':'f:hr.employee:private_phone',
                'T':'f:hr.employee:work_phone','Z':'f:hr.contract:trial_date_end','AN':'f:hr.employee:identification_id',
                'AO':'f:hr.employee:permit_no','AP':'f:hr.employee:passport_id','AR':'f:hr.contract:tax_identification_number',
                'BD':'f:hr.employee:private_street','BE':'f:hr.employee:emergency_phone','BF':'f:hr.employee:emergency_contact',
                'CH':'f:hr.employee:work_email','CP':'b:bank_name','CQ':'b:acc_number','CR':'b:acc_holder_name','CS':'b:bank_bic',
                'CT':'f:hr.employee:study_school','CU':'f:hr.employee:study_field'}
            excluded=set('A D G H L Q AM AW BL BM BN BO BP BT BU BV BW BX'.split())
            datecols=set('K X Y Z AB AC AL BY CA CB CL CN'.split())
            moneycols=set('BJ BK BQ BR'.split())
            bycol['AS']['note']='Saved full-month contract basic salary. Payroll report supplies the August baseline, including the separate new-hire/leaver calculation schedule. Paid days and period days are monthly inputs.'
            for c in {'AX','AY','AZ','BA'}:bycol[c]['note']='Saved employee/contract setting inferred from the August contribution treatment. Review when coverage, tax contract or statutory policy changes.'
        for col,destination in destinations.items():
            bycol[col]['destination']=destination
            bycol[col]['required']=col in ({'A','C','H','AF','J'} if key=='sg' else {'D','E','AS','J','AX','AY','AZ','BA'})
        for col in time|monthly:
            bycol[col]['tab']='Time' if col in time else 'Monthly adjustments'
            bycol[col]['required']=col not in {'AG','AR'}
            if not bycol[col]['note']:bycol[col]['note']='Selected-month input. Supply a reviewed value, including an explicit zero when none. Never saved as a recurring salary amount.'
        definitions=({
            'T':'Agreed standard working days for the selected Singapore month, after public holidays. Must be greater than zero; August uses 20. This is not the Indonesia calendar-day divisor.',
            'U':'Unpaid or inactive working days, including an approved joiner/leaver proration. Enter zero when none; cannot exceed standard working days.',
            'V':'Approved arrear days at the current salary rate. Use Other income for separately priced backpay. Enter zero when none.',
            'I':'Reviewed full-month employee CPF basis used by the supplied net-before-proration method. Confirm residency, age, PR year and wage ceilings separately. Enter a reviewed zero for an exemption; do not carry an old month forward.',
            'K':'Reviewed full-month self-help contribution basis used by the supplied net-before-proration method. Confirm the fund and any valid opt-out separately; zero must be explicit.',
            'L':'Actual employer CPF payable for this month, reviewed separately. Added to employer cost without a second proration. Enter zero only when confirmed.',
            'Y':'Approved cash medical allowance for the selected month. Added after salary proration. A reimbursement needs its correct separate treatment.',
            'Z':'Approved other cash income for the selected month, added after proration. Explain nonzero bonuses or backpay in Remarks.',
        } if key=='sg' else {
            'J':'Client PTKP codes: S/0, S/1, S/2, S/3, M/0, M/1, M/2 or M/3. S means single, M married; the number is approved tax dependants. Review changes in Records Desk; do not infer the tax code from a name.',
            'AT':'Approved payable calendar days for this month. A full month normally equals Period days; August joiners use 29 or 22 and the leaver uses four. Must be between zero and Period days.',
            'AU':'Calendar-day salary divisor for the selected month, normally 28, 29, 30 or 31. August uses 31. This is not Singapore standard working days.',
            'AV':'Approved signed salary/backpay adjustment in IDR, added after the basic-salary day calculation. Enter an explicit zero when none.',
            'AW':'Saved annual private-insurance premium in IDR. The formula divides it by 12 and rounds to whole rupiah. Review the saved amount when insured people or premiums change; zero means no premium.',
            'AX':'Saved BPJS manpower coverage: 1 covered, 0 not covered. Determines JKM, JKK, JHT and pension calculations. A registration number alone does not select coverage.',
            'AY':'Saved BPJS healthcare coverage: 1 covered, 0 not covered. Contributions use the linked wage floor, cap and rates.',
            'AZ':'Saved reimbursement setting: 1 the employer reimburses employee JHT, pension and healthcare contributions through BPJS allowance; 0 those contributions remain employee deductions.',
            'BA':'Saved tax allowance mode: 0 none, 1 gross-up of regular insurance-premium tax, 2 full monthly net-contract gross-up. A final period uses the separately approved annual allowance instead.',
            'BB':'Final tax period: 1 for December or the final employment-period reconciliation, otherwise 0. Final periods use annual tax less prior withholding, including a signed refund where applicable.',
            'BC':'Taxable gross from January through the month before this one. Include the applicable taxable benefits/allowances; exclude the current month and separately qualifying severance. Supply reviewed YTD values for a final period.',
            'BD':'Deductible employee JHT plus pension from January through the prior month. Exclude employer contributions, employee healthcare and this month. Used in final annual reconciliation.',
            'BE':'Job-expense eligible months including the current month, normally 1 to 12 for the actual employment period. Review this for joiners/leavers and final annual reconciliation; do not carry August eight into December.',
            'BF':'PPh21 already withheld from January through the prior month. Exclude the current month. Final-period payable tax is annual liability less this reviewed amount.',
            'BU':'Approved leave days being paid out this month, not the total HR leave balance. The formula uses saved basic salary divided by Period days. Enter zero when none.',
            'AJ':'THR already paid separately this month, deducted once from current take-home pay to prevent a second payment. Enter zero when none.',
            'AK':'Incentive already paid separately this month, deducted once from current take-home pay. The earning remains in gross; the earlier bank payment is not paid again.',
        })
        for col,note in definitions.items():bycol[col]['note']=note
        master_specs={}
        for col,label in headers.items():
            if col in aliases or col in excluded:continue
            kind='date' if col in datecols else 'money' if col in moneycols else 'quantity' if col in {'CR','CS','CT','CU','CW','CX','CY','CZ'} and col not in native else 'identifier' if col in {'D','E','AM','AN','AO','AP','AQ','AR','AS','AT','AU','AV','CL','CN','CQ','CS'} else 'text'
            if col in native and native[col] in {'b:bank_name','b:acc_holder_name','b:bank_bic','f:hr.employee:emergency_contact'}:kind='text' if col not in {'CN','CS'} else 'identifier'
            dest=native.get(col,'c:amount' if kind in {'money','quantity'} else 'c:text')
            name='Master '+label.split('(')[0].strip().replace('\n',' ')
            if name in {s['name'] for s in specs}:name+=' (HR)'
            s=add(prefix+'HC'+col,name,kind,dest,note='Source employee master, '+('Singapore' if key=='sg' else 'ID')+'!'+col+'. Retained as HR context, not automatically paid. '+('Source money is preserved under its original heading; ambiguous annual/monthly labels do not override contract salary.' if kind=='money' else ''))
            master_specs[col]=s
        masterid=add(prefix+'MASTERID','Employee master ID (source)','identifier','c:text',note='Original HR identifier. Indonesia master repeats IDs and collides with other countries; payroll matching uses the unique country-prefixed local payroll number.')
        sourcecode=add(prefix+'PAYROLLID','Payroll employee number (source)','identifier','c:text',note='Original local payroll identifier for the approved August crosswalk. Reference only; not a secondary identity mapping.')
        sourcename=add(prefix+'PAYROLLNAME','Payroll name (source)','text','c:text',note='Original payroll name. Employee master name is preferred after a unique name or national-ID match.')
        sourcejoin=add(prefix+'PAYROLLJOIN','Payroll joining date (source)','date','c:text',note='Original August payroll joining date, retained for audit. The employee master supplies the saved join date when matched.')
        sourceend=add(prefix+'PAYROLLEND','Payroll resign date (source)','date','c:text',note='Original August payroll resign date. An actual payroll departure overrides a later fixed-term end date in the master.')
        state=add(prefix+'SOURCESTATE','Payroll population review','text','c:text',note='Source population decision. Master-only people without an August payroll row remain Draft for payroll until reviewed; inactive Singapore profiles are archived.')
        # Standard native second ID is separately mapped, so government IDs are
        # never populated with a payroll alias.
        nativeid=add(prefix+'NATIVEID','Employee ID','identifier','f:hr.employee:employee_id',required=True,note='Same unique canonical value as Employee Code/Badge ID.')
        if key=='id':
            bankcode=add('IDBANKCODE','Bank routing code (source)','identifier','c:text',note='Bank Transfer column E, matched by the unique local payroll number. Preserve leading zeros. This is a domestic bank routing code, not a SWIFT/BIC.')
            payrollnik=add('IDPAYROLLNIK','Payroll national ID (source)','identifier','c:text',note='August payroll NIK, retained as a separate audit value. The HC national ID remains authoritative when both sources provide conflicting values; the population-review field flags the disagreement.')
            for c in ('CP','CQ','CR'):
                master_specs[c]['note']='Employee-master bank field where available. Missing values are supplied by the August Bank Transfer sheet, using the unique local payroll number and matching employee name. Bank net amounts are controls, never payroll inputs.'
            master_specs['AN']['note']='HC government ID, with the matched payroll NIK used only when HC has no value. A conflicting payroll NIK stays in the separate source field and is flagged for HR review.'
            master_specs['AU']['note']='HC BPJS manpower number. Missing values are supplied by the matched August BPJS TK report. This reference identifier does not determine coverage or calculate a contribution.'
            bank_rows={text(r['values'][1]):r for r in data['payroll_id']['sheets']['Bank Transfer']['rows']
                if isinstance(r['values'][1],(int,float)) and r['values'][2]}
            tk_rows={text(r['values'][2]):r for r in data['payroll_id']['sheets']['BPJS TK']['rows']
                if isinstance(r['values'][2],(int,float)) and r['values'][3]}
            assert len(bank_rows)==122
        rules_bycode={s['code']:s for s in specs}
        records=[];fixtures=[]
        existing_people=live['people']
        links=[]
        if key=='sg':
            for r,m in [(7,3),(8,4)]:links.append(dict(payroll_row=r,master_row=m,canonical_id=masters[m]['values'][3],method='reviewed_name'))
            for m in (5,6):links.append(dict(payroll_row=None,master_row=m,canonical_id=masters[m]['values'][3],method='master_only'))
        else:
            links=list(cross['links'])
            for extra in cross['outside']:links.append(dict(payroll_row=None,master_row=extra['row'],canonical_id='IDHC'+extra['id'],method='master_only'))
        controls={r['source_row']:r for r in json.loads((root/(key+'-controls.json')).read_text())['controls']}
        for link in links:
            pr=link['payroll_row'];mr=link['master_row'];p=pays.get(pr);m=masters.get(mr)
            mv=m['values'] if m else []
            pv=p['values'] if p else []
            getm=lambda c:clean(mv[column_index_from_string(c)-1]) if m else None
            values={}
            for col,s in master_specs.items():
                v=getm(col)
                if s['kind']=='date':v=date(v)
                elif s['kind'] in {'money','quantity'}:v=money(v)
                else:v=text(v)
                if s['destination']=='f:hr.employee:country_id' and v:
                    # Native country dropdowns need country names, not a newly
                    # created country record named after a demonym.
                    v={'Canadian':'Canada','Indian':'India','Singaporean':'Singapore','Indonesian':'Indonesia'}.get(v,v)
                values[s['code']]=v
            source_master=text(getm('D'))
            canonical=link['canonical_id']
            if key=='id':canonical=re.sub(r'[^A-Za-z0-9]','',canonical)
            assert canonical.isascii() and canonical.isalnum() and len(canonical)<=18,'Native Odoo badge IDs must be alphanumeric and at most 18 characters.'
            name=clean(getm('F')) or (pv[1] if key=='sg' else pv[4])
            name=str(name).strip()
            assert not any(canonical in (x['barcode'],x['employee_id']) for x in existing_people),(canonical,'Canonical employee code already exists')
            active_pay=bool(p)
            status=('August payroll; master matched by '+link['method']) if p and m else 'August payroll; absent from employee master' if p else 'Inactive master profile' if key=='sg' else 'Master only; payroll status and settings require review'
            values.update({masterid['code']:source_master,sourcecode['code']:text(pv[3]) if p and key=='id' else None,
                sourcename['code']:str(pv[1] if key=='sg' else pv[4]).strip() if p else None,state['code']:status,nativeid['code']:canonical})
            values[sourcejoin['code']]=date(pv[4] if key=='sg' else pv[6]) if p else None
            values[sourceend['code']]=date(pv[7]) if p and key=='id' else None
            if key=='id':
                values[bankcode['code']]=None
                values[payrollnik['code']]=None
                if p:
                    local_id=text(pv[3]);bank=bank_rows[local_id]['values']
                    assert norm(bank[2])==norm(pv[4]),(local_id,'Bank recipient identity does not match payroll')
                    assert bank[7]==pv[42],(local_id,'Bank report does not reconcile to payroll net')
                    for c,i in [('CP',3),('CQ',5),('CR',6)]:
                        if not values.get(master_specs[c]['code']):values[master_specs[c]['code']]=text(bank[i])
                    values[bankcode['code']]=text(bank[4])
                    payroll_nik=text(pv[46]);master_nik=values.get(master_specs['AN']['code'])
                    values[payrollnik['code']]=payroll_nik
                    if not master_nik:values[master_specs['AN']['code']]=payroll_nik
                    elif payroll_nik and norm(master_nik)!=norm(payroll_nik):
                        status+='; National ID differs: HC retained, review before statutory filing'
                        values[state['code']]=status
                    tk=tk_rows.get(local_id)
                    if tk and not values.get(master_specs['AU']['code']):values[master_specs['AU']['code']]=text(tk['values'][5])
            if key=='sg':
                core={'A':canonical,'B':'SG','C':name,'D':getm('U'),'E':date(getm('K')),'F':date(getm('X')),
                    'G':'SGD','H':float(pv[7]) if p else money(getm('BF')),'J':'SINDA' if pr==8 else 'None',
                    'AF':getm('AP') or getm('AN') or 'Review',
                    'AB':getm('CK') or (clean(pv[26]) if p else None),
                    'AC':text(getm('CL')) or (text(pv[27]) if p else None),
                    'AD':getm('CM') or (clean(pv[28]) if p else None),
                    'AE':getm('CN') or (clean(pv[29]) if p else None)}
                if p:core.update(controls[pr]['inputs'],AG='August source comparison. Master identity and join date retained.')
                for col,val in core.items():values[bycol[col]['code']]=val
                expected={bycol[c]['code']:v for c,v in controls[pr]['expected'].items()} if p else {}
                if p:expected.update({bycol['AJ']['code']:float(pv[7]),bycol['AI']['code']:float(pv[8])+float(pv[9]),bycol['AK']['code']:float(pv[7])+float(pv[10])+float(pv[11])})
                join=core['F'];end=date(getm('AC'));wage=core['H']
            else:
                core={}
                if p:
                    for col in 'A B D E F G H I J K AR'.split():core[col]=clean(pv[column_index_from_string(col)-1])
                    core.update(controls[pr]['inputs'])
                    assert not pv[14],'A nonzero leave payout needs its approved source days; never invent them.'
                    assert not pv[24] and not pv[40],'Qualifying severance needs its actual payment history.'
                    core.update(BU=0,BW=0)  # Explicit August no-leave/no-severance controls.
                    core.update(D=canonical,E=name,A=getm('AH') or core['A'],F=getm('U') or core['F'],
                        G=date(getm('X')) or date(core['G']),H=date(core['H']) or date(getm('AC')))
                    core['AW']=controls[pr]['inputs']['AW']
                else:core=dict(D=canonical,E=name,A=getm('AH'),F=getm('U'),G=date(getm('X')),H=date(getm('AC')),
                    AS=None,J=None,AW=None,AX=None,AY=None,AZ=None,BA=None,I=getm('AD'),K='L',B=None)
                for col,val in core.items():
                    if col in {'D','B'}:val=text(val)
                    if col in {'G','H'}:val=date(val)
                    values[bycol[col]['code']]=val
                values.update(IDTAXROUND=-1 if pr==105 else 0,IDFINALALLOW=0,IDYTDREVIEW=1 if p and core['BB']==1 else 0,IDSEVELIG=0)
                expected={bycol[c]['code']:v for c,v in controls[pr]['expected'].items()} if p else {}
                join=core['G'];end=core['H'];wage=core['AS']
            assert join and re.match(r'^\d{4}-\d\d-\d\d$',join),(canonical,'Valid start date required')
            for s in specs:
                if s['type']=='input' and s['code'] not in values:values[s['code']]=None
            record=dict(code=canonical,name=name,master_row=mr,payroll_row=pr,values=values,
                payroll_eligible=active_pay,profile_active=active_pay if key=='sg' else True,
                contract_start=join,contract_end=end,wage=wage,source_status=status,
                canonical_match=link['method'])
            if key=='sg':record['bank_fill_from_payroll']=bool(p and not getm('CL'))
            records.append(record)
            if p:
                expected[prefix+'READY']='Ready'
                fixtures.append(dict(row=pr,code=canonical,values=values,expected=expected))
        assert len({r['code'] for r in records})==len(records)
        assert len({s['name'] for s in specs})==len(specs),[k for k,n in Counter(s['name'] for s in specs).items() if n>1]
        assert len(fixtures)==(2 if key=='sg' else 122)
        cfgmeta=next(x for x in live['configs'] if x['id']==cfgid)
        definition=dict(code='RIZE_'+('SINGAPORE' if key=='sg' else 'INDONESIA')+'_CLIENT_MASTER_2026',
            name='Rize '+('Singapore' if key=='sg' else 'Indonesia')+' Payroll Client Master 2026',
            country=country,company='Rize Farms',source_config_code=cfgmeta['code'],source_config_id=cfgid,
            source_hashes={k:v['sha256'] for k,v in data.items() if k.endswith('_'+key)},rules=specs,
            time_order=[s['code'] for s in specs if s['tab']=='Time'],
            adjustments_order=[s['code'] for s in specs if s['tab']=='Monthly adjustments'],
            monthly_period='2026-08',expected_payroll_count=len(fixtures),master_profile_count=len(records),
            primary_code=bycol['A' if key=='sg' else 'D']['code'],native_id_code=nativeid['code'],
            name_code=bycol['C' if key=='sg' else 'E']['code'],salary_code=bycol['H' if key=='sg' else 'AS']['code'],
            source_scope='Rize-only client configuration and employee records. No shared application code changes.',
            notes=([
                'August payroll is the calculation source; the Singapore master supplies canonical employee names, IDs and joining dates.',
                'One joining date differs between the Singapore master and payroll. The master date is retained; the original payroll date remains a separate audit field. August pay is unchanged.',
                'CPF, self-help contributions and employer CPF are reviewed monthly inputs. Never infer CPF residency or an opt-out from a foreign nationality or name.',
                'Proration follows the supplied August net-before-proration formula. SDL uses actual wages; employer remittance is rounded down after summing employee SDL.',
                'Future months require a new period-specific download, reviewed time/contribution inputs and current records. Inactive master profiles are not on the payroll roster.',
            ] if key=='sg' else [
                'Retains the six published Indonesia statutory editions and all 292 linked constants. Use statutory edition refresh when rates, wage floors or caps change; do not overwrite historical editions.',
                'Canonical employee codes use ID followed by the local payroll number, for example ID170. This satisfies the native alphanumeric badge rule. Repeated original HR IDs are preserved as references and never matched as government identifiers.',
                '118 payroll employees match the master (111 names, seven national IDs); four are payroll-only. Four master-only profiles remain Draft for payroll pending review.',
                'HC banking cells are empty. The matched August Bank Transfer report supplies bank name, account and holder for all 122 payroll employees; its net controls reconcile exactly. Domestic routing codes remain separate from SWIFT/BIC.',
                'Ten HC national IDs differ from the matched payroll NIK. HC values are retained as employee records, payroll NIKs remain separate audit values, and affected profiles are flagged for review before statutory filing. Missing HC identifiers can be supplied by matched payroll or August BPJS TK references.',
                'Known source contract end dates are retained. Ended contracts remain available for valid historical months; future months require a real renewal record.',
                'A visible -1 IDR regular allowance adjustment reproduces source row 105. August gross, deductions, tax and net controls then match exactly; future-month adjustment starts blank.',
                'Ordinary months use TER. December or the final employment period requires annual reconciliation and reviewed prior YTD values. Final net contracts use an explicitly approved annual gross-up allowance rather than a silent monthly approximation.',
                'Qualifying severance requires explicit eligibility confirmation. Source report tabs and accounting/bank views are output controls, never automatic payroll inputs.',
            ]))
        safe=repo/'docs/configuration'/('rize-'+('singapore' if key=='sg' else 'indonesia')+'-client-master-2026.json')
        safe.write_text(json.dumps(definition,ensure_ascii=False,indent=2)+'\n')
        (root/(key+'-definition.json')).write_text(json.dumps(definition,ensure_ascii=False,indent=2))
        (root/(key+'-record-plan.json')).write_text(json.dumps(dict(records=records),ensure_ascii=False,indent=2))
        (root/(key+'-fixtures.json')).write_text(json.dumps(fixtures,ensure_ascii=False,indent=2))
        print(json.dumps(dict(country=country,components=len(specs),inputs=sum(s['type']=='input' for s in specs),
            formulas=sum(s['type']=='formula' for s in specs),time=len(definition['time_order']),
            adjustments=len(definition['adjustments_order']),records=len(records),payroll=len(fixtures))))


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--private-dir',type=Path,required=True)
    parser.add_argument('--repo',type=Path,default=Path(__file__).resolve().parents[1])
    args=parser.parse_args()
    prepare(args.private_dir,args.repo)

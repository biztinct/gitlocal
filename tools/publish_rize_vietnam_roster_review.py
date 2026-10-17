"""Publish the owner's 70-person review roster without assigning live payroll.

Run in Rize Odoo shell with VN_HC_ROSTER_DIR containing the two updated
workbooks and roster-decision.json. Validates against the original attached HC
source and records the scope explicitly on the Draft configuration.
"""
import base64
import collections
import hashlib
import io
import json
import os
from pathlib import Path
import openpyxl

assert env.cr.dbname=='rize'
p=Path(os.environ['VN_HC_ROSTER_DIR'])
decision=json.loads((p/'roster-decision.json').read_text())
assert decision['scope']=='Include in review roster only'
assert not decision['activate_configuration'] and not decision['assign_live_population']
cfg=env['hr.formula.config'].search([('code','=',decision['configuration_code'])])
assert len(cfg)==1 and cfg.state=='draft'
source=openpyxl.load_workbook(io.BytesIO(base64.b64decode(cfg.import_sample_file)),data_only=True)['VN']
people=[(str(r[4]),r[5],r[36]) for r in source.iter_rows(min_row=3,values_only=True) if r[36] in decision['included_source_statuses']]
assert len(people)==decision['expected_review_people']==70
assert dict(collections.Counter(r[2] for r in people))==decision['expected_status_counts']
assert len({r[0] for r in people})==70
expected=[r[:2] for r in people]
names=['Vietnam_Configuration_Review.xlsx','Vietnam_Monthly_Input_Template.xlsx']
monthly=openpyxl.load_workbook(p/names[1],data_only=False)
assert monthly.sheetnames==['Time','Monthly adjustments','Instructions']
for tab in ['Time','Monthly adjustments']:
    ws=monthly[tab]
    actual=[(ws.cell(r,1).value,ws.cell(r,2).value) for r in range(2,72)]
    assert actual==expected,tab
    assert all(c.value is None for row in ws.iter_rows(min_row=2,min_col=3) for c in row),tab
    assert all(c.value is None for row in ws.iter_rows(min_row=72,max_col=2) for c in row),tab
review=openpyxl.load_workbook(p/names[0],data_only=True)
assert review['Review']['D18'].value=='Approved' and review['Review']['C30'].value==10
assert sum(r[5]=='Include' for r in review['Roster review'].iter_rows(min_row=2,values_only=True))==70
models=['hr.employee','hr.contract','hr.payslip','hr.payslip.run','hr.contract.advantage']
counts=lambda:{m:env[m].with_context(active_test=False).search_count([]) for m in models}
before=counts()
facade=env['pb.payroll.templates'].with_company(cfg.company_id)
assert facade.workspace(cfg.id,'2026-10')['people']==0
probe=env['hr.payroll.import.batch'].with_company(cfg.company_id).new(dict(formula_config_id=cfg.id,company_id=cfg.company_id.id,file_header_row=1,file_data_start_row=2))
parsed=probe._parse_source_file((p/names[1]).read_bytes(),names[1])
assert len(parsed['rows'])==70
parsed_keys=[]
for row in parsed['rows']:
    key=next((row[k] for k in ['Country employee ID','Time|Country employee ID','Monthly adjustments|Country employee ID'] if row.get(k)),None)
    parsed_keys.append(str(key))
assert len(set(parsed_keys))==70 and set(parsed_keys)=={r[0] for r in people}
prior_hash=hashlib.sha256(json.dumps(cfg.rule_ids.read(['code','excel_formula','net_role','net_role_detail','value_kind','is_contract_component']),sort_keys=True,default=str).encode()).hexdigest()
files={}
for name in names:
    attachment=env['ir.attachment'].search([('res_model','=','hr.formula.config'),('res_id','=',cfg.id),('name','=',name)])
    assert len(attachment)==1
    attachment.write({'datas':base64.b64encode((p/name).read_bytes())})
    digest=hashlib.sha256((p/name).read_bytes()).hexdigest()
    assert hashlib.sha256(base64.b64decode(attachment.datas)).hexdigest()==digest
    files[name]={'attachment_id':attachment.id,'sha256':digest}
note='Review roster confirmed on 17 October 2026: 68 Active and 2 Notice Period employees (70 people). Inclusion is for review only. Payroll activation, population assignment and record changes remain pending validation.'
if note not in (cfg.description or ''):
    cfg.description=(cfg.description or '')+'<p>'+note+'</p>'
assert counts()==before
assert cfg.state=='draft' and facade.workspace(cfg.id,'2026-10')['people']==0
assert prior_hash==hashlib.sha256(json.dumps(cfg.rule_ids.read(['code','excel_formula','net_role','net_role_detail','value_kind','is_contract_component']),sort_keys=True,default=str).encode()).hexdigest()
result=dict(database=env.cr.dbname,config_id=cfg.id,state=cfg.state,review_people=70,active=68,notice_period=2,joined_upload_rows=70,pending_review_decisions=10,live_population=0,files=files,before=before,after=counts())
(p/'roster-review-receipt.json').write_text(json.dumps(result,indent=2))
env.cr.commit()
print('VN_HC_ROSTER_REVIEW_PUBLISHED',json.dumps(result))

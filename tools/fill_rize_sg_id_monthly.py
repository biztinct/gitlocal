"""Plan fills and preserve the native signed monthly-workbook envelope.

Artifact Tool authors all monthly values. This helper first generates a fill
plan from the native application export, then verifies the authored values and
copies only those cells into the original ZIP envelope. It preserves receipts,
protection, styles and the read-only Employee records worksheet byte for byte.
Never edits the user's original source workbooks.
"""
import argparse
import copy
import hashlib
import json
from pathlib import Path
import posixpath
import shutil
import xml.etree.ElementTree as ET
import zipfile

import openpyxl


NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def sheet_paths(z):
    relationships=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    rel={r.attrib['Id']:posixpath.normpath(posixpath.join('xl',r.attrib['Target'])) if not r.attrib['Target'].startswith('/') else r.attrib['Target'].lstrip('/') for r in relationships}
    return {s.attrib['name']:rel[s.attrib['{'+NS['r']+'}id']] for s in ET.fromstring(z.read('xl/workbook.xml')).find('s:sheets',NS)}


def plan(root,output):
    for key,country in [('sg','Singapore'),('id','Indonesia')]:
        definition=json.loads((root/(key+'-definition.json')).read_text())
        people={r['code']:r for r in json.loads((root/(key+'-fixtures.json')).read_text())}
        catalog=json.loads((root/(key+'-catalog.json')).read_text())
        byid={c['id']:c for c in catalog}
        wb=openpyxl.load_workbook(root/(key+'-issued-template.xlsx'))
        changes=[];seen={};columns=set()
        for name in ('Time','Monthly adjustments'):
            ws=wb[name];fields=[];seen[name]=set()
            for cell in ws[1]:
                if cell.comment and cell.comment.text.startswith('field: rule:'):
                    field=cell.comment.text.splitlines()[0].split('field: ',1)[1].strip()
                    fields.append((cell.column,byid[field]['code']))
            for row in range(2,ws.max_row+1):
                code=ws.cell(row,1).value
                if not code:continue
                assert code in people,(country,code)
                assert code not in seen[name],(country,code,'duplicate')
                seen[name].add(code)
                for column,component in fields:
                    assert ws.cell(row,column).value is None,(country,name,row,column,'Monthly inputs must start blank')
                    value=people[code]['values'].get(component)
                    spec=next(s for s in definition['rules'] if s['code']==component)
                    assert value is not None or not spec['required'],(country,code,component,'Required source value missing')
                    columns.add(component)
                    if value is not None:
                        changes.append(dict(sheet=name,cell=ws.cell(row,column).coordinate,value=value,employee_code=code,code=component))
        assert all(keys==set(people) for keys in seen.values())
        assert len(columns)==len(definition['time_order'])+len(definition['adjustments_order'])
        assert wb['Employee records'].protection.sheet and wb['_Payobook'].sheet_state=='veryHidden'
        target=output/('Rize_'+country+'_August_2026_TEST_Monthly_Inputs.xlsx')
        (root/(key+'-fill-plan.json')).write_text(json.dumps(dict(output=str(target),cells=changes),indent=2))
        print(json.dumps(dict(country=country,employees=len(people),monthly_fields=len(columns),authored_cells=len(changes))))


def preserve(root,output):
    for key,country in [('sg','Singapore'),('id','Indonesia')]:
        fill=json.loads((root/(key+'-fill-plan.json')).read_text())
        authored_values=openpyxl.load_workbook(root/(key+'-artifact-monthly.xlsx'),data_only=True)
        target=Path(fill['output'])
        with zipfile.ZipFile(root/(key+'-issued-template.xlsx')) as original,zipfile.ZipFile(root/(key+'-artifact-monthly.xlsx')) as authored:
            old_paths=sheet_paths(original);updated={}
            for sheet in ('Time','Monthly adjustments'):
                old=ET.fromstring(original.read(old_paths[sheet]))
                oldcells={c.attrib['r']:c for c in old.findall('.//s:sheetData/s:row/s:c',NS)}
                for item in (x for x in fill['cells'] if x['sheet']==sheet):
                    value=authored_values[sheet][item['cell']].value
                    assert value==item['value'],(country,sheet,item['cell'],value,item['value'])
                    dest=oldcells[item['cell']]
                    for child in list(dest):
                        if child.tag in {'{'+NS['s']+'}v','{'+NS['s']+'}f','{'+NS['s']+'}is'}:dest.remove(child)
                    if isinstance(value,(int,float)):
                        dest.attrib['t']='n';ET.SubElement(dest,'{'+NS['s']+'}v').text=str(value)
                    else:
                        dest.attrib['t']='inlineStr'
                        text=ET.SubElement(ET.SubElement(dest,'{'+NS['s']+'}is'),'{'+NS['s']+'}t')
                        text.text=str(value)
                updated[old_paths[sheet]]=ET.tostring(old,encoding='utf-8',xml_declaration=True)
            with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as final:
                for info in original.infolist():final.writestr(copy.copy(info),updated.get(info.filename,original.read(info.filename)))
            with zipfile.ZipFile(target) as final:
                changed=[p for p in original.namelist() if original.read(p)!=final.read(p)]
                assert set(changed)==set(updated)
        target.chmod(0o600)
        monthly=output/('Rize_'+country+'_October_2026_Monthly_Inputs.xlsx')
        shutil.copyfile(root/(key+'-current-template.xlsx'),monthly);monthly.chmod(0o600)
        current=openpyxl.load_workbook(monthly,data_only=True)
        for name in ('Time','Monthly adjustments'):
            ws=current[name]
            for header in ws[1]:
                if header.comment and header.comment.text.startswith('field: rule:'):
                    assert all(ws.cell(r,header.column).value is None for r in range(2,ws.max_row+1)),(country,name,header.value,'No historical values may carry forward')
        report=dict(country=country,authored_cells=len(fill['cells']),changed_parts=changed,
            readonly_records_unchanged=True,issued_receipt_unchanged=True,future_month_inputs_blank=True,
            august_sha256=hashlib.sha256(target.read_bytes()).hexdigest(),october_sha256=hashlib.sha256(monthly.read_bytes()).hexdigest())
        (root/(key+'-native-preservation.json')).write_text(json.dumps(report,indent=2))
        print(json.dumps(report))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('mode',choices=['plan','preserve']);p.add_argument('--private-dir',type=Path,required=True);p.add_argument('--output-dir',type=Path,required=True);a=p.parse_args()
    (plan if a.mode=='plan' else preserve)(a.private_dir,a.output_dir)

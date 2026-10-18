"""Odoo shell: Rize-only master heading alignment; dry-run unless VN_LABELS_APPLY=1.

VN_LABELS_DIR contains the committed rize-vietnam-master-labels.json and the
unchanged source workbook. This changes labels and Excel bindings only. It
preserves the owner's current activation, all calculation expressions and all
employee, contract, dependent, monthly-input and payroll records.
"""
import base64
import hashlib
import io
import json
import math
import os
from pathlib import Path

import openpyxl

assert env.cr.dbname == 'rize', 'Client configuration belongs only to Rize.'
root=Path(os.environ['VN_LABELS_DIR'])
definition=json.loads((root/'rize-vietnam-master-labels.json').read_text())
content=(root/definition['filename']).read_bytes()
assert hashlib.sha256(content).hexdigest()==definition['sha256']
cfg=env['hr.formula.config'].search([('code','=',definition['code'])])
assert len(cfg)==1 and cfg.id==4760 and cfg.company_id.name=='Rize Farms'
e=env(context=dict(env.context,allowed_company_ids=cfg.company_id.ids))
studio=e['pb.formula.studio'].with_company(cfg.company_id)
facade=e['pb.payroll.templates'].with_company(cfg.company_id)
by={r.code:r for r in cfg.rule_ids}
assert len(by)==185

def digest(table):
    e.cr.execute('SELECT row_to_json(t)::text FROM '+table+' t ORDER BY id')
    return hashlib.sha256('\n'.join(r[0] for r in e.cr.fetchall()).encode()).hexdigest()

tables=['hr_employee','hr_contract','hr_contract_advantage','vietnam_employee_dependent',
        'hr_payslip','hr_payslip_run','pb_payroll_month_inputs',
        'hr_payslip_import_mapping','pb_payroll_template_layout']
baseline={t:digest(t) for t in tables}
before={r.code:(r.column_letter,r.column_type,r.excel_formula,r.value_kind,
                r.net_role,r.net_role_detail,r.is_contract_component,r.is_text_component)
        for r in cfg.rule_ids}
state=cfg.state
people_before=facade.workspace(cfg.id,'2026-10')['people']
tests=cfg.run_sample_tests()
assert tests['total']==tests['passed']==96
changes=[]
for spec in definition['rules']:
    rule=by[spec['code']]
    if rule.name != spec['name']:
        changes.append(dict(code=rule.code,old_name=rule.name,name=spec['name']))
        rule.with_context(formula_version_note='Owner requested original August master headings.').write({'name':spec['name']})
cfg.write(dict(import_sample_file=base64.b64encode(content),
               import_sample_filename=definition['filename'],
               import_sample_columns_json=json.dumps(definition['columns'],ensure_ascii=False)))
types=studio._import_column_types(cfg)
bindings=[]
for spec in definition['rules']:
    rule=by[spec['code']]
    if (rule.column_type!='input' or not spec['key']
            or studio._import_column_type(types,spec['key']).get('kind')=='formula'):
        continue
    old=rule.source_binding_key
    if old != spec['key']:
        rule.set_source_binding('excel',spec['key'],origin='user')
        rule.source_sheet_name='Salary'
        rule.original_column_letter=spec['column']
        bindings.append(dict(code=rule.code,old_key=old,key=spec['key']))

board=studio.import_mapping_data(cfg.id,'sample')
left={r['id']:r for r in board['left']}
visible={c['key'] for c in definition['columns'] if types[c['key']]['kind']!='formula'}
assert {key[2:] for key in left if key.startswith('c:')}==visible
assert not any(c['meta'].get('orphan') for c in board['left'])
assert not any(c['id'] in {r.id for r in cfg.rule_ids if r.column_type=='formula'} for c in board['right'])
assert sum(bool(c['meta'].get('collapsedGroup')=='Constants') for c in board['right'])==5
assert all(w['leftId'] in left for w in board['wires'])
for spec in definition['rules']:
    rule=by[spec['code']]
    if rule.column_type=='input' and spec['key'] in visible:
        assert any(w['leftId']=='c:'+spec['key'] and w['rightId']==rule.id and w['kind']=='mapping' for w in board['wires']),spec['code']
        if spec['code']!='VNINSBASE':
            assert left['c:'+spec['key']]['label']==rule.name,spec['code']
assert not {'c:Salary|Y','c:Salary|BC','c:Salary|BJ'} & set(left)
assert 'c:Salary|M' in left and left['c:Salary|M']['meta']['badge']=='Mixed'

journey=studio.journey_data(cfg.id)
file=next(c for c in journey['lanes']['systems'] if c['id']=='file')
assert all(r['id'][2:] in visible and r['state']!='gone' for r in file['rows'])
assert not any(c['id']=='dangling' for c in journey['lanes']['scheme'][0]['chips'])
assert not any(r['label']=='Blood Type' for r in file['rows'])
assert not any(r['label']=='Blood Type' for r in journey['lanes']['scheme'][0]['rows'])
assert any(r['label']=='Blood Type' for r in journey['lanes']['scheme'][0]['folded']['rows'])

# Re-reading the exact source keeps reviewed positions instead of reverting to
# colliding USD/VND/Others headings or its numbered guide row.
reread=studio.import_mapping_read_headers(cfg.id,base64.b64encode(content).decode(),definition['filename'])
assert {c['id'] for c in reread['left']}==set(left)
assert {w['id'] for w in reread['wires'] if w['kind']=='mapping'}=={w['id'] for w in board['wires'] if w['kind']=='mapping'}

# Exercise the real loader without creating a pay run. Both physical Others
# columns and all three currency pairs must keep their own cell values.
Batch=e['hr.payroll.import.batch']
probe=Batch.new(dict(formula_config_id=cfg.id,company_id=cfg.company_id.id,source_type='excel'))
parsed=probe._parse_source_file(content,definition['filename'])
wb=openpyxl.load_workbook(io.BytesIO(content),data_only=True)
salary=wb['Salary']
row=next(r for r in parsed['rows'] if r.get('Salary|B')==salary['B13'].value)
for col in ['K','L','BB','BC','BI','BJ','U','AF','AP','AU','AX','AY']:
    assert row.get('Salary|'+col)==salary[col+'13'].value,(col,row.get('Salary|'+col),salary[col+'13'].value)
assert not any(r.get('Salary|B')==2 for r in parsed['rows']), 'Column guide became an employee.'

workspace=facade.workspace(cfg.id,'2026-10')
assert workspace['people']==people_before
workbooks=[]
for mode in ('template','records'):
    download=facade.download(cfg.id,'2026-10',mode)
    data=base64.b64decode(download['file_b64'])
    book=openpyxl.load_workbook(io.BytesIO(data))
    assert book['Employee records'].max_row==people_before+1
    if mode=='template':
        assert [s.title for s in book if s.sheet_state=='visible']==['Time','Monthly adjustments','Employee records']
        assert book['Employee records'].protection.sheet
        assert all(c.protection.locked for row in book['Employee records'] for c in row)
    else:
        review=facade.review(cfg.id,'2026-10','records',download['file_b64'],download['filename'])
        assert not review['issues'] and not review['changes'],(review['issues'],len(review['changes']))
    workbooks.append(dict(mode=mode,rows=download['rows']))

assert cfg.run_sample_tests()['passed']==96
assert before=={r.code:(r.column_letter,r.column_type,r.excel_formula,r.value_kind,
                       r.net_role,r.net_role_detail,r.is_contract_component,r.is_text_component)
               for r in cfg.rule_ids}
assert cfg.state==state
assert baseline=={t:digest(t) for t in tables}, 'Business records or native mappings changed.'
receipt=dict(database=e.cr.dbname,configuration=cfg.id,state=state,label_changes=changes,
             binding_changes=bindings,source_inputs=len(visible),source_formulas_hidden=board['formula_columns_hidden'],
             visible_mappings=board['wired'],other_inputs=sum(c['meta'].get('collapsedGroup')=='Other scheme inputs' for c in board['right']),
             journey_file_fields=len(file['rows']),workbooks=workbooks,scenarios=tests,
             business_unchanged=baseline,formulas_unchanged=True)
apply=os.getenv('VN_LABELS_APPLY')=='1'
(root/('applied.json' if apply else 'dry-run.json')).write_text(json.dumps(receipt,ensure_ascii=False,indent=2))
if apply:
    e.cr.commit()
else:
    e.cr.rollback()
print('VN_LABELS_RESULT='+json.dumps({k:v for k,v in receipt.items() if k not in ('business_unchanged','label_changes','binding_changes')},ensure_ascii=False))

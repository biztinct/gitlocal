"""Align VN HC display names and delivery attachments, preserving live mappings.

Run in the Rize Odoo shell with VN_LABELS_DIR containing definition.json and
both reviewed workbooks. Dry run unless VN_LABELS_APPLY=1. Source keys, codes,
formulas, roles, destinations, population and employee values remain unchanged.
"""
import base64
import hashlib
import json
import os
from pathlib import Path
import openpyxl

assert env.cr.dbname == 'rize'
root = Path(os.environ['VN_LABELS_DIR'])
d = json.loads((root / 'definition.json').read_text())
cfg = env['hr.formula.config'].search([('code', '=', d['code'])])
assert len(cfg) == 1 and cfg.id == 4639 and cfg.state == 'draft'
by = {r.code: r for r in cfg.rule_ids}
plan = {r['code']: r['name'] for r in d['rules'] if r['source_column']}
assert len(plan) == 65 and set(plan) <= set(by)

def snapshot(records, excluded=()):
    fields = [k for k, f in records._fields.items()
              if f.store and not f.compute and f.type != 'binary'
              and k not in {'create_date', 'create_uid', 'write_date', 'write_uid', *excluded}]
    rows = records.read(fields)
    for row in rows:
        for k in fields:
            if records._fields[k].type == 'many2one':
                row[k] = row[k][0] if row[k] else False
    return rows

def protected_state():
    return {
        'rules': snapshot(cfg.rule_ids, ('name',)),
        'sources': snapshot(cfg.rule_ids.mapped('source_ids')),
        'mappings': snapshot(env['hr.payslip.import.mapping'].search([('salary_structure_id', '=', cfg.id)])),
        'config': snapshot(cfg),
        'old_configuration': snapshot(env['hr.formula.config'].browse(3942).rule_ids),
        'business_counts': {m: env[m].with_context(active_test=False).search_count([])
                            for m in ['hr.employee', 'hr.contract', 'hr.payslip', 'hr.payslip.run', 'hr.contract.advantage']},
    }

before = protected_state()
(root / 'protected-before.json').write_text(json.dumps(before, indent=2, default=str))
old_names = {code: by[code].name for code in plan}
studio = env['pb.formula.studio'].with_company(cfg.company_id)
changes = []
for code, name in plan.items():
    if by[code].name != name:
        result = studio.save_component(by[code].id, {'name': name})
        assert result.get('ok'), (code, result)
        changes.append({'code': code, 'before': old_names[code], 'after': name})
assert all(by[code].name == name for code, name in plan.items())
assert protected_state() == before, 'A protected field changed during the label update'
tests = cfg.run_sample_tests()
assert tests['total'] == tests['passed'] == 30 and tests['failed'] == 0, tests
assert protected_state() == before
facade = env['pb.payroll.templates'].with_company(cfg.company_id)
assert facade.workspace(cfg.id, '2026-10')['people'] == 0

# The review roster is still 70 people; all actual monthly pay cells stay blank.
files = ['Vietnam_Configuration_Review.xlsx', 'Vietnam_Monthly_Input_Template.xlsx']
monthly = openpyxl.load_workbook(root / files[1], data_only=False)
assert monthly.sheetnames == ['Time', 'Monthly adjustments', 'Instructions']
identities = None
for tab in ['Time', 'Monthly adjustments']:
    sheet = monthly[tab]
    assert [sheet.cell(1, c).value for c in (1, 2)] == [plan['VNHCE'], plan['VNHCF']]
    people = [(sheet.cell(r, 1).value, sheet.cell(r, 2).value) for r in range(2, 72)]
    assert len({p[0] for p in people}) == 70 and all(p[0] and p[1] for p in people)
    if identities is None:
        identities = people
    assert people == identities
    assert all(c.value is None for row in sheet.iter_rows(min_row=2, min_col=3) for c in row)
review = openpyxl.load_workbook(root / files[0], data_only=True)
assert review['Review']['D18'].value == 'Approved' and review['Review']['C30'].value == 10
for tab in ['Payroll rules', 'Contract components']:
    for row in review[tab].iter_rows(min_row=2, values_only=True):
        if row[0] in plan:
            assert row[1] == plan[row[0]], (tab, row[0])
probe = env['hr.payroll.import.batch'].with_company(cfg.company_id).new(dict(
    formula_config_id=cfg.id, company_id=cfg.company_id.id, file_header_row=1, file_data_start_row=2))
parsed = probe._parse_source_file((root / files[1]).read_bytes(), files[1])
assert len(parsed['rows']) == 70
keys = [str(next((row[k] for k in [plan['VNHCE'], 'Time|' + plan['VNHCE'],
                                  'Monthly adjustments|' + plan['VNHCE']] if row.get(k)), '')) for row in parsed['rows']]
assert len(set(keys)) == 70 and set(keys) == {str(p[0]) for p in identities}
attachments = {}
for name in files:
    attachment = env['ir.attachment'].search([
        ('res_model', '=', 'hr.formula.config'), ('res_id', '=', cfg.id), ('name', '=', name)])
    assert len(attachment) == 1
    data = (root / name).read_bytes()
    attachment.write({'datas': base64.b64encode(data)})
    digest = hashlib.sha256(data).hexdigest()
    assert hashlib.sha256(base64.b64decode(attachment.datas)).hexdigest() == digest
    attachments[name] = {'id': attachment.id, 'sha256': digest}
assert protected_state() == before
result = dict(database=env.cr.dbname, configuration=cfg.id, state=cfg.state,
              renamed=len(changes), changes=changes, tests=tests, parsed_review_people=70,
              protected_fields_unchanged=True, business_counts=before['business_counts'], attachments=attachments)
(root / 'label-receipt.json').write_text(json.dumps(result, indent=2))
if os.environ.get('VN_LABELS_APPLY') == '1':
    env.cr.commit()
    print('VN_LABELS_DEPLOYED', json.dumps({k: v for k, v in result.items() if k != 'changes'}))
else:
    env.cr.rollback()
    print('VN_LABELS_DRY_RUN', json.dumps({k: v for k, v in result.items() if k != 'changes'}))

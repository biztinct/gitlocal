"""Extract anonymous numerical controls from the client's active payroll tab.

The output is private working data, not a repository fixture. Historical tabs
are inventoried only; their stale references do not drive the new template.
"""
import argparse
import json
from pathlib import Path

import openpyxl


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('source')
    parser.add_argument('output')
    args = parser.parse_args()
    values = openpyxl.load_workbook(args.source, data_only=True)
    formulas = openpyxl.load_workbook(args.source, data_only=False)
    sheet = values['Payroll Input - August26']
    controls = []
    for row in range(7, sheet.max_row + 1):
        if not sheet.cell(row, 2).value:
            continue
        controls.append({
            'source_row': row,
            'inputs': {new: sheet[f'{old}{row}'].value or 0 for new, old in {
                'H': 'H', 'I': 'I', 'K': 'J', 'L': 'K', 'T': 'S',
                'U': 'T', 'V': 'U', 'Y': 'X', 'Z': 'Y',
            }.items()},
            'expected': {new: sheet[f'{old}{row}'].value for new, old in {
                'M': 'L', 'N': 'M', 'O': 'N', 'P': 'O', 'Q': 'P',
                'W': 'V', 'X': 'W', 'AA': 'Z',
            }.items()},
        })
    inventory = [{
        'sheet': sheet.title,
        'formula_count': sum(cell.data_type == 'f' for row in sheet for cell in row),
        'broken_references': sum(cell.data_type == 'f' and '#REF!' in cell.value
                                 for row in sheet for cell in row),
    } for sheet in formulas]
    Path(args.output).write_text(json.dumps({'controls': controls, 'inventory': inventory}))
    print(json.dumps({'employees': len(controls), 'inventory': inventory}))


if __name__ == '__main__':
    main()

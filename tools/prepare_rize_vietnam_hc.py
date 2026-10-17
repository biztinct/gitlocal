"""Build a non-personal mapping definition and private source review from VN HC.

Read-only extraction; does not edit the supplied workbook or employee records.
Usage: python prepare_rize_vietnam_hc.py SOURCE WORKDIR
"""
import collections
import datetime as dt
import hashlib
import json
import sys
from pathlib import Path

import openpyxl

source, root = Path(sys.argv[1]), Path(sys.argv[2])
root.mkdir(parents=True, exist_ok=True)
wb = openpyxl.load_workbook(source, data_only=True)
wf = openpyxl.load_workbook(source, data_only=False)
s, sf = wb['VN'], wf['VN']
rows = [r for r in range(3, s.max_row + 1) if s.cell(r, 6).value]
rules, mapping, issues = [], [], []

def add(code, name, kind='text', group='Employee details', dest='', col='', note='', formula='', role='info', detail=False, tab='', required=False, value=None):
    rule = dict(code=code, name=name, kind=kind, group=group, destination=dest,
                source_column=col, note=note, formula=formula, role=role,
                detail=detail, tab=tab, required=required, value=value)
    rule['type'] = 'formula' if formula else 'constant' if value is not None else 'input'
    if col:
        rule['source_key'] = 'VN|' + str(sf[f'{col}2'].value)
    elif tab:
        rule['source_key'] = tab + '|' + name
    rules.append(rule)
    return rule

# Native destinations are used only where the source's meaning is unambiguous.
# Unmatched HR context is retained as contract text/amount, not forced into a
# selection field, tax status, related model or employment activation flag.
specs = {
 'B':('Country','text','c:text'), 'C':('Position ID','identifier','c:text'),
 'D':('Global employee ID','identifier','c:text'),
 'E':('Country employee ID','identifier','f:hr.employee:employee_id'),
 'F':('Employee name','text','f:hr.employee:name'),
 'G':('First name','text','c:text'), 'H':('Last name','text','c:text'),
 'I':('Gender as supplied','text','c:text'),
 'J':('Nationality as supplied','text','c:text'),
 'K':('Date of birth','date','f:hr.employee:birthday'),
 'M':('Place of birth','text','f:hr.employee:place_of_birth'),
 'P':('Marital status as supplied','text','c:text'),
 'Q':('HC dependent count','quantity','c:amount'),
 'R':('Personal email','text','f:hr.employee:private_email'),
 'S':('Personal phone','identifier','f:hr.employee:private_phone'),
 'T':('Work phone','identifier','f:hr.employee:work_phone'),
 'U':('Job title','text','f:hr.employee:job_title_text'),
 'V':('Country of work','text','c:text'), 'W':('City of work','text','c:text'),
 'X':('Contract start date','date','f:hr.contract:date_start'),
 'Y':('Probation start date','date','c:text'),
 'Z':('Probation end date','date','f:hr.contract:trial_date_end'),
 'AA':('Probation outcome','text','c:text'),
 'AB':('Labour contract signing date','date','c:text'),
 'AC':('Contract end date as supplied','text','c:text'),
 'AD':('Employment type as supplied','text','c:text'),
 'AE':('Reporting manager name','text','c:text'),
 'AF':('Reporting manager email','text','c:text'),
 'AG':('Hiring manager','text','c:text'),
 'AH':('Department','text','f:hr.employee:department_id'),
 'AI':('Subdepartment','text','c:text'), 'AJ':('Job level','text','c:text'),
 'AK':('Source employment status','text','c:text'),
 'AN':('Government ID','identifier','c:text'),
 'AO':('Work permit number','identifier','f:hr.employee:vietnam_work_permit_number'),
 'AP':('Passport number','identifier','f:hr.employee:vietnam_passport_number'),
 'AQ':('Visa status as supplied','text','c:text'),
 'AR':('Vietnam tax code','identifier','f:hr.employee:vietnam_tax_code'),
 'AW':('Social insurance number','identifier','f:hr.employee:vietnam_social_insurance_number'),
 'BD':('Current address','text','f:hr.employee:vietnam_temporary_address'),
 'BE':('Emergency contact phone','identifier','f:hr.employee:vietnam_emergency_contact_phone'),
 'BF':('Emergency contact name','text','f:hr.employee:vietnam_emergency_contact_name'),
 'BG':('Emergency relationship','text','f:hr.employee:vietnam_emergency_contact_relationship'),
 'BH':('Salary currency','text','c:text'), 'BI':('Gross or net basis','text','c:text'),
 'BK':('Monthly fixed salary VND','money','f:hr.contract:wage'),
 'BL':('Annual variable target rate','rate','c:amount'),
 'BQ':('HC bonus reference VND','money','c:amount'),
 'BS':('ESOP reference','text','c:text'),
 'BY':('Resignation date','date','c:text'),
 'BZ':('Notice period days','quantity','c:amount'),
 'CA':('Planned last working date','date','c:text'),
 'CB':('Actual last working date','date','c:text'),
 'CC':('Unserved notice days','quantity','c:amount'),
 'CD':('Exit reason','text','c:text'),
 'CE':('Exit interview status','text','c:text'),
 'CF':('Final settlement status','text','c:text'),
 'CG':('Clearance status','text','c:text'),
 'CH':('Work email','text','f:hr.employee:work_email'),
 'CP':('Bank name','text','b:bank_name'),
 'CQ':('Bank account number','identifier','b:acc_number'),
 'CR':('Bank beneficiary','text','b:acc_holder_name'),
 'CS':('Bank SWIFT code','identifier','b:bank_bic'),
 'CY':('Annual sick leave entitlement','quantity','c:amount'),
}
reference = {'A':'Row number is not an identity.', 'L':'Age is date-dependent; retain in HR reporting.',
 'AL':'Duplicate joining date derived from probation date. Use X for joining date.',
 'AM':'Service duration uses TODAY and is not a payroll input.',
 'BJ':'Annual salary = monthly fixed salary x 12. Reference, not monthly earnings.',
 'BM':'Annual target variable compensation; not an earned monthly bonus.',
 'BN':'Monthly target allocation; use approved actual bonus in monthly upload.',
 'BO':'Annual fixed plus target variable compensation; reference only.',
 'BP':'Monthly target gross includes unearned variable target; not payroll salary.',
 'BR':'13th-month estimate hardcodes 2026 and ignores exit date; upload approved actual payment.',
 'BT':'Employer-cost estimate assumes blanket eligibility; actual payroll calculates separately.',
 'BU':'Annual employer-cost reference; never an employee deduction.',
 'BV':'Annual target CTC; not payroll earnings.',
 'BW':'USD CTC reference uses DA2 FX rate; salary is paid in VND.',
 'BX':'USD annual CTC reference; not payroll earnings.',
 'CW':'Leave entitlement formula includes individual carryovers and 2026 dates. Review in leave module.',
 'CX':'Leave balance formula contains individual hardcoded usage. Review in leave module.',
 'CZ':'Sick leave balance is not paid days or a cash deduction.',
}
excluded = set('N O AS AT AU AV AX AY AZ BA BB BC CI CJ CK CL CM CN CO CT CU CV'.split())
for c in range(1,105):
    col = openpyxl.utils.get_column_letter(c)
    pop = sum(sf[f'{col}{r}'].value is not None for r in rows)
    if col in specs:
        name, kind, dest = specs[col]
        group = 'Compensation reference' if col in {'BH','BI','BK','BL','BQ','BS'} else 'Bank details' if col in {'CP','CQ','CR','CS'} else 'Employment context' if 24 <= c <= 39 or 77 <= c <= 85 else 'Employee details'
        note = 'Proposed mapping. Blank values keep existing records.'
        if col=='D': note='Not a match key: duplicate and missing IDs exist.'
        if col=='E': note='Unique across all 127 source people; proposed match key.'
        if col=='Q': note='HC count does not establish registered tax-dependent eligibility.'
        if col=='BK': note='VND gross only. Review salaries and effective date before applying.'
        if col=='AC': note='Mixed dates and Indefinite text. Preserve as text; review native end dates separately.'
        if col=='AK': note='Context only. Does not activate, terminate or select employees.'
        if col in {'BQ','BL'}: note='Target/reference only. Never automatically paid each month.'
        if kind == 'quantity':
            kind = 'integer'
        rule=add('VNHC'+col,name,kind,group,dest,col,note)
        status='Mapped'
    else:
        note=reference.get(col,'Not used by payroll. Retain in the relevant HR, document, asset or education process.')
        status='Reference only' if col in reference else 'Outside payroll'
        rule={}
    mapping.append(dict(column=col,source=str(sf[f'{col}2'].value),populated=pop,status=status,
                        code=rule.get('code',''),name=rule.get('name',''),kind=rule.get('kind',''),
                        destination=rule.get('destination',''),note=note))

# Joining date is a different destination from contract start, with the same source.
add('VNJOINED','Date of joining','date','Employment context','f:hr.employee:date_of_joining','X')

# New recurring controls are deliberately separate from HC's target/reference values.
for code,name,kind,note in [
 ('VNTAXMODE','Tax calculation mode','text','Set Resident or Manual after tax-residency review. No inference from nationality or contract type.'),
 ('VNREGISTER','Registered tax dependents','integer','Confirm registered eligible dependents; do not copy HC children count blindly.'),
 ('VNINSBASE','Insurance salary VND','money','Confirm statutory contribution base separately from gross salary.'),
 ('VNSIJOIN','SI participation','text','Set Yes or No after eligibility review.'),
 ('VNHIJOIN','HI participation','text','Set Yes or No after eligibility review.'),
 ('VNUIJOIN','UI participation','text','Set Yes or No after eligibility review, including foreign-worker rules.'),
 ('VNUIREGCAP','UI salary cap VND','money','Review work region. 106200000 is an example for region I, not a default for everyone.'),
 ('VNTUJOIN','Employer union participation','text','Set Yes or No after employer union contribution review.'),
 ('VNTRANSPORT','Fixed transport allowance VND','money','Confirm policy. Not inferred from job titles.'),
 ('VNPHONE','Fixed phone allowance VND','money','Confirm recurring amount and tax treatment.'),
 ('VNOTHERFIX','Other fixed allowance VND','money','Confirm contractual amount and tax treatment.'),
]:
    add(code,name,kind,'Recurring payroll settings','c:text' if kind=='text' else 'c:amount',note=note)

for code,name,kind,tab in [
 ('VNSTDDAYS','Standard working days','quantity','Time'),
 ('VNPAIDDAYS','Paid days','quantity','Time'),
 ('VNUNPAID','Unpaid leave days reference','quantity','Time'),
 ('VNWEEKHRS','Weekend overtime hours reference','quantity','Time'),
 ('VNACTBONUS','Actual variable bonus VND','money','Monthly adjustments'),
 ('VNTHIRTEEN','Actual 13th month VND','money','Monthly adjustments'),
 ('VNINCENTIV','Other incentive VND','money','Monthly adjustments'),
 ('VNOVERTIME','Approved overtime pay VND','money','Monthly adjustments'),
 ('VNADDEARN','Other taxable earnings VND','money','Monthly adjustments'),
 ('VNTAXEXEM','Approved cash tax exemption VND','money','Monthly adjustments'),
 ('VNBENEFIT','Taxable noncash benefit VND','money','Monthly adjustments'),
 ('VNREIMBURS','Nontaxable reimbursement VND','money','Monthly adjustments'),
 ('VNADVANCE','Advance recovery VND','money','Monthly adjustments'),
 ('VNDEDUCT','Other net deduction VND','money','Monthly adjustments'),
 ('VNUNIONEE','Employee union dues VND','money','Monthly adjustments'),
 ('VNPITMAN','Reviewed manual PIT VND','money','Monthly adjustments'),
 ('VNPITREF','PIT refund VND','money','Monthly adjustments'),
]:
    add(code,name,kind,'Monthly pay inputs',tab=tab,note='Positive amount. Review approved monthly values. Blank variable amounts are zero for a new month.')
for code,name,value in [('VNSICAP','SI and HI salary cap VND',50600000),('VNPERSONAL','Resident personal deduction VND',15500000),('VNDEPAMT','Resident dependent deduction VND',6200000)]:
    add(code,name,'money','Statutory parameters',value=value,note='2026 edition. Review applicability for the selected pay month.')

# Formula placeholders resolve to immutable component column letters in installer.
formula_defs = [
 ('VNREADY','Payroll input validation','IF(OR({VNSTDDAYS}<=0,{VNPAIDDAYS}<0,{VNPAIDDAYS}>{VNSTDDAYS}),"Review paid days",IF(OR({VNHCBH}<>"VND",{VNHCBI}<>"Gross",{VNHCBK}<=0),"Review gross VND salary",IF(OR(AND({VNSIJOIN}<>"Yes",{VNSIJOIN}<>"No"),AND({VNHIJOIN}<>"Yes",{VNHIJOIN}<>"No"),AND({VNUIJOIN}<>"Yes",{VNUIJOIN}<>"No"),AND({VNTUJOIN}<>"Yes",{VNTUJOIN}<>"No")),"Review participation",IF(AND(OR({VNSIJOIN}="Yes",{VNHIJOIN}="Yes",{VNUIJOIN}="Yes",{VNTUJOIN}="Yes"),{VNINSBASE}<=0),"Review insurance base",IF(AND({VNUIJOIN}="Yes",{VNUIREGCAP}<=0),"Review UI cap",IF(AND({VNTAXMODE}<>"Resident",{VNTAXMODE}<>"Manual"),"Review tax mode",IF(OR({VNREGISTER}<0,MOD({VNREGISTER},1)<>0,{VNPITMAN}<0),"Review tax inputs",IF(OR({VNTAXEXEM}<0,{VNTAXEXEM}>SUM({VNBASIC},{VNFIXALLOW},{VNACTBONUS},{VNTHIRTEEN},{VNINCENTIV},{VNOVERTIME},{VNADDEARN})),"Review exemption","Ready"))))))))','info',False,'Payroll validation'),
 ('VNBASIC','Prorated fixed salary VND','IF(OR({VNSTDDAYS}<=0,{VNPAIDDAYS}<0,{VNPAIDDAYS}>{VNSTDDAYS}),"Review paid days",ROUND({VNHCBK}/{VNSTDDAYS}*{VNPAIDDAYS},0))','earning',True,'Earnings'),
 ('VNFIXALLOW','Prorated fixed allowances VND','IF({VNSTDDAYS}<=0,"Review standard days",ROUND(SUM({VNTRANSPORT},{VNPHONE},{VNOTHERFIX})/{VNSTDDAYS}*{VNPAIDDAYS},0))','earning',True,'Earnings'),
 ('VNCASHGROSS','Cash gross pay VND','SUM({VNBASIC},{VNFIXALLOW},{VNACTBONUS},{VNTHIRTEEN},{VNINCENTIV},{VNOVERTIME},{VNADDEARN})','earning',False,'Totals'),
 ('VNSIEMP','Employee SI VND','IF({VNSIJOIN}="Yes",ROUND(MIN({VNINSBASE},{VNSICAP})*8%,0),IF({VNSIJOIN}="No",0,"Review SI participation"))','deduction',True,'Employee deductions'),
 ('VNHIEMP','Employee HI VND','IF({VNHIJOIN}="Yes",ROUND(MIN({VNINSBASE},{VNSICAP})*1.5%,0),IF({VNHIJOIN}="No",0,"Review HI participation"))','deduction',True,'Employee deductions'),
 ('VNUIEMP','Employee UI VND','IF({VNUIJOIN}="Yes",IF({VNUIREGCAP}<=0,"Review UI cap",ROUND(MIN({VNINSBASE},{VNUIREGCAP})*1%,0)),IF({VNUIJOIN}="No",0,"Review UI participation"))','deduction',True,'Employee deductions'),
 ('VNINSDED','Employee insurance total VND','SUM({VNSIEMP},{VNHIEMP},{VNUIEMP})','info',False,'Tax calculation'),
 ('VNTAXABLE','Taxable income before deductions VND','MAX(0,{VNCASHGROSS}+{VNBENEFIT}-{VNTAXEXEM})','info',False,'Tax calculation'),
 ('VNASSESS','Resident assessable income VND','MAX(0,{VNTAXABLE}-{VNINSDED}-{VNPERSONAL}-{VNREGISTER}*{VNDEPAMT})','info',False,'Tax calculation'),
 ('VNPITRES','Resident progressive PIT VND','ROUND(IF({VNASSESS}<=10000000,{VNASSESS}*5%,IF({VNASSESS}<=30000000,{VNASSESS}*10%-500000,IF({VNASSESS}<=60000000,{VNASSESS}*20%-3500000,IF({VNASSESS}<=100000000,{VNASSESS}*30%-9500000,{VNASSESS}*35%-14500000)))),0)','info',False,'Tax calculation'),
 ('VNPITDUE','PIT withheld VND','IF({VNTAXMODE}="Resident",{VNPITRES},IF({VNTAXMODE}="Manual",{VNPITMAN},"Review tax mode"))','deduction',True,'Employee deductions'),
 ('VNDEDTOTAL','Total employee deductions VND','SUM({VNINSDED},{VNPITDUE},{VNADVANCE},{VNDEDUCT},{VNUNIONEE})','deduction',False,'Totals'),
 ('VNNETPAY','Net payment VND','{VNCASHGROSS}-{VNDEDTOTAL}+{VNREIMBURS}+{VNPITREF}','net',False,'Totals'),
 ('VNSIER','Employer SI including accident VND','IF({VNSIJOIN}="Yes",ROUND(MIN({VNINSBASE},{VNSICAP})*17.5%,0),IF({VNSIJOIN}="No",0,"Review SI participation"))','employer_cost',True,'Employer contributions'),
 ('VNHIERR','Employer HI VND','IF({VNHIJOIN}="Yes",ROUND(MIN({VNINSBASE},{VNSICAP})*3%,0),IF({VNHIJOIN}="No",0,"Review HI participation"))','employer_cost',True,'Employer contributions'),
 ('VNUIERR','Employer UI VND','IF({VNUIJOIN}="Yes",IF({VNUIREGCAP}<=0,"Review UI cap",ROUND(MIN({VNINSBASE},{VNUIREGCAP})*1%,0)),IF({VNUIJOIN}="No",0,"Review UI participation"))','employer_cost',True,'Employer contributions'),
 ('VNTUER','Employer union contribution VND','IF({VNTUJOIN}="Yes",ROUND(MIN({VNINSBASE},{VNSICAP})*2%,0),IF({VNTUJOIN}="No",0,"Review employer union"))','employer_cost',True,'Employer contributions'),
 ('VNERCONTR','Total employer contributions VND','SUM({VNSIER},{VNHIERR},{VNUIERR},{VNTUER})','employer_cost',False,'Totals'),
 ('VNEMPCTC','Total employer payroll cost VND','SUM({VNCASHGROSS},{VNERCONTR},{VNREIMBURS},{VNBENEFIT})','info',False,'Totals'),
]
for code,name,formula,role,detail,group in formula_defs:
    if code not in {'VNREADY','VNBASIC','VNFIXALLOW'}:
        formula='IF({VNREADY}<>"Ready","Review payroll inputs",'+formula+')'
    add(code,name,'text' if code=='VNREADY' else 'money',group,formula=formula,role=role,detail=detail)
by_code={r['code']:r for r in rules}
for code in ['VNREIMBURS','VNPITREF']:
    by_code[code].update(role='earning',detail=False,group='Net payment additions')
for code in ['VNADVANCE','VNDEDUCT','VNUNIONEE']:
    by_code[code].update(role='deduction',detail=True,group='Employee deductions')
for code in ['VNACTBONUS','VNTHIRTEEN','VNINCENTIV','VNOVERTIME','VNADDEARN']:
    by_code[code].update(role='earning',detail=True,group='Earnings')

def issue(row,field,message):
    issues.append(dict(row=row,key=str(s[f'E{row}'].value),field=field,issue=message,status='Open'))
for r in rows:
    status=s[f'AK{r}'].value
    if status in ('Active','Notice Period'):
        for col in ['BK','BH','BI','CQ','AR','AW']:
            if s[f'{col}{r}'].value is None: issue(r,col,'Missing value for proposed payroll population')
        if s[f'BH{r}'].value!='VND': issue(r,'BH','Salary currency is not VND')
        if s[f'BI{r}'].value!='Gross': issue(r,'BI','Gross salary formula needs a reviewed gross amount')
        if s[f'CB{r}'].value: issue(r,'CB','Active/notice employee has a last-working date; review pay-month eligibility')
        for col in ['Z','AB','AC']:
            val=s[f'{col}{r}'].value
            if isinstance(val,dt.datetime) and val<s[f'X{r}'].value: issue(r,col,'Date precedes joining date')
        end=s[f'AC{r}'].value
        if isinstance(end,dt.datetime) and end<dt.datetime(2026,10,1): issue(r,'AC','Contract end precedes October 2026 while source status remains Active/Notice')
        q=s[f'Q{r}'].value
        if q is not None and not isinstance(q,(int,float)): issue(r,'Q','Dependent count is text; confirm registered tax dependents')
    gid=s[f'D{r}'].value
    if gid and sum(s[f'D{rr}'].value==gid for rr in rows)>1: issue(r,'D','Global employee ID is duplicated; country employee ID remains unique')
    if gid=='#N/A': issue(r,'D','Global employee ID contains literal #N/A')

review = [
 ['Roster','Confirm 68 Active plus 2 Notice Period employees for the first pay month. Exclude 57 Inactive unless final pay is approved.','Pending'],
 ['Identity','Approve Country employee ID (VN column E). Resolve changes against existing employees before importing.','Pending'],
 ['Salary','Approve BK gross VND salary and effective date. Review net/foreign-currency exceptions.','Pending'],
 ['Contracts','Review date anomalies, Indefinite end dates, probation dates, department and manager names.','Pending'],
 ['Recurring amounts','Review fixed transport, phone, other allowance, insurance base, UI region cap and participation flags.','Pending'],
 ['Tax','Set Resident or Manual mode and registered dependents. Nonresident, short-term, net-to-gross and special tax cases use reviewed manual PIT.','Pending'],
 ['Tax exemptions','Approve overtime exemption, phone/transport treatment and reimbursements. No automatic blanket exemption.','Pending'],
 ['Bonus','Confirm actual bonus and 13th-month payment policy. HC targets and 2026 estimates do not create monthly earnings.','Pending'],
 ['Leave','Validate balances in leave module. Sheet4 has no headers or period; it is not a paid-days source.','Pending'],
 ['Dependents','Dependent sheet has names, inconsistent relationship labels and a placeholder row. Add country employee ID, dates and eligibility before import.','Pending'],
 ['Cutover','After validation, apply approved record changes through Records Desk, assign the new scheme, generate the populated monthly download, and compare a parallel draft.','Pending'],
]
definition=dict(name='Rize Vietnam Payroll HC Oct 2026',code='RIZE_VIETNAM_HC_202610',
 source=source.name,sha256=hashlib.sha256(source.read_bytes()).hexdigest(), rules=rules,mapping=mapping,
 review=review,sources=[source.name,'2026 resident PIT: https://xaydungchinhsach.chinhphu.vn/quy-dinh-moi-ve-khau-tru-thue-thu-nhap-ca-nhan-119260703150410707.htm',
 '2026 insurance reference/cap: https://baohiemxahoi.gov.vn/tintuc/Pages/hoat-dong-he-thong-bao-hiem-xa-hoi.aspx?CateID=52&ItemID=26780&OtItem=date',
 'Insurance rates and accident component: https://baohiemxahoi.gov.vn/chuyen-trang-bhxh-bhyt/tu-van-hoi-dap/Pages/default.aspx?ItemID=18093'])
(root/'definition.json').write_text(json.dumps(definition,ensure_ascii=False,indent=2))
(root/'source-review.json').write_text(json.dumps(dict(counts=dict(collections.Counter(s[f'AK{r}'].value for r in rows)),issues=issues,
 roster=[dict(row=r,key=str(s[f'E{r}'].value),name=s[f'F{r}'].value,status=s[f'AK{r}'].value) for r in rows]),ensure_ascii=False,default=str))
print(json.dumps(dict(source_people=len(rows),rules=len(rules),dispositions=len(mapping),mapped=sum(m['status']=='Mapped' for m in mapping),issues=len(issues))))

"""Isolated XLSX cell-type checks; fixtures contain no employee data."""
import importlib.util
import io
from pathlib import Path
import unittest
import openpyxl

path = Path(__file__).resolve().parents[2] / 'pb_formula_studio/models/spreadsheet_columns.py'
spec = importlib.util.spec_from_file_location('spreadsheet_columns', path)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class FormulaColumnTypes(unittest.TestCase):
    def test_physical_columns_keep_repeated_headers_and_ignore_numbering_and_footer(self):
        book = openpyxl.Workbook()
        ws = book.active
        ws.title = 'Salary'
        ws.append(['Group', None, 'Other group'])
        ws.append(['Code', 'Others', 'Others'])
        ws.append([1, 2, 3])
        ws.append(['E1', 100, '=B4*2'])
        ws.append(['E2', 200, '=B5*2'])
        ws.append(['Totals', '=SUM(B4:B5)', 'an unrelated note'])
        out = io.BytesIO()
        book.save(out)
        cols = tuple((f'Salary|{letter}', 'Salary', label, letter, True, letter, 4, 5)
                     for letter,label in [('A','Code'),('B','Others'),('C','Others')])
        types = module.classify_xlsx_columns(out.getvalue(), cols)
        self.assertEqual(types['Salary|B']['kind'], 'input')
        self.assertEqual(types['Salary|C']['kind'], 'formula')
        self.assertEqual(types['Salary|C']['letter'], 'C')
        self.assertEqual(types['Salary|C']['data_rows'], 2)

    def test_sections_blanks_mixed_values_and_later_formulas(self):
        book = openpyxl.Workbook()
        ws = book.active
        ws.title = 'People'
        ws.append(['HR section'])
        ws.append(['Code', 'Age', 'Salary', 'Entered text'])
        ws.append(['TEST1', None, 100, '=literal text'])
        ws['D3'].data_type = 's'
        ws.append(['TEST2', '=10+20', '=50*2', None])
        ws.append(['TEST3', '=10+30', 200, None])
        out = io.BytesIO()
        book.save(out)
        cols = tuple((f'People|{header}', 'People', header, letter, True)
                     for letter, header in [('A','Code'),('B','Age'),('C','Salary'),('D','Entered text')])
        types = module.classify_xlsx_columns(out.getvalue(), cols)
        self.assertEqual(types['People|Age']['kind'], 'formula')
        self.assertEqual(types['People|Age']['formula_cells'], 2)
        self.assertEqual(types['Age'], types['People|B'])
        self.assertEqual(types['People|Salary']['kind'], 'mixed')
        self.assertEqual(types['People|Entered text']['kind'], 'input')

    def test_sheet_qualified_aliases_do_not_confuse_equal_headings(self):
        book = openpyxl.Workbook()
        book.active.title = 'First'
        book.active.append(['ID', 'Age'])
        book.active.append(['TEST1', '=20+1'])
        second = book.create_sheet('Second')
        second.append(['ID', 'Age'])
        second.append(['TEST2', 30])
        out = io.BytesIO()
        book.save(out)
        cols = tuple((f'{sheet}|{header}', sheet, header, letter, True)
                     for sheet in ['First','Second'] for header,letter in [('ID','A'),('Age','B')])
        types = module.classify_xlsx_columns(out.getvalue(), cols)
        self.assertEqual(types['First|Age']['kind'], 'formula')
        self.assertEqual(types['Second|Age']['kind'], 'input')
        self.assertEqual(types['First|B']['kind'], 'formula')
        self.assertEqual(types['Second|B']['kind'], 'input')
        self.assertEqual(types['First|ID']['kind'], 'input')

    def test_generated_sequence_and_mixed_input_quality(self):
        book = openpyxl.Workbook()
        ws = book.active
        ws.title = 'People'
        headers = ['Sl No', 'Employee ID', 'Reporting Manager', 'Salary', 'Not a sequence']
        ws.append(headers)
        ws.append([1, 'TEST1', '=F2', 100, 1])
        ws.append(['=A2+1', '=#N/A', 'Entered manager', '=2500*2', '=E2+2'])
        ws.append(['=A3+1', None, '=F4', None, '=E3+1'])
        # Formatting/blank rows and a headingless global constant do not add
        # records to blank-cell counts.
        ws['J10'] = 26318.8
        ws['A20'].number_format = '0'
        out = io.BytesIO()
        book.save(out)
        cols = tuple((f'People|{header}', 'People', header,
                      openpyxl.utils.get_column_letter(i + 1), True)
                     for i, header in enumerate(headers))
        types = module.classify_xlsx_columns(out.getvalue(), cols)
        serial = types['Sl No']
        self.assertEqual(serial['kind'], 'formula')
        self.assertTrue(serial['generated_sequence'])
        self.assertEqual(types['People|A'], serial)
        self.assertEqual(types['Not a sequence']['kind'], 'mixed')
        for header in ['Employee ID', 'Reporting Manager', 'Salary']:
            self.assertEqual(types[header]['kind'], 'mixed')
            self.assertEqual(types[header]['data_rows'], 3)
        self.assertEqual(types['Employee ID']['error_cells'], 1)
        self.assertEqual(types['Employee ID']['error_addresses'], ['B3'])
        self.assertEqual(types['Employee ID']['blank_cells'], 1)
        self.assertEqual(types['Salary']['blank_cells'], 1)
        self.assertEqual(types['Reporting Manager']['blank_cells'], 0)

    def test_sequence_requires_a_single_seed_and_exact_previous_row_references(self):
        book = openpyxl.Workbook()
        ws = book.active
        ws.append(['Entered numbers', 'Wrong reference', 'Wrong seed', 'Gap'])
        ws.append([1, 1, 100, 1])
        ws.append([2, '=B1+1', '=C2+1', None])
        ws.append(['=A3+1', '=B3+1', '=C3+1', '=D2+1'])
        out = io.BytesIO()
        book.save(out)
        cols = tuple((header, '', header, openpyxl.utils.get_column_letter(i + 1), True)
                     for i, header in enumerate(['Entered numbers', 'Wrong reference', 'Wrong seed', 'Gap']))
        types = module.classify_xlsx_columns(out.getvalue(), cols)
        for header in ['Entered numbers', 'Wrong reference', 'Wrong seed', 'Gap']:
            self.assertEqual(types[header]['kind'], 'mixed')
            self.assertFalse(types[header]['generated_sequence'])


if __name__ == '__main__':
    unittest.main()

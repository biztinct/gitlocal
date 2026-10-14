"""Run in the Rize Odoo shell after importing the August formula workbook.

RIZE_SETUP_DIR must contain formula-template.xlsx, monthly-inputs.xlsx and
monthly-probe.xlsx (six fictional rows). Uses the application APIs, does not
create employees/payslips, and leaves the scheme and statutory editions draft.
"""
import base64
import io
import json
import math
import os
from pathlib import Path

import openpyxl

assert env.cr.dbname == 'rize', 'This installation belongs to the Rize tenant.'
root = Path(os.environ['RIZE_SETUP_DIR'])
config = env['hr.formula.config'].search([
    ('code', '=', 'RIZE_INDONESIA_AUGUST_2026'),
    ('company_id.name', '=', 'Rize Farms'),
])
assert len(config) == 1 and config.state in ('draft', 'testing')
assert config.country_code == 'ID' and config.currency_id.name == 'IDR'
assert len(config.rule_ids) == 367
by_col = {r.column_letter: r for r in config.rule_ids}
employee_rules = config.rule_ids.filtered(lambda r: not r.tax_table_id)
assert len(employee_rules) == 75
updates = {}
earning_cols = set('L M N O P Q R S T U V W X'.split())
deduction_cols = set('AC AD AE AF AG AH AI AJ AK AL AM AN AO'.split())
profile_cols = set('A B C D E F G H I J K AR'.split())
quantity_cols = {'AT', 'AU', 'BU'}
integer_cols = {'C', 'AX', 'AY', 'AZ', 'BA', 'BB', 'BE'}
text_cols = {'A', 'E', 'F', 'I', 'J', 'K', 'AR'}
for rule in config.rule_ids:
    col = rule.column_letter
    group, role, detail, kind = 'Calculation steps', 'info', False, 'money'
    if rule.tax_table_id:
        code = rule.tax_table_id.code
        group = {'IDPARAMS': 'Statutory parameters', 'IDANNUAL': 'Annual tax bands',
                 'IDSEVERANCE': 'Severance tax bands'}.get(code, 'TER %s bands' % code[-1])
        if rule.tax_field == 'rate' or (rule.tax_parameter_code or '').endswith('RATE') or rule.tax_parameter_code in {'JHTER', 'JHTEE', 'JPER', 'JPEE', 'HIER', 'HIEE'}:
            kind = 'rate'
        elif rule.tax_parameter_code == 'PREMMONTHS':
            kind = 'integer'
    else:
        if col in profile_cols:
            group = 'Employee details'
        elif col in earning_cols:
            group, role, detail = 'Earnings', 'earning', True
        elif col == 'Y':
            group, role = 'Severance', 'earning'
        elif col in {'Z', 'AP', 'AQ'}:
            group = 'Totals'
            role = {'Z': 'earning', 'AP': 'deduction', 'AQ': 'net'}[col]
        elif col in {'AA', 'AB'}:
            group, role, detail = 'Employer contributions', 'employer_cost', True
        elif col in deduction_cols:
            group, role, detail = 'Deductions', 'deduction', True
        elif col in {'AS', 'AT', 'AU', 'AV', 'AW', 'AX', 'AY', 'AZ', 'BA', 'BB', 'BU'}:
            group = 'Payroll inputs'
        elif col in {'BC', 'BD', 'BE', 'BF'}:
            group = 'Year to date inputs'
        elif col == 'BT':
            group, role = 'Employer cost', 'employer_cost'
        elif col in {'BV', 'BW'}:
            group = 'Severance history'
        if col in text_cols:
            kind = 'text'
        elif col in {'B', 'D'}:
            kind = 'identifier'
        elif col in {'G', 'H'}:
            kind = 'date'
        elif col in quantity_cols:
            kind = 'quantity'
        elif col in integer_cols:
            kind = 'integer'
        elif col == 'BR':
            kind = 'rate'
    updates[rule.code] = {'group': group, 'pay_role': role, 'rollup': detail, 'kind': kind, 'signal': False}
config.set_component_setup(updates)
# Keep these reviewed treatment decisions stable if the classifier is rerun.
config.rule_ids.write({
    'net_role_source': 'user', 'net_role_confidence': 'certain',
    'net_role_reason': 'Reviewed for the Rize Indonesia August 2026 installation.',
    'value_kind_source': 'user',
    'value_kind_reason': 'Reviewed against the monthly input field definition.',
})
for rule in config.rule_ids:
    show = rule.column_letter in earning_cols | deduction_cols | {'Y', 'Z', 'AA', 'AB', 'AP', 'AQ', 'BT'} and not rule.tax_table_id
    if rule.appears_on_payslip != show:
        rule.appears_on_payslip = show
    if rule.column_type == 'input':
        # The compact monthly workbook has different letters. Bind by heading.
        rule.set_source_binding('excel', 'Payroll|%s' % rule.name)
for col, role in {'C': 'reference', 'E': 'identity', 'G': 'contract', 'H': 'contract'}.items():
    by_col[col].write({'column_role': role, 'column_role_source': 'user'})

source = openpyxl.load_workbook(root / 'formula-template.xlsx', data_only=False)['Payroll']
cached = openpyxl.load_workbook(root / 'formula-template.xlsx', data_only=True)['Payroll']
input_cols = [c for c in range(1, 76) if source.cell(30, c).data_type != 'f']
formula_cols = [c for c in range(1, 76) if source.cell(30, c).data_type == 'f']
expected_headers = [source.cell(29, c).value for c in input_cols]
blank = openpyxl.load_workbook(root / 'monthly-inputs.xlsx', data_only=False)
assert blank.sheetnames == ['Payroll', 'Instructions']
assert [c.value for c in blank['Payroll'][1]] == expected_headers
assert not any(c.value is not None for row in blank['Payroll'].iter_rows(min_row=2) for c in row)
assert not any(c.data_type == 'f' for sheet in blank for row in sheet for c in row)
assert 'Tax Map' not in blank.sheetnames

Batch = env['hr.payroll.import.batch']
probe = Batch.new({'formula_config_id': config.id, 'company_id': config.company_id.id,
                   'file_header_row': 1, 'file_data_start_row': 2,
                   'file_sheet_name': 'Payroll'})
parsed = probe._parse_source_file((root / 'monthly-probe.xlsx').read_bytes(), 'monthly-probe.xlsx')
assert parsed['multisheet'] and len(parsed['rows']) == 6
formula_checks = 0
for i, raw in enumerate(parsed['rows']):
    provenance = {}
    inputs = probe._transform_data_to_formula_inputs(raw, provenance=provenance)
    for c in input_cols:
        name = source.cell(29, c).value
        rule = by_col[source.cell(30, c).column_letter]
        original = cached.cell(30 + i, c).value
        if original is not None:
            expected = '000001' if i == 0 and name == 'Employee ID' else original
            assert inputs.get(rule.code) == expected, (i, name, inputs.get(rule.code), expected)
    expected = {by_col[source.cell(30, c).column_letter].code: cached.cell(30+i, c).value for c in formula_cols}
    name = 'Template example %s' % (i+1)
    sample = config.sample_data_ids.filtered(lambda s: s.name == name)
    vals = {'config_id': config.id, 'name': name, 'sequence': i+1,
            'source_type': 'import', 'source_date': '2026-08-01',
            'description': cached.cell(30+i, 44).value,
            'input_values_json': json.dumps(inputs),
            'expected_inputs_json': json.dumps(inputs),
            'expected_values_json': json.dumps(expected)}
    if sample:
        sample.write(vals)
    else:
        sample = env['hr.formula.sample.data'].create(vals)
    errors = {}
    values = sample._evaluate_rules_with_dependencies(inputs, readonly=True, errors=errors)
    assert not errors, errors
    for code, value in expected.items():
        assert isinstance(values.get(code), (int, float)) and math.isclose(values[code], value, rel_tol=0, abs_tol=1e-6), (i, code, values.get(code), value)
        formula_checks += 1
    # The selected treatment must account for net pay exactly once.
    earned = sum(values.get(r.code, 0) for r in config.rule_ids if r.net_role == 'earning' and not r.net_role_detail)
    deducted = sum(values.get(r.code, 0) for r in config.rule_ids if r.net_role == 'deduction' and not r.net_role_detail)
    assert math.isclose(earned-deducted, values[by_col['AQ'].code], rel_tol=0, abs_tol=1e-6)
    costs = sum(values.get(r.code, 0) for r in config.rule_ids if r.net_role == 'employer_cost' and not r.net_role_detail)
    assert costs == values[by_col['BT'].code]

# The guided setup initially supplies a generic BASIC/STDDAYS example. Those
# codes are absent from this imported schema; retain it archived for history.
placeholder = config.sample_data_ids.filtered(lambda s: s.name == 'Sample employee'
    and not json.loads(s.expected_values_json or '{}')
    and 'BASIC' in json.loads(s.input_values_json or '{}'))
placeholder.write({'active': False})
test_results = config.run_sample_tests()
assert test_results['passed'] == 6 and not test_results['failed']
board = config.value_kind_board()
assert not board['review_count'] and not board['role_conflict_count']
assert len(config.rule_ids.filtered('tax_table_id')) == 292
assert all(config.rule_ids.filtered('tax_table_id').mapped('tax_aligned'))
health = env['pb.formula.studio'].get_problems(config.id)
assert health['counts']['error'] == 0 and health['counts']['warning'] == 0, health['problems']
attachment = env['ir.attachment'].search([('res_model', '=', 'hr.formula.config'),
    ('res_id', '=', config.id), ('name', '=', 'Rize_Indonesia_Monthly_Inputs.xlsx')], limit=1)
attachment_vals = {'name': 'Rize_Indonesia_Monthly_Inputs.xlsx', 'type': 'binary',
    'res_model': 'hr.formula.config', 'res_id': config.id,
    'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'datas': base64.b64encode((root / 'monthly-inputs.xlsx').read_bytes())}
if attachment:
    attachment.write(attachment_vals)
else:
    attachment = env['ir.attachment'].create(attachment_vals)
env.cr.commit()
print(json.dumps({'config_id': config.id, 'state': config.state,
    'components': len(config.rule_ids), 'inputs': len(input_cols),
    'mapped_constants': 292, 'monthly_rows_tested': 6,
    'formula_checks': formula_checks, 'tests': test_results,
    'health': health['counts'], 'treatment_review': board['review_count'],
    'type_conflicts': board['role_conflict_count'], 'input_attachment': attachment.id}))

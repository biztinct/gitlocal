#!/usr/bin/env python3
"""Read a decrypted client workbook and emit local, anonymous test controls.

Never emit employee names, government IDs, bank details, or an opening
password. The JSON is temporary analysis input, not a repository fixture.
"""
import argparse
import json
from pathlib import Path

import openpyxl


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('source')
    parser.add_argument('output')
    args = parser.parse_args()
    w = openpyxl.load_workbook(args.source, data_only=True)
    f = openpyxl.load_workbook(args.source, data_only=False)
    s, t = w['0826 Payroll Report'], w['TER A B C']
    bands = {key: [[t.cell(r, c + 2).value or 0, t.cell(r, c + 3).value]
                   for r in range(10, 54) if isinstance(t.cell(r, c).value, (int, float))]
             for key, c in [('A', 2), ('B', 11), ('C', 20)]}
    p = w['Premi Insurance']
    premiums = {p[f'B{r}'].value: p[f'H{r}'].value for r in range(2, p.max_row + 1)
                if isinstance(p[f'B{r}'].value, (int, float))}
    calc = w['Calculation']
    controls = []
    for r in range(9, 131):
        employee = s[f'D{r}'].value
        vals = {c: s[f'{c}{r}'].value or 0 for c in
                ['L', 'M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U', 'V', 'W', 'X', 'Y',
                 'Z', 'AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ',
                 'AK', 'AL', 'AM', 'AN', 'AO', 'AP', 'AQ']}
        inputs = {c: vals[c] for c in ['P', 'Q', 'R', 'S', 'X', 'Y', 'AJ', 'AK', 'AL', 'AM']}
        inputs.update(J=s[f'J{r}'].value, AS=vals['L'], AT=31, AU=31, AV=0,
                      AW=premiums.get(employee, 0), AX=int(vals['T'] > 0),
                      AY=int(vals['V'] > 0), AZ=int(vals['N'] > 0), BA=2 if vals['N'] else 1,
                      BB=0, BC=0, BD=0, BE=8, BF=0, BV=0)
        for cr in [6, 7]:
            if calc[f'A{cr}'].value == employee:
                inputs.update(AS=calc[f'H{cr}'].value, AT=calc[f'G{cr}'].value)
        for cr in range(19, 28):
            if calc[f'A{cr}'].value == employee and isinstance(calc[f'G{cr}'].value, (int, float)):
                inputs.update(AS=calc[f'G{cr}'].value, AV=calc[f'H{cr}'].value or 0)
        if calc['A13'].value == employee:
            inputs.update(AS=calc['H13'].value, AT=calc['F13'].value, BA=0, BB=1,
                          BC=51061356, BD=1319686, BE=4, BF=3925731)
        controls.append({'source_row': r, 'inputs': inputs, 'expected': vals})
    inventory = [{'sheet': sheet.title, 'formula_count': sum(c.data_type == 'f' for row in sheet for c in row)} for sheet in f]
    Path(args.output).write_text(json.dumps({'bands': bands, 'controls': controls, 'inventory': inventory}))
    print(json.dumps({'employees': len(controls), 'TER_bands': {k: len(v) for k, v in bands.items()}, 'inventory': inventory}))


if __name__ == '__main__':
    main()

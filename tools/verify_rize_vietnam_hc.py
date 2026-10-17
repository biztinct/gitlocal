"""Verify deployed VN HC definition, attach delivery files and test XLSX parsing.

Runs through the Rize Odoo shell. Does not import payroll or write people data.
"""
import base64
import hashlib
import json
import os
from pathlib import Path
import openpyxl

assert env.cr.dbname=='rize'
root=Path(os.environ['VN_HC_DIR'])
d=json.loads((root/'definition.json').read_text())
receipt=json.loads((root/'install-receipt.json').read_text())
cfg=env['hr.formula.config'].browse(receipt['config_id']).exists()
assert cfg.code==d['code'] and cfg.state=='draft'
by={r.code:r for r in cfg.rule_ids}
assert set(by)=={r['code'] for r in d['rules']}
for r in d['rules']:
    actual=by[r['code']]
    assert actual.value_kind==r['kind'],r['code']
    assert actual.component_type==r['group'] and actual.net_role==r['role'] and actual.net_role_detail==r['detail'],r['code']
    if r['type']=='input':
        assert actual.source_ids.filtered(lambda s:s.kind=='excel').key==(r.get('source_key') or 'Recurring settings|'+r['name'])
    if r['formula']:
        assert actual.excel_formula=='='+r['formula'].format(**{k:v.column_letter+'2' for k,v in by.items()})
tests=cfg.run_sample_tests()
assert tests['passed']==30 and tests['failed']==0,tests
facade=env['pb.payroll.templates'].with_company(cfg.company_id)
w=facade.workspace(cfg.id,'2026-10')
assert w['people']==0
assert len(w['layout']['time'])==4 and len(w['layout']['adjustments'])==13
blank=openpyxl.load_workbook(root/'Vietnam_Monthly_Input_Template.xlsx',data_only=False)
assert blank.sheetnames==['Time','Monthly adjustments','Instructions']
for tab in ['Time','Monthly adjustments']:
    expected=['Country employee ID','Employee name']+[r['name'] for r in d['rules'] if r['tab']==tab]
    assert [c.value for c in blank[tab][1]]==expected
    assert all(c.value is None for row in blank[tab].iter_rows(min_row=2) for c in row)
probe=env['hr.payroll.import.batch'].with_company(cfg.company_id).new(dict(formula_config_id=cfg.id,company_id=cfg.company_id.id,file_header_row=1,file_data_start_row=2))
parsed=probe._parse_source_file((root/'monthly-probe.xlsx').read_bytes(),'monthly-probe.xlsx')
assert len(parsed['rows'])==1,(len(parsed['rows']),parsed.keys())
raw=parsed['rows'][0]
transformed=probe._transform_data_to_formula_inputs(raw)
assert transformed['VNSTDDAYS']==22 and transformed['VNPAIDDAYS']==22 and transformed['VNACTBONUS']==1000000,transformed
assert str(transformed['VNHCE'])=='VNTEST001',transformed
attachment_ids=[]
for name in ['Vietnam_Configuration_Review.xlsx','Vietnam_Monthly_Input_Template.xlsx']:
    path=root/name
    vals=dict(name=name,type='binary',res_model='hr.formula.config',res_id=cfg.id,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',datas=base64.b64encode(path.read_bytes()))
    a=env['ir.attachment'].search([('res_model','=','hr.formula.config'),('res_id','=',cfg.id),('name','=',name)])
    assert len(a)<=1
    if a: a.write(vals)
    else: a=env['ir.attachment'].create(vals)
    assert hashlib.sha256(base64.b64decode(a.datas)).hexdigest()==hashlib.sha256(path.read_bytes()).hexdigest()
    attachment_ids.append(a.id)
counts={m:env[m].with_context(active_test=False).search_count([]) for m in receipt['after']}
assert counts==receipt['after'],(counts,receipt['after'])
result=dict(database=env.cr.dbname,config_id=cfg.id,state=cfg.state,tests=tests,parsed_rows=1,monthly_inputs=17,layout_time=4,layout_adjustments=13,attachments=attachment_ids,people=w['people'],business_counts=counts)
(root/'verification-receipt.json').write_text(json.dumps(result,indent=2))
env.cr.commit()
print('VN_HC_VERIFIED',json.dumps(result))

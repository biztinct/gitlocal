# Builds "Rize HRMS Must Have - Traceability Matrix.xlsx" from RIZE/HRMS Proposal - Must Have.xlsx.
# Same look as build_traceability_matrix.py (the Recruit one). Row answers come from JSON files
# (one list of {tab,row,supported,solution,try_as,clicks,comments} per research batch).
# Client-facing: plain words, screen names, no internal codes, never the word "Odoo".
import glob
import json
import os
import sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SRC = '/Users/adity/Documents/GitHub/gitlocal/RIZE/HRMS Proposal - Must Have.xlsx'
OUT = '/Users/adity/Documents/GitHub/gitlocal/RIZE/Rize HRMS Must Have - Traceability Matrix.xlsx'
HERE = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'musthave_data')
DATA = sys.argv[1:] or sorted(glob.glob(os.path.join(HERE, 'tm_[A-Z].json')))
RECRUIT_TAB = 'Hiring & Recruitment'
RECRUIT_FILE = 'Rize Recruit - Traceability Matrix.xlsx'

ANS = {}
for f in DATA:
    for o in json.load(open(f)):
        ANS[(o['tab'], int(o['row']))] = o

WRAP = Alignment(wrap_text=True, vertical='top')
FILLS = {'Yes': 'E3EFE7', 'Partly': 'FBEFD5', 'No': 'F6E0DD', 'Ready, switched off': 'EEF0F3',
         'Rize action': 'EEF0F3'}
HEAD_FILL = PatternFill('solid', fgColor='3F7D5C')
HEAD_FONT = Font(name='Arial', size=9, bold=True, color='FFFFFF')
BODY = Font(name='Arial', size=9)
THIN = Side(style='thin', color='C9D1CA')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NEW = [('Supported?', 14, 'supported'), ('Solution type', 20, 'solution'), ('Try it as', 20, 'try_as'),
       ('Clicks to reach it', 70, 'clicks'), ('Comments', 50, 'comments')]

def head(ws, row, col, text, width):
    c = ws.cell(row=row, column=col, value=text)
    c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical='center'), BOX
    ws.column_dimensions[get_column_letter(col)].width = width

def put(ws, row, col, val, tint=False):
    c = ws.cell(row=row, column=col, value=val)
    c.alignment, c.border, c.font = WRAP, BOX, BODY
    if tint and val in FILLS:
        c.fill = PatternFill('solid', fgColor=FILLS[val])

def fit(ws, r):
    n = 2
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        if isinstance(v, str):
            w = ws.column_dimensions[get_column_letter(c)].width or 10
            n = max(n, sum(1 + len(l) // int(w * 1.15) for l in v.split('\n')))
    ws.row_dimensions[r].height = 12.5 * n + 4

wb = openpyxl.load_workbook(SRC)
seen, counts = set(), {}
for ws in wb.worksheets:
    # split grouped column widths (e.g. D:H) so each new column keeps its own width
    for key, dim in list(ws.column_dimensions.items()):
        if dim.min and dim.max and dim.max > dim.min:
            w = dim.width
            for i in range(dim.min, dim.max + 1):
                L = get_column_letter(i)
                ws.column_dimensions[L].width = w
                ws.column_dimensions[L].min = ws.column_dimensions[L].max = i
    if ws.title == RECRUIT_TAB:
        put(ws, 1, 1, f'Hiring is traced line by line in its own workbook: “{RECRUIT_FILE}”.')
        ws.cell(1, 1).font = Font(name='Arial', size=10, bold=True, color='3F7D5C')
        ws.cell(1, 1).border = Border()
        continue
    # header rows: column A starts with "Module" (some tabs hold two tables)
    heads = [r for r in range(1, ws.max_row + 1)
             if str(ws.cell(r, 1).value or '').startswith('Module')]
    for hr in heads:
        start = 5 if ws.cell(hr, 4).value else 4
        # Onboarding keeps its 4th "Covered in Probation Form" column; label it
        if ws.cell(hr, 4).value is None and any(ws.cell(r, 4).value for r in range(hr + 1, ws.max_row + 1)):
            head(ws, hr, 4, 'Note', 22)
            start = 5
        for i, (t, w, _) in enumerate(NEW):
            head(ws, hr, start + i, t, w)
        r = hr + 1
        while r <= ws.max_row and ws.cell(r, 1).value not in (None, '') \
                and not str(ws.cell(r, 1).value).startswith('Module'):
            if ws.cell(r, 2).value in (None, '') and ws.cell(r, 3).value in (None, ''):
                break     # a sub-title such as "Role-Based Access & Delegation"
            o = ANS.get((ws.title, r))
            assert o, f'no answer for {ws.title} row {r}'
            seen.add((ws.title, r))
            counts[o['supported']] = counts.get(o['supported'], 0) + 1
            for i, (_, _, k) in enumerate(NEW):
                put(ws, r, start + i, o[k], tint=(i == 0))
            for c in range(1, start):
                ws.cell(r, c).alignment = WRAP
            fit(ws, r)
            r += 1
    ws.freeze_panes = None
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0

extra = set(ANS) - seen
assert not extra, f'answers with no matching row: {sorted(extra)}'

# ---------------------------------------------------------------- start sheet
st = wb.create_sheet('Start here', 0)
for L, w in zip('ABCD', (26, 36, 18, 60)):
    st.column_dimensions[L].width = w
r = 1
def line(a, b=None, c=None, d=None, bold=False, size=None, h=None, merge=False):
    global r
    for i, v in enumerate((a, b, c, d), 1):
        if v is not None:
            cell = st.cell(r, i, v); cell.alignment = WRAP
            cell.font = Font(name='Arial', bold=bold, size=size or 10)
    if h: st.row_dimensions[r].height = h
    if merge: st.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    r += 1

line('Rize HRMS Must Have — try every requirement yourself', bold=True, size=16, h=26)
st.merge_cells('A1:D1')
line('This workbook is the "HRMS Proposal - Must Have" list with five columns added to every tab: whether Payobook '
     'supports the line, how, which sign-in to use, the exact clicks to reach it, and comments. '
     f'The Hiring & Recruitment tab is traced in its own workbook, “{RECRUIT_FILE}”.', h=48)
st.merge_cells(start_row=r - 1, start_column=1, end_row=r - 1, end_column=4)
r += 1
line('1. Where to try it', bold=True, size=13)
line('Address', 'https://payobook.com/web/login', merge=True)
line('The data', 'Made-up demo data. People and records start with DEMO. Nothing you do here reaches a real person.', h=32, merge=True)
r += 1
line('2. Demo sign-ins', bold=True, size=13)
line('Sign in as', 'Email', 'Password', 'What you can try')
for c in range(1, 5):
    st.cell(r - 1, c).fill, st.cell(r - 1, c).font = HEAD_FILL, HEAD_FONT
for row in json.load(open(os.path.join(HERE, 'tm_logins.json'))):
    line(*row, h=44 if len(row[3]) > 110 else 30)
line('Tip', 'Use a private (incognito) window for each sign-in to compare two people side by side.', merge=True)
r += 1
line('3. The added columns', bold=True, size=13)
for a, b in [('Supported?', 'Yes, Partly, No, Ready, switched off (built, but its emails or switch are off on the demo), '
                            'or Rize action (needs something only Rize can give, such as real dates or a vendor login).'),
             ('Solution type', 'Built for Rize in Payobook, Already in Payobook, Built with the rest later, or Not built.'),
             ('Try it as', 'Which demo sign-in to use. “Automatic” means nobody presses anything: the system does it on a schedule.'),
             ('Clicks to reach it', 'Numbered steps from signing in to the screen. Words in “quotes” are the exact words on the screen.'),
             ('Comments', 'What to look for, what is switched off on the demo, and anything that waits on Rize.')]:
    line(a, b, h=30, merge=True)
r += 1
line('4. Summary', bold=True, size=13)
for k in ('Yes', 'Partly', 'Ready, switched off', 'Rize action', 'No'):
    if counts.get(k):
        line(k, f'{counts[k]} lines', merge=True)
        st.cell(r - 1, 1).fill = PatternFill('solid', fgColor=FILLS[k])
st.page_setup.orientation = 'landscape'
st.sheet_properties.pageSetUpPr.fitToPage = True
st.page_setup.fitToWidth, st.page_setup.fitToHeight = 1, 0
wb.active = 0
wb.save(OUT)

bad = []
for w in openpyxl.load_workbook(OUT).worksheets:
    for row in w.iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str) and 'odoo' in v.lower():
                bad.append((w.title, v[:60]))
print('saved', OUT, counts, 'odoo hits:', bad)

# -*- coding: utf-8 -*-
"""TIDY P3 — a month is a scope, equal to the year (ledger rule 13).

Every test here asks the same question from a different side: does the board
still ADD UP when the thing it is about is one month rather than twelve. The
fixtures are made and rolled back inside the test's own transaction, on a
department this database already has, so the suite says the same thing on a
customer's data as on the demo company's.

The frozen "today" matters more than it looks: `past`, `current` and `future`
are the whole of the difference between "under budget" and "has not started",
and a test that reads the wall clock passes in September and fails in January.
"""
from datetime import date

from odoo.tests.common import TransactionCase, tagged

#: The year these tests build. Far enough out that no live database has budget
#: rows in it, so the figures asserted are exactly the ones created here.
FY = 2031
#: The frozen calendar. June is halfway: five finished months behind it, six
#: months that have not started ahead of it, and itself in progress.
TODAY = date(FY, 6, 11)


@tagged('post_install', '-at_install')
class TestBudgetMonth(TransactionCase):

    def setUp(self):
        super().setUp()
        self.Line = self.env['pb.budget.line']
        # THE BOARD READS `env.companies`, NOT `env.company`, and on more than
        # one database here the two do not agree (the reader's own company is
        # not in the set they are entitled to). A fixture built in `env.company`
        # is then invisible to the very facade it exists to test — and the test
        # fails with an empty board and no clue why. Build in a company the
        # board can actually see.
        self.company = None
        self.dept = self.dept2 = None
        for company in self.env.companies:
            depts = self.env['hr.department'].search(
                [('company_id', '=', company.id)], limit=2)
            if depts:
                self.company = company
                self.dept = depts[0]
                self.dept2 = depts[1] if len(depts) > 1 else depts[0]
                break
        if not self.dept:
            self.skipTest('no department in a company this reader can see')

    # ------------------------------------------------------------- fixtures
    def _row(self, month, budget, spent, dept=None):
        return self.Line.create({
            'company_id': self.company.id,
            'department_id': (dept or self.dept).id,
            'period_month': date(FY, month, 1),
            'pb_budget_type': 'admin',
            'forecast_cost': budget,
            'actual_cost': spent,
        })

    def _board(self, month=None):
        return self.env['pb.budget'].get_board(FY, 'admin', 'local', None,
                                               month)

    # ================================================================== T1
    def test_t1_a_month_totals_only_that_months_rows(self):
        """The function's figures ARE that month's lines, the strip's twelve
        entries add up to the year, and the state of a month is decided by the
        calendar rather than by whether anybody spent anything."""
        self._row(3, 1000.0, 900.0)
        self._row(4, 1000.0, 1400.0)
        self._row(3, 500.0, 500.0, self.dept2)

        year = self._board()
        march = self._board('%s-03' % FY)

        # Whatever function these departments roll into, MARCH's figures are
        # the March rows and only those.
        self.assertTrue(march['functions'])
        m_budget = sum(f['budget'] for f in march['functions'])
        m_spent = sum(f['spent'] for f in march['functions'])
        self.assertEqual(round(m_budget - self._year_only(year, 3, 'budget'), 2),
                         0.0)
        self.assertEqual(round(m_spent - self._year_only(year, 3, 'spent'), 2),
                         0.0)

        # The strip is twelve entries and it is the WHOLE year, whichever
        # month is in scope.
        for board in (year, march):
            self.assertEqual(len(board['strip']), 12)
            self.assertAlmostEqual(
                round(sum(s['spent'] for s in board['strip']), 2),
                round(sum(f['year_spent'] for f in board['functions']), 2),
                places=2)
            self.assertAlmostEqual(
                round(sum(s['budget'] for s in board['strip']), 2),
                round(sum(f['year_budget'] for f in board['functions']), 2),
                places=2)

        # The year board's own total is unchanged by any of this.
        self.assertEqual(year['scope']['kind'], 'year')
        self.assertEqual(march['scope']['kind'], 'month')
        self.assertEqual(march['scope']['key'], '%s-03' % FY)

    def test_t1b_the_state_of_a_month_follows_the_calendar(self):
        facade = self.env['pb.budget']
        self.assertEqual(facade._month_state('%s-03' % FY, TODAY), 'past')
        self.assertEqual(facade._month_state('%s-06' % FY, TODAY), 'current')
        self.assertEqual(facade._month_state('%s-09' % FY, TODAY), 'future')
        # And the pace with it: finished, part way, not started.
        self.assertEqual(facade._month_pace('%s-03' % FY, TODAY), 100.0)
        self.assertEqual(facade._month_pace('%s-09' % FY, TODAY), 0.0)
        part = facade._month_pace('%s-06' % FY, TODAY)
        self.assertTrue(0.0 < part < 100.0, part)
        self.assertEqual(part, round(10 / 30 * 100, 1))

    def test_t1c_a_month_outside_the_year_is_the_whole_year(self):
        """A stale bookmark lands somewhere real (zero dead-ends)."""
        facade = self.env['pb.budget']
        self.assertEqual(facade._month_in_fy('%s-03' % FY, FY), '%s-03' % FY)
        self.assertEqual(facade._month_in_fy('%s-03' % (FY - 1), FY), '')
        self.assertEqual(facade._month_in_fy('rubbish', FY), '')
        self.assertEqual(facade._month_in_fy(None, FY), '')
        board = self._board('%s-03' % (FY - 1))
        self.assertEqual(board['scope']['kind'], 'year')

    def _year_only(self, year_board, month, key):
        want = '%s-%02d' % (FY, month)
        return round(sum(m[key] for f in year_board['functions']
                         for m in f['months'] if m['key'] == want), 2)

    # ================================================================== T2
    def test_t2_the_month_words_and_where_they_change(self):
        """±5% either side of the budget, and "Not yet" ONLY for a month that
        has not started with nothing spent on it."""
        facade = self.env['pb.budget']
        tone = facade._tone_month

        # Exactly on the budget, and both edges of the band, are all "close".
        self.assertEqual(tone({'budget': 1000.0, 'spent': 1000.0}, 'past'),
                         'onpace')
        self.assertEqual(tone({'budget': 1000.0, 'spent': 1050.0}, 'past'),
                         'onpace')
        self.assertEqual(tone({'budget': 1000.0, 'spent': 950.0}, 'past'),
                         'onpace')
        # One dong past either edge is not.
        self.assertEqual(tone({'budget': 1000.0, 'spent': 1050.01}, 'past'),
                         'over')
        self.assertEqual(tone({'budget': 1000.0, 'spent': 949.99}, 'past'),
                         'calm')
        # No budget is its own answer, spent or not.
        self.assertEqual(tone({'budget': 0.0, 'spent': 400.0}, 'past'), 'none')
        self.assertEqual(tone({'budget': 0.0, 'spent': 0.0}, 'past'), 'none')
        # "Not yet" is a month that has not started AND has nothing on it. A
        # future month that HAS been spent on is judged like any other.
        self.assertEqual(tone({'budget': 1000.0, 'spent': 0.0}, 'future'),
                         'notyet')
        self.assertEqual(tone({'budget': 0.0, 'spent': 0.0}, 'future'),
                         'notyet')
        self.assertEqual(tone({'budget': 1000.0, 'spent': 2000.0}, 'future'),
                         'over')
        # Every word is a word, and none of them is the year's.
        for key in ('over', 'onpace', 'calm', 'none', 'notyet'):
            label = facade._tone_label_month({'tone': key})
            self.assertTrue(label, key)
            self.assertNotIn('year', str(label).lower())

    # ================================================================== T3
    def test_t3_every_headline_form_is_reachable(self):
        """Five sentences, and each one produced at least once."""
        facade = self.env['pb.budget']
        base = {
            'currency': {'code': 'VND'},
            'months': [],
        }

        def say(scope_state, functions, kpis, name='March'):
            payload = dict(base)
            payload['functions'] = functions
            payload['kpis'] = kpis
            payload['scope'] = {'kind': 'month', 'key': '%s-03' % FY,
                                'label': '%s %s' % (name, FY), 'name': name,
                                'state': scope_state}
            return str(facade._headline_month(payload))

        over = {'id': 1, 'name': 'Retail', 'tone': 'over', 'budget': 100.0,
                'spent': 130.0, 'variance': 30.0, 'variance_pct': 30.0}
        under = {'id': 2, 'name': 'Logistics', 'tone': 'calm', 'budget': 100.0,
                 'spent': 60.0, 'variance': -40.0, 'variance_pct': -40.0}

        future = say('future', [], {'spent': 0.0, 'budget': 0.0, 'left': 0.0,
                                    'burn': 0.0, 'pace': 0.0, 'functions': 0})
        self.assertIn('has not started', future)

        none = say('past', [under],
                   {'spent': 500.0, 'budget': 0.0, 'left': -500.0, 'burn': 0.0,
                    'pace': 100.0, 'functions': 1})
        self.assertIn('no budget set', none)

        current = say('current', [under],
                      {'spent': 62.0, 'budget': 100.0, 'left': 38.0,
                       'burn': 62.0, 'pace': 71.0, 'functions': 1})
        self.assertIn('so far', current)
        self.assertIn('62', current)
        self.assertIn('71', current)

        hot = say('past', [over, under],
                  {'spent': 190.0, 'budget': 200.0, 'left': 10.0, 'burn': 95.0,
                   'pace': 100.0, 'functions': 2})
        self.assertIn('over budget', hot)
        self.assertIn('Retail', hot)

        cool = say('past', [under],
                   {'spent': 60.0, 'budget': 100.0, 'left': 40.0, 'burn': 60.0,
                    'pace': 100.0, 'functions': 1})
        self.assertIn('under budget', cool)

        # Five DIFFERENT sentences, not one sentence five times.
        self.assertEqual(len({future, none, current, hot, cool}), 5)
        # And none of them is the year's.
        for line in (future, none, current, hot, cool):
            self.assertNotIn('of the year', line)

    def test_t3b_the_short_money_matches_what_a_tile_prints(self):
        facade = self.env['pb.budget']
        self.assertEqual(facade._short(12_400_000_000), '12.4bn')
        self.assertEqual(facade._short(8_100_000_000), '8.1bn')
        self.assertEqual(facade._short(2_500_000), '2.5m')
        self.assertEqual(facade._short(4_000), '4k')
        self.assertEqual(facade._short(-12_400_000_000), '-12.4bn')

    # ================================================================== T4
    def test_t4_the_drill_of_a_month_adds_up_and_says_what_it_cannot_know(self):
        line = self._row(3, 1000.0, 900.0)
        self._row(3, 500.0, 620.0, self.dept2)
        self._row(2, 1000.0, 700.0)
        fid = line.pb_function_id.id or line.department_id.id or 0

        board = self._board('%s-03' % FY)
        func = next((f for f in board['functions'] if f['id'] == fid), None)
        self.assertTrue(func, 'the fixture should have made a function')

        drill = self.env['pb.budget'].get_function(
            func['id'], FY, 'admin', 'local', '%s-03' % FY)
        self.assertTrue(drill['ok'])
        self.assertEqual(drill['scope']['kind'], 'month')

        # The departments inside it sum to the function's own month figure.
        self.assertAlmostEqual(
            round(sum(d['spent'] for d in drill['function']['departments']), 2),
            drill['function']['spent'], places=2)
        self.assertAlmostEqual(
            round(sum(d['budget'] for d in drill['function']['departments']), 2),
            drill['function']['budget'], places=2)

        cmp = drill['compare']
        for key in ('this', 'last', 'last_year', 'average'):
            self.assertIn(key, cmp)
        self.assertTrue(cmp['this']['has'])
        self.assertEqual(cmp['this']['spent'], drill['function']['spent'])
        # There is no year before this one, so the year-ago cell is BLANK
        # rather than a zero that would read as a collapse.
        self.assertFalse(cmp['last_year']['has'])
        self.assertEqual(cmp['last_year']['spent'], 0.0)

        # And a month with a month before it finds it.
        self.assertTrue(cmp['last']['has'])
        self.assertEqual(cmp['last']['key'], '%s-02' % FY)

        # The twelve bars are still the whole year.
        self.assertEqual(len(drill['function']['months']), 12)

    # ================================================================== T5
    def test_t5_the_spreadsheet_of_a_month_is_that_month(self):
        self._row(3, 1000.0, 900.0)
        self._row(4, 1000.0, 1400.0)
        board = self._board('%s-03' % FY)
        try:
            import openpyxl                                    # noqa: F401
        except ImportError:
            self.skipTest('no spreadsheet library on this database')

        import base64
        import io
        import openpyxl

        res = self.env['pb.budget.export'].build(
            FY, 'admin', 'local', 'xlsx', '%s-03' % FY)
        self.assertTrue(res['ok'])
        book = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(res['file_b64'])))
        sheet = book.active
        # The sheet is named after what is ON it.
        self.assertIn(str(FY), sheet.title)
        self.assertIn(board['scope']['label'][:12], sheet.title)
        self.assertIn(board['scope']['label'], res['filename'])

        # And its total row is the board's total, to the digit.
        spent = round(sum(f['spent'] for f in board['functions']), 2)
        found = []
        for row in sheet.iter_rows(values_only=True):
            if row and row[0] and str(row[0]).strip() in ('Total', 'Tổng'):
                found.append(row)
        self.assertTrue(found, 'the sheet must carry a total row')
        self.assertAlmostEqual(float(found[-1][3] or 0.0), spent, places=0)

        # A month's sheet drops the twenty-four month columns.
        head = [c for c in next(sheet.iter_rows(min_row=4, max_row=4,
                                                values_only=True)) if c]
        self.assertLessEqual(len(head), 8, head)
        self.assertTrue(any('%' in str(h) for h in head))

    # ================================================================== T6
    def test_t6_the_year_board_is_untouched_by_any_of_this(self):
        """The regression that would matter: a month scope must not have
        changed what the year says."""
        board = self.env['pb.budget'].get_board()
        self.assertTrue(board['ok'])
        self.assertEqual(board['scope']['kind'], 'year')
        self.assertEqual(len(board['months']), 12)
        self.assertEqual(len(board['strip']), 12)
        for f in board['functions']:
            # In year scope the year figures and the scope figures are one and
            # the same, or something has quietly started filtering.
            self.assertEqual(f['budget'], f['year_budget'])
            self.assertEqual(f['spent'], f['year_spent'])
            self.assertIn(f['tone'],
                          ('over', 'watch', 'onpace', 'calm', 'none'))

    def test_t6b_the_strip_and_the_tiles_never_disagree(self):
        self._row(3, 1000.0, 900.0)
        self._row(7, 2000.0, 2600.0)
        board = self._board()
        for cell in board['strip']:
            wanted = round(sum(m['spent'] for f in board['functions']
                               for m in f['months'] if m['key'] == cell['key']),
                           2)
            self.assertAlmostEqual(cell['spent'], wanted, places=2)

    def test_t6c_the_month_scope_names_this_product_and_no_other(self):
        """White-label, on the strings this phase adds."""
        board = self._board('%s-06' % FY)
        blob = ' '.join([
            str(board['headline']),
            ' '.join(str(s['tone_label']) for s in board['strip']),
            ' '.join(str(f['tone_label']) for f in board['functions']),
            str(board['scope']['label']), str(board['scope']['name']),
        ])
        self.assertNotIn('odoo', blob.lower())

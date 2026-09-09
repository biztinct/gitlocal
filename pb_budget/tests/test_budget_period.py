# -*- coding: utf-8 -*-
"""LOOK P3 — a stretch of months is a scope, equal to the year and to a month.

THE FIRST TEST IN THIS FILE IS THE POINT OF THE FILE. TIDY P3 shipped the month
strip and it was validated on real data eight days before this phase started;
this phase then replaced the single month key that threaded through `_matrix`,
`_strip`, `_headline`, `_tone`, `_expenses` and `_rows` with an ordered LIST of
month keys, so that a month, a quarter, "March to June" and the whole year are
one code path with a different list. That refactor is the hot path of a screen
somebody is using. T1 proves it moved nothing:

  * the year board and each of the twelve month boards are what an independent
    reading of `pb.budget.line` says they should be, function by function and
    month by month;
  * a stretch of ONE month is byte-identical to that month asked for directly;
  * a stretch of TWELVE is byte-identical to the year asked for directly.

The last two are the collapse rules of R1 stated as equality of the whole
payload, which is the strongest form the assertion has: there is no second,
subtly different board hiding behind the same words.

The frozen "today" matters more than it looks: `past`, `current` and `future`
are the whole of the difference between "under budget" and "has not started",
and a test that reads the wall clock passes in September and fails in January.
"""
import json
from datetime import date

from odoo.tests.common import TransactionCase, tagged

#: The year these tests build. Far enough out that no live database has budget
#: rows in it, so the figures asserted are exactly the ones created here.
FY = 2032
#: The frozen calendar. June is halfway: five finished months behind it, six
#: months that have not started ahead of it, and itself in progress.
TODAY = date(FY, 6, 11)


@tagged('post_install', '-at_install')
class TestBudgetPeriod(TransactionCase):

    def setUp(self):
        super().setUp()
        self.Budget = self.env['pb.budget']
        self.Line = self.env['pb.budget.line']
        # T22 — the board reads `env.companies`, not `env.company`, and on more
        # than one database here the two do not agree. A fixture built in the
        # wrong one is invisible to the very facade it exists to test.
        self.company = None
        self.dept = self.dept2 = None
        for company in self.env.companies:
            depts = self.env['hr.department'].search(
                [('company_id', '=', company.id)], limit=2)
            if depts:
                self.company = company
                self.dept = depts[0]
                self.dept2 = depts[1] if len(depts) > 1 else depts[0]
                break
        if not self.dept:
            self.skipTest('no department in a company this reader can see')

    # ------------------------------------------------------------- fixtures
    def _row(self, month, budget, spent, dept=None):
        return self.Line.create({
            'company_id': self.company.id,
            'department_id': (dept or self.dept).id,
            'period_month': date(FY, month, 1),
            'pb_budget_type': 'admin',
            'forecast_cost': budget,
            'actual_cost': spent,
        })

    def _twelve(self):
        """One row a month on one department, one more on another in three of
        them, and one month deliberately empty — so the fixture has a shape
        rather than a straight line."""
        for m in range(1, 13):
            if m == 8:                       # August: nothing at all
                continue
            self._row(m, 1000.0 * m, 900.0 * m)
        self._row(3, 500.0, 620.0, self.dept2)
        self._row(6, 400.0, 400.0, self.dept2)
        self._row(11, 300.0, 100.0, self.dept2)

    def _board(self, period=None):
        return self.Budget.get_board(FY, 'admin', 'local', None, period)

    def _key(self, m):
        return '%s-%02d' % (FY, m)

    def _blob(self, board):
        return json.dumps(board, sort_keys=True, default=str)

    # ================================================================== T1
    def test_t1_the_refactor_moved_nothing(self):
        """R3, as a test and not a paragraph."""
        self._twelve()
        year = self._board()

        # --- the year adds up to the rows, read independently -------------
        rows = self.Line.search([
            ('pb_budget_type', '=', 'admin'),
            ('period_month', '>=', date(FY, 1, 1)),
            ('period_month', '<=', date(FY, 12, 1)),
            ('company_id', 'in', self.env.companies.ids),
        ])
        want_budget = round(sum(r.forecast_cost for r in rows), 2)
        want_spent = round(sum(r.actual_cost for r in rows), 2)
        self.assertAlmostEqual(
            round(sum(f['budget'] for f in year['functions']), 2),
            want_budget, places=2)
        self.assertAlmostEqual(
            round(sum(f['spent'] for f in year['functions']), 2),
            want_spent, places=2)
        self.assertEqual(year['scope']['kind'], 'year')
        self.assertEqual(len(year['scope']['keys']), 12)

        # --- each of the twelve months is its own rows, and only those -----
        for m in range(1, 13):
            key = self._key(m)
            board = self._board(key)
            want_b = round(sum(r.forecast_cost for r in rows
                               if r.period_month.month == m), 2)
            want_s = round(sum(r.actual_cost for r in rows
                               if r.period_month.month == m), 2)
            self.assertAlmostEqual(
                round(sum(f['budget'] for f in board['functions']), 2),
                want_b, places=2, msg=key)
            self.assertAlmostEqual(
                round(sum(f['spent'] for f in board['functions']), 2),
                want_s, places=2, msg=key)
            # And the twelve `months[]` on every tile are still the WHOLE
            # year, whatever the scope: the spark never narrows.
            for f in board['functions']:
                self.assertEqual(len(f['months']), 12, key)
            self.assertEqual(len(board['strip']), 12, key)

        # --- a stretch of one IS that month, to the byte (R1) --------------
        for m in (1, 3, 6, 8, 12):
            key = self._key(m)
            self.assertEqual(
                self._blob(self._board('%s..%s' % (key, key))),
                self._blob(self._board(key)),
                'a stretch of one month must BE that month: %s' % key)

        # --- a stretch of twelve IS the year, to the byte (R1) -------------
        self.assertEqual(
            self._blob(self._board('%s..%s' % (self._key(1), self._key(12)))),
            self._blob(year),
            'a stretch of the whole strip must BE the year board')

        # --- and asking backwards is the same stretch ----------------------
        self.assertEqual(
            self._blob(self._board('%s..%s' % (self._key(6), self._key(3)))),
            self._blob(self._board('%s..%s' % (self._key(3), self._key(6)))),
            'a stretch dragged right to left is the same stretch')

    # ================================================================== T2
    def test_t2_a_quarter_is_three_months_and_adds_up_to_them(self):
        self._twelve()
        q1 = self._board('Q1')
        self.assertEqual(q1['scope']['kind'], 'range')
        self.assertEqual(q1['scope']['quarter'], 'Q1')
        self.assertEqual(q1['scope']['keys'],
                         [self._key(1), self._key(2), self._key(3)])

        one_at_a_time_b = one_at_a_time_s = 0.0
        for m in (1, 2, 3):
            board = self._board(self._key(m))
            one_at_a_time_b += sum(f['budget'] for f in board['functions'])
            one_at_a_time_s += sum(f['spent'] for f in board['functions'])
        self.assertAlmostEqual(
            round(sum(f['budget'] for f in q1['functions']), 2),
            round(one_at_a_time_b, 2), places=2)
        self.assertAlmostEqual(
            round(sum(f['spent'] for f in q1['functions']), 2),
            round(one_at_a_time_s, 2), places=2)

        # The four quarters together are the year, and they do not overlap.
        every = []
        for q in self.Budget._fy_quarters(FY):
            self.assertEqual(len(q['keys']), 3, q['key'])
            every += q['keys']
        self.assertEqual(every, self.Budget._month_keys(FY))
        # And a quarter asked for by name is the same board as the same three
        # months asked for as a stretch.
        self.assertEqual(
            self._blob(self._board('Q1')),
            self._blob(self._board('%s..%s' % (self._key(1), self._key(3)))))

    # ================================================================== T4
    def test_t4_the_collapse_rules(self):
        """R1 — nothing downstream ever has two ways to say the same thing."""
        self._twelve()
        one = self._board('%s..%s' % (self._key(4), self._key(4)))
        self.assertEqual(one['scope']['kind'], 'month')
        self.assertEqual(one['scope']['key'], self._key(4))
        self.assertEqual(one['scope']['quarter'], '')

        twelve = self._board('%s..%s' % (self._key(1), self._key(12)))
        self.assertEqual(twelve['scope']['kind'], 'year')
        self.assertEqual(twelve['scope']['key'], str(FY))

        stretch = self._board('%s..%s' % (self._key(3), self._key(6)))
        self.assertEqual(stretch['scope']['kind'], 'range')
        self.assertEqual(stretch['scope']['key'],
                         '%s..%s' % (self._key(3), self._key(6)))
        self.assertEqual(len(stretch['scope']['keys']), 4)
        # It is not a quarter, and it does not pretend to be one.
        self.assertEqual(stretch['scope']['quarter'], '')

    # ================================================================== T5
    def test_t5_a_quarter_is_a_quarter_of_the_fiscal_year(self):
        """R5 — Q1 is fiscal months 1–3, whatever the calendar says."""
        param = self.env['ir.config_parameter'].sudo()
        before = param.get_param('pb_budget.fy_start_month')
        try:
            param.set_param('pb_budget.fy_start_month', '7')
            self.assertEqual(self.Budget._fy_start_month(), 7)
            quarters = self.Budget._fy_quarters(FY)
            self.assertEqual(quarters[0]['keys'],
                             ['%s-07' % FY, '%s-08' % FY, '%s-09' % FY])
            self.assertEqual(quarters[3]['keys'],
                             ['%s-04' % (FY + 1), '%s-05' % (FY + 1),
                              '%s-06' % (FY + 1)])
            # The year is still printed as a year that spans two.
            self.assertEqual(self.Budget._fy_label(FY),
                             '%s/%s' % (FY, str(FY + 1)[-2:]))
            # And the chip's own words name the months, so nobody is guessing.
            self.assertIn(str(FY), quarters[0]['name'])
            self.assertIn('/', quarters[0]['name'])
            title = str(quarters[0]['title'])
            self.assertTrue(title, 'a quarter must say which months it means')
            self.assertNotEqual(title, str(quarters[1]['title']))
            scope = self.Budget._scope(FY, quarters[0]['keys'], TODAY)
            self.assertEqual(scope['quarter'], 'Q1')
            self.assertIn('Q1', str(scope['label']))
        finally:
            param.set_param('pb_budget.fy_start_month', before or '1')
        self.assertEqual(self.Budget._fy_start_month(), 1)

    # ================================================================== T6
    def test_t6_the_pace_of_a_stretch(self):
        keys = self.Budget._month_keys(FY)
        pace = self.Budget._period_pace

        # A stretch of one month IS `_month_pace`, to the decimal.
        for m in range(1, 13):
            self.assertEqual(pace([keys[m - 1]], TODAY),
                             self.Budget._month_pace(keys[m - 1], TODAY),
                             keys[m - 1])

        self.assertEqual(pace(keys[0:3], TODAY), 100.0)      # wholly past
        self.assertEqual(pace(keys[8:12], TODAY), 0.0)       # wholly future
        part = pace(keys[3:9], TODAY)                        # straddles today
        self.assertTrue(0.0 < part < 100.0, part)
        # April to September is 183 days; 11 June is the 72nd of them.
        first = date(FY, 4, 1)
        total = (date(FY, 10, 1) - first).days
        self.assertEqual(part, round((TODAY - first).days / total * 100, 1))

        # The state follows the same calendar.
        self.assertEqual(self.Budget._period_state(keys[0:3], TODAY), 'past')
        self.assertEqual(self.Budget._period_state(keys[8:12], TODAY), 'future')
        self.assertEqual(self.Budget._period_state(keys[3:9], TODAY), 'current')

    # ================================================================== T7
    def test_t7_a_deep_link_never_lands_on_an_empty_board(self):
        """R6 / ledger rule 21 — a saved link is never an error."""
        self._twelve()
        keys = self.Budget._month_keys(FY)
        read = self.Budget._period_in_fy

        # The vocabulary that already exists keeps working, exactly.
        self.assertEqual(read('%s-03' % FY, FY), ['%s-03' % FY])
        self.assertEqual(read('', FY), keys)
        self.assertEqual(read(None, FY), keys)
        today_key = date.today().strftime('%Y-%m')
        want = [today_key] if today_key in keys else keys
        self.assertEqual(read('current', FY), want)

        # The new vocabulary, in the same prefix.
        self.assertEqual(read('%s-03..%s-06' % (FY, FY), FY),
                         ['%s-0%s' % (FY, m) for m in (3, 4, 5, 6)])
        self.assertEqual(read('Q2', FY), keys[3:6])
        self.assertEqual(read('q2', FY), keys[3:6])

        # HALF outside the year is CLAMPED to the year.
        self.assertEqual(read('%s-11..%s-03' % (FY - 1, FY), FY),
                         ['%s-01' % FY, '%s-02' % FY, '%s-03' % FY])
        self.assertEqual(read('%s-10..%s-04' % (FY, FY + 1), FY),
                         ['%s-%s' % (FY, m) for m in (10, 11, 12)])

        # WHOLLY outside it, and every shape of rubbish, is the whole year.
        for bad in ('%s-01..%s-06' % (FY - 2, FY - 2),
                    '%s-01..%s-06' % (FY + 3, FY + 3),
                    'rubbish', '..', 'Q9', '%s-13' % FY, '2026-03..',
                    '..%s-03' % FY, '%s-99..%s-99' % (FY, FY), 0, False):
            self.assertEqual(read(bad, FY), keys, repr(bad))

        # And every one of them draws a real board rather than a blank one.
        for asked in ('%s-03' % FY, 'current', '%s-03..%s-06' % (FY, FY), 'Q2',
                      '%s-11..%s-03' % (FY - 1, FY), 'rubbish',
                      '%s-01..%s-06' % (FY - 2, FY - 2)):
            board = self._board(asked)
            self.assertTrue(board['ok'], asked)
            self.assertTrue(board['functions'], asked)
            self.assertTrue(board['scope']['keys'], asked)
            self.assertTrue(str(board['headline']).strip(), asked)

    # ================================================================== T8
    def test_t8_the_exports_name_the_stretch_they_are_about(self):
        self._twelve()
        try:
            import openpyxl                                    # noqa: F401
        except ImportError:
            self.skipTest('no spreadsheet library on this database')

        import base64
        import io
        import openpyxl

        period = '%s-03..%s-06' % (FY, FY)
        board = self._board(period)
        res = self.env['pb.budget.export'].build(FY, 'admin', 'local', 'xlsx',
                                                 period)
        self.assertTrue(res['ok'])
        # The FILE says which stretch it is about, in its own name.
        self.assertIn(str(board['scope']['label']), res['filename'])

        book = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(res['file_b64'])))
        sheet = book.active
        title = str(sheet.cell(row=1, column=1).value or '')
        self.assertIn(str(board['scope']['label']), title)

        # And the total row is the board's total, to the digit.
        spent = round(sum(f['spent'] for f in board['functions']), 2)
        found = [r for r in sheet.iter_rows(values_only=True)
                 if r and r[0] and str(r[0]).strip() in ('Total', 'Tổng')]
        self.assertTrue(found, 'the sheet must carry a total row')
        self.assertAlmostEqual(float(found[-1][3] or 0.0), spent, places=0)

        # A stretch's sheet answers with the variance, not with the pace, and
        # drops the twenty-four month columns it is not about.
        head = [c for c in next(sheet.iter_rows(min_row=4, max_row=4,
                                                values_only=True)) if c]
        self.assertLessEqual(len(head), 8, head)

        # The printed page draws the whole year with the stretch marked.
        rec = self.env['pb.budget.export'].create({
            'fy': FY, 'budget_type': 'admin', 'currency_mode': 'local',
            'month': period,
        })
        bars = rec.month_bars(board)
        self.assertEqual(len(bars), 12)
        self.assertEqual([b['key'] for b in bars if b['in_scope']],
                         board['scope']['keys'])
        lines = rec.narrative(board)
        self.assertTrue(lines)
        for line in lines:
            self.assertIn(str(board['scope']['name']), str(line))

    # ================================================================== T11
    def test_t11_every_shape_of_stretch_has_words_for_itself(self):
        """§2d — no stretch is left with a blank name or a wrong sentence."""
        self._twelve()
        keys = self.Budget._month_keys(FY)

        # A stretch inside one calendar year carries the year once.
        march_to_june = self.Budget._scope(FY, keys[2:6], TODAY)
        self.assertIn(str(FY), str(march_to_june['label']))
        self.assertNotIn(str(FY), str(march_to_june['name']))

        # A quarter names itself AND its months.
        q2 = self.Budget._scope(FY, keys[3:6], TODAY)
        self.assertIn('Q2', str(q2['label']))
        self.assertIn('Q2', str(q2['name']))
        self.assertEqual(q2['short'], 'Q2')

        # One that crosses the turn of the calendar year carries it on both
        # ends, or "December to February" means two different things.
        param = self.env['ir.config_parameter'].sudo()
        before = param.get_param('pb_budget.fy_start_month')
        try:
            param.set_param('pb_budget.fy_start_month', '7')
            crossing = self.Budget._scope(
                FY, self.Budget._month_keys(FY)[5:9], TODAY)
            self.assertIn(str(FY), str(crossing['label']))
            self.assertIn(str(FY + 1), str(crossing['label']))
            self.assertIn(str(FY), str(crossing['name']))
            self.assertIn(str(FY + 1), str(crossing['name']))
        finally:
            param.set_param('pb_budget.fy_start_month', before or '1')

        # Nothing is ever blank, in any shape.
        for span in (keys[0:1], keys[0:2], keys[2:6], keys[3:6], keys[0:12],
                     keys[11:12], []):
            scope = self.Budget._scope(FY, span, TODAY)
            for field in ('kind', 'key', 'label', 'name', 'short', 'state'):
                self.assertTrue(str(scope[field]).strip(),
                                '%s empty for %s' % (field, span))
            self.assertIn(scope['kind'], ('year', 'month', 'range'))
            self.assertIn(scope['state'], ('past', 'current', 'future'))

    def test_t11b_a_stretch_with_nothing_in_it_still_answers(self):
        """A stretch wholly in the future, one with no budget at all, and one
        where a function spent nothing — three states, three sentences, and no
        tile disappears."""
        self._twelve()
        # August has no rows at all; July and September do.
        empty = self._board('%s-08..%s-08' % (FY, FY))
        self.assertTrue(empty['ok'])
        self.assertTrue(str(empty['headline']).strip())
        # THE TILES DO NOT REARRANGE. Every function the year has, a month with
        # nothing in it also has — "this function spent nothing" is an answer.
        year = self._board()
        self.assertEqual([f['id'] for f in empty['functions']],
                         [f['id'] for f in year['functions']])
        for f in empty['functions']:
            self.assertEqual(f['budget'], 0.0)
            self.assertEqual(f['spent'], 0.0)
            self.assertTrue(str(f['tone_label']).strip())

    # ================================================================= T12
    def test_t12_the_new_words_are_this_product_and_carry_no_stray_signs(self):
        """White-label, and L9: not one literal per cent sign reaches a
        screen. The board is MADE of percentages and the same sentence exists
        on both sides of this feature, where `_()` needs `%%` and `_t()` does
        not."""
        self._twelve()
        for period in (None, '%s-06' % FY, 'Q2', '%s-03..%s-06' % (FY, FY),
                       '%s-11..%s-02' % (FY, FY + 1)):
            board = self._board(period)
            blob = ' '.join([
                str(board['headline']),
                str(board['scope']['label']), str(board['scope']['name']),
                str(board['scope']['short']),
                ' '.join(str(q['name']) for q in board['quarters']),
                ' '.join(str(q['title']) for q in board['quarters']),
                ' '.join(str(s['tone_label']) for s in board['strip']),
                ' '.join(str(f['tone_label']) for f in board['functions']),
            ])
            self.assertNotIn('odoo', blob.lower(), repr(period))
            self.assertNotIn('%%', blob, repr(period))
            # No code vocabulary leaks into a sentence a person reads.
            for word in ('mkeys', 'scope[', 'kind', 'YYYY'):
                self.assertNotIn(word, blob, repr(period))

    def test_t12b_a_running_stretch_gets_its_own_sentence(self):
        """The one sentence of the five that could not be shared: a month says
        "the month's budget" out loud and a stretch may not."""
        keys = self.Budget._month_keys(FY)
        payload = {
            'currency': {'code': 'VND'}, 'months': [],
            'functions': [{'id': 1, 'name': 'Retail', 'tone': 'calm',
                           'budget': 100.0, 'spent': 60.0, 'variance': -40.0,
                           'variance_pct': -40.0}],
            'kpis': {'spent': 62.0, 'budget': 100.0, 'left': 38.0,
                     'burn': 62.0, 'pace': 41.0, 'functions': 1},
        }
        payload['scope'] = self.Budget._scope(FY, keys[3:9], TODAY)
        payload['scope']['state'] = 'current'
        line = str(self.Budget._headline_month(payload))
        self.assertIn('62', line)
        self.assertIn('41', line)
        self.assertNotIn('%%', line)
        self.assertNotIn("the month's", line)
        self.assertIn(str(payload['scope']['name']), line)

    # ================================================================= T14
    def test_t14_the_drill_of_a_stretch_adds_up_and_compares_like_for_like(self):
        self._twelve()
        period = '%s-04..%s-06' % (FY, FY)
        board = self._board(period)
        func = board['functions'][0]
        drill = self.Budget.get_function(func['id'], FY, 'admin', 'local',
                                         period)
        self.assertTrue(drill['ok'])
        self.assertEqual(drill['scope']['kind'], 'range')
        self.assertAlmostEqual(
            round(sum(d['spent'] for d in drill['function']['departments']), 2),
            drill['function']['spent'], places=2)
        # The twelve bars are still the whole year.
        self.assertEqual(len(drill['function']['months']), 12)

        # "Compared with what" compares a stretch with a stretch of the SAME
        # length, not with one month.
        cmp = drill['compare']
        self.assertEqual(cmp['this']['spent'], drill['function']['spent'])
        # Three months earlier is January to March, which the fixture has.
        self.assertTrue(cmp['last']['has'])
        jan_mar = self.Budget.get_function(
            func['id'], FY, 'admin', 'local', '%s-01..%s-03' % (FY, FY))
        self.assertAlmostEqual(cmp['last']['spent'],
                               jan_mar['function']['spent'], places=2)
        # There is no year before this one, so that cell is BLANK rather than a
        # zero that would read as a collapse.
        self.assertFalse(cmp['last_year']['has'])

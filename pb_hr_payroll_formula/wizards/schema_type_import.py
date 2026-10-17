# -*- coding: utf-8 -*-
"""Read bounded, real schema rows instead of the truncated preview display."""
import base64
import io

from odoo import models


class MultisheetSchemaTypes(models.TransientModel):
    _inherit = 'hr.formula.multisheet.import.wizard'

    def _schema_import_evidence(self, components):
        self.ensure_one()
        import openpyxl
        from openpyxl.utils import column_index_from_string
        from openpyxl.utils.cell import coordinate_from_string
        from ..integrations import ExcelConnector

        content = base64.b64decode(self.import_file)
        connector = ExcelConnector(None)
        connector.load_workbook_multisheet(content, include_formulas=True)
        formulas = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
        starts, evidence = {}, {}
        for comp in components:
            if not comp.column_letter or comp.source_sheet not in formulas.sheetnames:
                continue
            data_sheet = connector.workbook[comp.source_sheet]
            formula_sheet = formulas[comp.source_sheet]
            if comp.source_sheet not in starts:
                if self._use_color_import():
                    info = self._get_color_coded_sheet_data(formula_sheet, formula_sheet)
                    start = info['sheet_data']['data_start_row'] + 1
                else:
                    info = connector.load_sheet_with_detection(comp.source_sheet)
                    start = info['data_start_row']
                starts[comp.source_sheet] = start
            start = starts[comp.source_sheet]
            source_column = (comp.constant_cell_ref if comp.column_type == 'constant' else False) \
                or comp.source_column_letter or comp.column_letter
            if comp.column_type == 'constant' and any(char.isdigit() for char in source_column):
                letter, row = coordinate_from_string(source_column)
                cell = data_sheet.cell(row, column_index_from_string(letter))
                evidence[comp.id] = self.env['hr.formula.rule']._schema_type_evidence(
                    [cell.value], [cell.number_format])
                continue
            column = column_index_from_string(source_column)
            values, formats = [], []
            for row in range(start, min(start + 40, data_sheet.max_row + 1)):
                cell = data_sheet.cell(row, column)
                value = cell.value
                # W41 exports include a code-header row. It is not employee data.
                if isinstance(value, str) and (value.strip() in (
                        comp.generated_code, comp.generated_name) or value.startswith('=')):
                    continue
                if value is not None:
                    values.append(value)
                formats.append(cell.number_format)
            evidence[comp.id] = self.env['hr.formula.rule']._schema_type_evidence(values, formats)
        formulas.close()
        return evidence

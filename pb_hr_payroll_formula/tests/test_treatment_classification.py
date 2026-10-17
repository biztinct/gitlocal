# -*- coding: utf-8 -*-
"""Schema types and reporting treatment must be ready before the first pay run."""
import base64
import datetime
import io
from unittest.mock import patch

from odoo.exceptions import UserError
from odoo.tests import TransactionCase, tagged


@tagged('post_install', '-at_install', 'treatment_classification')
class TestTreatmentClassification(TransactionCase):
    def setUp(self):
        super().setUp()
        self.config = self.env['hr.formula.config'].create({
            'name': 'Treatment fixture', 'country_code': 'VN',
        })
        specs = [
            ('A', 'BASICT', 'Base salary', 'input', ''),
            ('B', 'ALLOWT', 'Allowance', 'input', ''),
            ('C', 'GROSST', 'Gross pay', 'formula', '=A5+B5'),
            ('D', 'TAXT', 'Income tax', 'input', ''),
            ('E', 'NETPAYT', 'Net pay', 'formula', '=C5-D5'),
            ('F', 'NAMET', 'Full name', 'input', ''),
            ('G', 'DATET', 'Commencing date', 'input', ''),
            ('H', 'HOURST', 'Working hours', 'input', ''),
        ]
        self.rules = {}
        for letter, code, name, ctype, formula in specs:
            self.rules[code] = self.env['hr.formula.rule'].create({
                'config_id': self.config.id, 'code': code, 'name': name,
                'forced_column_letter': letter, 'original_column_letter': letter,
                'column_type': ctype, 'excel_formula': formula,
            })
        self.rules['NAMET'].write(self.rules['NAMET']._schema_type_evidence(['Lan', 'Mai']))
        self.rules['DATET'].write(self.rules['DATET']._schema_type_evidence(
            [datetime.date(2026, 1, 1)], ['dd/mm/yyyy']))

    def test_01_all_axes_ready_and_subtotals_follow_graph(self):
        self.rules['BASICT'].imported_component_type = 'Imported section'
        result = self.config.reclassify_from_formulas()
        self.assertGreater(result['changed'], 0)
        for code, group in [('BASICT', 'Earnings'), ('TAXT', 'Deductions'),
                            ('NETPAYT', 'Net pay'), ('NAMET', 'Information')]:
            self.assertEqual(self.rules[code].component_type, group, code)
        self.assertEqual(self.rules['BASICT'].imported_component_type, 'Imported section')
        self.assertTrue(self.rules['BASICT'].net_role_detail)
        self.assertFalse(self.rules['GROSST'].net_role_detail)
        self.assertEqual(self.rules['NAMET'].value_kind, 'text')
        self.assertEqual(self.rules['DATET'].value_kind, 'date')
        self.assertEqual(self.rules['HOURST'].value_kind, 'quantity')
        self.assertFalse(self.env['hr.payroll.import.batch'].search_count(
            [('formula_config_id', '=', self.config.id)]))

    def test_02_manual_choices_survive_independently(self):
        self.config.reclassify_from_formulas()
        self.config.set_component_setup({
            'ALLOWT': {'group': 'Benefits', 'pay_role': 'info', 'rollup': False},
            'NAMET': {'kind': 'identifier'},
        })
        self.config.reclassify_from_formulas()
        self.assertEqual(self.rules['ALLOWT'].component_type, 'Benefits')
        self.assertEqual(self.rules['ALLOWT'].component_type_source, 'user')
        self.assertEqual(self.rules['ALLOWT'].net_role, 'info')
        self.assertFalse(self.rules['ALLOWT'].net_role_detail)
        self.assertEqual(self.rules['NAMET'].value_kind, 'identifier')

    def test_03_repeat_is_stable_and_offers_standard_groups(self):
        self.config.reclassify_from_formulas()
        self.assertEqual(self.config.reclassify_from_formulas()['changed'], 0)
        board = self.config.value_kind_board()
        self.assertTrue({'Earnings', 'Deductions', 'Employer costs', 'Net pay',
                         'Information'}.issubset(set(board['groups'])))
        self.assertTrue(all(row['group_source'] == 'auto' for row in board['rows']))

    def test_04_safety_rejection_rolls_back_type_changes_too(self):
        self.rules['NAMET'].write({'value_kind': 'money', 'value_kind_source': 'auto'})
        with patch.object(type(self.config), '_net_role_would_worsen',
                          return_value={'gap_before': 0, 'gap_after': 100}):
            with self.assertRaises(UserError):
                self.config.reclassify_from_formulas()
        self.assertEqual(self.rules['NAMET'].value_kind, 'money')
        self.assertFalse(self.rules['NAMET'].component_type)

    def _workbook(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Payroll'
        ws.append(['Employee Code', 'Full name', 'Commencing date',
                   'Base salary', 'Income tax', 'Net pay', 'Policy rate'])
        for number in range(2, 7):
            ws.append(['001' + str(number), 'Lan ' + str(number),
                       datetime.date(2026, 1, number), 1000000, 10000,
                       '=D%d-E%d' % (number, number), .05])
            ws.cell(number, 3).number_format = 'dd/mm/yyyy'
            ws.cell(number, 7).number_format = '0%'
        data = io.BytesIO()
        wb.save(data)
        return base64.b64encode(data.getvalue())

    def test_05_standard_import_classifies_without_pay_data(self):
        config = self.env['hr.formula.config'].create({
            'name': 'Standard schema import', 'country_code': 'VN',
            'use_color_coded_excel_import': False})
        self.env['hr.formula.import.wizard'].create({
            'config_id': config.id, 'import_source': 'excel',
            'import_file': self._workbook(), 'import_filename': 'schema.xlsx',
        }).action_import()
        rules = {r.name: r for r in config.rule_ids}
        self.assertEqual(rules['Full name'].value_kind, 'text')
        self.assertEqual(rules['Commencing date'].value_kind, 'date')
        self.assertEqual(rules['Policy rate'].value_kind, 'rate')
        self.assertEqual(rules['Employee Code'].value_kind, 'identifier')
        self.assertEqual(rules['Base salary'].component_type, 'Earnings')
        self.assertEqual(rules['Income tax'].component_type, 'Deductions')
        self.assertEqual(rules['Net pay'].component_type, 'Net pay')
        # Persist type counts only, not names or sample cell contents.
        self.assertNotIn('Lan', str(rules['Full name'].schema_value_shape))

    def test_06_color_schema_populates_types_before_net_is_defined(self):
        from .test_column_roles import _build_color_workbook
        config = self.env['hr.formula.config'].create({
            'name': 'Colour schema import', 'country_code': 'VN',
            'use_color_coded_excel_import': True})
        self.env['hr.formula.import.wizard'].create({
            'config_id': config.id, 'import_source': 'excel',
            'import_file': _build_color_workbook(), 'import_filename': 'schema.xlsx',
        }).action_import()
        rules = {r.name: r for r in config.rule_ids}
        self.assertEqual(rules['Cost Centre Ref'].value_kind, 'text')
        self.assertEqual(rules['Base Salary'].value_kind, 'money')
        self.assertTrue(rules['Cost Centre Ref'].schema_value_shape['texty'])

    def test_07_multisheet_execute_uses_source_columns_not_appended_positions(self):
        config = self.env['hr.formula.config'].create({
            'name': 'Multisheet schema import', 'country_code': 'VN',
            'use_color_coded_excel_import': False})
        wizard = self.env['hr.formula.multisheet.import.wizard'].create({
            'config_id': config.id, 'import_file': self._workbook(),
            'import_filename': 'schema.xlsx',
        })
        for letter, source, code, name, kind, formula in [
            ('A', 'B', 'FULLNAME', 'Full name', 'input', ''),
            ('B', 'C', 'STARTDATE', 'Commencing date', 'input', ''),
            ('C', 'D', 'SALARY', 'Base salary', 'input', ''),
            ('D', 'E', 'TAX', 'Income tax', 'input', ''),
            ('E', 'F', 'NETPAY', 'Net pay', 'formula', '=C2-D2'),
        ]:
            self.env['hr.formula.multisheet.component.preview'].create({
                'wizard_id': wizard.id, 'source_sheet': 'Payroll',
                'column_letter': letter, 'source_column_letter': source,
                'generated_code': code, 'generated_name': name,
                'column_type': kind, 'excel_formula': formula, 'include_in_import': True,
            })
        wizard.action_execute_import()
        rules = {r.code: r for r in config.rule_ids}
        self.assertEqual(rules['FULLNAME'].value_kind, 'text')
        self.assertEqual(rules['STARTDATE'].value_kind, 'date')
        self.assertEqual(rules['SALARY'].component_type, 'Earnings')
        self.assertEqual(rules['TAX'].component_type, 'Deductions')

    def test_08_clearing_manual_group_returns_to_automatic(self):
        self.config.reclassify_from_formulas()
        self.config.set_component_setup({'BASICT': {'group': 'Custom'}})
        self.config.set_component_setup({'BASICT': {'group': ''}})
        self.assertEqual(self.rules['BASICT'].component_type, 'Earnings')
        self.assertEqual(self.rules['BASICT'].component_type_source, 'auto')

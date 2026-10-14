"""Country tax tables and explicit links to imported payroll constants.

Published editions are immutable. A dated refresh changes only a draft/testing
scheme through the existing statutory approval route. Payroll arithmetic still
uses the scheme's versioned constants: changing the library never rewrites a
historical payroll or silently changes an active scheme.
"""
import hashlib
import json
import math

from markupsafe import Markup, escape
from odoo import _, api, fields, models
from odoo.exceptions import UserError, ValidationError

from .formula_legislation import _COUNTRY
from .statutory_approval import PbStatutoryProposal

_APPLY_KEY = 'payroll_tax_apply'
_APPLY_TOKEN = object()


def stamp(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, default=str).encode()).hexdigest()


def same_number(left, right):
    # Odoo float rounding adds a few binary ULPs at billion-rupiah magnitudes.
    # Keep the tolerance much smaller than a currency unit or a rate digit.
    return math.isclose(left, right, rel_tol=1e-14, abs_tol=1e-9)


def notice(answer):
    return {'type': 'ir.actions.client', 'tag': 'display_notification', 'params': {
        'title': _('Statutory tables'), 'message': answer['message'],
        'type': 'success' if answer['applied'] else 'info', 'sticky': not answer['applied']}}


class PayrollTaxTable(models.Model):
    _name = 'hr.payroll.tax.table'
    _description = 'Country Statutory Table Edition'
    _order = 'country_code, code, effective_from desc, id desc'

    name = fields.Char(required=True)
    code = fields.Char(required=True, index=True)
    company_id = fields.Many2one('res.company', required=True, default=lambda s: s.env.company, index=True)
    country_code = fields.Selection(_COUNTRY, required=True, index=True)
    version = fields.Char(required=True, default=lambda s: fields.Date.context_today(s).strftime('%Y-%m'))
    effective_from = fields.Date(required=True, default=fields.Date.today)
    effective_to = fields.Date()
    method = fields.Selection([('progressive', 'Progressive marginal bands'),
                               ('whole_income', 'Selected rate on whole income (TER)'),
                               ('parameters', 'Statutory constants')], required=True, default='progressive')
    basis = fields.Selection([('monthly', 'Monthly'), ('annual', 'Annual'), ('daily', 'Daily')], required=True, default='monthly')
    state = fields.Selection([('draft', 'Draft'), ('published', 'Published')], default='draft', required=True, copy=False)
    authority = fields.Char(string='Source / regulation')
    slab_ids = fields.One2many('hr.payroll.tax.band', 'table_id', copy=True, string='Tax bands')
    parameter_ids = fields.One2many('hr.payroll.tax.parameter', 'table_id', copy=True, string='Statutory values')
    binding_ids = fields.One2many('hr.formula.rule', 'tax_table_id', string='Mapped payroll constants')

    _edition_unique = models.Constraint('unique(company_id,country_code,code,version)',
                                         'Use a unique version for this country and table code.')

    @api.constrains('effective_from', 'effective_to', 'code')
    def _check_metadata(self):
        import re
        for table in self:
            if not re.fullmatch(r'[A-Z][A-Z0-9]*', table.code or ''):
                raise ValidationError(_('Table codes must be uppercase letters and digits, starting with a letter.'))
            if table.effective_to and table.effective_to < table.effective_from:
                raise ValidationError(_('Effective to must be on or after effective from.'))

    def _check_mutable(self):
        if self.env.context.get(_APPLY_KEY) is _APPLY_TOKEN:
            return
        if any(t.state == 'published' for t in self):
            raise UserError(_('Published editions cannot be edited. Duplicate the edition with a new version and effective date.'))

    @api.model_create_multi
    def create(self, vals_list):
        if self.env.context.get(_APPLY_KEY) is not _APPLY_TOKEN and any(v.get('state') == 'published' for v in vals_list):
            raise UserError(_('Publish this edition through its approval action.'))
        return super().create(vals_list)

    def write(self, vals):
        self._check_mutable()
        if 'state' in vals and self.env.context.get(_APPLY_KEY) is not _APPLY_TOKEN:
            raise UserError(_('Publish this edition through its approval action.'))
        return super().write(vals)

    def unlink(self):
        self._check_mutable()
        if self.binding_ids:
            raise UserError(_('This edition is mapped to payroll constants and cannot be deleted.'))
        return super().unlink()

    def copy(self, default=None):
        self.ensure_one()
        return super().copy(dict({'version': self.version + ' (copy)', 'state': 'draft'}, **(default or {})))

    def _validate_edition(self):
        for t in self:
            if t.method == 'parameters':
                if not t.parameter_ids or t.slab_ids:
                    raise ValidationError(_('A constants table needs statutory values and no tax bands.'))
            else:
                rows = t.slab_ids.sorted('sequence')
                if not rows or t.parameter_ids:
                    raise ValidationError(_('A tax table needs tax bands and no statutory constants.'))
                if rows.mapped('sequence') != list(range(1, len(rows) + 1)):
                    raise ValidationError(_('Number the tax bands consecutively, starting at 1.'))
                previous = 0
                for i, row in enumerate(rows):
                    if (i == len(rows) - 1 and row.income_to != 0) or (i != len(rows) - 1 and row.income_to <= previous):
                        raise ValidationError(_('Band upper bounds must increase; the final band must use 0 for unlimited.'))
                    previous = row.income_to
            clash = self.search([('id', '!=', t.id), ('company_id', '=', t.company_id.id),
                                 ('country_code', '=', t.country_code), ('code', '=', t.code),
                                 ('effective_from', '=', t.effective_from), ('state', '=', 'published')], limit=1)
            if clash:
                raise ValidationError(_('A published edition already starts on this date. Choose a new effective date.'))
        return True

    def _values(self):
        self.ensure_one()
        if self.method == 'parameters':
            return {(p.code, 'parameter'): p.value for p in self.parameter_ids}
        return {(b.sequence, key): value for b in self.slab_ids
                for key, value in [('upper', b.income_to), ('rate', b.tax_rate)]}

    def _stamp(self):
        self.ensure_one()
        return stamp({'metadata': self.read(['code', 'country_code', 'company_id', 'version',
                                             'effective_from', 'effective_to', 'method', 'basis', 'state'])[0],
                      'values': sorted((str(k[0]), k[1], v) for k, v in self._values().items())})

    def action_publish(self):
        self.ensure_one()
        self.check_access('write')
        self._validate_edition()
        if self.state != 'draft':
            raise UserError(_('This edition is already published.'))
        proposal = self.env['pb.statutory.proposal'].propose(
            'country_tax_publish', _('Publish %s', self.display_name),
            payload={'table_id': self.id,
                     'values': [{'slot': str(k[0]), 'field': k[1], 'value': v}
                                for k, v in sorted(self._values().items(), key=lambda x: str(x[0]))]},
            snapshot={'stamp': self._stamp()},
            facts={'effective_date': {'value': str(self.effective_from), 'unit': ''},
                   'rows_changed': {'value': len(self.slab_ids) + len(self.parameter_ids), 'unit': ''}}, target=self)
        return notice(proposal.answer())

    def action_export_mapping(self):
        """Export current DB values to the same workbook mapping format.

        The workbook is a constants/mapping patch for the specified payroll
        cells, not a replacement payroll schema. Download happens through an
        attachment on this record, respecting its normal access rules.
        """
        import base64
        import io
        import openpyxl
        from openpyxl.styles import Font
        from ..formula_engine.tax_workbook import HEADERS
        self.ensure_one()
        wb = openpyxl.Workbook()
        meta = wb.active
        meta.title = 'Tax Map'
        meta.append(HEADERS)
        values = self._values()
        for rule in self.binding_ids.sorted('id'):
            if not rule.tax_source_cell or not rule.tax_source_sheet:
                continue
            slot = rule.tax_parameter_code if rule.tax_field == 'parameter' else rule.tax_band_sequence
            key = (slot, rule.tax_field)
            if key not in values:
                raise UserError(_('A mapped slot is missing from this table.'))
            source = rule.tax_source_sheet
            sheet = wb[source] if source in wb.sheetnames else wb.create_sheet(source)
            cell = sheet[rule.tax_source_cell]
            cell.value = values[key]
            cell.font = Font(color='0000FF')
            if rule.tax_field == 'rate':
                cell.number_format = '0.00%'
            meta.append([self.code, self.country_code, self.version, str(self.effective_from),
                         str(self.effective_to or ''), self.method, self.basis, slot,
                         rule.tax_field, source, rule.tax_source_cell, rule.name])
        if meta.max_row == 1:
            raise UserError(_('No payroll cells have been mapped to this edition.'))
        out = io.BytesIO()
        wb.save(out)
        attachment = self.env['ir.attachment'].create({
            'name': '%s_%s_constants.xlsx' % (self.code, self.version),
            'datas': base64.b64encode(out.getvalue()), 'res_model': self._name, 'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'})
        return {'type': 'ir.actions.act_url', 'url': '/web/content/%s?download=true' % attachment.id, 'target': 'self'}


class PayrollTaxBand(models.Model):
    _name = 'hr.payroll.tax.band'
    _description = 'Country Tax Band'
    _order = 'sequence, id'
    table_id = fields.Many2one('hr.payroll.tax.table', required=True, ondelete='cascade', index=True)
    sequence = fields.Integer(required=True, default=1)
    income_from = fields.Float(compute='_compute_from', string='Above')
    income_to = fields.Float(string='Up to (inclusive; 0 = unlimited)', required=True)
    tax_rate = fields.Float(string='Rate', digits=(16, 8), required=True)

    @api.depends('table_id.slab_ids.income_to', 'table_id.slab_ids.sequence')
    def _compute_from(self):
        for band in self:
            previous = band.table_id.slab_ids.filtered(lambda b: b.sequence < band.sequence).sorted('sequence')
            band.income_from = previous[-1].income_to if previous else 0

    @api.constrains('income_to', 'tax_rate', 'sequence')
    def _check_values(self):
        for band in self:
            if not math.isfinite(band.income_to) or band.income_to < 0 or not math.isfinite(band.tax_rate) or not 0 <= band.tax_rate <= 1 or band.sequence < 1:
                raise ValidationError(_('Tax bands need non-negative bounds and rates between 0 and 100%.'))

    @api.model_create_multi
    def create(self, vals_list):
        self.env['hr.payroll.tax.table'].browse([v['table_id'] for v in vals_list])._check_mutable()
        return super().create(vals_list)

    def write(self, vals):
        self.table_id._check_mutable()
        if vals.get('table_id'):
            self.env['hr.payroll.tax.table'].browse(vals['table_id'])._check_mutable()
        return super().write(vals)

    def unlink(self):
        self.table_id._check_mutable()
        return super().unlink()


class PayrollTaxParameter(models.Model):
    _name = 'hr.payroll.tax.parameter'
    _description = 'Country Statutory Parameter'
    _order = 'code, id'
    table_id = fields.Many2one('hr.payroll.tax.table', required=True, ondelete='cascade', index=True)
    code = fields.Char(required=True)
    name = fields.Char(required=True)
    value = fields.Float(required=True, digits=(20, 8))
    _slot_unique = models.Constraint('unique(table_id,code)', 'Parameter codes must be unique within an edition.')

    @api.constrains('value')
    def _check_value(self):
        if any(not math.isfinite(p.value) for p in self):
            raise ValidationError(_('Statutory values must be finite numbers.'))

    @api.model_create_multi
    def create(self, vals_list):
        self.env['hr.payroll.tax.table'].browse([v['table_id'] for v in vals_list])._check_mutable()
        return super().create(vals_list)

    def write(self, vals):
        self.table_id._check_mutable()
        if vals.get('table_id'):
            self.env['hr.payroll.tax.table'].browse(vals['table_id'])._check_mutable()
        return super().write(vals)

    def unlink(self):
        self.table_id._check_mutable()
        return super().unlink()


class TaxMappedRule(models.Model):
    _inherit = 'hr.formula.rule'
    tax_table_id = fields.Many2one('hr.payroll.tax.table', ondelete='restrict', string='Statutory edition')
    tax_field = fields.Selection([('upper', 'Band upper bound'), ('rate', 'Band rate'), ('parameter', 'Statutory value')])
    tax_band_sequence = fields.Integer(string='Band')
    tax_parameter_code = fields.Char(string='Parameter')
    tax_source_sheet = fields.Char(string='Workbook sheet')
    tax_source_cell = fields.Char(string='Workbook cell')
    tax_expected_value = fields.Float(compute='_compute_tax_value', digits=(20, 8), string='Database value')
    tax_aligned = fields.Boolean(compute='_compute_tax_value', string='Aligned')

    @api.depends('tax_table_id.slab_ids.income_to', 'tax_table_id.slab_ids.tax_rate',
                 'tax_table_id.parameter_ids.value', 'constant_value', 'tax_field',
                 'tax_band_sequence', 'tax_parameter_code')
    def _compute_tax_value(self):
        for rule in self:
            slot = rule.tax_parameter_code if rule.tax_field == 'parameter' else rule.tax_band_sequence
            value = rule.tax_table_id._values().get((slot, rule.tax_field)) if rule.tax_table_id else None
            rule.tax_expected_value = value or 0
            rule.tax_aligned = value is not None and same_number(rule.constant_value, value)

    @api.constrains('tax_table_id', 'config_id', 'column_type', 'tax_field', 'tax_band_sequence', 'tax_parameter_code')
    def _check_tax_binding(self):
        for rule in self.filtered('tax_table_id'):
            table = rule.tax_table_id
            slot = rule.tax_parameter_code if rule.tax_field == 'parameter' else rule.tax_band_sequence
            if rule.column_type != 'constant' or rule.config_id.country_code != table.country_code or rule.config_id.company_id != table.company_id:
                raise ValidationError(_('Statutory constants must match the payroll scheme company and country.'))
            if (slot, rule.tax_field) not in table._values():
                raise ValidationError(_('The mapped statutory slot does not exist.'))

    def _version_snapshot(self):
        result = super()._version_snapshot()
        result['tax_edition'] = self.tax_table_id.id or False
        result['tax_version'] = self.tax_table_id.version or ''
        return result


class TaxMappedConfig(models.Model):
    _inherit = 'hr.formula.config'
    statutory_asof = fields.Date(string='Statutory values as of', copy=True)

    @api.constrains('country_code', 'company_id')
    def _check_mapped_country(self):
        for config in self:
            if config.rule_ids.filtered(lambda r: r.tax_table_id and (
                    r.tax_table_id.country_code != config.country_code or
                    r.tax_table_id.company_id != config.company_id)):
                raise ValidationError(_('Mapped statutory editions must match the configuration company and country.'))

    def _tax_refresh_plan(self, asof):
        self.ensure_one()
        self.check_access('write')
        if self.state not in ('draft', 'testing'):
            raise UserError(_('Refresh a draft or testing scheme. Copy or branch the active scheme for the new payroll period first.'))
        asof = fields.Date.to_date(asof)
        plan = []
        for rule in self.rule_ids.filtered('tax_table_id'):
            old = rule.tax_table_id
            candidates = self.env['hr.payroll.tax.table'].search([
                ('company_id', '=', self.company_id.id), ('country_code', '=', self.country_code),
                ('code', '=', old.code), ('method', '=', old.method), ('basis', '=', old.basis),
                ('state', '=', 'published'), ('effective_from', '<=', asof),
                '|', ('effective_to', '=', False), ('effective_to', '>=', asof)],
                order='effective_from desc, id desc', limit=1)
            if not candidates:
                raise UserError(_('No published edition of %s covers this date.', old.code))
            slot = rule.tax_parameter_code if rule.tax_field == 'parameter' else rule.tax_band_sequence
            value = candidates._values().get((slot, rule.tax_field))
            if value is None:
                raise UserError(_('The new edition changes the band structure. Re-import or remap the formulas before refreshing %s.', old.code))
            # A newly added band also changes formula structure; never merely
            # update the old slots and pretend the new rate is being used.
            if set(old._values()) != set(candidates._values()):
                raise UserError(_('The new edition changes the band structure. Re-import the updated template for %s.', old.code))
            plan.append({'rule_id': rule.id, 'name': rule.name, 'before': rule.constant_value,
                         'field': rule.tax_field, 'slot': str(slot),
                         'old_table_id': old.id, 'table_id': candidates.id, 'after': value,
                         'source_stamp': candidates._stamp()})
        if not plan:
            raise UserError(_('This scheme has no mapped statutory constants. Import a workbook with a Tax Map first.'))
        return plan

    def _apply_tax_refresh(self, asof):
        if self.env.context.get(_APPLY_KEY) is not _APPLY_TOKEN:
            raise UserError(_('Refresh statutory values through the review action.'))
        plan = self._tax_refresh_plan(asof)
        changed = 0
        for row in plan:
            if not same_number(row['before'], row['after']) or row['old_table_id'] != row['table_id']:
                self.env['hr.formula.rule'].browse(row['rule_id']).with_context(
                    formula_version_reason='legislation').write({
                        'constant_value': row['after'], 'tax_table_id': row['table_id']})
                changed += 1
        self.statutory_asof = asof
        Ver = self.env['hr.formula.rule.version']
        last = Ver.search([('config_id', '=', self.id)], order='id desc', limit=1)
        self.env['hr.formula.config.milestone'].create({
            'config_id': self.id, 'name': _('Statutory refresh · %s', asof),
            'version_hwm': last.id if last else 0})
        return {'config_id': self.id, 'changed': changed, 'asof': str(asof)}

    def action_activate(self):
        for config in self:
            rules = config.rule_ids.filtered('tax_table_id')
            if rules.filtered(lambda r: r.tax_table_id.state != 'published' or not r.tax_aligned):
                raise UserError(_('Publish the mapped statutory editions and align their constants before activating this scheme.'))
        return super().action_activate()


class TaxRefreshWizard(models.TransientModel):
    _name = 'hr.payroll.tax.refresh.wizard'
    _description = 'Review Statutory Constant Refresh'
    config_id = fields.Many2one('hr.formula.config', required=True,
                                domain="[('state', 'in', ['draft','testing'])]")
    asof = fields.Date(required=True, default=fields.Date.today, string='Rules effective on')
    preview_html = fields.Html(compute='_compute_preview', sanitize=False)

    @api.depends('config_id', 'asof')
    def _compute_preview(self):
        for wizard in self:
            wizard.preview_html = ''
            if not wizard.config_id or not wizard.asof:
                continue
            try:
                rows = wizard.config_id._tax_refresh_plan(wizard.asof)
                wizard.preview_html = Markup('<p>%s</p><table class="table"><thead><tr><th>%s</th><th>%s</th><th>%s</th></tr></thead><tbody>%s</tbody></table>') % (
                    escape(_('Review the database values before updating this draft. Active schemes and completed payrolls are preserved.')),
                    escape(_('Constant')), escape(_('Current')), escape(_('Database')),
                    Markup('').join(Markup('<tr><td>%s</td><td>%s</td><td>%s</td></tr>') %
                                    (escape(r['name']), r['before'], r['after']) for r in rows))
            except UserError as exc:
                wizard.preview_html = Markup('<p class="text-danger">%s</p>') % escape(str(exc))

    def action_refresh(self):
        self.ensure_one()
        plan = self.config_id._tax_refresh_plan(self.asof)
        proposal = self.env['pb.statutory.proposal'].propose(
            'country_tax_refresh', _('Refresh statutory constants · %s', self.config_id.name),
            payload={'config_id': self.config_id.id, 'asof': str(self.asof), 'changes': plan},
            snapshot={'stamp': stamp(plan)}, target=self.config_id,
            facts={'rows_changed': {'value': len(plan), 'unit': ''},
                   'configs_affected': {'value': 1, 'unit': ''},
                   'effective_date': {'value': str(self.asof), 'unit': ''}})
        return notice(proposal.answer())


class CountryTaxProposal(models.Model):
    _inherit = 'pb.statutory.proposal'
    _proposal_kind_labels = {**PbStatutoryProposal._proposal_kind_labels,
                             'country_tax_publish': 'Publish country statutory table',
                             'country_tax_refresh': 'Refresh mapped payroll constants'}

    def _live_snapshot(self):
        if self.kind == 'country_tax_publish':
            return {'stamp': self.env['hr.payroll.tax.table'].browse(self.payload()['table_id'])._stamp()}
        if self.kind == 'country_tax_refresh':
            data = self.payload()
            return {'stamp': stamp(self.env['hr.formula.config'].browse(data['config_id'])._tax_refresh_plan(data['asof']))}
        return super()._live_snapshot()

    def _proposal_rows(self):
        if self.kind == 'country_tax_refresh':
            data = self.payload()
            return [(_('Rules effective on'), '', data['asof'])] + [
                (r['name'], r['before'], r['after']) for r in data.get('changes', [])
                if r['before'] != r['after'] or r['old_table_id'] != r['table_id']]
        if self.kind == 'country_tax_publish':
            table = self.env['hr.payroll.tax.table'].browse(self.payload()['table_id'])
            return [(_('Country'), '', table.country_code), (_('Version'), '', table.version),
                    (_('Effective from'), '', str(table.effective_from)),
                    (_('Method'), '', table.method), (_('Basis'), '', table.basis)] + [
                ('%s %s' % (r['slot'], r['field']), '', r['value'])
                for r in self.payload().get('values', [])]
        return super()._proposal_rows()

    def _apply_country_tax_publish(self):
        table = self.env['hr.payroll.tax.table'].browse(self.payload()['table_id'])
        table._validate_edition()
        table.with_context(**{_APPLY_KEY: _APPLY_TOKEN}).write({'state': 'published'})
        return {'table_id': table.id}

    def _apply_country_tax_refresh(self):
        data = self.payload()
        return self.env['hr.formula.config'].browse(data['config_id']).with_context(
            **{_APPLY_KEY: _APPLY_TOKEN})._apply_tax_refresh(data['asof'])
